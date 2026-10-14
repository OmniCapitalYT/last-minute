#!/usr/bin/env python3
"""Peer medians of year-on-year revenue growth by industry (IND.grw), revision 3 (October 2026): Chase's small-business
median plus a measured size bridge for the level; Census and BLS industry deviations, pooled with the Fiserv Small
Business Index where Fiserv has the industry's own cell, shrunk by group-specific persistence; every SE carries the
uncertainty common to all industries (the level, the size bridge and selection into sale).

Run:  python3 -I /tmp/claude-0/fit6/fitscripts/peer_growth.py          (about 65 s)
Out:  /tmp/claude-0/fit6/out/peer_growth/RESULT.json (key js_table = MODELS.fitted.peer_grw; js_line = the exact line to
      embed), js_line.txt, estimates.csv, observations.csv, bridge.csv, persistence.csv, online_sources.csv,
      sbcs_all_firms.csv, fiserv_annual.csv, fiserv_cells.csv, fiserv_persistence.csv, btos_nb.csv, size_bridge.csv,
      nonstore_454_mismatch.csv. Earlier rounds are kept beside it, unchanged: fitscripts/peer_growth.v1.py with
      out/peer_growth.v1/ (round 1) and fitscripts/peer_growth.v2.py with out/peer_growth.v2/ (revision 2). Both copies
      still name out/peer_growth/ as their output directory, so run them only from a copy with OUT changed.

Why
  The console z-scores the user's nominal year-on-year revenue growth against a per-industry peer median,
  z.grw = (grw - IND.grw)/.13; the z-score feeds the quality composite (and through it model value, price, closing odds
  and both scores), the Growth sub-score and the buyer fit functions. The 16 medians were set by hand (mean 8.6%).
  No source publishes the median growth of small businesses that are being sold, so the medians are built from the
  closest published quantities, each step with a stated, tested bridge. Round 1 shipped M + 0.43 * D with M = 5% from one
  rounded Chase figure. Revision 2 applied the analyst's five decisions on the completeness critique: Fiserv's Small
  Business Index tested M (it does not support 5% at face value: +1.2% YoY in Feb 2026) and was pooled into six industry
  deviations; group lambdas replaced the pooled one; saas and ecom ship only if their population-relevant sources agree
  within 2 SE (saas did not, ecom did); restaurants keep Chase's direct median; a BTOS size bridge S (+1.7 pts) was
  applied. Its adversarial review accepted it with changes. Revision 3 applies the analyst's seven decisions on that review:
  1. Level SE. M stays 5% (Chase's median of continuing firms is the construct the console needs; Fiserv's index measures
     aggregate point-of-sale sales including entry and exit), but SE_M is widened from 1.25 to 2.0 pts to cover the Fiserv
     conflict (Fiserv 3.8 pts below M: z -3.1 on 1.25 pts, -1.9 on 2.0), and selection into sale (+-2 pts, judgment, common
     to all industries) is added in quadrature. With S's SE (0.65) the common SE is 2.9 pts; shipped SEs are now .030-.037.
  2. lambda, non-manufacturing = 0.48: the mean of the three adjacent pre-COVID AIES windows (2014-16 -> 2017-19 .53,
     2015-17 -> 2018-19 .49, 2015-16 -> 2017-18 .43). SE = .10 (analyst; the bootstrap SE of the three-window mean is .06) and
     the SD of all six windows (.057) in quadrature = .12. Revision 2's window (2017-19 -> 2023-24, .56) spans COVID.
     Manufacturing stays .12 (2017-19 -> 2023-24); its SE also adds the SD of its six windows (.13): .09 -> .15.
  3. Staffing: no override. The formula value ships: M + S + lambda D = -0.1% (rounds to .000, as the old override). Its SE,
     .037, exceeds .03, so it is flagged weak. (Its own SE, .023, is the largest of the 15 shipped industries; the common
     terms alone are .029, so every industry's SE is now near or above .03.)
  4. Empire Flippers grouping: seller types matched exactly and case-insensitively. E-commerce = the gm fit's product
     sellers (Amazon FBA, Amazon FBM, eCommerce, DropShipping, Subscription-Box): -0.5% (n 184), not -2.0% (n 169);
     revision 2 missed 'DropShipping' (case) and 'Amazon FBM'. Software = SaaS and (non-box) Subscription.
  5. ecom: Empire Flippers is measured against M, without the size bridge (its sellers are small: median revenue $342k);
     BTOS 454 carries a 2.5-pt cell-mismatch SE (the analyst's 2-3 pts; 454 also holds vending and direct selling, 35% of its
     firms). The agreement test compares the sources' levels at the console's size and counts the SEs of M, S and selection
     only where two sources differ in them. They agree (z -0.6): ecom ships at .055, flagged weak. The unrounded .0527 is
     .0002 above a rounding edge: a 3-pt term gives .050 (2 pts and none give .055).
  6. Fiserv error variance: the Fiserv-AIES gaps of the 51 matched cells correlate .56 (bootstrap SE .11) between 2023 and
     2024 (.51 for Fiserv's own error, AIES's sampling error removed), so a 4-year mean is not V_F2 / 2. With errors
     equicorrelated across years, V_F4 = V_F2 (1 + 3 rho) / (2 (1 + rho)) = .84 V_F2: error SD 7.9 pts, not 6.1. Fiserv's
     pool weights fall from 3-15% to 2-9%.
  7. Size-bridge note reworded: S is conservative, because Chase's median customer ($125k revenue in 2025, JPMCI Denver
     snapshot page) may be a nonemployer, smaller than the 1-4-employee class that BTOS starts at.
  saas stays judgment at .25 (max |z| 5.8; 4.5 with a 2.5-pt mismatch term on BTOS 511 too).
  Values that moved against revision 2: vet .085 -> .080, wash .090 -> .085, fit .080 -> .075, ecom .050 -> .055; staff
  ships the formula (.000) instead of the override (.000); the SEs rose from .015-.021 (staff's override .030) to .030-.037.

What ships (MODELS.fitted.peer_grw)
  grw_i = M + S + lambda_g * D_i, rounded to the nearest 0.005:
  M = 5.0% (SE 2.0 pts): median YoY revenue growth of Chase small-business customers, Feb 2026 (JPMCI Denver snapshot).
  S = +1.7 pts (SE 0.65): size bridge from Chase's typical customer (the 1-4-employee class, or smaller) to the console's
      (10-49 employees); conservative.
  D_i = deviation of industry i from the all-industry aggregate (Census AIES, QSS, MRTS; BLS QCEW bridged), round 1's
        estimator, pooled with Fiserv's 2023-26 deviation for medical, vet, dist, wash, auto and fit.
  lambda_g = 0.48 (non-manufacturing), 0.12 (manufacturing).
  Restaurants: M + S + lambda * (2.0% - M). Staffing: the formula value, weak. ecom: M + S + lambda * (pooled Empire
  Flippers and BTOS 454 deviation), weak. saas: not shipped (judgment .25, evidence in not_shipped).
  SE^2 = (lambda se_D)^2 + (D se_lambda)^2 + (lambda * 1.5 pts)^2 + SE_M^2 + SE_S^2 + SE_SEL^2 (the last three common).

What it does
  0. Data. Reads every download directly (paths below; zip members are read in memory; the Fiserv API responses were
     saved as JSON; nothing downloaded is executed). The snapshot's index.html is read for the current IND.grw.
  1. Level M = 5.0% (Chase), SE 2.0 pts (analyst, revision 3; 1.25 before); selection into sale SE 2.0 pts (analyst, judgment).
  2. Comparators, 3. payroll-to-revenue bridge, 5. Census deviations: round 1's code.
  4. AIES persistence: six windows for each group (non-manufacturing leaf industries, manufacturing six-digit industries).
     Non-manufacturing lambda = the mean of the three adjacent pre-COVID windows (bootstrap SE of the mean, industries
     resampled); manufacturing = the 2017-19 -> 2023-24 window (bootstrap and HC1 SEs). Both SEs add the six-window SD.
  6. Fiserv. National: monthly SA YoY (the published figure; the API reproduces the Sep 2026 report's Figure 2 to 0.05 pt)
     and annual NSA growth (sums of the monthly index; a cell-year with a missing month is dropped) -> the M test (on the
     1.25-pt SE it was decided on, and on the widened 2.0) and the restaurant offset. Industries: annual NSA growth 2020-25
     and Jan-Sep 2026; console deviations (2023-26 mean against Fiserv's all-industry). Persistence: 2023-24 -> 2025-26 and
     lags for the six-digit leaves (bootstrap SE) and the subsectors. Agreement with Census: Fiserv cells matched to AIES
     codes (same years 2023-24; AIES 2023-24 predicting Fiserv 2025-26) and the console industries (Census D against
     Fiserv D). Fiserv's error variance = var(Fiserv - AIES) over the matched cells minus AIES's own sampling variance
     (2-year mean), carried to the 4-year mean with the measured year-to-year correlation of the gaps (revision 3).
  7. BTOS. Revenue net balance (share increased minus share decreased; two weeks against the two weeks before) by
     employment size, sector, subsector and sector x size. Calibration: within-year sector deviations of QCEW payroll
     growth x 1.10 on BTOS sector net balance, 2024 and 2025 stacked, clustered by sector (checks: by year, leave one
     sector out, against Fiserv). Size bridge S = slope x (mean of 10-19 and 20-49 employees minus 1-4 employees), 2025-26;
     applied if S > 2 SE. Checks: S against all employer firms, by sector, in 2023-24; SBCS 12-month revenue change by
     employment size (2023-25 surveys, ordered-probit latent medians); Chase payroll net shares by revenue size (Jan 2026
     Pulse chart, mapped to medians with Chase's own payroll sd); BTOS-implied industry deviations against Census D.
  8. Pooling and forward values (see What ships). A 2026 nowcast (lambda = 1) is reported, not shipped.
  9. Restaurants: the Chase rule, re-checked against Fiserv's 722 deviation in Aug 2026 (single-month sd from the last 12
     months) and BTOS 722; the Census-pooled alternative is stored. The out-of-sample check (Census nowcast against Chase +
     S) leaves selection out of its SE: both sides measure all continuing firms.
 10. Staffing: the formula value (no override); weak if its SE exceeds .03 (analyst). The deviation of the latest readings
     alone (QCEW 2025, QSS 2026H1) and the value it would give are stored.
 11. saas and ecom: pairwise z of the population-relevant sources on their levels carried to the console's size (Empire
     Flippers + S, since it is measured against M; Acquire and SaaS Capital as published; BTOS M + S + D); the variance
     counts each source's own SE and the SEs of M, S and selection only where the two sources differ in them. If all
     |z| <= 2, ship M + S + lambda * the inverse-variance pooled deviation (own SEs only; the common SEs are added once),
     else keep the hand-set value as judgment. 11b: evidence on the 454 mismatch (SUSB 2022 firm shares; growth of 454
     less 4541 in AIES 2013-24 and MRTS 2025-26) and ecom sensitivities (mismatch 0/2/3 pts, EF base, EF grouping).
 12. Checks (round 1's plus fiserv and btos), sensitivity (no S, S against all employer firms, Census-only D, Fiserv error
     variance not reduced or halved, lambda .56 (revision 2) and .43 (round 1), lambda +-1 SE, Fiserv's lambda 0.16),
     rounding edges (values within 0.0005 of a rounding boundary), tests (assertions printed and stored).
 13. Output. Values rounded to the nearest 0.005 (half up); SEs to 0.001; lambdas to 0.01.

Data and fetch (each source in its own directory under /tmp/claude-0/fit6/data/B-grw/; read with python3 -I)
  round 1 (unchanged):
  census_aies/ts/AIESTIMESERIES.zip   https://www2.census.gov/programs-surveys/aies/data/2024/API_Datasets/AIESTIMESERIES.zip
  census_aies/z2023/AIES00BASIC.zip   https://www2.census.gov/programs-surveys/aies/data/2023/AIES00BASIC.zip
  census_aies/z2024/AIES00BASIC.zip   https://www2.census.gov/programs-surveys/aies/data/2024/AIES00BASIC.zip
  census_qss/qss-all-current.xlsx     https://www.census.gov/services/qss/qss-all-current.xlsx   (2026Q2 release)
  census_mrts/mrtssales92-present.xlsx https://www.census.gov/retail/mrts/www/mrtssales92-present.xlsx
  census_susb/us_state_6digitnaics_{2017,2022}.xlsx  https://www2.census.gov/programs-surveys/susb/tables/{2017,2022}/
  bls_qcew/{year}_{code}.csv          https://data.bls.gov/cew/data/api/{2022..2025}/a/industry/{code}.csv (send a
                                      descriptive User-Agent; BLS answers 403 otherwise)
  fed_sbcs/*appendix*.xlsx            https://www.fedsmallbusiness.org (report pages 2017-2026, "appendix" workbooks)
  jpmci/denver-snapshot.pdf           https://www.jpmorganchase.com/content/dam/jpmorganchase/institute/documents/denver-snapshot.pdf
  chase_sb_pulse/{smb-pulse-january-2026,april-sb-pulse2026,may-sb-pulse2026,sb-pulse-sep-2026}.pdf
                                      https://media.chase.com/content/dam/chase/media-center/intelligence/<name>.pdf
  lit_relbd/carra-wp-2017-03.pdf      https://www.census.gov/content/dam/Census/library/working-papers/2017/adrm/carra-wp-2017-03.pdf
  saas_capital/RB33WS5-2025.pdf       https://www.saas-capital.com/wp-content/uploads/2025/06/RB33WS5-2025-B2B-Benchmarking-Private-SaaS-Company-Growth-Rates.pdf
  local: /tmp/claude-0/data5/empireflippers/clean/{ef_listings,ef_sold_derived}.csv; /tmp/claude-0/data5/kaggle/clean/
         acquire_saas_listings_2023.csv
  new in revision 2:
  fiserv_sbi/api/{adjusted,nsa}_US_all.json and {adjusted,nsa}_US_{23,42,44,45,48,49,51,53,54,56,61,62,71,72,81}_n6.json
      the public data API behind fiserv.com/en/fiserv-small-business-index (no key; the page's config says APIGEE off):
      curl -X POST -H 'Content-Type: application/json' -H 'Accept: application/json'
        --data '{"geo":"US","naicsCode2":"81","startPeriod":"2019-01-01","endPeriod":"2026-09-01","isInflation":false,
                 "includeNaics6":true}' https://fsbi.fiservapps.com/index-data/api/v1/index-details/adjusted
      (seasonally adjusted; .../index-details without /adjusted = not adjusted; omit naicsCode2 for the national file)
  fiserv_sbi/fsbi-september-2026-monthly-overview-1026.pdf   https://www.fiserv.com/content/dam/fiserv-ent/final-files/
      marketing-collateral/documents-handouts-flyers/fsbi-september-2026-monthly-overview-1026.pdf (methodology, Figure 2)
  census_btos/{National,Employment Size Class,Sector,Subsector,Sector by Employment Size Class}.xlsx
      https://www.census.gov/hfp/btos/downloads/<name>.xlsx (spaces as %20), data through collection period 202620
  used in revision 3 (downloaded in the research round):
  jpmci/institute_all-topics_business-growth-and-entrepreneurship_denver-small-business-snapshot.html
      https://www.jpmorganchase.com/institute/all-topics/business-growth-and-entrepreneurship/denver-small-business-snapshot
      (Figure 2 text: "the median small business generated $125,200 in revenues during 2025"; the size-bridge note)
  Revision 3 also reads more of files already used: MRTS rows 454 (nonstore retailers) and 45431 (fuel dealers), AIES
  4542/4543 and SUSB 2022 454 sub-industry firm counts (the 454 cell-mismatch evidence).
  Numbers read off charts are constants below with file, page and figure. Text quotes are confirmed in the PDF text with
  pdftotext (poppler) when it is installed, and in downloaded pages as text (tags stripped).
"""
import glob, hashlib, io, json, math, os, re, shutil, subprocess, zipfile
import numpy as np
import pandas as pd
import openpyxl
import statsmodels.api as sm
from scipy import stats

OUT = '/tmp/claude-0/fit6/out/peer_growth'
DATA = '/tmp/claude-0/fit6/data/B-grw'
AIES_TS = DATA + '/census_aies/ts/AIESTIMESERIES.zip'
AIES_BASIC = {2023: DATA + '/census_aies/z2023/AIES00BASIC.zip', 2024: DATA + '/census_aies/z2024/AIES00BASIC.zip'}
QCEW_DIR = DATA + '/bls_qcew'
QSS_XLSX = DATA + '/census_qss/qss-all-current.xlsx'
MRTS_XLSX = DATA + '/census_mrts/mrtssales92-present.xlsx'
SUSB = {2017: DATA + '/census_susb/us_state_6digitnaics_2017.xlsx', 2022: DATA + '/census_susb/us_state_6digitnaics_2022.xlsx'}
SBCS_DIR = DATA + '/fed_sbcs'
DENVER_PDF = DATA + '/jpmci/denver-snapshot.pdf'
DENVER_HTML = DATA + '/jpmci/institute_all-topics_business-growth-and-entrepreneurship_denver-small-business-snapshot.html'
PULSE_PDF = {'jan': DATA + '/chase_sb_pulse/smb-pulse-january-2026.pdf', 'apr': DATA + '/chase_sb_pulse/april-sb-pulse2026.pdf',
             'may': DATA + '/chase_sb_pulse/may-sb-pulse2026.pdf', 'sep': DATA + '/chase_sb_pulse/sb-pulse-sep-2026.pdf'}
CARRA_PDF = DATA + '/lit_relbd/carra-wp-2017-03.pdf'
SAASCAP_PDF = DATA + '/saas_capital/RB33WS5-2025.pdf'
EF_LISTINGS = '/tmp/claude-0/data5/empireflippers/clean/ef_listings.csv'
EF_SOLD = '/tmp/claude-0/data5/empireflippers/clean/ef_sold_derived.csv'
ACQUIRE = '/tmp/claude-0/data5/kaggle/clean/acquire_saas_listings_2023.csv'
FISERV_API = DATA + '/fiserv_sbi/api'
FISERV_PDF = DATA + '/fiserv_sbi/fsbi-september-2026-monthly-overview-1026.pdf'
FISERV_FAQ = DATA + '/fiserv_sbi/api/faqs.json'   # GET https://fsbi.fiservapps.com/index-data/api/faqs (methodology notes)
FISERV_CODES2 = ['23', '42', '44', '45', '48', '49', '51', '53', '54', '56', '61', '62', '71', '72', '81']
BTOS_DIR = DATA + '/census_btos'
BTOS_FILES = {'national': 'National.xlsx', 'size': 'Employment Size Class.xlsx', 'sector': 'Sector.xlsx',
              'subsector': 'Subsector.xlsx', 'sector_size': 'Sector by Employment Size Class.xlsx'}
INDEX = '/tmp/claude-0/fit6snap/index.html'          # frozen snapshot, read only

# ---------------------------------------------------------------- documented constants (judgment or read by eye)
M = 0.050            # Chase national median YoY revenue growth, Feb 2026 (Denver snapshot p. 2: "slightly below the national rate of 5%")
M_DENVER = 0.040     # same sentence: "Denver small businesses saw median revenue growth of 4% year-over-year as of February 2026"
SE_M = 0.020         # analyst (revision 3): widened from 1.25 pts to cover the Fiserv conflict (Fiserv all-industry +1.2% YoY in Feb 2026)
SE_M_R2 = 0.0125     # rounds 1-2 judgment (definition: deposit inflows; rounding of the quoted 5%); kept for the M-test comparison
SE_SEL = 0.020       # analyst (revision 3): selection into sale (businesses being sold against all continuing businesses), +-2 pts,
                     # judgment, common to all industries, added in quadrature
WEAK_SE = 0.030      # analyst (revision 3): staffing ships the formula value, flagged weak if its SE exceeds .03
SE_BRIDGE = 0.015    # aggregate-to-median deviation bridge; checked against Chase medians in check 'median_vs_aggregate'
FLOOR = 0.015        # floor on the SE of one AIES industry-year growth rate (frame and vintage breaks)
SE_AIES_NOCV = 0.025 # SE of an AIES growth rate with no published CV pair (the 2023 observations)
QSS_EXTRA = 0.015    # added to QSS SEs: half-year YoY against annual growth
SE_MRTS = 0.020      # MRTS kind-of-business against the console industry, and month-sum against annual
REC = {2023: 0.5, 2024: 1.0, 2025: 2.0, 2026: 2.0}   # recency weights, nowcast only
# Chase Small Business Pulse, Sep 2026, p. 3, chart "Median inflow growth: August 2023 vs. August 2026" (read off the
# dot plot, +-0.1 pt): Restaurants and Other Eating Places 5.9% (Aug 2023) -> 2.0% (Aug 2026).
CHASE_REST = {'2023-08': 0.059, '2026-08': 0.020}
SE_CHASE_DEV = 0.010 # judgment: restaurants read in Aug 2026 against the national median of Feb 2026
STAFF_OVERRIDE_R2 = (0.000, 0.030)   # rounds 1-2 override, dropped in revision 3 (analyst: no override; comparison only)
CURRENT_JUDGMENT_KEYS = ('saas', 'ecom')   # the two industries that ship only if their population-relevant sources agree
SE_454_MISMATCH = 0.025   # analyst (revision 3): cell-mismatch SE on the BTOS 454 deviation, 2-3 pts; the midpoint is used because the
                          # measured evidence (11b) points both ways: 35% of 454 firms are not e-commerce (SUSB 2022), and that share
                          # times their growth gap to 4541 had an RMS of 3.5-3.8 pts in AIES 2013-24 (fuel dealers), but is under
                          # 0.3 pt in the BTOS window (MRTS 2025-26); the definitional gap (online sellers vs 4541) is not measured
# Empire Flippers seller types, matched case-insensitively. EF_PROD is the gm fit's product-seller set (peer_gross_margin.py,
# EF_PROD: Amazon FBA, eCommerce, DropShipping, Amazon FBM, Subscription-Box); EF_SAAS: software and (non-box) subscriptions
EF_PROD = ('amazon fba', 'ecommerce', 'dropshipping', 'amazon fbm', 'subscription-box')
EF_SAAS = ('saas', 'subscription')
SAASCAP = {'median_all': 0.25, 'median_bootstrapped': 0.23, 'median_equity': 0.25, 'share_flat_or_negative': 0.069, 'n_min': 1000,
           'fig2': {'<$1M': (0.20, 0.40, 1.00), '$1-3M': (0.14, 0.28, 0.50), '$3-5M': (0.10, 0.24, 0.40),
                    '$5-10M': (0.13, 0.24, 0.42), '$10-20M': (0.08, 0.20, 0.37), '>$20M': (0.10, 0.20, 0.31)}}
RELBD_FIG5 = {'5': (0.38, -0.31), '6': (0.38, -0.31), '7': (0.35, -0.30), '8': (0.34, -0.29), '9': (0.35, -0.27),
              '10': (0.34, -0.27), '11': (0.31, -0.27), '12': (0.31, -0.26), '13': (0.31, -0.25), '14': (0.29, -0.245),
              '15': (0.32, -0.25), '16+': (0.23, -0.17)}
CHASE_PAYROLL = {'apr': (0.06, 0.23), 'may': (0.07, 0.25), 'sep': (0.05, 0.18)}
# Chase Small Business Pulse, Jan 2026, p. 3, chart "Net share of employer businesses increasing payroll YOY" (2025
# points, Jan..Dec, read off a 200-dpi render, +-0.3 pt): larger (> $1M yearly revenue), all, smaller (< $1M).
CHASE_PAYROLL_SIZE_2025 = {'larger': [33.5, 29.1, 30.7, 30.5, 26.4, 30.0, 30.8, 27.9, 30.2, 27.4, 27.1, 29.0],
                           'all': [26.3, 23.0, 24.3, 24.0, 20.6, 23.8, 24.3, 22.0, 24.0, 21.8, 21.1, 23.0],
                           'smaller': [22.1, 19.6, 20.6, 20.3, 17.2, 20.3, 20.6, 18.6, 20.6, 18.6, 17.7, 19.8]}
SPREAD_KEPT = 0.13
SPREAD_CENTRAL = 0.18
# analyst (revision 3): non-manufacturing lambda = mean of the adjacent pre-COVID AIES windows (the main window spans COVID);
# manufacturing stays at the main window (.12). SE = sampling SE (analyst: .10 for non-manufacturing; the measured window SE
# for manufacturing) plus the across-window SD of the six windows, in quadrature.
LAM_WINDOWS_NONMFG = ('2014-16 -> 2017-19 (pre-COVID, adjacent)', '2015-17 -> 2018-19 (pre-COVID)', '2015-16 -> 2017-18 (2-year, adjacent)')
SE_LAM_SAMPLING_NONMFG = 0.10   # analyst (revision 3); the bootstrap SE of the three-window mean is computed as a check
LAM_R2 = {'non_manufacturing': 0.56, 'manufacturing': 0.12}         # revision 2 (comparison only)
LAM_POOLED_R1 = 0.43 # round 1's pooled lambda (comparison only)
BR_DECISION = 1.10   # the analyst's decision (round 1): the payroll-to-revenue deviation slope
Z_AGREE = 2.0        # analyst: saas and ecom ship only if their population-relevant sources agree within 2 SE
Z_APPLY = 2.0        # analyst: apply the size bridge if it is clearly non-zero (> 2 SE)
SIZE_REF = ('C', 'D')   # BTOS 10-19 and 20-49 employees: the console's typical business (SDE about $0.7M, ~$2-4M revenue)
SIZE_CHASE = 'A'        # BTOS 1-4 employees: the closest employer class to Chase's median customer ($125k revenue)
BTOS_WINDOW = (202501, 202699)    # collection periods 2025-26 (BTOS data end at 202620; 202521-23 were not collected)
FISERV_YEARS = [2023, 2024, 2025, 2026]   # Fiserv deviation window; 2026 = Jan-Sep against Jan-Sep 2025
# console industry -> Fiserv cell. POOL: the industry's own NAICS cell (or the same sector cell Census uses, dist = 42;
# medical = 621, which holds both physicians and dentists). REPORT_ONLY: a cell shared by two console industries or a
# catch-all, reported but not pooled. mfg and ecom have no Fiserv cell (Fiserv covers no manufacturing and no nonstore).
FISERV_POOL = {'medical': '621', 'vet': '541940', 'dist': '42', 'wash': '811192', 'auto': '811111', 'fit': '713940'}
FISERV_REPORT_ONLY = {'home': '238000', 'constr': '238000', 'msp': '541OTH', 'prof': '541OTH', 'staff': '561000',
                      'land': '561000', 'rest': '722511', 'saas': '51XXXX'}
# Fiserv cell (NAICS 2022) -> AIES time-series code (NAICS 2017): every Fiserv cell with a matching AIES series
FISERV_AIES = {'423430': '42343', '423450': '42345', '424000': '424', '444140': '44413', '445110': '44511', '449121': '44221',
               '449210': '443142', '455000': '452', '456130': '44613', '457000': '447', '458110': '4481', '458310': '44831',
               '459110': '45111', '459120': '45112', '459420': '45322', '459991': '453991', '484110': '48411', '485310': '48531',
               '488410': '48841', '493110': '49311', '51XXXX': '51', '532000': '532', '541213': '541213', '541810': '54181',
               '541820': '54182', '541910': '54191', '541940': '54194', '561000': '561', '621310': '62131', '621320': '62132',
               '713940': '71394', '721110': '72111', '722320': '72232', '722511': '722511', '722513': '722513',
               '811111': '811111', '811121': '811121', '811192': '811192', '811412': '811412', '812112': '812112',
               '812210': '81221', '812320': '81232', '812930': '81293', '441': '441', '444': '444', '445': '445', '621': '621',
               '722': '722', '811': '811', '812': '812', '813': '813', '42': '42', '423': '423'}
FISERV_SERVICES = ('48', '49', '53', '54', '56', '61', '62', '71', '72', '81')
# console industry -> BTOS cell (subsector, or sector where Census uses the sector) for the ordering check
BTOS_CELL = {'home': ('subsector', '238'), 'constr': ('subsector', '238'), 'medical': ('subsector', '621'), 'vet': ('subsector', '541'),
             'msp': ('subsector', '541'), 'prof': ('subsector', '541'), 'mfg': ('sector', '31'), 'dist': ('sector', '42'),
             'staff': ('subsector', '561'), 'land': ('subsector', '561'), 'wash': ('subsector', '811'), 'auto': ('subsector', '811'),
             'fit': ('subsector', '713'), 'rest': ('subsector', '722'), 'saas': ('subsector', '511'), 'ecom': ('subsector', '454')}
BTOS_QCEW = {'23': '23', '31': '31-33', '42': '42', '44': '44-45', '54': '54', '56': '56', '62': '62', '71': '71', '72': '72', '81': '81'}
MFG2 = ('31', '32', '33')
B_BOOT = 2000
SEED = 20261014

INDS = ['home', 'medical', 'vet', 'msp', 'saas', 'prof', 'mfg', 'dist', 'ecom', 'staff', 'land', 'wash', 'auto', 'constr', 'fit', 'rest']
MAIN = [k for k in INDS if k not in CURRENT_JUDGMENT_KEYS]
GROUP = {k: ('manufacturing' if k == 'mfg' else 'non_manufacturing') for k in INDS}
IND = {'home': (['238*'], ['2382'], None, None), 'medical': (['6211'], ['6211'], ['6211'], None),
       'vet': (['54194'], ['541940'], None, None), 'msp': (['5415'], ['5415'], ['5415'], None),
       'saas': (['5112'], ['5132'], ['5112'], None), 'prof': (['541211', '54111'], ['541211', '541110'], ['5412', '5411'], None),
       'mfg': (['31-33'], ['31-33'], None, None), 'dist': (['42'], ['42'], None, None), 'ecom': (['4541'], [], None, '4541'),
       'staff': (['56132'], ['561320'], ['5613'], None), 'land': (['56173'], ['561730'], None, None),
       'wash': (['811192'], ['811192'], None, None), 'auto': (['811111'], ['811111'], ['8111'], None),
       'constr': (['238*'], ['2389'], None, None), 'fit': (['71394'], ['713940'], ['7139'], None),
       'rest': (['722511'], ['722511'], None, '722')}
CVMAP = {'54194': '541940', '56132': '561320', '56173': '561730', '71394': '713940', '54111': '541110', '5112': '511210'}
SUSB_CODE = {'home': '238220', 'medical': '621111', 'vet': '541940', 'msp': '5415', 'saas': '511210', 'prof': '541211',
             'mfg': '31-33', 'dist': '42', 'ecom': '454110', 'staff': '561320', 'land': '561730', 'wash': '811192',
             'auto': '811111', 'constr': '238990', 'fit': '713940', 'rest': '722511'}
SECTORS = ['21', '22', '23', '31-33', '42', '44-45', '48-49', '51', '52', '53', '54', '55', '56', '61', '62', '71', '72', '81']
QCEW_SECTORS = ['23', '31-33', '42', '44-45', '54', '56', '62', '71', '72', '81']
QSS_SECTORS = ['54', '56', '62', '71', '81']
BRIDGE_PAIRS = {'6211': '6211', '6212': '6212', '54194': '541940', '5415': '5415', '541511': '541511', '541512': '541512',
                '5112': '5132', '541211': '541211', '5412': '5412', '54111': '541110', '31-33': '31-33', '339950': '339950',
                '42': '42', '423': '423', '56132': '561320', '56173': '561730', '811192': '811192', '811111': '811111',
                '811121': '811121', '71394': '713940', '722511': '722511', '722513': '722513', '44-45': '44-45', '72': '72',
                '81': '81', '62': '62', '71': '71', '56': '56'}


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def sha_files(paths):
    h = hashlib.sha256()
    for p in sorted(paths):
        with open(p, 'rb') as f:
            h.update(f.read())
    return h.hexdigest()


def r005(x):
    """nearest 0.005, half up; returned with 3 decimals (no negative zero)"""
    v = math.floor(x * 200 + 0.5) / 200
    return round(v, 3) + 0.0


def clean(o):
    """JSON-safe, deterministic: numpy scalars to Python, NaN to None, floats to 6 decimals"""
    if isinstance(o, dict):
        return {str(k): clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [clean(v) for v in o]
    if isinstance(o, (np.integer,)):
        return int(o)
    if isinstance(o, (float, np.floating)):
        f = float(o)
        return None if not math.isfinite(f) else round(f, 6) + 0.0
    if isinstance(o, np.bool_):
        return bool(o)
    return o


def pdf_text(path, first=None, last=None, layout=False):
    """pdftotext (poppler) in reading order, whitespace collapsed (layout=True keeps lines); None when poppler is missing"""
    exe = shutil.which('pdftotext')
    if exe is None:
        return None
    cmd = [exe] + (['-f', str(first), '-l', str(last)] if first else []) + (['-layout'] if layout else []) + [path, '-']
    r = subprocess.run(cmd, capture_output=True, timeout=120, check=False)
    t = r.stdout.decode('utf-8', 'replace')
    return t if layout else re.sub(r'\s+', ' ', t)


def quote_check(path, patterns, first=None, last=None):
    t = pdf_text(path, first, last)
    if t is None:
        return {'checked': False, 'reason': 'pdftotext not installed'}
    return {'checked': True, 'found': {p: bool(re.search(p, t)) for p in patterns}}


def text_check(path, patterns):
    """regex search in a downloaded text file (JSON/HTML read as text, tags stripped, whitespace collapsed)"""
    t = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', open(path, encoding='utf-8', errors='replace').read()))
    return {'checked': True, 'found': {p: bool(re.search(p, t)) for p in patterns}}


def read_pipe_zip(path):
    """the .dat member of a Census bulk zip (pipe-delimited, header starts with '#'), read in memory"""
    z = zipfile.ZipFile(path)
    member = [n for n in z.namelist() if n.endswith('.dat')]
    assert len(member) == 1, member
    d = pd.read_csv(io.BytesIO(z.read(member[0])), sep='|', dtype=str, keep_default_na=False, na_values=[''])
    d.columns = [c.lstrip('#') for c in d.columns]
    return d[(d.GEOTYPE == '01') & (d.TYPOP == '00') & (d.TAXSTAT == '00')].copy()


def ols_hc1(x, y):
    X = sm.add_constant(np.asarray(x, float))
    f = sm.OLS(np.asarray(y, float), X).fit(cov_type='HC1')
    return float(f.params[1]), float(f.bse[1]), f


def boot_slope(x, y, seed=SEED, B=B_BOOT):
    rng = np.random.default_rng(seed); x = np.asarray(x, float); y = np.asarray(y, float); n = len(x); bs = np.empty(B)
    for i in range(B):
        idx = rng.integers(0, n, n)
        bs[i] = np.cov(x[idx], y[idx], ddof=1)[0, 1] / np.var(x[idx], ddof=1)
    return float(bs.std(ddof=1)), [float(np.quantile(bs, .025)), float(np.quantile(bs, .975))]


def agreement(x, y):
    """r, Spearman and both OLS slopes (with HC1 SEs) between two deviation vectors"""
    x = np.asarray(x, float); y = np.asarray(y, float)
    b_yx, se_yx, _ = ols_hc1(x, y); b_xy, se_xy, _ = ols_hc1(y, x)
    return dict(n=int(len(x)), r=float(np.corrcoef(x, y)[0, 1]), spearman=float(stats.spearmanr(x, y)[0]),
                slope_y_on_x=b_yx, se_y_on_x=se_yx, slope_x_on_y=b_xy, se_x_on_y=se_xy)


# ================================================================ 0. data
def load_aies():
    d = read_pipe_zip(AIES_TS)
    d['rev'] = pd.to_numeric(d.RCPT_TOT_VAL, errors='coerce'); d['YEAR'] = d.YEAR.astype(int)
    p = d.pivot_table(index='NAICS', columns='YEAR', values='rev', aggfunc='first')
    p = p.reindex(columns=range(int(p.columns.min()), int(p.columns.max()) + 1))
    labels = d.drop_duplicates('NAICS').set_index('NAICS').NAICS_LABEL
    basic = {}
    for y, path in AIES_BASIC.items():
        b = read_pipe_zip(path)
        b['rev'] = pd.to_numeric(b.RCPT_TOT_VAL, errors='coerce'); b['cv'] = pd.to_numeric(b.RCPT_TOT_CV, errors='coerce')
        basic[y] = b.set_index('NAICS2017')[['NAICS2017_LABEL', 'INDLEVEL', 'rev', 'cv']]
    j = basic[2023].join(basic[2024], lsuffix='23', rsuffix='24', how='inner')
    j['g'] = j.rev24 / j.rev23 - 1
    j['se_g'] = (1 + j.g) * np.sqrt((j.cv23 / 100) ** 2 + (j.cv24 / 100) ** 2)   # independent samples: an upper bound
    return p, labels, j, len(d)


def load_qcew():
    rows, files = [], sorted(glob.glob(os.path.join(QCEW_DIR, '*.csv')))
    for f in files:
        d = pd.read_csv(f, dtype={'area_fips': str, 'industry_code': str, 'own_code': str}, low_memory=False)
        d = d[(d.area_fips == 'US000') & (d.own_code == '5')]
        if d.empty:
            continue
        r = d.iloc[0]
        rows.append(dict(ind=r.industry_code, year=int(r.year), wages=float(r.total_annual_wages), emp=float(r.annual_avg_emplvl)))
    Q = pd.DataFrame(rows)
    P = Q.pivot_table(index='ind', columns='year', values='wages')
    out = pd.DataFrame({'w23': P[2023] / P[2022] - 1, 'w24': P[2024] / P[2023] - 1, 'w25': P[2025] / P[2024] - 1})
    return out, len(files), sha_files(files), len(Q)


def load_qss():
    wb = openpyxl.load_workbook(QSS_XLSX, read_only=True, data_only=True)
    t1 = [s for s in wb.sheetnames if s.startswith('table1a')][0]; t2 = [s for s in wb.sheetnames if s.startswith('table2a')][0]
    hdr1 = list(wb[t1].iter_rows(min_row=5, max_row=5, values_only=True))[0]
    hdr2 = list(wb[t2].iter_rows(min_row=6, max_row=6, values_only=True))[0]
    c2 = [i for i, h in enumerate(hdr1) if h and re.sub(r'\s+', ' ', str(h)).startswith('2Q 2026 from 2Q 2025')][0]
    c1 = [i for i, h in enumerate(hdr1) if h and re.sub(r'\s+', ' ', str(h)).startswith('1Q 2026 from 1Q 2025')][0]
    s2 = [i for i, h in enumerate(hdr2) if h and re.sub(r'\s+', ' ', str(h)).startswith('2Q 2026 from 2Q 2025')][0]
    s1 = [i for i, h in enumerate(hdr2) if h and re.sub(r'\s+', ' ', str(h)).startswith('1Q 2026 from 1Q 2025')][0]
    out = {}
    for sheet, cols, keys in [(t1, (c2, c1), ('yoy_2q26', 'yoy_1q26')), (t2, (s2, s1), ('se_2q26', 'se_1q26'))]:
        for r in wb[sheet].iter_rows(values_only=True):
            if r[0] is None:
                continue
            code = str(r[0]).strip()
            if not re.fullmatch(r'\d{2,6}', code):
                continue
            rec = out.setdefault(code, {})
            if keys[0] in rec:
                continue                                    # first row of a code only
            for c, k in zip(cols, keys):
                rec[k] = float(r[c]) if isinstance(r[c], (int, float)) else None
    return out, t1


def load_mrts():
    wb = openpyxl.load_workbook(MRTS_XLSX, read_only=True, data_only=True)
    want = {'Retail and food services sales, total': 'total', 'Electronic shopping and mail-order houses': '4541',
            'Food services and drinking places': '722', 'Nonstore retailers': '454', 'Fuel dealers': '45431'}
    out = {}
    for ws in wb.worksheets:
        yr = ws.title
        rows = list(ws.iter_rows(values_only=True))
        hdr = None
        for r in rows[:8]:
            if r and any(isinstance(c, str) and ('TOTAL' in c or 'CY CUM' in c) for c in r if c):
                hdr = r
        if hdr is None:
            continue
        hdr = [str(c) if c is not None else '' for c in hdr]
        adj = False
        for r in rows:
            if isinstance(r[1], str) and r[1].strip().startswith('ADJUSTED'):
                adj = True                                   # only the NOT ADJUSTED block (annual totals live there)
            if adj or not (isinstance(r[1], str) and r[1].strip() in want):
                continue
            k = want[r[1].strip()]
            if k in out.get(yr, {}):
                continue
            months = [c for c, h in zip(r, hdr) if re.match(r'^[A-Z][a-z]{2}\.? \d{4}', h)]
            tot = [c for c, h in zip(r, hdr) if h.strip() == 'TOTAL']
            out.setdefault(yr, {})[k] = (tot[0] if tot else None, months)
    g, lv = {}, {}
    for k in ['total', '4541', '722', '454', '45431']:
        for yr in ['2023', '2024', '2025']:
            g[(k, int(yr))] = out[yr][k][0] / out[str(int(yr) - 1)][k][0] - 1
            lv[(k, int(yr))] = out[yr][k][0]
        lv[(k, 2022)] = out['2022'][k][0]
        m26 = [x for x in out['2026'][k][1] if isinstance(x, (int, float))]
        m25 = out['2025'][k][1][:len(m26)]
        g[(k, 2026)] = sum(m26) / sum(m25) - 1
        lv[(k, 2026)] = sum(m26); lv[(k, '2025ytd')] = sum(m25)
    n_months_2026 = len([x for x in out['2026']['total'][1] if isinstance(x, (int, float))])
    return g, n_months_2026, lv


def load_susb(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    hdr, rows = None, []
    for row in wb.worksheets[0].iter_rows(values_only=True):
        if hdr is None:
            if row and row[0] == 'State':
                hdr = [str(c).replace('\n', ' ') for c in row]
            continue
        if row[0] == '00':
            rows.append(row)
        elif rows:
            break                                            # the US rows come first
    d = pd.DataFrame(rows, columns=hdr)
    d['size'] = d['Enterprise Size'].astype(str).str.slice(0, 2)
    for c in ['Firms', 'Receipts ($1,000)']:
        d[c] = pd.to_numeric(d[c], errors='coerce')
    return d


def load_sbcs():
    """SBCS employer firms, 'Revenue change, prior 12 months' (Decreased / No change / Increased): all firms, the eight
    industry groups and (revision 2) employment-size groups (logic of research sbcs_extract.py; 2016 and 2020 not used)"""
    Qn = 'Revenue change, prior 12 months'
    rows = []

    def add(year, gtype, group, resp, val, n=None):
        resp = {'Decreased': 'dec', 'No change': 'nc', 'Increased': 'inc'}.get(str(resp).strip())
        if resp is None or val in (None, '', ' '):
            return
        v = float(str(val).replace('%', '')) if not isinstance(val, (int, float)) else float(val)
        rows.append(dict(survey_year=year, gtype=gtype, group=str(group).strip(), resp=resp, val=v, n=n))

    def norm_ind(s):
        return re.sub(r'\s+', ' ', s.replace('Industry=', '').replace('&', 'and').strip())

    def norm_size(s):
        return re.sub(r'\s+', ' ', str(s).replace('–', '-').replace('—', '-').replace('Number of employees=', '').strip())
    for fn, yr in [('2018_sbcs-employer-firms-appendix.xlsx', 2017), ('2019_sbcs-employer-firms-appendix.xlsx', 2018)]:
        wb = openpyxl.load_workbook(os.path.join(SBCS_DIR, fn), read_only=True, data_only=True)
        for sh, gt in [('EmployerFirms', 'all'), ('EmployerFirms-by-Industry', 'industry')]:
            R = list(wb[sh].iter_rows(values_only=True))
            hdr, nrow = R[1], R[2]
            for i, r in enumerate(R):
                if r[0] == Qn:
                    for j in range(i + 1, i + 4):
                        for k in range(1, len(hdr)):
                            if hdr[k] is None and gt != 'all':
                                continue
                            if gt == 'all' and k > 1:
                                continue
                            grp = 'All employer firms' if gt == 'all' else norm_ind(hdr[k])
                            try:
                                n = int(str(nrow[k]).replace('N=', '').replace(',', ''))
                            except ValueError:
                                n = None
                            add(yr, gt, grp, R[j][0], R[j][k], n)
                    break
    wb = openpyxl.load_workbook(os.path.join(SBCS_DIR, '2019_sbcs-employer-firms-appendix-2019.xlsx'), read_only=True, data_only=True)
    for r in wb['NationalDataAndCrossTabs'].iter_rows(values_only=True):
        if r[2] != Qn:
            continue
        u = str(r[0])
        gt = 'all' if u == 'All firms' else 'industry' if u.startswith('Industry=') else None
        if gt:
            add(2019, gt, 'All employer firms' if gt == 'all' else norm_ind(u), r[3], r[4] / 100, r[5])
    wb = openpyxl.load_workbook(os.path.join(SBCS_DIR, '2021_sbcs-employer-firms-appendix-2021.xlsx'), read_only=True, data_only=True)
    for sh, gt in [('Employer firms', 'all'), ('Industry', 'industry')]:
        for r in wb[sh].iter_rows(values_only=True):
            if r[1] == Qn:
                add(2021, gt, 'All employer firms' if gt == 'all' else norm_ind(r[0]), r[2], float(r[3]) / 100, r[4])
    for fn, yr in [('2023_sbcs-employer-firms-appendix-2022.xlsx', 2022), ('2024_sbcs-employer-firms-appendix-2023.xlsx', 2023),
                   ('2025_sbcs-employer-firms-appendix-2024.xlsx', 2024), ('2026_sbcs-employer-firms-appendix-2025.xlsx', 2025)]:
        wb = openpyxl.load_workbook(os.path.join(SBCS_DIR, fn), read_only=True, data_only=True)
        allsh = 'Employer firms' if 'Employer firms' in wb.sheetnames else 'Employer Firms'
        for sh, gt in [(allsh, 'all'), ('Industry', 'industry'), ('Firm size (employees)', 'employees')]:
            for r in wb[sh].iter_rows(values_only=True):
                if r[1] == Qn:
                    grp = 'All employer firms' if gt == 'all' else (norm_ind(r[0]) if gt == 'industry' else norm_size(r[0]))
                    add(yr, gt, grp, r[2], float(r[3]), r[4])
    L = pd.DataFrame(rows)
    W = L.pivot_table(index=['survey_year', 'gtype', 'group'], columns='resp', values='val', aggfunc='first').reset_index()
    N = L.groupby(['survey_year', 'gtype', 'group']).n.first().reset_index()
    W = W.merge(N, on=['survey_year', 'gtype', 'group']).sort_values(['survey_year', 'gtype', 'group']).reset_index(drop=True)
    S = W[['dec', 'nc', 'inc']].sum(axis=1)
    for c in ['dec', 'nc', 'inc']:
        W[c] = np.where(S > 50, W[c] / 100, W[c])
    S = W[['dec', 'nc', 'inc']].sum(axis=1)
    for c in ['dec', 'nc', 'inc']:
        W[c] = W[c] / S                                       # renormalised shares
    z1 = stats.norm.ppf(1 - W.inc); z2 = stats.norm.ppf(W.dec)
    W['mu_s'] = -(z1 + z2) / 2                                # ordered probit: median / sd
    W['t_s'] = (z1 - z2) / 2                                  # half-width of the 'no change' band / sd
    nn = pd.to_numeric(W.n, errors='coerce')                  # delta-method SE of mu_s (shares treated as independent)
    W['se_mu_s'] = 0.5 * np.sqrt(W.inc * (1 - W.inc) / nn / stats.norm.pdf(z1) ** 2 + W.dec * (1 - W.dec) / nn / stats.norm.pdf(z2) ** 2)
    files = sorted(glob.glob(os.path.join(SBCS_DIR, '20*_sbcs-*appendix*.xlsx')))
    return W, files


def current_ind():
    html = open(INDEX, encoding='utf-8').read()
    block = html[html.index('const IND = {'):]
    block = block[:block.index('};')]
    cur = {}
    for k in INDS:
        m = re.search(r'\b' + k + r'\s*:\s*\{[^}]*?\bgrw:\s*(-?[\d.]+)', block)
        cur[k] = float(m.group(1))
    return cur


def load_fiserv():
    """Fiserv SBI API responses: monthly sales index (2019 = 100) by NAICS cell, SA and NSA, Jan 2019 - Sep 2026.
    Returns {'sa': DataFrame code x period, 'nsa': ...}, level and description by code, the published SA YoY (API field)
    and the file list. A cell present in several files must carry identical values (asserted)."""
    tabs, lev, desc, api_yoy, files = {}, {}, {}, {}, []
    for adj, pre in (('sa', 'adjusted'), ('nsa', 'nsa')):
        rows = []
        fl = [os.path.join(FISERV_API, pre + '_US_all.json')] + [os.path.join(FISERV_API, f'{pre}_US_{c}_n6.json') for c in FISERV_CODES2]
        for f in fl:
            files.append(f)
            for x in json.load(open(f, encoding='utf-8')):
                lv = 6 if x['naicsFinal'] != '0' else (3 if x['naicsCode3'] != '0' else (2 if x['naicsCode2'] != '0' else 0))
                code = {6: x['naicsFinal'], 3: x['naicsCode3'], 2: x['naicsCode2'], 0: 'ALL'}[lv]
                lev[code] = lv
                desc[code] = {6: x['naicsFinalDesc'], 3: x['naicsShortDesc3'], 2: x['naicsShortDesc2'], 0: 'All small business'}[lv]
                yv = x.get('salesYoyDiff')
                rows.append((code, x['period'], float(x['salesIndex']), float(yv) if yv is not None else np.nan))
        d = pd.DataFrame(rows, columns=['code', 'period', 'sales', 'yoy_api'])
        spread = d.groupby(['code', 'period']).sales.agg(lambda s: s.max() - s.min())
        assert spread.max() == 0, 'a Fiserv cell differs between files'
        d = d.drop_duplicates(['code', 'period'])
        d['period'] = pd.to_datetime(d.period)
        tabs[adj] = d.pivot_table(index='code', columns='period', values='sales', aggfunc='first').sort_index()
        api_yoy[adj] = d.pivot_table(index='code', columns='period', values='yoy_api', aggfunc='first').sort_index()
    return tabs, lev, desc, api_yoy, sorted(set(files))


def load_btos():
    """BTOS question 4 (revenue, two weeks against the two weeks before): shares increased / decreased with SEs by
    group and collection period; net balance nb = increased - decreased (fraction). Also the national revenue index."""
    rows = []
    for gtype, fname in BTOS_FILES.items():
        wb = openpyxl.load_workbook(os.path.join(BTOS_DIR, fname), read_only=True, data_only=True)
        keys = {'national': [], 'size': ['Empsize'], 'sector': ['Sector'], 'subsector': ['Subsector'], 'sector_size': ['Sector', 'Empsize']}[gtype]
        for sh, kind in (('Response Estimates', 'est'), ('Response Standard Errors', 'se')):
            it = wb[sh].iter_rows(values_only=True)
            hdr = [str(h) for h in next(it)]
            qi, ai = hdr.index('Question ID'), hdr.index('Answer')
            pcols = [(i, int(h)) for i, h in enumerate(hdr) if re.fullmatch(r'\d{6}', h)]
            ki = [hdr.index(k) for k in keys]
            for r in it:
                if r is None or r[qi] is None or str(r[qi]).strip() != '4':
                    continue
                ans = str(r[ai]).strip()
                if ans not in ('Increased', 'Decreased'):
                    continue
                key = '|'.join(str(r[i]) for i in ki) if ki else 'US'
                for i, p in pcols:
                    v = r[i]
                    try:
                        v = float(str(v).strip().replace('%', '')) / 100
                    except (TypeError, ValueError):
                        v = np.nan
                    rows.append((gtype, key, p, kind + '_' + ans[:3].lower(), v))
    L = pd.DataFrame(rows, columns=['gtype', 'key', 'period', 'var', 'val'])
    W = L.pivot_table(index=['gtype', 'key', 'period'], columns='var', values='val', aggfunc='first').reset_index()
    W['nb'] = W.est_inc - W.est_dec
    W['se_nb'] = np.sqrt(W.se_inc ** 2 + W.se_dec ** 2)          # upper bound (the two shares covary negatively)
    W = W.dropna(subset=['nb']).sort_values(['gtype', 'key', 'period']).reset_index(drop=True)
    wb = openpyxl.load_workbook(os.path.join(BTOS_DIR, BTOS_FILES['national']), read_only=True, data_only=True)
    it = wb['Index Estimates'].iter_rows(values_only=True)
    hdr = [str(h) for h in next(it)]
    idx = {}
    for r in it:
        if r and str(r[0]).strip() == 'Revenues':
            for i, h in enumerate(hdr):
                if re.fullmatch(r'\d{6}', h):
                    try:
                        idx[int(h)] = float(r[i])
                    except (TypeError, ValueError):
                        pass                                     # '.' = not collected (Oct-Nov 2025 shutdown)
    return W, idx


print('reading data ...')
TS, AIES_LABELS, CVB, n_aies_ts_rows = load_aies()
QW, n_qcew_files, qcew_cat_sha, n_qcew_rows = load_qcew()
QSS, qss_sheet = load_qss()
MRTS, mrts_months_2026, MRTS_LV = load_mrts()
MRTS_YTD = 'Jan-' + ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec'][mrts_months_2026 - 1] + ' 2026'
SU = {y: load_susb(p) for y, p in SUSB.items()}
SBCS, sbcs_files = load_sbcs()
CUR = current_ind()
FIS, FIS_LEVEL, FIS_DESC, FIS_API_YOY, fiserv_files = load_fiserv()
BT, BT_REV_INDEX = load_btos()

# firm-count weights: SUSB 2022, firms under 500 employees, NAICS sectors
s22 = SU[2022][SU[2022]['size'] == '08'].set_index('NAICS')
FIRMS = {k: int(s22.loc[k, 'Firms']) for k in SECTORS}


def fw(vals):
    v = {k: x for k, x in vals.items() if k in FIRMS and x is not None and np.isfinite(x)}
    return sum(FIRMS[k] * x for k, x in v.items()) / sum(FIRMS[k] for k in v), sorted(v)


def ts_growth(code, y):
    if code not in TS.index:
        return np.nan
    a, b = TS.loc[code, y - 1], TS.loc[code, y]
    return b / a - 1 if np.isfinite(a) and np.isfinite(b) and a > 0 else np.nan


def qss_h1(code):
    q = QSS.get(code, {})
    if not all(isinstance(q.get(k), float) for k in ('yoy_1q26', 'yoy_2q26', 'se_1q26', 'se_2q26')):
        return None
    return (q['yoy_1q26'] + q['yoy_2q26']) / 200, (q['se_1q26'] + q['se_2q26']) / 200


# ================================================================ 2. comparators
C, C_sectors = {}, {}
C['aies2023'], C_sectors['aies2023'] = fw({k: ts_growth(k, 2023) for k in SECTORS})
C['aies2024'], C_sectors['aies2024'] = fw({k: CVB.g.get(k) for k in SECTORS})
for col in ['w23', 'w24', 'w25']:
    C['qcew_' + col], C_sectors['qcew_' + col] = fw({k: QW.loc[k, col] for k in QCEW_SECTORS})
C['qss2026'], C_sectors['qss2026'] = fw({k: qss_h1(k)[0] for k in QSS_SECTORS})
C['mrts2025'] = MRTS[('total', 2025)]; C['mrts2026'] = MRTS[('total', 2026)]
print('comparators', {k: round(v, 4) for k, v in C.items()})

# ================================================================ 3. payroll-to-revenue bridge
brows = []
for a, b in BRIDGE_PAIRS.items():
    if a not in TS.index or b not in QW.index:
        continue
    for y, col in [(2023, 'w23'), (2024, 'w24')]:
        brows.append(dict(aies=a, qcew=b, year=y, rev_g=ts_growth(a, y), pay_g=QW.loc[b, col]))
BD = pd.DataFrame(brows).dropna().reset_index(drop=True)
grp = pd.factorize(BD.aies)[0]
m_pool = sm.OLS(BD.rev_g, sm.add_constant(BD.pay_g)).fit(cov_type='cluster', cov_kwds={'groups': grp})
BD['rev_dev'] = BD.rev_g - BD.groupby('year').rev_g.transform('mean')
BD['pay_dev'] = BD.pay_g - BD.groupby('year').pay_g.transform('mean')
m_dev = sm.OLS(BD.rev_dev, sm.add_constant(BD.pay_dev)).fit(cov_type='cluster', cov_kwds={'groups': grp})
BRIDGE = dict(slope=float(m_dev.params.pay_dev), se=float(m_dev.bse.pay_dev), r=float(BD.rev_dev.corr(BD.pay_dev)),
              resid_sd=float(np.std(m_dev.resid, ddof=2)), n=len(BD), industries=int(BD.aies.nunique()),
              pooled=dict(intercept=float(m_pool.params.const), slope=float(m_pool.params.pay_g), se=float(m_pool.bse.pay_g),
                          r=float(BD.rev_g.corr(BD.pay_g)), resid_sd=float(np.std(m_pool.resid, ddof=2))))
BR = BR_DECISION
RES = round(BRIDGE['resid_sd'], 3)
print('bridge', {k: (round(v, 4) if isinstance(v, float) else v) for k, v in BRIDGE.items() if k != 'pooled'})


# ================================================================ 4. persistence (AIES)
def deviations(codes):
    """growth (clipped to +-50%) minus the cross-industry median of the same year, within the group"""
    g = (TS.loc[codes] / TS.loc[codes].shift(1, axis=1) - 1).clip(-0.5, 0.5)
    return g - g.median(axis=0)


def window_slope(DEV, wa, wb, boot=False, seed=SEED):
    a = DEV[wa].mean(axis=1); b = DEV[wb].mean(axis=1)
    ok = a.notna() & b.notna()
    x, y = a[ok].values, b[ok].values
    X = sm.add_constant(x)
    fit = sm.OLS(y, X).fit(); hc = sm.OLS(y, X).fit(cov_type='HC1')
    out = dict(windows=[list(wa), list(wb)], n=int(ok.sum()), slope=float(fit.params[1]), se_ols=float(fit.bse[1]),
               se_hc1=float(hc.bse[1]), r=float(np.corrcoef(x, y)[0, 1]), resid_sd=float(np.std(fit.resid, ddof=2)))
    if boot:
        out['se_boot'], out['boot_ci95'] = boot_slope(x, y, seed)
    return out, (a[ok], b[ok])


six = [c for c in TS.index if len(c) == 6 and c.isdigit()]
DEV6 = deviations(six)
nm = [c for c in TS.index if c.isdigit() and 4 <= len(c) <= 6 and c[:2] not in MFG2]
NONMFG_LEAF = [c for c in nm if not any(o != c and o.startswith(c) for o in nm)]
MFG6 = [c for c in six if c[:2] in MFG2]
GROUP_DEV = {'all_six_digit': DEV6, 'manufacturing_six_digit': deviations(MFG6), 'non_manufacturing_leaf': deviations(NONMFG_LEAF)}
MAIN_W = ([2017, 2018, 2019], [2023, 2024])
WINDOWS = [('2017-19 -> 2023-24 (main)', *MAIN_W), ('2014-16 -> 2017-19 (pre-COVID, adjacent)', [2014, 2015, 2016], [2017, 2018, 2019]),
           ('2015-17 -> 2018-19 (pre-COVID)', [2015, 2016, 2017], [2018, 2019]), ('2014-16 -> 2023-24 (long gap)', [2014, 2015, 2016], [2023, 2024]),
           ('2018-19 -> 2023-24 (2-year)', [2018, 2019], [2023, 2024]), ('2015-16 -> 2017-18 (2-year, adjacent)', [2015, 2016], [2017, 2018])]
PERS, pers_pairs = window_slope(DEV6, *MAIN_W, boot=True)
BY_GROUP = {}
for g, DV in GROUP_DEV.items():
    wins = {lab: window_slope(DV, wa, wb)[0] for lab, wa, wb in WINDOWS}
    sl = [v['slope'] for v in wins.values()]
    BY_GROUP[g] = dict(windows=wins, range=[min(sl), max(sl)], n_codes=len(DV))
GMAIN = {'non_manufacturing': window_slope(GROUP_DEV['non_manufacturing_leaf'], *MAIN_W, boot=True)[0],
         'manufacturing': window_slope(GROUP_DEV['manufacturing_six_digit'], *MAIN_W, boot=True)[0]}


def boot_window_mean(DV, labels, seed=SEED, B=B_BOOT):
    """bootstrap SE of the mean of several window slopes (industries resampled, the same draw for every window)"""
    wins = [(wa, wb) for lab, wa, wb in WINDOWS if lab in labels]
    A = [DV[wa].mean(axis=1).values for wa, wb in wins]; Bv = [DV[wb].mean(axis=1).values for wa, wb in wins]
    rng = np.random.default_rng(seed); n = len(DV); out = np.empty(B)
    for i in range(B):
        idx = rng.integers(0, n, n)
        s = []
        for a, b in zip(A, Bv):
            x, y = a[idx], b[idx]
            ok = np.isfinite(x) & np.isfinite(y)
            s.append(np.cov(x[ok], y[ok], ddof=1)[0, 1] / np.var(x[ok], ddof=1))
        out[i] = np.mean(s)
    return float(out.std(ddof=1)), [float(np.quantile(out, .025)), float(np.quantile(out, .975))]


LAM_MFG, LAM_NONMFG_MAIN = GMAIN['manufacturing']['slope'], GMAIN['non_manufacturing']['slope']
NM_WIN = BY_GROUP['non_manufacturing_leaf']['windows']
LAM_NONMFG = float(np.mean([NM_WIN[lab]['slope'] for lab in LAM_WINDOWS_NONMFG]))     # analyst: adjacent pre-COVID mean
WIN_SD = {g: float(np.std([v['slope'] for v in BY_GROUP[gk]['windows'].values()], ddof=1))
          for g, gk in (('non_manufacturing', 'non_manufacturing_leaf'), ('manufacturing', 'manufacturing_six_digit'))}
NM_BOOT_SE, NM_BOOT_CI = boot_window_mean(GROUP_DEV['non_manufacturing_leaf'], LAM_WINDOWS_NONMFG)
SE_LAM_SAMPLING = {'non_manufacturing': max(SE_LAM_SAMPLING_NONMFG, NM_BOOT_SE),
                   'manufacturing': max(GMAIN['manufacturing']['se_hc1'], GMAIN['manufacturing']['se_boot'])}
SE_LAM_RAW = {g: math.sqrt(SE_LAM_SAMPLING[g] ** 2 + WIN_SD[g] ** 2) for g in ('non_manufacturing', 'manufacturing')}
SE_LAM = {g: round(v, 2) for g, v in SE_LAM_RAW.items()}
LAM = {'non_manufacturing': round(LAM_NONMFG, 2), 'manufacturing': round(LAM_MFG, 2)}
LAM_INFO = dict(non_manufacturing=dict(rule='mean of the adjacent pre-COVID windows', windows={lab: NM_WIN[lab]['slope'] for lab in LAM_WINDOWS_NONMFG},
                                       mean=LAM_NONMFG, se_sampling_lead=SE_LAM_SAMPLING_NONMFG, se_boot_of_mean=NM_BOOT_SE, boot_ci95=NM_BOOT_CI,
                                       across_window_sd=WIN_SD['non_manufacturing'], se=SE_LAM_RAW['non_manufacturing'],
                                       main_window=LAM_NONMFG_MAIN, main_window_se=max(GMAIN['non_manufacturing']['se_hc1'], GMAIN['non_manufacturing']['se_boot']),
                                       all_windows={lab: v['slope'] for lab, v in NM_WIN.items()}),
                manufacturing=dict(rule='main window 2017-19 -> 2023-24 (analyst: stays .12)', slope=LAM_MFG, se_sampling=SE_LAM_SAMPLING['manufacturing'],
                                   across_window_sd=WIN_SD['manufacturing'], se=SE_LAM_RAW['manufacturing'],
                                   all_windows={lab: v['slope'] for lab, v in BY_GROUP['manufacturing_six_digit']['windows'].items()}),
                revision2=LAM_R2)
lag_rows = []
for gname in ('all_six_digit', 'non_manufacturing_leaf'):
    DV = GROUP_DEV[gname]
    for lag in (1, 2, 3, 5):
        for y in range(2015, 2025):
            if y - lag < 2014:
                continue
            a, b = DV[y - lag], DV[y]
            ok = a.notna() & b.notna()
            if ok.sum() < 50:
                continue
            lag_rows.append(dict(group=gname, lag=lag, year=y, n=int(ok.sum()), r=float(np.corrcoef(a[ok], b[ok])[0, 1]),
                                 slope=float(np.polyfit(a[ok], b[ok], 1)[0])))
LAGS = pd.DataFrame(lag_rows)
covid = LAGS.year.isin([2020, 2021, 2022]) | (LAGS.year - LAGS.lag).isin([2020, 2021])
lag_summary = {f'{gname}_lag{int(l)}': dict(median_slope_all=float(g.slope.median()), median_slope_excl_covid=float(g[~covid[g.index]].slope.median()),
                                            median_r_excl_covid=float(g[~covid[g.index]].r.median()), years=int(len(g)))
               for (gname, l), g in LAGS.groupby(['group', 'lag'])}
print('persistence: pooled', round(PERS['slope'], 4), '| non-manufacturing adjacent pre-COVID mean', round(LAM_NONMFG, 3), 'se', round(SE_LAM_RAW['non_manufacturing'], 3),
      '(boot of mean', round(NM_BOOT_SE, 3), '; window sd', round(WIN_SD['non_manufacturing'], 3), '; main window', round(LAM_NONMFG_MAIN, 3), ')',
      '| manufacturing', round(LAM_MFG, 3), 'se', round(SE_LAM_RAW['manufacturing'], 3))

# ================================================================ 5. Census deviations per industry (round 1)
OBS, ROWS = [], []
for k in INDS:
    ac, qc, qsc, mc = IND[k]
    obs = []                                   # (label, year, deviation, variance)
    for code in ac:
        if code == '238*':
            obs.append(('aies2024_238', 2024, CVB.loc['238', 'g'] - C['aies2024'], max(CVB.loc['238', 'se_g'], FLOOR) ** 2))
            continue
        for y in (2023, 2024):
            g = ts_growth(code, y)
            bc = CVMAP.get(code, code)
            se = CVB.loc[bc, 'se_g'] if (y == 2024 and bc in CVB.index and np.isfinite(CVB.loc[bc, 'se_g'])) else SE_AIES_NOCV
            obs.append((f'aies{y}_{code}', y, g - C[f'aies{y}'], max(se, FLOOR) ** 2))
    for code in qc:
        for col, y in [('w23', 2023), ('w24', 2024), ('w25', 2025)]:
            obs.append((f'qcew{y}_{code}', y, BR * (QW.loc[code, col] - C['qcew_' + col]), RES ** 2))
    for code in (qsc or []):
        v, se = qss_h1(code)
        obs.append((f'qss2026_{code}', 2026, v - C['qss2026'], max(se, FLOOR) ** 2 + QSS_EXTRA ** 2))
    if mc:
        for y in (2025, 2026):
            obs.append((f'mrts{y}_{mc}', y, MRTS[(mc, y)] - C[f'mrts{y}'], SE_MRTS ** 2))
    yv = np.array([o[1] for o in obs]); d = np.array([o[2] for o in obs], float); v = np.array([o[3] for o in obs], float)
    assert np.all(np.isfinite(d)), (k, obs)
    a = np.array([REC[t] for t in yv]) / v
    D_now = float((a * d).sum() / a.sum())
    se_D_now = float(max(np.sqrt((a ** 2 * v).sum()) / a.sum(), np.sqrt(np.cov(d, aweights=a)) / np.sqrt(max(len(d) / 2, 1))))
    b = 1 / v
    D_avg = float((b * d).sum() / b.sum())
    se_D_avg = float(max(np.sqrt(1 / b.sum()), np.std(d, ddof=1) / np.sqrt(max(len(d) / 2, 1))))
    for lab, yy, dd, vv in obs:
        OBS.append(dict(ind=k, obs=lab, year=int(yy), deviation=dd, se=math.sqrt(vv), weight_forward=1 / vv / b.sum(),
                        weight_nowcast=REC[int(yy)] / vv / a.sum()))
    ROWS.append(dict(ind=k, current=CUR[k], n_obs=len(obs), D_avg=D_avg, se_D_avg=se_D_avg, D_now=D_now, se_D_now=se_D_now))
E = pd.DataFrame(ROWS).set_index('ind')

# ================================================================ 6. Fiserv Small Business Index
SA, NSA = FIS['sa'], FIS['nsa']
PER = list(SA.columns)


def f_yoy(P, code, ym):
    t = pd.Timestamp(ym + '-01')
    return float(P.loc[code, t] / P.loc[code, t - pd.DateOffset(years=1)] - 1)


def f_annual(P, code, y):
    """calendar-year growth of the summed index; 2026 = Jan-Sep against Jan-Sep 2025"""
    last = max(c for c in PER if c.year == 2026).month if y == 2026 else 12
    a = [c for c in PER if c.year == y and c.month <= last]; b = [c for c in PER if c.year == y - 1 and c.month <= last]
    return float(P.loc[code, a].sum(min_count=len(a)) / P.loc[code, b].sum(min_count=len(b)) - 1)   # NaN if a month is missing


FY = list(range(2020, 2027))
FG = pd.DataFrame({y: [f_annual(NSA, c, y) for c in NSA.index] for y in FY}, index=NSA.index)
FG['level'] = [FIS_LEVEL[c] for c in FG.index]; FG['desc'] = [FIS_DESC[c] for c in FG.index]
last_month = max(PER)
# 6a. the M test: national YoY (SA, the published monthly figure) and annual NSA growth
nat_m = {m.strftime('%Y-%m'): f_yoy(SA, 'ALL', m.strftime('%Y-%m')) for m in PER if m >= pd.Timestamp('2020-01-01')}
last12 = [m.strftime('%Y-%m') for m in PER[-12:]]
FIS_NAT = dict(feb2026_sa=nat_m['2026-02'], sep2026_sa=nat_m['2026-09'], mean_last12_sa=float(np.mean([nat_m[m] for m in last12])),
               last12=last12[0] + '..' + last12[-1], annual_nsa={y: float(FG.loc['ALL', y]) for y in (2022, 2023, 2024, 2025)},
               ytd2026_nsa=float(FG.loc['ALL', 2026]), cagr_2019_to_last=float((SA.loc['ALL', last_month] / SA.loc['ALL', [c for c in PER if c.year == 2019]].mean())
                                                                               ** (12 / ((last_month.year - 2019) * 12 + last_month.month - 6.5)) - 1))
FIS_NAT['gap_to_M_feb2026'] = FIS_NAT['feb2026_sa'] - M
FIS_NAT['z_vs_SE_M_round2'] = FIS_NAT['gap_to_M_feb2026'] / SE_M_R2     # the test as decided, on M's SE before revision 3
FIS_NAT['supports_M'] = bool(abs(FIS_NAT['z_vs_SE_M_round2']) <= 2)       # at face value
FIS_NAT['z_vs_SE_M'] = FIS_NAT['gap_to_M_feb2026'] / SE_M                 # revision 3: SE_M widened to cover this gap (not a test)
FIS_NAT['within_widened_SE_M'] = bool(abs(FIS_NAT['z_vs_SE_M']) <= 2)
# the API reproduces the Sep 2026 report (Figure 2, 'Sales Index - YoY', Sep 25 .. Sep 26, one decimal)
fig2 = None
t2 = pdf_text(FISERV_PDF, 2, 2, layout=True)
if t2 is not None:
    for line in t2.splitlines():
        if re.match(r'^\s*Sales Index\s*.\s*YoY', line):
            fig2 = [float(x) / 100 for x in re.findall(r'(-?\d+\.\d)%', line)]
            break
fig2_months = [m.strftime('%Y-%m') for m in PER[-13:]]
api_nat = {m: float(FIS_API_YOY['sa'].loc['ALL', pd.Timestamp(m + '-01')]) / 100 for m in fig2_months}   # the API's own YoY field
FIG2 = dict(checked=fig2 is not None, months=fig2_months, report=fig2, api=[round(api_nat[m], 5) for m in fig2_months],
            computed=[round(nat_m[m], 5) for m in fig2_months],
            max_abs_diff=(None if fig2 is None else float(max(abs(a - api_nat[m]) for a, m in zip(fig2, fig2_months)))),
            max_abs_diff_computed_vs_api=float(max(abs(nat_m[m] - api_nat[m]) for m in fig2_months)))
# restaurant offsets: Chase restaurant median minus Fiserv 722 in the same month, against the national offset
REST_OFFSET = dict(chase_rest_aug2026=CHASE_REST['2026-08'], fiserv_722_aug2026=f_yoy(SA, '722', '2026-08'),
                   offset_rest_aug2026=CHASE_REST['2026-08'] - f_yoy(SA, '722', '2026-08'),
                   offset_national_feb2026=M - FIS_NAT['feb2026_sa'])
print('Fiserv national: Feb 2026', round(FIS_NAT['feb2026_sa'], 4), 'last 12 mean', round(FIS_NAT['mean_last12_sa'], 4),
      'annual', {y: round(v, 4) for y, v in FIS_NAT['annual_nsa'].items()}, 'YTD 2026', round(FIS_NAT['ytd2026_nsa'], 4),
      '| Figure 2 max diff', FIG2['max_abs_diff'])

# 6b. console industries: Fiserv deviations against Fiserv's all-industry, 2023-26 mean
FMAP = dict(FISERV_POOL); FMAP.update(FISERV_REPORT_ONLY)
FDEV = FG[FY].sub(FG.loc['ALL', FY], axis=1)
F_IND = {}
for k in INDS:
    c = FMAP.get(k)
    if c is None:
        continue
    vals = [float(FDEV.loc[c, y]) for y in FISERV_YEARS]
    F_IND[k] = dict(code=c, desc=FIS_DESC[c], pooled=k in FISERV_POOL, D=float(np.mean(vals)), by_year=dict(zip(FISERV_YEARS, vals)),
                    growth=dict((y, float(FG.loc[c, y])) for y in FISERV_YEARS))

# 6c. persistence of Fiserv's own industry deviations (80 leaf industries, NSA annual growth)
leaf_codes = [c for c in FG.index if FIS_LEVEL[c] == 6]
LEAF = FG.loc[leaf_codes, FY].astype(float).clip(-0.5, 0.5)
FLDEV = LEAF - LEAF.median(axis=0)


def f_window(wa, wb, boot=False, seed=SEED):
    ok = FLDEV[wa + wb].notna().all(axis=1)                            # leaves with every year of both windows
    a = FLDEV.loc[ok, wa].mean(axis=1); b = FLDEV.loc[ok, wb].mean(axis=1)
    x, y = a.values, b.values
    s, se, f = ols_hc1(x, y)
    out = dict(windows=[wa, wb], n=int(len(x)), slope=s, se_hc1=se, r=float(np.corrcoef(x, y)[0, 1]), resid_sd=float(np.std(f.resid, ddof=2)))
    if boot:
        out['se_boot'], out['boot_ci95'] = boot_slope(x, y, seed)
    return out


FPERS_MAIN = f_window([2023, 2024], [2025, 2026], boot=True)
FPERS = {'2023-24 -> 2025-26 (main)': FPERS_MAIN, '2023 -> 2025-26': f_window([2023], [2025, 2026]),
         '2022-23 -> 2025-26': f_window([2022, 2023], [2025, 2026]), '2022-24 -> 2025-26': f_window([2022, 2023, 2024], [2025, 2026]),
         '2023 -> 2024 (lag 1)': f_window([2023], [2024]), '2024 -> 2025 (lag 1)': f_window([2024], [2025]),
         '2025 -> 2026 (lag 1)': f_window([2025], [2026]), '2023 -> 2025 (lag 2)': f_window([2023], [2025]),
         '2024 -> 2026 (lag 2)': f_window([2024], [2026]), '2023 -> 2026 (lag 3)': f_window([2023], [2026])}
sub_codes = [c for c in FG.index if FIS_LEVEL[c] == 3]
SUB = FG.loc[sub_codes, FY].astype(float).clip(-0.5, 0.5)
SUBDEV = (SUB - SUB.median(axis=0)).dropna(subset=[2023, 2024, 2025, 2026])
_x, _y = SUBDEV[[2023, 2024]].mean(axis=1), SUBDEV[[2025, 2026]].mean(axis=1)
_s, _se, _f = ols_hc1(_x, _y)
FPERS['subsectors 2023-24 -> 2025-26'] = dict(n=int(len(SUBDEV)), slope=_s, se_hc1=_se, r=float(np.corrcoef(_x, _y)[0, 1]))
SE_LF = max(FPERS_MAIN['se_hc1'], FPERS_MAIN['se_boot'])
FPERS_VS = dict(fiserv=FPERS_MAIN['slope'], se_fiserv=SE_LF, aies_non_manufacturing=LAM_NONMFG, se_aies=SE_LAM_RAW['non_manufacturing'],
                aies_basis='adjacent pre-COVID mean (shipped lambda), SE with the across-window SD',
                z=(FPERS_MAIN['slope'] - LAM_NONMFG) / math.sqrt(SE_LF ** 2 + SE_LAM_RAW['non_manufacturing'] ** 2),
                aies_main_window=LAM_NONMFG_MAIN,
                z_vs_main_window=(FPERS_MAIN['slope'] - LAM_NONMFG_MAIN) / math.sqrt(SE_LF ** 2 + max(GMAIN['non_manufacturing']['se_hc1'], GMAIN['non_manufacturing']['se_boot']) ** 2),
                aies_lag1_non_manufacturing=lag_summary['non_manufacturing_leaf_lag1']['median_slope_excl_covid'],
                fiserv_lag1_median=float(np.median([FPERS[k]['slope'] for k in FPERS if 'lag 1' in k])),
                sd_leaf_dev_by_year={y: float(FLDEV[y].std()) for y in FY})
print('Fiserv persistence (2023-24 -> 2025-26):', round(FPERS_MAIN['slope'], 3), 'se', round(SE_LF, 3), '| AIES non-mfg', round(LAM_NONMFG, 3))

# 6d. agreement with Census at the cell level: Fiserv cells matched to AIES codes
crows = []
for fc, acode in FISERV_AIES.items():
    if fc not in FG.index or acode not in TS.index:
        continue
    se24 = CVB.loc[acode, 'se_g'] if (acode in CVB.index and np.isfinite(CVB.loc[acode, 'se_g'])) else SE_AIES_NOCV
    crows.append(dict(fiserv=fc, aies=acode, desc=FIS_DESC[fc], service=fc[:2] in FISERV_SERVICES, leaf=FIS_LEVEL[fc] == 6,
                      f23=FG.loc[fc, 2023], f24=FG.loc[fc, 2024], f25=FG.loc[fc, 2025], f26=FG.loc[fc, 2026],
                      a23=ts_growth(acode, 2023), a24=ts_growth(acode, 2024), se_a23=max(SE_AIES_NOCV, FLOOR), se_a24=max(se24, FLOOR)))
CELL = pd.DataFrame(crows).dropna(subset=['f23', 'f24', 'f25', 'f26', 'a23', 'a24']).reset_index(drop=True)
for c in ['f23', 'f24', 'f25', 'f26', 'a23', 'a24']:
    CELL[c + 'd'] = CELL[c].astype(float).clip(-0.5, 0.5) - CELL[c].astype(float).clip(-0.5, 0.5).median()
CELL['F2'] = CELL[['f23d', 'f24d']].mean(axis=1); CELL['A2'] = CELL[['a23d', 'a24d']].mean(axis=1)
CELL['F_next'] = CELL[['f25d', 'f26d']].mean(axis=1)


def cell_stats(D_):
    same = agreement(D_.A2, D_.F2)
    pred_s, pred_se, _ = ols_hc1(D_.A2, D_.F_next)
    own_s, own_se, _ = ols_hc1(D_.F2, D_.F_next)
    return dict(n=int(len(D_)), same_2023_24=same, same_2023=agreement(D_.a23d, D_.f23d), same_2024=agreement(D_.a24d, D_.f24d),
                predict_fiserv_2025_26_from_aies_2023_24=dict(slope=pred_s, se=pred_se, r=float(np.corrcoef(D_.A2, D_.F_next)[0, 1])),
                predict_fiserv_2025_26_from_fiserv_2023_24=dict(slope=own_s, se=own_se),
                sd_fiserv_2yr=float(D_.F2.std()), sd_aies_2yr=float(D_.A2.std()), sd_diff_2yr=float((D_.F2 - D_.A2).std()))


CELLS = {'all': cell_stats(CELL), 'services': cell_stats(CELL[CELL.service]), 'leaf': cell_stats(CELL[CELL.leaf])}
V_A2 = float(((CELL.se_a23 ** 2 + CELL.se_a24 ** 2) / 4).mean())       # AIES sampling variance of a 2-year mean growth
V_F2 = max(CELLS['all']['sd_diff_2yr'] ** 2 - V_A2, V_A2)             # Fiserv error variance of a 2-year mean deviation
# revision 3 (review): the Fiserv-AIES gaps of a cell are correlated across years (persistent cell and coverage differences),
# so a 4-year mean does not halve V_F2. Moments over the matched cells: single-year Fiserv error variance V_F1 = var(gap_t)
# minus AIES's sampling variance (mean of 2023 and 2024); C_F = cov(gap_2023, gap_2024), AIES sampling errors taken as
# independent across years; rho = C_F / V_F1 (clipped to 0..1). Errors equicorrelated across years (a persistent cell
# component plus independent annual noise): V_Fn = V_F1 (rho + (1 - rho) / n), so V_F4 = V_F2 (1 + 3 rho) / (2 (1 + rho)).
GAP23 = (CELL.f23d - CELL.a23d).values.astype(float); GAP24 = (CELL.f24d - CELL.a24d).values.astype(float)
RHO_GAP = float(np.corrcoef(GAP23, GAP24)[0, 1])
_rng = np.random.default_rng(SEED); _bs = np.empty(B_BOOT)
for _i in range(B_BOOT):
    _idx = _rng.integers(0, len(GAP23), len(GAP23))
    _bs[_i] = np.corrcoef(GAP23[_idx], GAP24[_idx])[0, 1]
RHO_GAP_SE = float(_bs.std(ddof=1)); RHO_GAP_CI = [float(np.quantile(_bs, .025)), float(np.quantile(_bs, .975))]
V_A1 = {2023: float((CELL.se_a23 ** 2).mean()), 2024: float((CELL.se_a24 ** 2).mean())}
V_F1 = float(np.mean([np.var(GAP23, ddof=1) - V_A1[2023], np.var(GAP24, ddof=1) - V_A1[2024]]))
C_F = float(np.cov(GAP23, GAP24, ddof=1)[0, 1])
RHO_F = float(min(max(C_F / V_F1, 0.0), 1.0))
V_F4 = V_F2 * (1 + 3 * RHO_F) / (2 * (1 + RHO_F))                     # 4-year mean (2023-26), equicorrelated errors
V_F4_R2 = V_F2 / 2                                                     # revision 2: errors independent across years
V_F4_AR1 = V_F1 / 16 * (4 + 2 * (3 * RHO_F + 2 * RHO_F ** 2 + RHO_F ** 3))   # alternative: AR(1) errors (check only)
FIS_ERR = dict(model='equicorrelated (persistent cell component plus independent annual noise)', n_cells=int(len(GAP23)),
               rho_gap=RHO_GAP, rho_gap_se_boot=RHO_GAP_SE, rho_gap_ci95=RHO_GAP_CI, spearman_gap=float(stats.spearmanr(GAP23, GAP24)[0]),
               v_aies_1yr=V_A1, v_fiserv_1yr=V_F1, cov_fiserv=C_F, rho_fiserv=RHO_F, v_fiserv_2yr=V_F2, v_fiserv_2yr_from_moments=(V_F1 + C_F) / 2,
               v_fiserv_4yr=V_F4, ratio_4yr_to_2yr=V_F4 / V_F2, v_fiserv_4yr_round2_halved=V_F4_R2, v_fiserv_4yr_ar1=V_F4_AR1,
               sd_fiserv_4yr=math.sqrt(V_F4), sd_fiserv_4yr_round2=math.sqrt(V_F4_R2))
print('Fiserv vs AIES cells: n', CELLS['all']['n'], 'r(2023-24)', round(CELLS['all']['same_2023_24']['r'], 3), 'services r',
      round(CELLS['services']['same_2023_24']['r'], 3), '| sd(F-A)', round(CELLS['all']['sd_diff_2yr'], 4), '| gap corr 2023 vs 2024', round(RHO_GAP, 3),
      'se', round(RHO_GAP_SE, 3), 'rho_F', round(RHO_F, 3), '-> Fiserv error sd (4-yr mean)', round(math.sqrt(V_F4), 4), '(halved:', round(math.sqrt(V_F4_R2), 4), ')')

# 6e. agreement at the console level: Census D against Fiserv D
cons = [k for k in INDS if k in F_IND and k != 'saas']
uniq = [k for k in cons if k not in ('constr', 'prof', 'land')]       # one console industry per Fiserv cell
CONS_AGREE = {'all_mapped': agreement([E.loc[k, 'D_avg'] for k in cons], [F_IND[k]['D'] for k in cons]),
              'all_mapped_without_staff': agreement([E.loc[k, 'D_avg'] for k in cons if k != 'staff'], [F_IND[k]['D'] for k in cons if k != 'staff']),
              'unique_cells': agreement([E.loc[k, 'D_avg'] for k in uniq], [F_IND[k]['D'] for k in uniq]),
              'own_cells': agreement([E.loc[k, 'D_avg'] for k in list(FISERV_POOL) + ['rest']], [F_IND[k]['D'] for k in list(FISERV_POOL) + ['rest']]),
              'industries': {k: dict(census=float(E.loc[k, 'D_avg']), fiserv=F_IND[k]['D'], fiserv_cell=F_IND[k]['code'], pooled=F_IND[k]['pooled']) for k in cons}}
for lab in ('all_mapped', 'own_cells'):
    CONS_AGREE[lab]['x'] = 'Census D'; CONS_AGREE[lab]['y'] = 'Fiserv D'
print('Census D vs Fiserv D: all mapped r', round(CONS_AGREE['all_mapped']['r'], 3), 'slope F on C', round(CONS_AGREE['all_mapped']['slope_y_on_x'], 2),
      '| own cells r', round(CONS_AGREE['own_cells']['r'], 3), 'slope', round(CONS_AGREE['own_cells']['slope_y_on_x'], 2))

# ================================================================ 7. BTOS
BTW = BT[(BT.period >= BTOS_WINDOW[0]) & (BT.period <= BTOS_WINDOW[1])]


def bt_series(gtype, key, lo=BTOS_WINDOW[0], hi=BTOS_WINDOW[1]):
    s = BT[(BT.gtype == gtype) & (BT.key == key) & (BT.period >= lo) & (BT.period <= hi)].set_index('period').nb
    return s


def bt_mean(gtype, key, lo=BTOS_WINDOW[0], hi=BTOS_WINDOW[1]):
    s = bt_series(gtype, key, lo, hi)
    return float(s.mean()), float(s.std(ddof=1) / math.sqrt(len(s))), int(len(s))


NAT_NB = bt_mean('national', 'US')
SIZE_NB = {k: bt_mean('size', k) for k in 'ABCDEFG'}
SIZE_NB_2324 = {k: bt_mean('size', k, 202300, 202499)[0] for k in 'ABCDEFG'}
# the published revenue index is 50 + 50 * nb (BTOS index = share increased + half the share unchanged): parsing check
nat_s = bt_series('national', 'US', 0, 999999)
IDX_CHECK = float(max(abs(BT_REV_INDEX[p] - (50 + 50 * nat_s[p])) for p in nat_s.index if p in BT_REV_INDEX))
# 7a. calibration: within-year sector deviations of QCEW payroll growth x 1.10 on BTOS sector net balance, 2024 and 2025
crow = []
for y, col in ((2024, 'w24'), (2025, 'w25')):
    for bs, qs in BTOS_QCEW.items():
        nb_y, _, n_y = bt_mean('sector', bs, y * 100, y * 100 + 99)
        crow.append(dict(year=y, sector=bs, qcew=qs, nb=nb_y, n_periods=n_y, g=BR * float(QW.loc[qs, col])))
CAL = pd.DataFrame(crow)
CAL['nb_dev'] = CAL.nb - CAL.groupby('year').nb.transform('mean'); CAL['g_dev'] = CAL.g - CAL.groupby('year').g.transform('mean')
m_cal = sm.OLS(CAL.g_dev, sm.add_constant(CAL.nb_dev)).fit(cov_type='cluster', cov_kwds={'groups': pd.factorize(CAL.sector)[0]})
CALIB = dict(slope=float(m_cal.params.nb_dev), se=float(m_cal.bse.nb_dev), r=float(np.corrcoef(CAL.nb_dev, CAL.g_dev)[0, 1]),
             resid_sd=float(np.std(m_cal.resid, ddof=2)), n=int(len(CAL)), sectors=int(CAL.sector.nunique()),
             by_year={int(y): dict(zip(('slope', 'se'), ols_hc1(g.nb, g.g)[:2])) for y, g in CAL.groupby('year')},
             leave_one_sector_out=[float(sm.OLS(CAL[CAL.sector != s].g_dev, sm.add_constant(CAL[CAL.sector != s].nb_dev)).fit().params.nb_dev)
                                   for s in sorted(CAL.sector.unique())])
CALIB['loso_range'] = [min(CALIB['leave_one_sector_out']), max(CALIB['leave_one_sector_out'])]
# checks of the calibration against Fiserv (small-business sales, 2025-26): sector and subsector
fs_rows = []
for bs, fc in {'23': '23', '42': '42', '44': '44-45', '48': '48-49', '51': '51', '53': '53', '54': '54', '56': '56', '61': '61',
               '62': '62', '71': '71', '72': '72', '81': '81'}.items():
    fs_rows.append(dict(cell=bs, nb=bt_mean('sector', bs)[0], g=float(FG.loc[fc, [2025, 2026]].mean())))
FSC = pd.DataFrame(fs_rows).dropna()
ss_rows = []
for fc in [c for c in FG.index if FIS_LEVEL[c] == 3]:
    bc = {'48X': '484', '71X': '713'}.get(fc, fc)
    s = bt_series('subsector', bc)
    if len(s) >= 20:
        ss_rows.append(dict(cell=fc, btos=bc, nb=float(s.mean()), g=float(FG.loc[fc, [2025, 2026]].mean())))
FSS = pd.DataFrame(ss_rows).dropna()
CALIB['check_fiserv_sector'] = dict(zip(('slope', 'se'), ols_hc1(FSC.nb, FSC.g)[:2]), n=int(len(FSC)), r=float(np.corrcoef(FSC.nb, FSC.g)[0, 1]))
CALIB['check_fiserv_subsector'] = dict(zip(('slope', 'se'), ols_hc1(FSS.nb, FSS.g)[:2]), n=int(len(FSS)), r=float(np.corrcoef(FSS.nb, FSS.g)[0, 1]))
BSL, BSL_SE = CALIB['slope'], CALIB['se']
print('BTOS calibration: slope', round(BSL, 3), 'se', round(BSL_SE, 3), 'r', round(CALIB['r'], 3), '| Fiserv checks',
      round(CALIB['check_fiserv_sector']['slope'], 3), round(CALIB['check_fiserv_subsector']['slope'], 3))


# 7b. size bridge
def size_gap(ref, base, lo=BTOS_WINDOW[0], hi=BTOS_WINDOW[1], gtype='size', prefix=''):
    """per-period mean of the reference classes minus the base (a class or 'national'); mean and SE over periods"""
    R = pd.concat([bt_series(gtype, prefix + c, lo, hi) for c in ref], axis=1).mean(axis=1)
    Bs = bt_series('national', 'US', lo, hi) if base == 'national' else bt_series(gtype, prefix + base, lo, hi)
    d = (R - Bs).dropna()
    return float(d.mean()), float(d.std(ddof=1) / math.sqrt(len(d))), int(len(d))


def to_growth(gap, se_gap):
    v = BSL * gap
    return v, math.sqrt((BSL_SE * gap) ** 2 + (BSL * se_gap) ** 2)


GAP_A = size_gap(SIZE_REF, SIZE_CHASE)
GAP_NAT = size_gap(SIZE_REF, 'national')
S_VAL, S_SE = to_growth(GAP_A[0], GAP_A[1])
S_NAT, S_NAT_SE = to_growth(GAP_NAT[0], GAP_NAT[1])
S_Z = S_VAL / S_SE
S_APPLIED = bool(S_Z > Z_APPLY)
S = S_VAL if S_APPLIED else 0.0
SE_S = S_SE if S_APPLIED else 0.0
# by sector (heterogeneity) and in 2023-24
sec_gap = {}
for sec in sorted(set(k.split('|')[0] for k in BT[BT.gtype == 'sector_size'].key)):
    try:
        g_, se_, n_ = size_gap(SIZE_REF, SIZE_CHASE, gtype='sector_size', prefix=sec + '|')
    except (ValueError, ZeroDivisionError):
        continue
    if n_ >= 20 and np.isfinite(g_):
        sec_gap[sec] = dict(nb_gap=g_, se=se_, n=n_, growth=BSL * g_)
GAP_A_2324 = size_gap(SIZE_REF, SIZE_CHASE, 202300, 202499)
# SBCS check: 12-month revenue change by employment size, latent median (mu / s), surveys 2023-25
sb = SBCS[(SBCS.gtype.isin(['employees', 'all'])) & SBCS.survey_year.isin([2023, 2024, 2025])].copy()
sbcs_size = {}
for Y, g in sb.groupby('survey_year'):
    gg = g.set_index('group')
    ref = float(np.mean([gg.loc['10-19 employees', 'mu_s'], gg.loc['20-49 employees', 'mu_s']]))
    sbcs_size[int(Y)] = dict(ref_10_49=ref, emp_1_4=float(gg.loc['1-4 employees', 'mu_s']), all=float(gg.loc['All employer firms', 'mu_s']),
                             d_vs_1_4=ref - float(gg.loc['1-4 employees', 'mu_s']), d_vs_all=ref - float(gg.loc['All employer firms', 'mu_s']),
                             se_d_vs_1_4=float(math.sqrt(np.mean([gg.loc['10-19 employees', 'se_mu_s'] ** 2, gg.loc['20-49 employees', 'se_mu_s'] ** 2]) / 2
                                                         + gg.loc['1-4 employees', 'se_mu_s'] ** 2)))
sb_d14 = float(np.mean([v['d_vs_1_4'] for v in sbcs_size.values()])); sb_dall = float(np.mean([v['d_vs_all'] for v in sbcs_size.values()]))
sb_se14 = float(math.sqrt(sum(v['se_d_vs_1_4'] ** 2 for v in sbcs_size.values())) / len(sbcs_size))
SBCS_SIZE = dict(by_survey=sbcs_size, d_vs_1_4_s_units=sb_d14, se_s_units=sb_se14, d_vs_all_s_units=sb_dall,
                 pts_vs_1_4={f'{s:.2f}': sb_d14 * s for s in (SPREAD_KEPT, SPREAD_CENTRAL)},
                 pts_vs_all={f'{s:.2f}': sb_dall * s for s in (SPREAD_KEPT, SPREAD_CENTRAL)})
# Chase check: payroll net shares by revenue size (2025), mapped to median payroll growth with the normal sd implied by
# Chase's own median and net share (CHASE_PAYROLL)
sd_pay = float(np.mean([med / stats.norm.ppf((1 + net) / 2) for med, net in CHASE_PAYROLL.values()]))
ns = {k: float(np.mean(v)) / 100 for k, v in CHASE_PAYROLL_SIZE_2025.items()}
med_pay = {k: sd_pay * stats.norm.ppf((1 + v) / 2) for k, v in ns.items()}
CHASE_SIZE = dict(net_share_2025=ns, sd_normal=sd_pay, median_payroll_growth=med_pay, larger_minus_all=med_pay['larger'] - med_pay['all'],
                  larger_minus_smaller=med_pay['larger'] - med_pay['smaller'])
print('size bridge: nb gap C,D - A', round(GAP_A[0], 4), '->', round(S_VAL, 4), 'se', round(S_SE, 4), 'z', round(S_Z, 2), 'applied', S_APPLIED,
      '| vs national', round(S_NAT, 4), '| SBCS pts', {k: round(v, 4) for k, v in SBCS_SIZE['pts_vs_1_4'].items()},
      '| Chase payroll larger-all', round(CHASE_SIZE['larger_minus_all'], 4))
# 7c. BTOS-implied industry deviations (ordering check) against Census D
BT_IND = {}
for k, (gt, key) in BTOS_CELL.items():
    s = bt_series(gt, key)
    if len(s) < 10:
        continue
    gap = float(s.mean() - bt_series('national', 'US').reindex(s.index).mean())
    BT_IND[k] = dict(cell=gt + ' ' + key, nb=float(s.mean()), nb_gap=gap, D=BSL * gap, n=int(len(s)),
                     se=math.sqrt((BSL_SE * gap) ** 2 + CALIB['resid_sd'] ** 2))
bu = [k for k in MAIN if k in BT_IND and k not in ('constr', 'msp', 'prof', 'land', 'auto')]   # one console industry per BTOS cell
BT_AGREE = dict(unique_cells=agreement([E.loc[k, 'D_avg'] for k in bu], [BT_IND[k]['D'] for k in bu]), cells=bu,
                with_fiserv=agreement([F_IND[k]['D'] for k in bu if k in F_IND], [BT_IND[k]['D'] for k in bu if k in F_IND]))
bu2 = [k for k in bu if k != 'staff']                                  # 561 holds staffing (5613) and much else: a broad cell
BT_AGREE['unique_cells_without_staff'] = agreement([E.loc[k, 'D_avg'] for k in bu2], [BT_IND[k]['D'] for k in bu2])
BT_AGREE['fiserv_vs_census_same_cells_without_staff'] = agreement([E.loc[k, 'D_avg'] for k in bu2 if k in F_IND], [F_IND[k]['D'] for k in bu2 if k in F_IND])
BT_AGREE['cells_without_staff'] = bu2

# ================================================================ 8. pooling and forward values
E['F_D'] = [F_IND[k]['D'] if k in F_IND else np.nan for k in E.index]
E['F_pooled'] = [k in FISERV_POOL for k in E.index]


def fiserv_pool(V):
    """pooled deviation, its SE and Fiserv's weight when Fiserv's 2023-26 deviation has error variance V (FISERV_POOL only)"""
    D, se, w = E.D_avg.astype(float).copy(), E.se_D_avg.astype(float).copy(), pd.Series(0.0, index=E.index)
    for k in FISERV_POOL:
        vC = float(E.loc[k, 'se_D_avg']) ** 2
        w[k] = (1 / V) / (1 / vC + 1 / V)
        D[k] = (1 - w[k]) * float(E.loc[k, 'D_avg']) + w[k] * F_IND[k]['D']
        se[k] = math.sqrt(1 / (1 / vC + 1 / V))
    return D, se, w


E['D_pool'], E['se_D_pool'], E['w_fiserv'] = fiserv_pool(V_F4)
E['D_pool_vf2'] = fiserv_pool(V_F2)[0]           # sensitivity: no reduction for the 4-year mean
E['D_pool_r2'], _, E['w_fiserv_r2'] = fiserv_pool(V_F4_R2)   # revision 2: V_F2 halved
E['lam'] = [LAM[GROUP[k]] for k in E.index]
E['se_lam'] = [SE_LAM[GROUP[k]] for k in E.index]
LEVEL = M + S
SE_COMMON = math.sqrt(SE_M ** 2 + SE_S ** 2 + SE_SEL ** 2)    # common to every industry: M, S and selection into sale
E['forward'] = LEVEL + E.lam * E.D_pool
E['se_forward_rel'] = np.sqrt((E.lam * E.se_D_pool) ** 2 + (E.D_pool * E.se_lam) ** 2 + (E.lam * SE_BRIDGE) ** 2)
E['se_forward'] = np.sqrt(E.se_forward_rel ** 2 + SE_COMMON ** 2)
E['nowcast'] = LEVEL + E.D_now
E['se_nowcast'] = np.sqrt(E.se_D_now ** 2 + SE_BRIDGE ** 2 + SE_COMMON ** 2)
E['value'] = E.forward; E['se'] = E.se_forward
E['basis'] = ['forward: M + S + lambda * D (Census/BLS' + (' pooled with Fiserv ' + FISERV_POOL[k] if k in FISERV_POOL else '') + ')' for k in E.index]

# ================================================================ 9. restaurants: Chase's direct small-restaurant median, re-checked
D_chase = CHASE_REST['2026-08'] - M
lam_r, se_lam_r = LAM['non_manufacturing'], SE_LAM['non_manufacturing']
rest_val = LEVEL + lam_r * D_chase
rest_gap = float(E.loc['rest', 'forward'] - rest_val)
rest_se = math.sqrt(SE_COMMON ** 2 + (lam_r * SE_CHASE_DEV) ** 2 + (D_chase * se_lam_r) ** 2 + (rest_gap / 2) ** 2)
v_c = SE_CHASE_DEV ** 2; v_k = E.loc['rest', 'se_D_avg'] ** 2 + SE_BRIDGE ** 2
D_pool_r = (D_chase / v_c + E.loc['rest', 'D_avg'] / v_k) / (1 / v_c + 1 / v_k)
# Fiserv re-check: the 722 (restaurants and bars) deviation from Fiserv's all-industry, same month and trailing 12 months
f722 = {m: f_yoy(SA, '722', m) - nat_m[m] for m in last12 + ['2023-08']}
f722_last12 = [f722[m] for m in last12]
f_dev_aug26 = f722['2026-08']
f_sd_month = float(np.std(f722_last12, ddof=1))
REST_FISERV = dict(fiserv_722_dev_aug2026=f_dev_aug26, fiserv_722_dev_last12_mean=float(np.mean(f722_last12)), fiserv_722_dev_month_sd=f_sd_month,
                   fiserv_722511_dev_aug2026=f_yoy(SA, '722511', '2026-08') - nat_m['2026-08'],
                   fiserv_722513_dev_aug2026=f_yoy(SA, '722513', '2026-08') - nat_m['2026-08'],
                   fiserv_722_dev_aug2023=f722['2023-08'], fiserv_722_change_aug2023_aug2026=f_yoy(SA, '722', '2026-08') - f_yoy(SA, '722', '2023-08'),
                   fiserv_national_change_aug2023_aug2026=nat_m['2026-08'] - nat_m['2023-08'],
                   chase_change_aug2023_aug2026=CHASE_REST['2026-08'] - CHASE_REST['2023-08'],
                   chase_dev_aug2026=D_chase, z=(D_chase - f_dev_aug26) / math.sqrt(SE_CHASE_DEV ** 2 + f_sd_month ** 2),
                   fiserv_722_D_2023_26=float(np.mean([FDEV.loc['722', y] for y in FISERV_YEARS])),
                   btos_722_D=BT_IND['rest']['D'], btos_722_se=BT_IND['rest']['se'])
REST_FISERV['rule_confirmed'] = bool(abs(REST_FISERV['z']) <= 2)
REST = dict(chase_aug2026=CHASE_REST['2026-08'], chase_aug2023=CHASE_REST['2023-08'], D_chase=D_chase, lam=lam_r, value=rest_val, se=rest_se,
            census_forward=float(E.loc['rest', 'forward']), census_forward_se=float(E.loc['rest', 'se_forward']),
            gap_census_minus_chase=rest_gap, pooled_alternative=dict(D=float(D_pool_r), value=float(LEVEL + lam_r * D_pool_r),
                                                                   value_rounded=r005(LEVEL + lam_r * D_pool_r)),
            fiserv_recheck=REST_FISERV)
E.loc['rest', 'value'] = rest_val; E.loc['rest', 'se'] = rest_se
E.loc['rest', 'basis'] = "Chase direct small-restaurant median (Aug 2026): M + S + lambda * (2.0% - M); re-checked against Fiserv 722"

# ================================================================ 10. staffing: the formula value (revision 3, analyst: no override)
# The formula already uses the latest readings (QCEW 2025 annual payroll, QSS 2026H1); the 2023-24 observations still carry
# most of the precision weight. Flagged weak if its SE exceeds WEAK_SE (analyst).
STAFF_SE = float(E.loc['staff', 'se_forward'])
STAFF_WEAK = bool(STAFF_SE > WEAK_SE)
staff_obs = [o for o in OBS if o['ind'] == 'staff']
staff_latest = [o for o in staff_obs if o['year'] >= 2025]               # QCEW 2025 and QSS 2026H1
D_STAFF_LATEST = float(np.average([o['deviation'] for o in staff_latest], weights=[1 / o['se'] ** 2 for o in staff_latest]))
STAFF = dict(rule=f'formula value, no override (analyst); flagged weak if its SE exceeds {WEAK_SE:.2f}', value=float(E.loc['staff', 'forward']),
             rounded=r005(float(E.loc['staff', 'forward'])), se=STAFF_SE, se_relative=float(E.loc['staff', 'se_forward_rel']), se_common=SE_COMMON,
             weak=STAFF_WEAK, D=float(E.loc['staff', 'D_pool']), se_D=float(E.loc['staff', 'se_D_pool']), lam=LAM['non_manufacturing'],
             observations={o['obs']: dict(year=o['year'], deviation=o['deviation'], se=o['se'], weight_forward=o['weight_forward']) for o in staff_obs},
             D_latest_2025_26=D_STAFF_LATEST, value_at_latest_D=LEVEL + LAM['non_manufacturing'] * D_STAFF_LATEST,
             nowcast=float(E.loc['staff', 'nowcast']), qss_5613_2026h1=qss_h1('5613')[0],
             se_round2_terms=float(math.sqrt(float(E.loc['staff', 'se_forward_rel']) ** 2 + SE_M_R2 ** 2 + SE_S ** 2)),
             override_round2=dict(value=STAFF_OVERRIDE_R2[0], se=STAFF_OVERRIDE_R2[1]))
E.loc['staff', 'basis'] = ('forward: M + S + lambda * D (Census/BLS); no override (revision 3)'
                           + (f'; flagged weak (SE {STAFF_SE:.3f} > {WEAK_SE:.2f})' if STAFF_WEAK else ''))

# ================================================================ 11. saas and ecom: ship only if population-relevant sources agree
ef = pd.read_csv(EF_LISTINGS, low_memory=False); efs = pd.read_csv(EF_SOLD, low_memory=False)
ef = ef.merge(efs[['listing_number', 'sold_date_est']], on='listing_number', how='left')
ef['yr'] = pd.to_datetime(ef.sold_date_est, errors='coerce').dt.year
ef['g'] = pd.to_numeric(ef.gross_revenue_trend_percent, errors='coerce') / 100
ef = ef[ef.g.notna() & (ef.gross_revenue_trend_percent_months == 12) & (ef.listing_status == 'Sold')]


def ef_group(m):
    """revision 3 (review): exact, case-insensitive seller types; the product sellers are the gm fit's EF_PROD
    (Amazon FBA, Amazon FBM, eCommerce, DropShipping, Subscription-Box); SaaS and (non-box) Subscription are software"""
    m = str(m).strip().lower()
    if m in EF_SAAS:
        return 'saas_sub'
    if m == 'amazon fba':
        return 'amazon_fba'
    if m == 'amazon fbm':
        return 'amazon_fbm'
    if m in EF_PROD:
        return 'ecom_other'                                   # eCommerce, DropShipping, Subscription-Box
    return 'content_other'


def ef_group_r2(m):
    """revision 2's grouping (comparison only): 'Dropship' matched case-sensitively, so DropShipping and Amazon FBM fell
    into content_other, and Subscription-Box went to saas_sub"""
    m = str(m)
    if 'SaaS' in m or 'Subscription' in m:
        return 'saas_sub'
    if 'Amazon FBA' in m:
        return 'amazon_fba'
    if 'Ecommerce' in m or 'Dropship' in m or 'eCommerce' in m or 'Shopify' in m:
        return 'ecom_other'
    return 'content_other'


ef['grp'] = ef.monetization_primary.map(ef_group)
ef['grp_r2'] = ef.monetization_primary.map(ef_group_r2)
ECOM_GRPS = ['amazon_fba', 'amazon_fbm', 'ecom_other']


def med_stats(x):
    x = np.asarray(x, float); iqr = np.quantile(x, .75) - np.quantile(x, .25)
    return dict(n=int(len(x)), median=float(np.median(x)), se_median=float(1.2533 * (iqr / 1.349) / math.sqrt(max(len(x), 1))),
                p25=float(np.quantile(x, .25)), p75=float(np.quantile(x, .75)), share_negative=float((x < 0).mean()))


recent = ef[ef.yr >= 2024]
EF = {g: med_stats(recent[recent.grp == g].g) for g in ['amazon_fba', 'amazon_fbm', 'ecom_other', 'saas_sub', 'content_other']}
_ecom = recent[recent.grp.isin(ECOM_GRPS)]
EF['ecom_pooled'] = med_stats(_ecom.g)
EF['ecom_pooled'].update(types={str(t): int(c) for t, c in _ecom.monetization_primary.value_counts().sort_index().items()},
                         median_annual_revenue=float(pd.to_numeric(_ecom.average_annual_gross_revenue, errors='coerce').median()))
EF['saas_sub']['types'] = {str(t): int(c) for t, c in recent[recent.grp == 'saas_sub'].monetization_primary.value_counts().sort_index().items()}
EF['ecom_pooled_round2_grouping'] = med_stats(recent[recent.grp_r2.isin(['amazon_fba', 'ecom_other'])].g)
EF_ALL = {g: med_stats(ef[ef.grp == g].g) for g in ['amazon_fba', 'amazon_fbm', 'ecom_other', 'saas_sub']}
EF_ALL['ecom_pooled'] = med_stats(ef[ef.grp.isin(ECOM_GRPS)].g)
aq = pd.read_csv(ACQUIRE)
aq_stated = pd.to_numeric(aq.annual_growth_pct, errors='coerce').dropna() / 100
aq_zero = pd.to_numeric(aq.growth_annual_raw_pct, errors='coerce').dropna() / 100
ACQ = dict(stated=med_stats(aq_stated), unstated_as_zero=dict(n=int(len(aq_zero)), median=float(aq_zero.median())))
sc_sd = float(np.median([(p75 - p25) / 1.349 for p25, _, p75 in SAASCAP['fig2'].values()]))
SC_SE = 1.2533 * sc_sd / math.sqrt(SAASCAP['n_min'])


def lvl_src(label, level, se_level, pop, base, sold):
    """a source that measures the median level directly, in its own population and size class. Its deviation (pooled when
    the sources agree) is taken against the level of its size class: M for Empire Flippers' small online sellers (no size
    bridge; revision 3), M + S for console-sized firms. Its SE is its own (sampling) error. G = the level carried to the
    console's size (+ S when the base is M), the quantity the agreement test compares."""
    on_m = base == 'M'
    return dict(source=label, value=level, se=se_level, dev=level - (M if on_m else LEVEL), se_dev=se_level, G=level + (S if on_m else 0.0),
                a_M=0, a_S=int(on_m), sold=int(sold), base=base, population=pop)


def btos_src(label, k, pop, mismatch=0.0):
    """BTOS-implied source: deviation = calibrated net balance; its own SE = slope, calibration residual and, for a cell
    that is not the console's industry, a cell-mismatch term. Level G = M + S + deviation; 'se' adds M's and S's SEs"""
    b = BT_IND[k]
    se_dev = math.sqrt(b['se'] ** 2 + mismatch ** 2)
    return dict(source=label, value=LEVEL + b['D'], se=math.sqrt(SE_M ** 2 + SE_S ** 2 + se_dev ** 2), dev=b['D'], se_dev=se_dev, G=LEVEL + b['D'],
                a_M=1, a_S=1, sold=0, base='M+S', population=pop, mismatch=mismatch, se_btos=b['se'])


def pair_z(a, b):
    """z of two sources' console-size levels G. The variance counts each source's own error, and the SEs of M, S and
    selection into sale only where the two sources differ in them (BTOS carries M, a level source does not; a source on
    base M is carried by S; a sold-business source against an all-firm one)"""
    v = (a['se_dev'] ** 2 + b['se_dev'] ** 2 + (a['a_M'] - b['a_M']) ** 2 * SE_M ** 2 + (a['a_S'] - b['a_S']) ** 2 * SE_S ** 2
         + (a['sold'] - b['sold']) ** 2 * SE_SEL ** 2)
    return (a['G'] - b['G']) / math.sqrt(v), math.sqrt(v)


def pool_online(srcs):
    """inverse-variance pooled deviation of agreeing sources (own SEs only; the common terms are added once to the value)"""
    wts = np.array([1 / s_['se_dev'] ** 2 for s_ in srcs])
    Dk = float(sum(w * s_['dev'] for w, s_ in zip(wts, srcs)) / wts.sum())
    return Dk, float(1 / math.sqrt(wts.sum())), (wts / wts.sum()).tolist()


EF_ECOM_LABEL = 'Empire Flippers e-commerce sellers (Amazon FBA/FBM, eCommerce, DropShipping) sold 2024-26, 12-month trend (median)'
POP = {'saas': [lvl_src('Empire Flippers SaaS/subscription sold 2024-26, 12-month trend (median)', EF['saas_sub']['median'], EF['saas_sub']['se_median'],
                        f"online businesses sold, n={EF['saas_sub']['n']}", base='M', sold=True),
                lvl_src('Acquire.com SaaS listings (ask >= $1M, 2021-23), stated growth (median)', ACQ['stated']['median'], ACQ['stated']['se_median'],
                        f"SaaS businesses listed for sale that state growth, n={ACQ['stated']['n']}", base='M+S', sold=True),
                lvl_src('SaaS Capital 2025 survey, median 2024 growth', SAASCAP['median_all'], SC_SE, 'private B2B SaaS firms, >1,000 (practitioner survey)',
                        base='M+S', sold=False),
                btos_src('BTOS 511 software publishers, 2025-26: M + S + calibrated net balance', 'saas', 'employer firms, firm-weighted (to 202615)')],
       'ecom': [lvl_src(EF_ECOM_LABEL, EF['ecom_pooled']['median'], EF['ecom_pooled']['se_median'],
                        f"online businesses sold, n={EF['ecom_pooled']['n']} (median revenue ${EF['ecom_pooled']['median_annual_revenue'] / 1e3:.0f}k)",
                        base='M', sold=True),
                btos_src('BTOS 454 nonstore retailers, 2025-26: M + S + calibrated net balance', 'ecom',
                         'employer firms, firm-weighted (to 202615; 454 also holds vending and direct selling)', mismatch=SE_454_MISMATCH)]}
NOT_POP = {'saas': [dict(source='Census forward (AIES 5112, QCEW 5132, QSS 5112), M + S + lambda D', value=float(E.loc['saas', 'forward']),
                         se=float(E.loc['saas', 'se_forward']), population='all software publishers (aggregate, large-firm dominated)'),
                    dict(source='QSS 5112 software publishers, 2026H1 YoY', value=qss_h1('5112')[0], se=qss_h1('5112')[1], population='aggregate'),
                    dict(source='Fiserv 51 Information, 2023-26 deviation + M + S', value=LEVEL + F_IND['saas']['D'], se=None,
                         population='all information industries (POS)')],
           'ecom': [dict(source='Census forward (AIES 4541, MRTS 4541), M + S + lambda D', value=float(E.loc['ecom', 'forward']),
                         se=float(E.loc['ecom', 'se_forward']), population='electronic shopping, aggregate (Amazon-dominated)'),
                    dict(source=f'MRTS 4541, {MRTS_YTD} YTD', value=MRTS[('4541', 2026)], se=None, population='aggregate'),
                    dict(source='Fiserv: no nonstore or online series (NAICS 2022 retail by product; POS data)', value=None, se=None, population='-')]}
def pairs_of(srcs):
    out = []
    for i in range(len(srcs)):
        for j in range(i + 1, len(srcs)):
            a_, b_ = srcs[i], srcs[j]
            z, sd = pair_z(a_, b_)
            out.append(dict(a=a_['source'], b=b_['source'], diff=a_['G'] - b_['G'], se=sd, z=z))
    return out


ONLINE = {}
for k, srcs in POP.items():
    pairs = pairs_of(srcs)
    agree = all(abs(p['z']) <= Z_AGREE for p in pairs)
    rec = dict(sources=srcs, pairs=pairs, max_abs_z=max(abs(p['z']) for p in pairs), agree=agree, current=CUR[k], not_population_relevant=NOT_POP[k],
               z_rule='z of the console-size levels G (EF median + S; Acquire and SaaS Capital as published; BTOS M + S + D); the variance counts '
                      'own SEs and the SEs of M, S and selection into sale only where the two sources differ in them')
    if agree:
        Dk, se_Dk, wts = pool_online(srcs)
        lam_k, se_lam_k = LAM['non_manufacturing'], SE_LAM['non_manufacturing']
        val = LEVEL + lam_k * Dk
        se_val = math.sqrt((lam_k * se_Dk) ** 2 + (Dk * se_lam_k) ** 2 + SE_COMMON ** 2)
        alt = {}                                              # the shipped value if one source alone set the deviation
        for s_ in srcs:
            alt[s_['source']] = LEVEL + lam_k * s_['dev']
        rec.update(shipped=True, D=Dk, se_D=se_Dk, weights=wts, value=val, se=se_val, se_relative=math.sqrt((lam_k * se_Dk) ** 2 + (Dk * se_lam_k) ** 2),
                   single_source_values=alt)
        E.loc[k, 'value'] = val; E.loc[k, 'se'] = se_val
        E.loc[k, 'basis'] = 'population-relevant sources agree within 2 SE: M + S + lambda * pooled deviation (weak)'
    else:
        rec.update(shipped=False, kept=CUR[k])
        E.loc[k, 'value'] = np.nan; E.loc[k, 'se'] = np.nan
        E.loc[k, 'basis'] = f'judgment kept ({CUR[k]:.2f}): population-relevant sources disagree (max |z| {rec["max_abs_z"]:.1f} > 2)'
    ONLINE[k] = rec
SHIP = [k for k in INDS if np.isfinite(E.loc[k, 'value'])]
print('online:', {k: (v['agree'], round(v['max_abs_z'], 2), round(v.get('value', float('nan')), 4)) for k, v in ONLINE.items()})

# 11b. evidence on the BTOS 454 cell mismatch (the SE term is the analyst's 2-3 pts; SE_454_MISMATCH = the midpoint). 454 holds
# electronic shopping (4541, the console's e-commerce), vending (4542) and direct selling (4543, incl. fuel dealers). A
# firm-weighted net balance of 454 differs from 4541's by (1 - firm share of 4541) x (the other firms' growth - 4541's).
_s22 = SU[2022][SU[2022]['size'] == '08'].drop_duplicates('NAICS').set_index('NAICS')['Firms'].astype(float)
SH_4541 = float(_s22['4541'] / _s22['454'])
SH_OTHER = {c: float(_s22[c] / _s22['454']) for c in ('4542', '4543', '45431')}
mm_rows = []
for y in range(2013, 2025):                                  # AIES revenue (aggregate within each sub-industry)
    g41 = ts_growth('4541', y)
    go = (SH_OTHER['4542'] * ts_growth('4542', y) + SH_OTHER['4543'] * ts_growth('4543', y)) / (SH_OTHER['4542'] + SH_OTHER['4543'])
    mm_rows.append(dict(source='AIES', year=y, g_4541=g41, g_other=go, mismatch=(1 - SH_4541) * (go - g41)))
for y, prev in ((2025, 2024), (2026, '2025ytd')):          # MRTS sales: other = 454 - 4541 (the BTOS window)
    go = (MRTS_LV[('454', y)] - MRTS_LV[('4541', y)]) / (MRTS_LV[('454', prev)] - MRTS_LV[('4541', prev)]) - 1
    mm_rows.append(dict(source='MRTS' + (f' ({MRTS_YTD})' if y == 2026 else ''), year=y, g_4541=MRTS[('4541', y)], g_other=go,
                        mismatch=(1 - SH_4541) * (go - MRTS[('4541', y)])))
MMD = pd.DataFrame(mm_rows)
_a = MMD[(MMD.source == 'AIES') & (MMD.year != 2020)]
_two = [float(MMD[(MMD.source == 'AIES') & MMD.year.isin(p)].mismatch.mean()) for p in ((2013, 2014), (2015, 2016), (2017, 2018), (2021, 2022), (2023, 2024))]
_w = MMD[MMD.source != 'AIES'].mismatch
MISMATCH_454 = dict(used=SE_454_MISMATCH, lead_range=[0.02, 0.03], firm_share_4541_susb2022_lt500=SH_4541, firm_shares_other=SH_OTHER,
                    aies_annual_2013_24_excl_2020=dict(mean=float(_a.mismatch.mean()), rms=float(np.sqrt((_a.mismatch ** 2).mean())), n=int(len(_a))),
                    aies_two_year_means=dict(windows=['2013-14', '2015-16', '2017-18', '2021-22', '2023-24'], values=_two,
                                             rms=float(np.sqrt(np.mean(np.square(_two))))),
                    btos_window_mrts_2025_26=dict(values=[float(x) for x in _w], mean=float(_w.mean())),
                    fuel_dealers_mrts=dict(g2025=MRTS[('45431', 2025)], g2026_ytd=MRTS[('45431', 2026)]),
                    reading=(f'{(1 - SH_4541) * 100:.0f}% of 454 firms are not e-commerce; that share times their revenue growth gap to '
                             f'4541 had an RMS of {np.sqrt((_a.mismatch ** 2).mean()) * 100:.1f} pts a year in AIES 2013-24 (2020 left out; '
                             f'fuel-price swings in fuel dealers) and {np.sqrt(np.mean(np.square(_two))) * 100:.1f} pts for 2-year means, but '
                             f'{_w.abs().max() * 100:.1f} pt at most in MRTS 2025-26 (the BTOS window); aggregate sub-industry growth stands '
                             'in for the typical firm; the definitional gap between online businesses for sale and 4541 is not measured. '
                             'The midpoint of the analyst\'s 2-3 pts is used.'),
                    rows=MMD.round(6).to_dict('records'))


def ecom_variant(mismatch=SE_454_MISMATCH, ef_base='M', ef_stats=None):
    """ecom value under alternative choices (cell-mismatch term, EF base, EF grouping): value, rounded, EF weight, z"""
    st = ef_stats or EF['ecom_pooled']
    srcs = [lvl_src('EF', st['median'], st['se_median'], '-', base=ef_base, sold=True), btos_src('BTOS 454', 'ecom', '-', mismatch=mismatch)]
    Dk, _, wts = pool_online(srcs)
    v = LEVEL + LAM['non_manufacturing'] * Dk
    return dict(value=v, rounded=r005(v), w_ef=wts[0], z=pair_z(srcs[0], srcs[1])[0])


ECOM_SENS = {'shipped: mismatch 2.5 pts, EF against M, EF_PROD set': ecom_variant(),
             'mismatch 2.0 pts': ecom_variant(mismatch=0.02), 'mismatch 3.0 pts': ecom_variant(mismatch=0.03),
             'no mismatch term': ecom_variant(mismatch=0.0), 'EF against M + S (revision 2 base)': ecom_variant(ef_base='M+S'),
             'EF revision-2 grouping (n 169)': ecom_variant(ef_stats=EF['ecom_pooled_round2_grouping']),
             'revision 2 choices (no mismatch, EF against M + S, old grouping) at the revision-3 lambda':
                 ecom_variant(mismatch=0.0, ef_base='M+S', ef_stats=EF['ecom_pooled_round2_grouping'])}
# saas (not shipped): the same rule with a 2.5-pt cell-mismatch term on BTOS 511 too (robustness of the judgment verdict)
_saas_mm = [s_ if s_.get('mismatch') is None else btos_src(s_['source'], 'saas', s_['population'], mismatch=SE_454_MISMATCH) for s_ in POP['saas']]
SAAS_MAXZ_WITH_511_MISMATCH = max(abs(p['z']) for p in pairs_of(_saas_mm))
print('ecom sensitivity:', {k: (round(v['value'], 4), v['rounded'], round(v['w_ef'], 3), round(v['z'], 2)) for k, v in ECOM_SENS.items()})
print('454 mismatch evidence: share 4541', round(SH_4541, 3), '| AIES rms', round(MISMATCH_454['aies_annual_2013_24_excl_2020']['rms'], 4),
      '2-yr rms', round(MISMATCH_454['aies_two_year_means']['rms'], 4), '| MRTS 2025-26', [round(x, 4) for x in _w])

# ================================================================ shipped values
E['shipped'] = [r005(x) if np.isfinite(x) else np.nan for x in E.value]
E['shipped_se'] = [round(float(x), 3) if np.isfinite(x) else np.nan for x in E.se]
E['forward_rounded'] = [r005(x) for x in E.forward]
# each shipped industry's own SE (the shipped SE without the common terms); staffing's rank among them
SE_OWN = pd.Series({k: math.sqrt(max(float(E.loc[k, 'se']) ** 2 - SE_COMMON ** 2, 0.0)) for k in SHIP}).sort_values(ascending=False)
STAFF['se_relative_rank'] = dict(rank=int(list(SE_OWN.index).index('staff')) + 1, of=int(len(SE_OWN)),
                                 next_largest=float(SE_OWN.drop('staff').iloc[0]), next_largest_ind=str(SE_OWN.drop('staff').index[0]))


def ship_at(lam_of=lambda k: LAM[GROUP[k]], level=LEVEL, D_col='D_pool'):
    """rounded values of the 14 formula industries (staff included, revision 3) and the restaurant rule under alternative
    choices (online industries excluded: see ECOM_SENS)"""
    out = {}
    for k in MAIN:
        out[k] = r005(level + lam_of(k) * (D_chase if k == 'rest' else float(E.loc[k, D_col])))
    return out


R1 = json.load(open('/tmp/claude-0/fit6/out/peer_growth.v1/RESULT.json')) if os.path.exists('/tmp/claude-0/fit6/out/peer_growth.v1/RESULT.json') else None
R2 = json.load(open('/tmp/claude-0/fit6/out/peer_growth.v2/RESULT.json')) if os.path.exists('/tmp/claude-0/fit6/out/peer_growth.v2/RESULT.json') else None
SENS = {'no_size_bridge': dict(values=ship_at(level=M)),
        'size_bridge_vs_all_employer_firms': dict(S=S_NAT, values=ship_at(level=M + S_NAT)),
        'census_only_deviations': dict(values=ship_at(D_col='D_avg')),
        'fiserv_error_variance_V_F2_no_reduction': dict(values=ship_at(D_col='D_pool_vf2')),
        'fiserv_error_variance_halved_revision2': dict(values=ship_at(D_col='D_pool_r2')),
        'lambda_main_window_0.56_revision2': dict(values=ship_at(lam_of=lambda k: LAM_R2[GROUP[k]])),
        'lambda_pooled_round1_0.43': dict(values=ship_at(lam_of=lambda k: LAM_POOLED_R1)),
        'lambda_minus_1se': dict(values=ship_at(lam_of=lambda k: LAM[GROUP[k]] - SE_LAM[GROUP[k]])),
        'lambda_plus_1se': dict(values=ship_at(lam_of=lambda k: LAM[GROUP[k]] + SE_LAM[GROUP[k]])),
        'lambda_fiserv_0.16_non_manufacturing': dict(values=ship_at(lam_of=lambda k: FPERS_MAIN['slope'] if GROUP[k] == 'non_manufacturing' else LAM['manufacturing']))}
for s in SENS.values():
    s['changed'] = {k: [float(E.loc[k, 'shipped']), v] for k, v in s['values'].items() if abs(v - float(E.loc[k, 'shipped'])) > 1e-9}
SENS['ecom'] = ECOM_SENS
SENS['staff_override_round2'] = STAFF_OVERRIDE_R2[0]
SENS['round1_shipped'] = R1['js_table']['by_industry'] if R1 else None
SENS['round2_shipped'] = R2['js_table']['by_industry'] if R2 else None
# rounding edges: shipped values within 0.0005 of a rounding boundary
EDGES = {k: float(E.loc[k, 'value']) for k in SHIP if abs((float(E.loc[k, 'value']) * 200) % 1 - 0.5) < 0.1}
print(E[['current', 'D_avg', 'F_D', 'w_fiserv', 'D_pool', 'lam', 'forward', 'value', 'shipped', 'shipped_se']].round(4).to_string())

# ================================================================ 12. checks (round 1's, kept)
rn, rse = float(E.loc['rest', 'nowcast']), float(E.loc['rest', 'se_nowcast'])
rse_cmp = math.sqrt(rse ** 2 - SE_SEL ** 2)       # both sides measure all continuing firms: selection into sale is not compared
REST_OOS = dict(census_nowcast=rn, se=rse, se_comparison=rse_cmp, census_forward=float(E.loc['rest', 'forward']),
                chase_direct_aug2026_plus_S=CHASE_REST['2026-08'] + S, z=(CHASE_REST['2026-08'] + S - rn) / rse_cmp,
                z_round2_se_M=(CHASE_REST['2026-08'] + S - rn) / math.sqrt(rse_cmp ** 2 - SE_M ** 2 + SE_M_R2 ** 2))
mva = [('national, Feb 2026 vs mean of AIES 2024 revenue and QCEW 2025 payroll (firm-weighted)', M, (C['aies2024'] + C['qcew_w25']) / 2),
       ('restaurants, Aug 2023 vs MRTS 722 calendar 2023', CHASE_REST['2023-08'], MRTS[('722', 2023)]),
       (f'restaurants, Aug 2026 vs MRTS 722 {MRTS_YTD}', CHASE_REST['2026-08'], MRTS[('722', 2026)])]
gaps = np.array([a - b for _, a, b in mva])
MVA = dict(rows=[dict(comparison=l, chase_median=a, census_aggregate=b, gap=a - b) for l, a, b in mva],
           gap_mean=float(gaps.mean()), gap_sd=float(gaps.std(ddof=1)), n=len(gaps), se_bridge_used=SE_BRIDGE)
shipped_keys = [k for k in INDS if np.isfinite(E.loc[k, 'shipped'])]
cur = np.array([CUR[k] for k in shipped_keys])
SPEAR = dict(shipped=float(stats.spearmanr(cur, E.loc[shipped_keys, 'shipped'])[0]), forward=float(stats.spearmanr(cur, E.loc[shipped_keys, 'forward'])[0]),
             nowcast=float(stats.spearmanr(cur, E.loc[shipped_keys, 'nowcast'])[0]), n=len(shipped_keys),
             mean_current=float(cur.mean()), mean_shipped=float(E.loc[shipped_keys, 'shipped'].mean()))
sz = {}
for y in (2017, 2022):
    t = SU[y][SU[y]['size'] == '08'].drop_duplicates('NAICS').set_index('NAICS')
    sz[y] = t['Receipts ($1,000)']
cagr_all = (sz[2022]['--'] / sz[2017]['--']) ** 0.2 - 1
srows = []
for k in INDS:
    c = SUSB_CODE[k]
    if c in sz[2017].index and c in sz[2022].index and sz[2017][c] > 0 and sz[2022][c] > 0:
        g = (sz[2022][c] / sz[2017][c]) ** 0.2 - 1
        srows.append(dict(ind=k, naics=c, cagr_small_2017_22=g, dev=g - cagr_all, D_avg=float(E.loc[k, 'D_avg'])))
SS = pd.DataFrame(srows)
SUSB_CHECK = dict(cagr_all_small_firms=float(cagr_all), n=len(SS), r=float(SS.dev.corr(SS.D_avg)),
                  spearman=float(stats.spearmanr(SS.dev, SS.D_avg)[0]), slope=float(np.polyfit(SS.dev, SS.D_avg, 1)[0]))
SECT = {'Manufacturing': ['31-33'], 'Retail': ['44-45'], 'Leisure and hospitality': ['71', '72'], 'Finance and insurance': ['52'],
        'Healthcare and education': ['61', '62'], 'Professional services and real estate': ['53', '54'],
        'Business support and consumer services': ['56', '81'], 'Non-manufacturing goods production and associated services': ['23', '42', '48-49', '22', '21']}


def aies_sector_g(code, y):
    return CVB.loc[code, 'g'] if (y == 2024 and code in CVB.index) else ts_growth(code, y)


sv = []
for Y in (2023, 2024, 2025):
    allrow = SBCS[(SBCS.survey_year == Y) & (SBCS.gtype == 'all')]
    for grp_name, codes in SECT.items():
        s = SBCS[(SBCS.survey_year == Y) & (SBCS.gtype == 'industry') & (SBCS.group == grp_name)]
        if s.empty or allrow.empty:
            continue
        if Y in (2023, 2024):
            cen, _ = fw({c: aies_sector_g(c, Y) for c in codes}) if any(np.isfinite(aies_sector_g(c, Y)) for c in codes) else (np.nan, [])
            tot = C[f'aies{Y}']
        else:
            ok = [c for c in codes if c in QW.index]
            cen = fw({c: QW.loc[c, 'w25'] for c in ok})[0] if ok else np.nan
            tot = C['qcew_w25']
        sv.append(dict(year=Y, sector=grp_name, sbcs_dev=float(s.mu_s.iloc[0] - allrow.mu_s.iloc[0]), census_dev=cen - tot, n=float(s.n.iloc[0])))
SV = pd.DataFrame(sv).dropna(subset=['census_dev'])
m_sv = sm.WLS(SV.sbcs_dev, sm.add_constant(SV.census_dev), weights=SV.n).fit()
SBCS_CHECK = dict(r=float(SV.sbcs_dev.corr(SV.census_dev)), slope_s_units=float(m_sv.params.census_dev), se=float(m_sv.bse.census_dev), n=len(SV))
LN5 = math.log(5)
relbd = {age: 1.139 * (p90 - p10) / (2 * LN5) for age, (p90, p10) in RELBD_FIG5.items()}
sb_all = SBCS[(SBCS.gtype == 'all') & SBCS.survey_year.isin([2017, 2018, 2019, 2022, 2023, 2024, 2025])].set_index('survey_year')
sbcs_s = {int(y): dict(t_s=float(r.t_s), s_band_2_5=0.025 / r.t_s, s_band_5=0.05 / r.t_s) for y, r in sb_all.iterrows()}
chase_pay = {}
for k, (med, net) in CHASE_PAYROLL.items():
    pinc = (1 + net) / 2
    b_ = -med / math.log(2 * (1 - pinc))
    chase_pay[k] = dict(median=med, net_share=net, sd_normal=med / stats.norm.ppf(pinc), sd_laplace_pct=1.139 * b_)
young = [relbd[a] for a in ['5', '6', '7', '8', '9', '10']]
SPREAD = dict(kept=SPREAD_KEPT, re_lbd=dict(range=[min(relbd.values()), max(relbd.values())], age16plus=relbd['16+'], age5_10=[min(young), max(young)]),
              sbcs=dict(range=[min(v['s_band_2_5'] for v in sbcs_s.values()), max(v['s_band_5'] for v in sbcs_s.values())], by_year=sbcs_s),
              chase_payroll=dict(range=[min(min(v['sd_normal'], v['sd_laplace_pct']) for v in chase_pay.values()),
                                        max(max(v['sd_normal'], v['sd_laplace_pct']) for v in chase_pay.values())], by_issue=chase_pay),
              central_guess=SPREAD_CENTRAL)

QUOTES = {
    'denver_snapshot': quote_check(DENVER_PDF, [r'median revenue growth of 4% year-? ?over-year as of February 2026', r'national rate of 5%',
                                                r'Direct Deposit Account and less than \$20M']),
    'denver_snapshot_page_median_revenue': text_check(DENVER_HTML, [r'the median small business generated \$125,200 in revenues during 2025']),
    'pulse_jan_size_chart': quote_check(PULSE_PDF['jan'], [r'Net share of employer businesses increasing payroll YOY',
                                                           r'Smaller businesses \(less than \$1 million in yearly revenue\)',
                                                           r'Larger businesses \(more than \$1 million in yearly revenue\)'], 3, 3),
    'pulse_apr': quote_check(PULSE_PDF['apr'], [r'Net share increasing payroll y/y, 23%', r'Median payroll increase y/y, 6%']),
    'pulse_may': quote_check(PULSE_PDF['may'], [r'Net share increasing payroll y/y, 25%', r'Median payroll increase y/y, 7%']),
    'pulse_sep': quote_check(PULSE_PDF['sep'], [r'unchanged in August \(18%\)', r'median payroll.{0,60}growth came in at 5%',
                                                r'has trended higher since mid-2024']),
    'pulse_sep_page3_chart_title': quote_check(PULSE_PDF['sep'], [r'Median inflow growth: August 2023 vs\. August 2026'], 3, 3),
    'fiserv_sep2026_report': quote_check(FISERV_PDF, [r'Overall sales rose \+2\.2% YoY', r'Full-?\s?Service Restaurants held steady at \+0\.8% YoY',
                                                      r'creation and closure of businesses', r'aggregated merchant sales from credit card authorizations, check']),
    'fiserv_faq_method_change': text_check(FISERV_FAQ, [r'In December 2024 we changed how we use the quarterly establishment counts',
                                                        r'2M\+ businesses extrapolated to represent census population by state and industry']),
    'carra_page50_figure5': quote_check(CARRA_PDF, [r'Figure 5 Output Growth Distribution for Continuing Firms, 1996-2013'], 50, 50),
    'saas_capital': quote_check(SAASCAP_PDF, [r'overall median growth rate for all companies in the survey registered 25%',
                                              r'Bootstrapped companies report median growth of 23%',
                                              r'6\.9% of the companies reported flat or negative growth'])}
quotes_ok = all(all(q['found'].values()) for q in QUOTES.values() if q.get('checked'))

# ================================================================ tests (assertions on the shipped table)
WEAK = [k for k in INDS if (k in CURRENT_JUDGMENT_KEYS and k in SHIP) or (k == 'staff' and STAFF_WEAK)]
_ef_types = set(t.strip().lower() for t in EF['ecom_pooled']['types'])
TESTS = {}
TESTS['shipped values finite, on the 0.005 grid, SEs >= the common SE (M, S, selection)'] = bool(all(
    np.isfinite(E.loc[k, 'shipped']) and abs(E.loc[k, 'shipped'] * 200 - round(E.loc[k, 'shipped'] * 200)) < 1e-9
    and E.loc[k, 'shipped_se'] >= round(SE_COMMON, 3) - 1e-9 for k in SHIP))
TESTS['the 14 main-street industries ship'] = all(k in SHIP for k in MAIN)
TESTS['lambdas follow the analyst (non-mfg: adjacent pre-COVID mean; mfg: main window), within 0.01'] = (
    abs(LAM_NONMFG - LAM['non_manufacturing']) <= 0.01 and abs(LAM_MFG - LAM['manufacturing']) <= 0.01
    and LAM['manufacturing'] == LAM_R2['manufacturing'])
TESTS['lambda SEs include the across-window SD (and the analyst\'s .10 sampling SE)'] = (
    SE_LAM_RAW['non_manufacturing'] >= math.hypot(SE_LAM_SAMPLING_NONMFG, WIN_SD['non_manufacturing']) - 1e-12
    and SE_LAM_RAW['manufacturing'] >= WIN_SD['manufacturing'])
TESTS['staff ships the formula value (no override); weak exactly when its SE > .03'] = bool(
    abs(float(E.loc['staff', 'value']) - float(E.loc['staff', 'forward'])) < 1e-12 and (('staff' in WEAK) == (STAFF_SE > WEAK_SE)))
TESTS['Fiserv 4-year error variance uses the measured gap correlation (not halved)'] = bool(
    abs(V_F4 - V_F2 * (1 + 3 * RHO_F) / (2 * (1 + RHO_F))) < 1e-15 and (V_F4 > V_F4_R2 if RHO_F > 0 else True)
    and abs(FIS_ERR['v_fiserv_2yr_from_moments'] - V_F2) < 1e-9)
TESTS['EF e-commerce set = the gm fit product sellers (case-insensitive; FBM and DropShipping in)'] = bool(
    _ef_types <= set(EF_PROD) and EF['ecom_pooled']['n'] == int(recent.monetization_primary.astype(str).str.strip().str.lower().isin(EF_PROD).sum()))
TESTS['ecom: BTOS 454 carries the cell-mismatch term; EF measured against M (no size bridge)'] = bool(
    abs(POP['ecom'][1]['mismatch'] - SE_454_MISMATCH) < 1e-15 and abs(POP['ecom'][0]['dev'] - (EF['ecom_pooled']['median'] - M)) < 1e-15
    and 0.02 <= SE_454_MISMATCH <= 0.03)
TESTS['measured bridge slope rounds to the decision (1.10)'] = round(BRIDGE['slope'], 2) == BR_DECISION
TESTS['bridge slope significant (|t| > 3)'] = abs(BRIDGE['slope'] / BRIDGE['se']) > 3
TESTS['Fiserv API reproduces the Sep 2026 report Figure 2 (13 months, 0.05 pt)'] = (not FIG2['checked']) or FIG2['max_abs_diff'] <= 0.0005 + 1e-9
TESTS['BTOS revenue index = 50 + 50 * net balance (parsing)'] = IDX_CHECK <= 0.11
TESTS['size bridge decision follows the 2-SE rule'] = S_APPLIED == (S_Z > Z_APPLY)
TESTS['saas/ecom decisions follow the 2-SE agreement rule'] = all(ONLINE[k]['agree'] == (k in SHIP) for k in ONLINE)
TESTS['restaurant OOS (Census nowcast vs Chase + S) within 2 SE'] = abs(REST_OOS['z']) < 2
TESTS['every PDF quote found (when pdftotext is installed)'] = bool(quotes_ok)
for t, ok in TESTS.items():
    print(('PASS ' if ok else 'FAIL ') + t)
# outcomes that are findings, not assertions
OUTCOMES = {'fiserv_supports_M': FIS_NAT['supports_M'], 'fiserv_gap_within_widened_SE_M': FIS_NAT['within_widened_SE_M'],
            'fiserv_persistence_differs_from_aies_non_mfg': bool(abs(FPERS_VS['z']) > 2),
            'size_bridge_applied': S_APPLIED, 'restaurant_rule_confirmed_by_fiserv': REST_FISERV['rule_confirmed'],
            'saas_ships': ONLINE['saas']['agree'], 'ecom_ships': ONLINE['ecom']['agree'], 'staff_weak': STAFF_WEAK,
            'fiserv_gap_correlation_positive': bool(RHO_GAP - 2 * RHO_GAP_SE > 0)}

# ================================================================ 13. output
src = [
    dict(name='JPMorganChase internal data via JPMCI Small business snapshot: Denver (May 2026)', url='https://www.jpmorganchase.com/content/dam/jpmorganchase/institute/documents/denver-snapshot.pdf',
         file=DENVER_PDF, years='Feb 2026', n='Chase small-business customers with a direct deposit account and < $20M revenue (count not given)', used='M = 5% (national), Denver 4%'),
    dict(name='JPMCI Small business snapshot: Denver, web page (Figure 2 text)', url='https://www.jpmorganchase.com/institute/all-topics/business-growth-and-entrepreneurship/denver-small-business-snapshot',
         file=DENVER_HTML, years='2025', n='-', used="median 2025 revenue of the sample, $125,200 (size of Chase's median customer: the size-bridge note)"),
    dict(name='Fiserv Small Business Index, public data API (index-details, SA and NSA, all NAICS cells incl. NAICS-6)',
         url='https://fsbi.fiservapps.com/index-data/api/v1/index-details/adjusted (POST; and .../index-details for NSA)', file=FISERV_API + '/{adjusted,nsa}_US_*.json',
         years='Jan 2019 - Sep 2026 (monthly)', n=f'{len(fiserv_files)} files; {len(FG)} NAICS cells ({len(leaf_codes)} six-digit leaves); about 2M small businesses (POS)',
         used='M test, industry deviations (pooled for 6), persistence, agreement with AIES, restaurant re-check', sha256_note='hash of the files concatenated in name order'),
    dict(name='Fiserv Small Business Index Monthly Overview, September 2026', url='https://www.fiserv.com/content/dam/fiserv-ent/final-files/marketing-collateral/documents-handouts-flyers/fsbi-september-2026-monthly-overview-1026.pdf',
         file=FISERV_PDF, years='Sep 2025 - Sep 2026', n='12 pages', used='methodology (aggregate incl. creation and closure), Figure 2 reproduction'),
    dict(name='Census Business Trends and Outlook Survey (BTOS), revenue question by employment size, sector, subsector, sector x size',
         url='https://www.census.gov/hfp/btos/downloads/{National,Employment Size Class,Sector,Subsector,Sector by Employment Size Class}.xlsx',
         file=BTOS_DIR + '/*.xlsx', years='collection periods 202319-202620 (2025-26 used)', n=f'{int(BT.period.nunique())} periods; {int(BT.key.nunique())} groups',
         used='size bridge, sector calibration, ordering check, BTOS 454 and 511 for ecom and saas', sha256_note='hash of the five workbooks concatenated in name order'),
    dict(name='Census AIES time series (released 2026-09-10)', url='https://www2.census.gov/programs-surveys/aies/data/2024/API_Datasets/AIESTIMESERIES.zip',
         file=AIES_TS, years='1992-2024 (services from 2013)', n=f'{n_aies_ts_rows} US rows; {TS.shape[0]} NAICS; {len(six)} six-digit', used='revenue growth 2023-24, persistence, bridge, Fiserv cell comparison'),
    dict(name='Census AIES 2023 basic table', url='https://www2.census.gov/programs-surveys/aies/data/2023/AIES00BASIC.zip', file=AIES_BASIC[2023], years='2023', n=f'{len(CVB)} NAICS matched to 2024', used='CVs, 238 and sector 2024 growth'),
    dict(name='Census AIES 2024 basic table', url='https://www2.census.gov/programs-surveys/aies/data/2024/AIES00BASIC.zip', file=AIES_BASIC[2024], years='2024', n=f'{len(CVB)} NAICS matched to 2023', used='CVs, 238 and sector 2024 growth'),
    dict(name='BLS QCEW annual averages, US private (own_code 5)', url='https://data.bls.gov/cew/data/api/{2022..2025}/a/industry/{code}.csv',
         file=QCEW_DIR + '/{year}_{code}.csv', years='2022-2025', n=f'{n_qcew_files} files, {n_qcew_rows} industry-years', used='payroll growth 2023-25, BTOS calibration', sha256_note='hash of the files concatenated in name order'),
    dict(name='Census Quarterly Services Survey, 2026Q2 release', url='https://www.census.gov/services/qss/qss-all-current.xlsx', file=QSS_XLSX, years='2025Q1-2026Q2',
         n=f'{sum(1 for v in QSS.values() if isinstance(v.get("yoy_2q26"), float))} codes with YoY', used='2026H1 YoY revenue and SEs'),
    dict(name='Census Monthly Retail Trade', url='https://www.census.gov/retail/mrts/www/mrtssales92-present.xlsx', file=MRTS_XLSX, years=f'2022-2026 ({MRTS_YTD})',
         n='total, 722, 4541, 454, 45431', used='restaurant and e-commerce sales growth; 454 cell-mismatch evidence (454 less 4541, fuel dealers)'),
    dict(name='Census SUSB 2022, US x 6-digit NAICS x employment size', url='https://www2.census.gov/programs-surveys/susb/tables/2022/us_state_6digitnaics_2022.xlsx', file=SUSB[2022], years='2022',
         n=f'{len(SU[2022])} US rows', used='sector weights (firms < 500 employees); persistence check; firm shares inside 454'),
    dict(name='Census SUSB 2017, US x 6-digit NAICS x employment size', url='https://www2.census.gov/programs-surveys/susb/tables/2017/us_state_6digitnaics_2017.xlsx', file=SUSB[2017], years='2017',
         n=f'{len(SU[2017])} US rows', used='persistence check'),
    dict(name='Chase Small Business Pulse, Sep 2026', url='https://media.chase.com/content/dam/chase/media-center/intelligence/sb-pulse-sep-2026.pdf', file=PULSE_PDF['sep'], years='Aug 2023, Aug 2026',
         n='Chase business deposit accounts', used='restaurant median 2.0% / 5.9% (p. 3 chart); payroll median 5%, net share 18%'),
    dict(name='Chase Small Business Pulse, Jan 2026', url='https://media.chase.com/content/dam/chase/media-center/intelligence/smb-pulse-january-2026.pdf', file=PULSE_PDF['jan'], years='2024-2025',
         n='Chase business deposit accounts (employer businesses)', used='payroll net share by revenue size (p. 3 chart, size-bridge check)'),
    dict(name='Chase Small Business Pulse, Apr 2026', url='https://media.chase.com/content/dam/chase/media-center/intelligence/april-sb-pulse2026.pdf', file=PULSE_PDF['apr'], years='2026', n='-', used='payroll median 6%, net share 23%'),
    dict(name='Chase Small Business Pulse, May 2026', url='https://media.chase.com/content/dam/chase/media-center/intelligence/may-sb-pulse2026.pdf', file=PULSE_PDF['may'], years='2026', n='-', used='payroll median 7%, net share 25%'),
    dict(name='Fed Small Business Credit Survey, employer-firm appendix workbooks (surveys 2017-2025)', url='https://www.fedsmallbusiness.org (report pages, appendix .xlsx)',
         file=SBCS_DIR + '/20*_sbcs-*appendix*.xlsx', years='2017-2025', n=f"{int(SBCS[SBCS.gtype == 'all'].n.sum())} firm responses (all-firm rows)", used='revenue change shares (spread, sector ordering, size-bridge check)'),
    dict(name='Haltiwanger, Jarmin, Kulick, Miranda, CARRA WP 2017-03 (revenue-enhanced LBD)', url='https://www.census.gov/content/dam/Census/library/working-papers/2017/adrm/carra-wp-2017-03.pdf',
         file=CARRA_PDF, years='1996-2013', n='continuing firms (published percentiles)', used='Figure 5, p. 50 (spread check)'),
    dict(name='SaaS Capital, 2025 B2B Benchmarking Private SaaS Company Growth Rates', url='https://www.saas-capital.com/wp-content/uploads/2025/06/RB33WS5-2025-B2B-Benchmarking-Private-SaaS-Company-Growth-Rates.pdf',
         file=SAASCAP_PDF, years='2024 growth', n='>1,000 private B2B SaaS firms', used='saas source (not shipped)'),
    dict(name='Empire Flippers sold listings (public marketplace API)', url='https://api.empireflippers.com/api/v1/listings/list', file=EF_LISTINGS, years='sold 2024-2026',
         n=f"{int(len(recent))} sold listings with a 12-month trend", used='ecom and saas sources'),
    dict(name='Empire Flippers sold dates (derived)', url='https://api.empireflippers.com/api/v1/listings/list', file=EF_SOLD, years='2014-2026', n=f'{len(efs)} sold listings', used='sale year'),
    dict(name='Acquire.com SaaS listings, ask >= $1M (Kaggle extract)', url='local: /tmp/claude-0/data5/kaggle', file=ACQUIRE, years='listed 2021-06..2023-04',
         n=f"{len(aq)} listings ({ACQ['stated']['n']} state growth)", used='saas source (not shipped)'),
    dict(name='Console snapshot (current hand-set IND.grw)', url='-', file=INDEX, years='HEAD 56ba230', n='16', used='current values')]
for s in src:
    if s['file'].endswith('{code}.csv'):
        s['sha256'] = qcew_cat_sha[:16]
    elif s['file'].startswith(FISERV_API):
        s['sha256'] = sha_files(fiserv_files)[:16]
    elif s['file'].startswith(BTOS_DIR):
        s['sha256'] = sha_files([os.path.join(BTOS_DIR, f) for f in BTOS_FILES.values()])[:16]
    elif '*' in s['file']:
        s['sha256'] = sha_files(sbcs_files)[:16]; s['files'] = [os.path.basename(f) for f in sbcs_files]
    else:
        s['sha256'] = sha(s['file'])[:16]

pct = lambda x: f'{x * 100:.1f}%'
W_LO, W_HI = float(E.loc[list(FISERV_POOL), 'w_fiserv'].min()), float(E.loc[list(FISERV_POOL), 'w_fiserv'].max())
EO = ONLINE['ecom']
NM_WIN_TXT = ', '.join(f"{NM_WIN[lab]['slope']:.2f}" for lab in LAM_WINDOWS_NONMFG)
MFG_RANGE = BY_GROUP['manufacturing_six_digit']['range']
ONLINE_TXT = '; '.join(
    (f"{k} ships flagged weak (population-relevant sources agree, max |z| {ONLINE[k]['max_abs_z']:.1f}"
     + (f"; Empire Flippers measured against M, BTOS 454 with a {SE_454_MISMATCH * 100:.1f}-pt cell-mismatch SE)" if k == 'ecom' else ')'))
    if ONLINE[k]['agree'] else
    f"{k} stays judgment at {CUR[k]:.2f} (population-relevant sources disagree, max |z| {ONLINE[k]['max_abs_z']:.1f})" for k in CURRENT_JUDGMENT_KEYS)
method = ('grw_i = M + S + lambda_g * D_i. M = 5.0% nominal (Chase small-business median YoY revenue growth, Feb 2026), the median of continuing '
          f'firms the console needs; its SE is widened to {SE_M * 100:.1f} pts to cover the conflict with Fiserv\'s all-industry sales '
          f"({FIS_NAT['feb2026_sa'] * 100:+.1f}% YoY in Feb 2026; Fiserv measures aggregate point-of-sale sales, entry and exit included). "
          f'S = {S * 100:+.1f} pts: size bridge from Chase\'s typical customer to the console\'s typical business (Census BTOS revenue net balance of '
          '10-49- minus 1-4-employee firms, 2025-26, x the cross-sector slope of QCEW payroll growth on the BTOS net balance); S is conservative, '
          'because Chase\'s median customer ($125k revenue in 2025) may be a nonemployer, smaller than the 1-4 class. D_i = precision-weighted '
          'deviation of industry i from the firm-weighted all-industry aggregate (Census AIES 2023-24, BLS QCEW payroll 2023-25 x 1.10, QSS 2026H1, '
          'MRTS 2025-26), pooled by inverse variance with the Fiserv Small Business Index deviation 2023-26 for medical, vet, dist, wash, auto and fit '
          f"(Fiserv's error variance measured from its disagreement with AIES over {CELLS['all']['n']} matched cells; its 4-year mean uses the "
          f"measured year-to-year correlation of Fiserv's errors, {RHO_F:.2f}, not independence; weights {W_LO * 100:.0f}-{W_HI * 100:.0f}%). "
          f"lambda_g = {LAM['non_manufacturing']:.2f} (non-manufacturing: mean of the three adjacent pre-COVID AIES windows) and "
          f"{LAM['manufacturing']:.2f} (manufacturing: 2017-19 -> 2023-24), the persistence of multi-year industry deviations; their SEs add the "
          'across-window SD. Restaurants use Chase\'s direct small-restaurant median (2.0%, Aug 2026). Staffing ships the formula value (no override)'
          + (f', flagged weak (SE above {WEAK_SE:.2f})' if STAFF_WEAK else '') + '. ' + ONLINE_TXT
          + f'. Values rounded to the nearest 0.005. Every SE includes three terms common to all industries: M {SE_M * 100:.1f} pts, S {SE_S * 100:.2f} '
          f'and selection into sale {SE_SEL * 100:.1f} (judgment), together {SE_COMMON * 100:.1f} pts.')
source = ('fit/peer_growth.py (revision 3): JPMorganChase internal data (median YoY revenue growth of Chase small-business customers, Feb 2026, '
          'quoted in the JPMCI Denver snapshot, May 2026); Census BTOS 2025-26 (size bridge, ecom); Census AIES, QSS, MRTS, SUSB 2022; BLS QCEW 2022-25; '
          'Fiserv Small Business Index 2019-Sep 2026 (test of M, industry deviations, persistence); Fed SBCS 2023-25 and Chase Pulse Jan 2026 (size checks); '
          'Chase Small Business Pulse Sep 2026 (restaurants); Empire Flippers sold listings 2024-26 (ecom). Nominal growth.')
by_ind = {k: float(E.loc[k, 'shipped']) for k in INDS if k in SHIP}
se_ind = {k: float(E.loc[k, 'shipped_se']) for k in INDS if k in SHIP}
weak = list(WEAK)
not_shipped = {}
for k in CURRENT_JUDGMENT_KEYS:
    o = ONLINE[k]
    if o['agree']:
        continue
    not_shipped[k] = dict(kept=CUR[k], tier='judgment', rule='ship only if every pair of population-relevant sources agrees within 2 SE',
                          max_abs_z=round(o['max_abs_z'], 2),
                          sources={s['source']: [round(s['value'], 3), round(s['se'], 3)] for s in o['sources']},
                          other={s['source']: round(s['value'], 3) for s in o['not_population_relevant'] if s['value'] is not None})
    if k == 'saas':
        not_shipped[k]['max_abs_z_with_511_mismatch'] = round(SAAS_MAXZ_WITH_511_MISMATCH, 2)
z_r2, z_r3 = FIS_NAT['z_vs_SE_M_round2'], FIS_NAT['z_vs_SE_M']
notes = {'rest': f"Chase's direct small-restaurant median (2.0% Aug 2026) replaces the Census deviation: M + S + lambda (2.0% - M); "
                 f"Fiserv's restaurant deviation (722, Aug 2026) is {f_dev_aug26 * 100:+.1f} pts against Chase's {D_chase * 100:+.1f} "
                 f"(z {REST_FISERV['z']:+.1f}); the Census forward value is {E.loc['rest', 'forward'] * 100:.1f}%",
         'staff': (f"formula value, no override: M + S + lambda D with D {E.loc['staff', 'D_pool'] * 100:+.1f} pts (Census AIES 2023-24, QCEW 2023-25, "
                   f"QSS 2026H1) gives {E.loc['staff', 'forward'] * 100:+.1f}%; "
                   + (f"flagged weak: SE {STAFF_SE:.3f} > {WEAK_SE:.2f} (its own part {STAFF['se_relative']:.3f}, "
                      + (f"the largest of the {STAFF['se_relative_rank']['of']} shipped" if STAFF['se_relative_rank']['rank'] == 1 else
                         f"rank {STAFF['se_relative_rank']['rank']} of {STAFF['se_relative_rank']['of']}") + ')' if STAFF_WEAK else f"SE {STAFF_SE:.3f}")
                   + f"; the latest readings alone (QCEW 2025, QSS 2026H1) put D at {D_STAFF_LATEST * 100:+.1f} pts"),
         'fiserv': (f"Fiserv deviations pooled for {', '.join(FISERV_POOL)} with weight {W_LO * 100:.0f}-{W_HI * 100:.0f}%: its industry series do not "
                    f"agree with AIES for the same cells and years (r {CELLS['all']['same_2023_24']['r']:.2f}, n {CELLS['all']['n']}); the error "
                    f"variance of its 4-year mean uses the year-to-year correlation of the Fiserv-AIES gaps ({RHO_GAP:.2f}, SE {RHO_GAP_SE:.2f}; "
                    f"{RHO_F:.2f} for Fiserv's own error once AIES's sampling error is removed), not independence"),
         'mfg': (f"lambda {LAM['manufacturing']:.2f} (manufacturing deviations barely persist; the six windows run {MFG_RANGE[0]:.2f} to "
                 f"{MFG_RANGE[1]:.2f}, so its SE adds their SD); Fiserv has no manufacturing series")}
if EO['agree']:
    notes['ecom'] = (f"Empire Flippers e-commerce sellers sold 2024-26 ({EF['ecom_pooled']['median'] * 100:+.1f}%, n {EF['ecom_pooled']['n']}, median "
                     f"revenue ${EF['ecom_pooled']['median_annual_revenue'] / 1e3:.0f}k) measured against M (no size bridge), pooled with BTOS 454 "
                     f"({EO['sources'][1]['dev'] * 100:+.1f} pts, with a {SE_454_MISMATCH * 100:.1f}-pt cell-mismatch SE: 454 also holds vending and "
                     f"direct selling); weights {EO['weights'][0] * 100:.0f}/{EO['weights'][1] * 100:.0f}%, z {EO['pairs'][0]['z']:+.1f}; weak")
js = {'by_industry': by_ind, 'se': se_ind,
      'anchor': {'median': M, 'se': SE_M, 'selection_se': SE_SEL, 'se_common': round(SE_COMMON, 4),
                 'se_note': (f"SE of M widened from {SE_M_R2 * 100:.2f} to {SE_M * 100:.1f} pts to cover the Fiserv conflict (Fiserv all-industry "
                             f"{FIS_NAT['feb2026_sa'] * 100:+.1f}% YoY in Feb 2026, {FIS_NAT['gap_to_M_feb2026'] * 100:+.1f} pts from M: z {z_r2:.1f} on "
                             f"{SE_M_R2 * 100:.2f} pts, {z_r3:.1f} on {SE_M * 100:.1f}); selection into sale ({SE_SEL * 100:.1f} pts, judgment) added in "
                             f"quadrature; with S's SE the common SE is {SE_COMMON * 100:.1f} pts, inside every industry's se"),
                 'source': 'JPMorganChase internal data: median YoY revenue growth of Chase small-business customers '
                           '(direct deposit account, < $20M revenue), national, Feb 2026; quoted in the JPMCI Denver snapshot (May 2026, Denver 4%)',
                 'size_bridge': round(S, 4), 'size_bridge_se': round(SE_S, 4), 'level': round(LEVEL, 4),
                 'size_bridge_note': ('Census BTOS 2025-26: revenue net balance of firms with 10-49 employees (the console\'s typical business) '
                                      f'minus 1-4 employees (the smallest employer class BTOS covers) = {GAP_A[0] * 100:+.1f} pts x {BSL:.3f} '
                                      f'(QCEW payroll on BTOS, sectors 2024-25) = {S_VAL * 100:+.1f} pts (SE {S_SE * 100:.2f}); applied because > 2 SE. '
                                      'Conservative: Chase\'s median customer ($125k revenue in 2025) may be a nonemployer, smaller than the 1-4 class, '
                                      'so the gap from it to the console\'s typical business is likely larger')},
      'lambda': LAM['non_manufacturing'], 'lambda_se': SE_LAM['non_manufacturing'],
      'lambda_note': (f"mean of the adjacent pre-COVID AIES windows ({NM_WIN_TXT}); SE = {SE_LAM_SAMPLING['non_manufacturing']:.2f} sampling and "
                      f"{WIN_SD['non_manufacturing']:.3f} across-window SD in quadrature; the 2017-19 -> 2023-24 window ({LAM_NONMFG_MAIN:.2f}) spans COVID"),
      'lambda_mfg': LAM['manufacturing'], 'lambda_mfg_se': SE_LAM['manufacturing'], 'bridge_slope': BR,
      'weak': weak,
      'notes': notes,
      'not_shipped': not_shipped,
      'refresh': 're-pull AIES, QSS, MRTS, QCEW, BTOS and Fiserv each year (and the latest Chase median); industry deviations are mostly transient',
      'checks': {
          'bridge': {'slope': round(BRIDGE['slope'], 2), 'se': round(BRIDGE['se'], 2), 'r': round(BRIDGE['r'], 3), 'n': BRIDGE['n'],
                     'industries': BRIDGE['industries'], 'resid_sd': round(BRIDGE['resid_sd'], 3)},
          'persistence': {'non_manufacturing': round(LAM_NONMFG, 3), 'non_manufacturing_se': round(SE_LAM_RAW['non_manufacturing'], 3),
                          'non_manufacturing_rule': 'mean of the adjacent pre-COVID windows',
                          'non_manufacturing_windows': {lab.split(' (')[0]: round(NM_WIN[lab]['slope'], 3) for lab in LAM_WINDOWS_NONMFG},
                          'non_manufacturing_se_boot_of_mean': round(NM_BOOT_SE, 3), 'non_manufacturing_across_window_sd': round(WIN_SD['non_manufacturing'], 3),
                          'non_manufacturing_main_window': round(LAM_NONMFG_MAIN, 3),
                          'non_manufacturing_range': [round(x, 2) for x in BY_GROUP['non_manufacturing_leaf']['range']],
                          'manufacturing': round(LAM_MFG, 3), 'manufacturing_se': round(SE_LAM_RAW['manufacturing'], 3),
                          'manufacturing_across_window_sd': round(WIN_SD['manufacturing'], 3),
                          'manufacturing_range': [round(x, 2) for x in BY_GROUP['manufacturing_six_digit']['range']],
                          'pooled': round(PERS['slope'], 3), 'n_non_manufacturing': GMAIN['non_manufacturing']['n'], 'n_manufacturing': GMAIN['manufacturing']['n'],
                          'windows': 'non-manufacturing: 2014-16 -> 2017-19, 2015-17 -> 2018-19, 2015-16 -> 2017-18; manufacturing: 2017-19 -> 2023-24',
                          'susb_r': round(SUSB_CHECK['r'], 2)},
          'restaurant_oos': {'census_nowcast': round(rn, 3), 'se': round(rse, 3), 'se_comparison': round(rse_cmp, 3),
                             'chase_median_plus_S': round(CHASE_REST['2026-08'] + S, 3), 'z': round(REST_OOS['z'], 2)},
          'median_vs_aggregate': {'gap_mean': round(MVA['gap_mean'], 3), 'gap_sd': round(MVA['gap_sd'], 3), 'n': MVA['n']},
          'spearman_current': {'rho': round(SPEAR['shipped'], 2), 'n': SPEAR['n'], 'mean_current': round(SPEAR['mean_current'], 3), 'mean_shipped': round(SPEAR['mean_shipped'], 3)},
          'spread': {'kept': SPREAD_KEPT, 're_lbd': [round(x, 2) for x in SPREAD['re_lbd']['range']], 'sbcs': [round(x, 2) for x in SPREAD['sbcs']['range']],
                     'chase_payroll': [round(x, 2) for x in SPREAD['chase_payroll']['range']], 'central': SPREAD_CENTRAL},
          'fiserv': {'national_yoy_feb2026': round(FIS_NAT['feb2026_sa'], 4), 'national_yoy_last12_mean': round(FIS_NAT['mean_last12_sa'], 4),
                     'national_2025': round(FIS_NAT['annual_nsa'][2025], 4), 'national_ytd2026': round(FIS_NAT['ytd2026_nsa'], 4),
                     'gap_to_M': round(FIS_NAT['gap_to_M_feb2026'], 4), 'supports_M': FIS_NAT['supports_M'],
                     'z_vs_se_m_round2': round(z_r2, 2), 'z_vs_se_m': round(z_r3, 2), 'within_widened_se_m': FIS_NAT['within_widened_SE_M'],
                     'industries_D_2023_26': {k: round(v['D'], 4) for k, v in F_IND.items()},
                     'pooled': sorted(FISERV_POOL), 'pool_weight': {k: round(float(E.loc[k, 'w_fiserv']), 3) for k in FISERV_POOL},
                     'pool_weight_round2': {k: round(float(E.loc[k, 'w_fiserv_r2']), 3) for k in FISERV_POOL},
                     'error_sd_4yr': round(math.sqrt(V_F4), 4), 'error_sd_4yr_round2_halved': round(math.sqrt(V_F4_R2), 4),
                     'gap_rho_2023_2024': round(RHO_GAP, 3), 'gap_rho_se': round(RHO_GAP_SE, 3), 'rho_fiserv_errors': round(RHO_F, 3),
                     'agreement_console': {'r': round(CONS_AGREE['all_mapped']['r'], 2), 'slope_fiserv_on_census': round(CONS_AGREE['all_mapped']['slope_y_on_x'], 2),
                                           'n': CONS_AGREE['all_mapped']['n'], 'r_own_cells': round(CONS_AGREE['own_cells']['r'], 2),
                                           'slope_own_cells': round(CONS_AGREE['own_cells']['slope_y_on_x'], 2), 'n_own_cells': CONS_AGREE['own_cells']['n']},
                     'agreement_cells': {'r_2023_24': round(CELLS['all']['same_2023_24']['r'], 2), 'n': CELLS['all']['n'],
                                         'r_services': round(CELLS['services']['same_2023_24']['r'], 2), 'n_services': CELLS['services']['n'],
                                         'aies_predicts_fiserv_slope': round(CELLS['all']['predict_fiserv_2025_26_from_aies_2023_24']['slope'], 2),
                                         'aies_predicts_fiserv_se': round(CELLS['all']['predict_fiserv_2025_26_from_aies_2023_24']['se'], 2)},
                     'persistence': {'slope': round(FPERS_MAIN['slope'], 3), 'se': round(SE_LF, 3), 'n': FPERS_MAIN['n'], 'windows': '2023-24 -> 2025-26',
                                     'z_vs_aies_non_manufacturing': round(FPERS_VS['z'], 2), 'lag1_median': round(FPERS_VS['fiserv_lag1_median'], 2),
                                     'subsector_slope': round(FPERS['subsectors 2023-24 -> 2025-26']['slope'], 3),
                                     'subsector_se': round(FPERS['subsectors 2023-24 -> 2025-26']['se_hc1'], 3)},
                     'restaurants': {'dev_722_aug2026': round(f_dev_aug26, 4), 'chase_dev_aug2026': round(D_chase, 4), 'z': round(REST_FISERV['z'], 2),
                                     'rule_confirmed': REST_FISERV['rule_confirmed']}},
          'btos': {'net_balance_by_size_2025_26': {k: round(v[0], 4) for k, v in SIZE_NB.items()}, 'national': round(NAT_NB[0], 4),
                   'gap_10_49_minus_1_4': round(GAP_A[0], 4), 'gap_se': round(GAP_A[1], 4),
                   'calibration_slope': round(BSL, 3), 'calibration_se': round(BSL_SE, 3), 'calibration_r': round(CALIB['r'], 2), 'calibration_n': CALIB['n'],
                   'calibration_loso_range': [round(x, 3) for x in CALIB['loso_range']],
                   'size_bridge': round(S_VAL, 4), 'size_bridge_se': round(S_SE, 4), 'z': round(S_Z, 2), 'applied': S_APPLIED,
                   'vs_all_employer_firms': round(S_NAT, 4), 'sectors_positive': f"{sum(1 for v in sec_gap.values() if v['nb_gap'] > 0)}/{len(sec_gap)}",
                   'sbcs_check_pts': {k: round(v, 4) for k, v in SBCS_SIZE['pts_vs_1_4'].items()},
                   'chase_payroll_larger_minus_all': round(CHASE_SIZE['larger_minus_all'], 4),
                   'ordering_vs_census': {'r': round(BT_AGREE['unique_cells']['r'], 2), 'n': BT_AGREE['unique_cells']['n'],
                                          'r_without_staff': round(BT_AGREE['unique_cells_without_staff']['r'], 2),
                                          'n_without_staff': BT_AGREE['unique_cells_without_staff']['n'],
                                          'r_btos_vs_fiserv': round(BT_AGREE['with_fiserv']['r'], 2), 'n_btos_vs_fiserv': BT_AGREE['with_fiserv']['n']}},
          'ecom': {'ef_median': round(EF['ecom_pooled']['median'], 4), 'ef_se': round(EF['ecom_pooled']['se_median'], 4), 'ef_n': EF['ecom_pooled']['n'],
                   'ef_base': 'M (no size bridge)', 'ef_dev': round(EO['sources'][0]['dev'], 4),
                   'btos_454_dev': round(EO['sources'][1]['dev'], 4), 'btos_454_se': round(EO['sources'][1]['se_dev'], 4),
                   'cell_mismatch_se': SE_454_MISMATCH, 'z': round(EO['pairs'][0]['z'], 2),
                   'weights': [round(x, 3) for x in EO.get('weights', [])],
                   'mismatch_evidence': {'firm_share_4541': round(SH_4541, 3),
                                         'aies_rms_2013_24': round(MISMATCH_454['aies_annual_2013_24_excl_2020']['rms'], 4),
                                         'aies_two_year_rms': round(MISMATCH_454['aies_two_year_means']['rms'], 4),
                                         'mrts_2025_26': [round(x, 4) for x in MISMATCH_454['btos_window_mrts_2025_26']['values']]},
                   'sensitivity': {k: v['rounded'] for k, v in ECOM_SENS.items()}},
          'staff': {'value': round(STAFF['value'], 4), 'se': round(STAFF_SE, 4), 'se_relative': round(STAFF['se_relative'], 4), 'weak': STAFF_WEAK,
                    'weak_rule_se': WEAK_SE, 'D': round(STAFF['D'], 4), 'D_latest_2025_26': round(D_STAFF_LATEST, 4),
                    'value_at_latest_D': round(STAFF['value_at_latest_D'], 4), 'override_round2': STAFF_OVERRIDE_R2[0]}},
      'source': source, 'method': method}
js_line = 'MODELS.fitted.peer_grw=' + json.dumps(clean(js), separators=(',', ':'), ensure_ascii=True) + ';'

os.makedirs(OUT, exist_ok=True)
est = E.reset_index()[['ind', 'current', 'n_obs', 'D_avg', 'se_D_avg', 'F_D', 'F_pooled', 'w_fiserv', 'w_fiserv_r2', 'D_pool', 'se_D_pool', 'D_pool_vf2',
                       'D_pool_r2', 'lam', 'se_lam', 'D_now', 'se_D_now', 'nowcast', 'se_nowcast', 'forward', 'se_forward', 'se_forward_rel',
                       'forward_rounded', 'value', 'se', 'shipped', 'shipped_se', 'basis']]
est.round(6).to_csv(OUT + '/estimates.csv', index=False)
pd.DataFrame(OBS).round(6).to_csv(OUT + '/observations.csv', index=False)
BD.round(6).to_csv(OUT + '/bridge.csv', index=False)
pd.DataFrame({'naics': pers_pairs[0].index, 'label': AIES_LABELS.reindex(pers_pairs[0].index).values, 'dev_2017_19': pers_pairs[0].values,
              'dev_2023_24': pers_pairs[1].values}).round(6).to_csv(OUT + '/persistence.csv', index=False)
pd.DataFrame([dict(ind=k, population_relevant=True, **s_) for k, o in ONLINE.items() for s_ in o['sources']] +
             [dict(ind=k, population_relevant=False, **s_) for k, o in ONLINE.items() for s_ in o['not_population_relevant']]
             ).round(6).to_csv(OUT + '/online_sources.csv', index=False)
SBCS.round(6).to_csv(OUT + '/sbcs_all_firms.csv', index=False)
FG.reset_index().rename(columns={'index': 'code'}).round(6).to_csv(OUT + '/fiserv_annual.csv', index=False)
CELL.round(6).to_csv(OUT + '/fiserv_cells.csv', index=False)
FLDEV.reset_index().rename(columns={'index': 'code'}).round(6).to_csv(OUT + '/fiserv_persistence.csv', index=False)
BT.round(6).to_csv(OUT + '/btos_nb.csv', index=False)
pd.DataFrame([dict(group='size', key=k, nb_2025_26=v[0], se=v[1], n=v[2], nb_2023_24=SIZE_NB_2324[k]) for k, v in SIZE_NB.items()] +
             [dict(group='sector_gap_10_49_minus_1_4', key=k, nb_2025_26=v['nb_gap'], se=v['se'], n=v['n'], nb_2023_24=np.nan) for k, v in sec_gap.items()]
             ).round(6).to_csv(OUT + '/size_bridge.csv', index=False)
MMD.round(6).to_csv(OUT + '/nonstore_454_mismatch.csv', index=False)
with open(OUT + '/js_line.txt', 'w') as f:
    f.write(js_line + '\n')

CHECKS = {
    'fiserv_M_test': dict(outcome=(('does not support M = 5% at face value (on the original 1.25-pt SE)' if not FIS_NAT['supports_M'] else 'supports M')
                                   + ('; SE_M widened to 2.0 pts to cover the gap (revision 3), so it now lies within 2 SE by construction'
                                      if FIS_NAT['within_widened_SE_M'] else '; outside 2 SE even on the widened SE_M')), **FIS_NAT,
                          se_M=SE_M, se_M_round2=SE_M_R2, figure2_reproduction=FIG2, restaurant_offset=REST_OFFSET),
    'fiserv_persistence': dict(outcome=('lower than AIES non-manufacturing (|z| > 2); attributed to noise in Fiserv industry series (see agreement), '
                                        'group lambdas kept') if abs(FPERS_VS['z']) > 2 else
                               'within 2 SE of the shipped AIES non-manufacturing lambda (adjacent pre-COVID mean); group lambdas kept', main=FPERS_MAIN,
                               windows=FPERS, comparison=FPERS_VS),
    'fiserv_vs_census_cells': dict(outcome=f"weak: r {CELLS['all']['same_2023_24']['r']:.2f} (services {CELLS['services']['same_2023_24']['r']:.2f}) for the same cells and years",
                                   **CELLS, v_aies_2yr=V_A2, v_fiserv_2yr=V_F2, v_fiserv_4yr=V_F4, error_correlation=FIS_ERR),
    'fiserv_vs_census_console': dict(outcome=f"r {CONS_AGREE['all_mapped']['r']:.2f} (n {CONS_AGREE['all_mapped']['n']}); own cells r {CONS_AGREE['own_cells']['r']:.2f} "
                                             f"(n {CONS_AGREE['own_cells']['n']}), slope {CONS_AGREE['own_cells']['slope_y_on_x']:.2f}", **CONS_AGREE),
    'btos_size_bridge': dict(outcome=f"applied: {S_VAL * 100:+.2f} pts (SE {S_SE * 100:.2f}, z {S_Z:.1f})" if S_APPLIED else f"not applied (z {S_Z:.1f})",
                             net_balance_by_size=SIZE_NB, net_balance_by_size_2023_24=SIZE_NB_2324, national=NAT_NB, gap_vs_1_4=GAP_A, gap_vs_national=GAP_NAT,
                             gap_vs_1_4_2023_24=GAP_A_2324, bridge=S_VAL, se=S_SE, z=S_Z, bridge_vs_national=S_NAT, se_vs_national=S_NAT_SE,
                             by_sector=sec_gap, calibration=CALIB, calibration_rows=CAL.round(6).to_dict('records'), sbcs_check=SBCS_SIZE, chase_payroll_check=CHASE_SIZE,
                             index_parsing_max_abs_diff=IDX_CHECK),
    'btos_ordering': dict(outcome=f"BTOS-implied deviations vs Census D: r {BT_AGREE['unique_cells']['r']:.2f} (n {BT_AGREE['unique_cells']['n']})",
                          industries=BT_IND, agreement=BT_AGREE),
    'restaurant_rule': dict(outcome='confirmed by Fiserv (within 2 SE)' if REST_FISERV['rule_confirmed'] else 'not confirmed by Fiserv', **REST),
    'restaurant_oos': dict(outcome='within 2 SE' if abs(REST_OOS['z']) < 2 else 'outside 2 SE', **REST_OOS),
    'median_vs_aggregate': dict(outcome=f"gap mean {MVA['gap_mean'] * 100:+.1f} pts, sd {MVA['gap_sd'] * 100:.1f}", **MVA),
    'sbcs_sector_ordering': dict(outcome=f"r {SBCS_CHECK['r']:.2f}", **SBCS_CHECK),
    'susb_persistence': dict(outcome=f"r {SUSB_CHECK['r']:.2f}", **SUSB_CHECK),
    'spearman_current': dict(outcome=f"rho {SPEAR['shipped']:.2f}", **SPEAR),
    'spread': dict(outcome='.13 kept (check only); evidence centres near .18', **SPREAD),
    'online': ONLINE, 'ecom_sensitivity': ECOM_SENS, 'nonstore_454_mismatch': MISMATCH_454,
    'staff': dict(outcome=('formula value shipped, flagged weak' if STAFF_WEAK else 'formula value shipped'), **STAFF),
    'lambda': dict(outcome=f"non-manufacturing {LAM['non_manufacturing']:.2f} (adjacent pre-COVID mean, SE {SE_LAM['non_manufacturing']:.2f}); "
                           f"manufacturing {LAM['manufacturing']:.2f} (SE {SE_LAM['manufacturing']:.2f})", **LAM_INFO),
    'rounding_edges': EDGES, 'tests': TESTS, 'outcomes': OUTCOMES}
RES = dict(
    name='peer_growth', revision=3, js_key='peer_grw', tier='aggregate', js_table=js, js_line=js_line,
    shipped={k: dict(value=float(E.loc[k, 'shipped']), se=float(E.loc[k, 'shipped_se']), current=CUR[k], change=float(E.loc[k, 'shipped']) - CUR[k],
                     round1=(R1['js_table']['by_industry'].get(k) if R1 else None), round2=(R2['js_table']['by_industry'].get(k) if R2 else None),
                     round2_se=(R2['js_table']['se'].get(k) if R2 else None), unrounded=float(E.loc[k, 'value']),
                     se_relative=math.sqrt(max(float(E.loc[k, 'se']) ** 2 - SE_COMMON ** 2, 0.0)),
                     weak=k in WEAK,
                     forward=float(E.loc[k, 'forward']), se_forward=float(E.loc[k, 'se_forward']), se_forward_relative=float(E.loc[k, 'se_forward_rel']),
                     nowcast=float(E.loc[k, 'nowcast']), D_census=float(E.loc[k, 'D_avg']), se_D_census=float(E.loc[k, 'se_D_avg']),
                     D_fiserv=(F_IND[k]['D'] if k in F_IND else None), w_fiserv=float(E.loc[k, 'w_fiserv']), D_pooled=float(E.loc[k, 'D_pool']),
                     lam=float(E.loc[k, 'lam']), basis=E.loc[k, 'basis']) for k in INDS if k in SHIP},
    not_shipped={k: dict(kept=CUR[k], tier='judgment', evidence=ONLINE[k]) for k in CURRENT_JUDGMENT_KEYS if k not in SHIP},
    parameters=dict(M=M, se_M=SE_M, se_M_round2=SE_M_R2, se_selection=SE_SEL, se_common=SE_COMMON, size_bridge=S, size_bridge_se=SE_S,
                    size_bridge_measured=S_VAL, level=LEVEL, lambda_shipped=LAM, lambda_se=SE_LAM,
                    lambda_measured={'non_manufacturing': LAM_NONMFG, 'non_manufacturing_main_window': LAM_NONMFG_MAIN, 'manufacturing': LAM_MFG,
                                     'pooled': PERS['slope']}, lambda_se_unrounded=SE_LAM_RAW, lambda_se_sampling=SE_LAM_SAMPLING,
                    lambda_across_window_sd=WIN_SD, weak_se_rule=WEAK_SE, se_454_mismatch=SE_454_MISMATCH, ef_product_sellers=list(EF_PROD),
                    ef_software=list(EF_SAAS),
                    bridge_slope_shipped=BR, bridge_resid_sd_used=RES, se_bridge_median_vs_aggregate=SE_BRIDGE, floor_aies=FLOOR, se_aies_no_cv=SE_AIES_NOCV,
                    qss_extra=QSS_EXTRA, se_mrts=SE_MRTS, recency_weights_nowcast=REC, fiserv_error_var_4yr=V_F4, fiserv_error_var_4yr_round2=V_F4_R2,
                    fiserv_error_rho=RHO_F, btos_calibration_slope=BSL,
                    rounding='values to the nearest 0.005 (half up); SEs to 0.001; lambdas to 0.01'),
    checks=CHECKS, comparators=dict(values=C, sectors=C_sectors, firm_weights_susb2022_lt500=FIRMS), bridge=BRIDGE,
    persistence=dict(main_pooled=PERS, groups=GMAIN, by_group=BY_GROUP, lags=lag_summary, susb_rows=SS.round(6).to_dict('records'),
                     sensitivity_not_shipped=SENS),
    fiserv=dict(industries=F_IND, national=FIS_NAT, error_correlation=FIS_ERR), staff=STAFF,
    online=dict(empire_flippers_2024_26=EF, empire_flippers_all_years=EF_ALL, acquire=ACQ, saas_capital=SAASCAP, saas_capital_se=SC_SE),
    quotes=QUOTES, tests=TESTS, sources=src)
with open(OUT + '/RESULT.json', 'w') as f:
    json.dump(clean(RES), f, indent=1, sort_keys=True, ensure_ascii=True)
    f.write('\n')
print(js_line)
print('wrote', OUT + '/RESULT.json')
