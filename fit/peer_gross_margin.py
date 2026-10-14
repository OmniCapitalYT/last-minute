#!/usr/bin/env python3
"""Peer gross-margin medians by industry and size (IND.gm -> MODELS.fitted.peer_gm). October 2026, revision 3.

Run:  python3 -I /tmp/claude-0/fit6/fitscripts/peer_gross_margin.py          (about 1-2 min)
      PGM_OUT=<dir> sends the outputs elsewhere (scratch re-runs only).
Out:  /tmp/claude-0/fit6/out/peer_gross_margin/RESULT.json (js_table = the table to embed; js_table_compact = the same on
      one line, also in peer_gm.json; change_from_round2), and tables: knots.csv, cells.csv, payroll_cut.csv,
      medical_materials.csv, change_from_round2.csv, owner_pay.csv, owner_pay_inputs.csv, median_gap_bach_derived.csv,
      ca_5_20M_check.csv, census_check.csv, us_vs_ca.csv, slope_test.csv, dealledger_check.csv,
      dealledger_stated_margins.csv, ef_listings_used.csv, definition_test.csv, spread_bach_derived.csv,
      research_replication.csv, recency.csv, oews_direct_share.csv. The page snapshot is read, never written; nothing is
      written outside the output directory. Earlier rounds are kept beside this file (peer_gross_margin.v1.py, .v2.py)
      with their outputs in /tmp/claude-0/fit6/out/peer_gross_margin.v1/ and .v2/; round 2's RESULT is read for the
      change table only.

Why
  The console z-scores the user's gross margin against a hand-set peer median per industry
  (z.gm = (gm - IND.gm)/.13, index.html model(); it feeds the quality composite, three category sub-scores, the buyer
  fit and the synergy pool). Round 1 replaced 14 of the 16 medians with size-varying medians from small-business tax
  data; round 2 fixed the definition, the Canadian owner pay, the surveys and e-commerce. Its review found the default
  definition (wages AND payroll costs of direct staff) applied only where the Census test covers it, too few materials
  in medical, optimistic saas SEs and a car-wash value without a definition. The analyst decided this revision:
  1. Payroll costs: the services and auto (medical, vet, msp, saas, prof, staff, land, fit, auto) load the payroll costs
     of the direct pay they load: the fit's own factor (SOI benefit and pension deductions over pay + employer FICA and
     FUTA) on the US direct pay, and on the Canadian direct pay drawn from labour and commissions (ISED glossary: that
     line is salaries and wages; cost-of-sales wages already include benefits). Half the cut is an SE term. Trades and
     manufacturing are unchanged (the Census 2022 test shows payroll costs already in their margin).
  2. Medical materials: physicians and dentists deduct the Census SAS 2022 supplies, other materials and purchased
     professional and technical services (6211, 6212; over Economic Census 2022 receipts), size-adjusted by the SOI
     $1-500K asset class's cost-of-goods share over the all-size share, in place of tax COGS less its labour where
     larger. A materials SE term carries three alternative readings.
  3. saas: the BACH median term is 0.5 correlated between the knots read off BACH's small class and those that blend in
     the medium class (the reviewer's reading; 1 before), so KBCM reaches the lower knots only partly; the top knot keeps
     the KBCM level. The same error model is used for every industry.
  4. Car wash: definition.wash = "Revenue less chemicals, water, utilities and supplies (car washes: wages are not in
     it)"; the hand-set .70 stays judgment on it (not_identified.wash).
  5. Recency: medical and prof (US 2004-13 cells alone) carry a level SE term of half the all-size SOI drift
     2004-13 -> 2019-22 on the labour-loaded definition.
  6. Auto: after the payroll cut the bridge and Ratchet+Wrench share the shipped basis (both counted direct pay without
     its payroll costs); they are combined by inverse variance, the survey moved by the full US-return cut at its size,
     the doubt about that cut shared (correlation 1), and a .02 judgment term for the survey's population (as KBCM's).
  7. mfg: the all-manufacturing cells load the all-manufacturing OEWS direct share (31-33); the Census 339950 row is
     the misc-manufacturing component, and the shipped mfg curve is tested against Census 31-33 (332710 as context).
  8. Metadata: by_industry[ind].tier (aggregate bridge, survey-combined, listing-level); checks.definition_test_1M is
     marked context only.

What ships (js_table.by_industry: knots at log10 revenue 5.5/6.0/6.5/7.0 in mid-2026 USD, SE, slope per decade, tier)
  15 industries. 12 aggregate bridges (the two-country tax-data bridge with the Canadian owner-pay fix; medical and prof
  on US cells alone); 2 survey-combined (saas with KBCM 2021, auto with Ratchet+Wrench 2026); 1 listing-level (ecom,
  Empire Flippers median, flat). wash stays judgment (definition.wash, not_identified.wash). No median shift. The
  z-scale .13 is kept. Each shipped value, its SE and the console's current value are in RESULT.shipped_values; what
  stays judgment, with its evidence, in RESULT.not_shipped; the change from round 2 with its parts in
  RESULT.change_from_round2.

What it does
  0. Units. Every revenue is in mid-2026 US dollars (CPI-U; CAD 2024 at the 2024 average rate, then CPI-U).
  1. Cells. US: IRS SOI Corporation Source Book 2004-13, Table 1, asset classes $1-500K to $10-25M: receipts, cost of
     goods, non-officer salaries and wages, officers' pay, pension and benefit deductions, returns. CA: ISED / Statistics
     Canada Financial Performance Data 2024, revenue quartiles of businesses with revenue CAD 30k-5M.
  2. Definition per industry. default: US 1 - [M + max(L, P) + f P]/receipts, L = lambda x COGS (SOI Table 12 sector
     labour share of COGS), P = d(SW + L) with d the industry's direct-occupation share of its wage bill (BLS OEWS May
     2025), M = COGS - L (physicians and dentists: max(COGS - L, m x receipts), m the size-adjusted Census share), f the
     payroll-cost factor (the services and auto; 0 elsewhere); officers' pay is never loaded. CA 1 - [COS + E(1 + f)]/rev
     with E = max(0, d(DW + IL - OP) - DW), OP below (physicians and dentists: COS - DW floored at m x rev).
     rest: 1 - (COGS - L)/receipts and 1 - (COS - DW)/revenue. dist: 1 - COGS/receipts, 1 - COS/revenue.
  3. Canadian owner pay. OP = s(x) x pi_q x revenue: s(x) = US officers' pay / receipts for the same industry at the
     same revenue (SOI asset classes, the same local-linear smoother, held flat beyond its cells +-0.25 decade); pi_q =
     the share of the quartile's firms that are incorporated, from ISED's reports by incorporation status. ISED's own
     reading - the incorporated minus the unincorporated labour share at matched revenue - is kappa = 0.7-6.2x the SOI
     share (median 1.9; held to 0.5-3); half the change it makes is the 'owner' SE term. Where OP is half or more of
     the Canadian labour bill (DW + IL) in the third or fourth quartile the Canadian cells are dropped: physicians,
     dentists, lawyers and accountants. medical and prof then use US cells alone; below the US cells (+0.25 decade) the
     US local trend is extended with a shape term of RMS slope gap x distance, and a size-slope 'tilt' term covers the
     second country they lack. saas has no Canadian cell (ISED does not publish 5112).
  4. Size curves and knots: US local-linear in log10 revenue (triangular kernel, half-width 1 decade, sqrt(returns)
     weights); CA quartiles joined linearly; a country carried past its cells (+0.25 decade) follows the other's shape.
     Knot = mean of the two countries; the page interpolates linearly and holds the end values.
  5. Medians. From BACH (sample 0, 2019-23, median over country-years) the median firm minus the ratio of sums, for the
     convention-matched margin. Over the 13 divisions the sign is not consistent (small 9+/4-, medium 8+/5-), so
     nothing is shifted: the RMS of the brackets enters the SE, one level within the small-class knots (to $3.2M) and
     one within those blending the medium class ($10M and beyond), 0.5 correlated between the two. Derived statistics
     only (BACH licence).
  6. Surveys. Inverse variance where the definition matches, a check where it does not. The survey enters at its own
     size; the other knots move through the bridge's own error covariance (the best linear update of the bridge given
     the survey; a shared adjustment enters as a correlated error). A survey's median is compared with the bridge's own
     median over the survey's revenue mix (normal within-industry spread at the survey's SD).
     auto - Ratchet+Wrench 2026, p. 9 overall shop gross profit (bins; 9% do not track; more than 430 responses, so
       n = 391), respondents' median revenue from p. 5. Shops count parts, sublet and technician pay without its
       payroll costs (FenderBender job costing), as the US bridge did: moved by technician pay x the payroll-cost factor
       at its size, SE = sampling, binning, .02 population (judgment) and half that cut (shared with the bridge). COMBINED.
     saas - KBCM 2021 total gross margin with support in COGS (GAAP, payroll-loaded, stock pay added back). $5-25M
       group, median 75% (n = 101); SE from the subscription-margin spread (n = 172), rounding, and 0.02 judgment for
       2020 and venture-backed respondents. COMBINED. DealLedger stated SaaS margins (n = 27): undefined, a check.
     ecom - Empire Flippers sold listings of product sellers with cogs_percent (COGS = units sold x landed cost): median,
       flat. SE = sampling (bootstrap) + seller mix + year + 0.01 definition + the slope SE beyond the 90th percentile.
     Each survey's within-industry SD is recorded against .13 (kept).
  7. SE at each knot = sqrt(diag of the summed covariance): time, source (half the US-CA gap; 0.04 with one country,
     plus the tilt), shape, bridge (d +- 0.15; restaurants lambda x0.5-1.5), proxy (vet, land .04; fit .05; saas .04),
     blend (binomial SE of the BIZCOMPS weight x component gap), owner, median, payroll (half the cut), materials (RMS
     of the alternative readings: purchased services left out, the pooled small classes at the knots as the size
     adjustment, no size adjustment) and recency. Slope SE from the same covariance. Shipped values rounded to 3
     decimals (Python round); RESULT keeps them unrounded.
  8. Checks (RESULT.check_outcomes, details in RESULT.checks): Census 2022 direct-cost decomposition of the shipped
     curves (238220, 238990, 31-33) on the decided wording and wages-only, the misc-manufacturing component (339950) and
     machine shops (332710, context); the payroll cut by country; the medical materials against Census SAS 2022; the
     recency drift and the small classes' own trends; US against CA at $1M (round 1, round 2, shipped); size slope
     across the countries; the surveys; DealLedger stated margins; BACH spread and the surveys' SDs against .13;
     Canadian incorporated cells above CAD 5M (not used); ADA HPI 2025 dental bound; research replication; Census
     AWTS/ARTS; definition test against the hand-set values (context only); materiality through the harness.

Occupation rules for d (OEWS May 2025, share of the industry's wage bill, detailed occupations; research round)
  trades 238220, 238900: 47-, 49-, 51-, 53-; auto 811110: 49-, 51-, 53-; car wash 811190: 49-, 53-;
  physicians 621100, dentists 621200: 29-, 31-; vets 541940: 29-, 31-, 39-; IT 541500: 15-;
  software 513200: 15-1231, 15-1232, 15-1244, 43-4051; accounting 541200: 13-2, 43-3011, 43-3031; legal 541100: 23-;
  all manufacturing 31-33 (natsector file) and misc. manufacturing 339900: 49-, 51-, 53-; temporary help 561320: all
  except 11-, 41-, 13-1071, 13-1141, 13-1161, 43-4161; landscaping 561730: 37-, 47-, 49-, 53-; fitness 713940: 39-.

Data and fetch commands (downloads are untrusted: run with python3 -I, never import or execute anything from them;
PDFs are read with poppler's pdftotext)
  /tmp/claude-0/fit6/data/A-gm/irs_soi_sourcebook/{04..07}sbflatfile.zip, {08..13}sbfltfile.zip
      curl -LO https://www.irs.gov/pub/irs-soi/13sbfltfile.zip   (and 12..08 sbfltfile; 07..04 sbflatfile)
  /tmp/claude-0/fit6/data/A-gm/irs_soi_corp/{19..22}co12ccr.xlsx, {19..22}co51ccr.xlsx
      curl -LO https://www.irs.gov/pub/irs-soi/22co12ccr.xlsx   (and 19-21; likewise co51ccr)
  /tmp/claude-0/fit6/data/A-gm/ised_canada/rev_<NAICS>.html, pm_<NAICS>.html
      curl -L -o rev_23822.html https://ised-isde.canada.ca/app/ixb/cis/performance/rev/23822   (and /performance/<NAICS>)
  /tmp/claude-0/fit6/data/peer_gross_margin/ised_fpd_incstatus/fpd_<NAICS>_bt<3|2|1>_rr<0|1>.html
      GET https://ised-isde.canada.ca/app/ixb/fpd-dpf/profilestart?lang=eng (cookie + _csrf), then POST to
      /app/ixb/fpd-dpf/start with reportCriteria.businessType=3|2|1 (all|incorporated|unincorporated),
      reportCriteria.revenueRange=0|1 (CAD 30k-5M | 5-20M, incorporated only), reportCriteria.distributedBy=0,
      reportCriteria.expenseUnit=1, reportCriteria.area=00, reportCriteria.year=2024, industrySearchKeyword and
      searchResultsSelectCode=<NAICS>, continueWithSearchCode=Create report (script: work/fit2-peer_gross_margin/fpd_fetch.sh)
  /tmp/claude-0/fit6/data/peer_gross_margin/ised_glossary/fpd_glossary.html   (new, round 3)
      curl -L -o fpd_glossary.html https://ised-isde.canada.ca/site/financial-performance-data/en/glossary
  /tmp/claude-0/data5/bls_oews/raw/oesm25in4.zip   https://www.bls.gov/oes/special-requests/oesm25in4.zip
  /tmp/claude-0/fit6/data/A-gm/census_ec22_constr/EC2223BASIC.zip, census_ec22_mfg/EC2231BASIC.zip,
  census_ec22_all/EC2200BASIC_US.dat (US rows of sector00/EC2200BASIC.zip)
      curl -LO https://www2.census.gov/programs-surveys/economic-census/data/2022/sector23/EC2223BASIC.zip (sector31, sector00)
  /tmp/claude-0/fit6/data/A-gm/census_sas/Table5.xlsx (Service Annual Survey 2022, selected expenses of employer firms)
      https://www2.census.gov/programs-surveys/sas/data/2022/ (downloaded by the research round)
  /tmp/claude-0/fit6/data/A-gm/census_awts/2022_awts_purchmarg_nomsbo_table4.xlsx, census_arts/gmper.xlsx
  /tmp/claude-0/fit6/data/A-gm/bach_docs/bach_docs.zip (licence: no redistribution; cite "BACH database: ECCBSO ...")
      curl -L -o bach_docs.zip 'https://www.bach.banque-france.fr/bach-ws-fo/ws/log-documentation/downloadocs?file_name=bach.zip'
  /tmp/claude-0/fit6/data/C-rec-ret/ratchetwrench/rw2026_industry_survey.pdf
      curl -LO https://img.ratchetandwrench.com/files/base/ebm/ratchetandwrench/document/2026/06/6a4277e469691969c387beca-2026_rwindustry_survey_fnl.pdf
  /tmp/claude-0/fit6/data/C-rec-ret/kbcm_2021/2021-KBCM-SaaS-Survey.pdf
      curl -LO https://www.forentrepreneurs.com/wp-content/uploads/2021/10/2021-KBCM-SaaS-Survey.pdf
  /tmp/claude-0/data5/empireflippers/clean/ef_listings.csv (public Empire Flippers marketplace API, see its MANIFEST.json)
  /tmp/claude-0/fit6/data/peer_gross_margin/ef_docs/ef_amazing.html      curl -L -o ef_amazing.html https://empireflippers.com/amazing/
  /tmp/claude-0/fit6/data/peer_gross_margin/fenderbender/fb_33026361.html
      curl -L https://www.fenderbender.com/running-a-shop/operations/article/33026361/basic-financial-knowledge-allows-shop-owners-to-measure-improve-performance
  /tmp/claude-0/fit6/data/peer_gross_margin/ada_hpi/hpidata_sdpi_2025.xlsx
      curl -LO https://www.ada.org/-/media/project/ada-organization/ada/ada-org/files/resources/research/hpi/hpidata_sdpi_2025.xlsx
  /tmp/claude-0/fit6/data/peer_gross_margin/pcd_carwash/pcd_winning_combinations.html
      curl -L -o pcd_winning_combinations.html https://www.carwash.com/winning-combinations-chemicals-equipment/
  /tmp/claude-0/dl/jeffsosville_dealledger/data/snapshots/<date>/listings.csv (CC0)   git clone --depth 1 https://github.com/jeffsosville/dealledger
  /tmp/claude-0/fit6/data/peer_gross_margin/fred_cpi_fx/CPIAUCNS.csv, AEXCAUS.csv
      curl -L -o CPIAUCNS.csv 'https://fred.stlouisfed.org/graph/fredgraph.csv?id=CPIAUCNS'   (and id=AEXCAUS)
  /tmp/claude-0/fit6snap/index.html (read only: the hand-set IND values, the z.gm scale and the BIZCOMPS NAICS-6 counts)
  /tmp/claude-0/fit6/work/A-gm/out/final_table.csv (the research round's table; compared in the replication check only)
  /tmp/claude-0/fit6/out/peer_gross_margin.v2/RESULT.json (round 2's outputs; the change table only)"""
import csv, glob, hashlib, html, io, json, os, re, shutil, subprocess, zipfile
import numpy as np
import pandas as pd
import openpyxl
from scipy import optimize, stats
import statsmodels.api as sm

OUT = os.environ.get('PGM_OUT', '/tmp/claude-0/fit6/out/peer_gross_margin')   # PGM_OUT only for scratch re-runs
DA = '/tmp/claude-0/fit6/data/A-gm'
DP = '/tmp/claude-0/fit6/data/peer_gross_margin'
DC = '/tmp/claude-0/fit6/data/C-rec-ret'
SOI_SB_DIR = DA + '/irs_soi_sourcebook'
SOI_CORP_DIR = DA + '/irs_soi_corp'
ISED_DIR = DA + '/ised_canada'
FPD_DIR = DP + '/ised_fpd_incstatus'
OEWS_ZIP = '/tmp/claude-0/data5/bls_oews/raw/oesm25in4.zip'
EC_CONSTR = DA + '/census_ec22_constr/EC2223BASIC.zip'
EC_MFG = DA + '/census_ec22_mfg/EC2231BASIC.zip'
EC_US = DA + '/census_ec22_all/EC2200BASIC_US.dat'
SAS5 = DA + '/census_sas/Table5.xlsx'
ISED_GLOSSARY = DP + '/ised_glossary/fpd_glossary.html'
PREV_RESULT = '/tmp/claude-0/fit6/out/peer_gross_margin.v2/RESULT.json'   # round 2's outputs: read for the change table only
AWTS = DA + '/census_awts/2022_awts_purchmarg_nomsbo_table4.xlsx'
ARTS = DA + '/census_arts/gmper.xlsx'
BACH_ZIP = DA + '/bach_docs/bach_docs.zip'
DL_SNAPS = '/tmp/claude-0/dl/jeffsosville_dealledger/data/snapshots'
CPI_CSV = DP + '/fred_cpi_fx/CPIAUCNS.csv'
FX_CSV = DP + '/fred_cpi_fx/AEXCAUS.csv'
RW_PDF = DC + '/ratchetwrench/rw2026_industry_survey.pdf'
KBCM_PDF = DC + '/kbcm_2021/2021-KBCM-SaaS-Survey.pdf'
EF_CSV = '/tmp/claude-0/data5/empireflippers/clean/ef_listings.csv'
EF_DOC = DP + '/ef_docs/ef_amazing.html'
FB_DOC = DP + '/fenderbender/fb_33026361.html'
ADA_XLSX = DP + '/ada_hpi/hpidata_sdpi_2025.xlsx'
PCD_DOC = DP + '/pcd_carwash/pcd_winning_combinations.html'
INDEX = '/tmp/claude-0/fit6snap/index.html'
HARNESS = '/tmp/claude-0/fit6/harness.js'      # loads the snapshot's model into a node sandbox (materiality check)
SDE_REF = 708000                               # the console's size centre ($708k SDE)
RESEARCH_TABLE = '/tmp/claude-0/fit6/work/A-gm/out/final_table.csv'
os.makedirs(OUT, exist_ok=True)

KNOTS = [5.5, 6.0, 6.5, 7.0]       # log10 revenue per firm, mid-2026 USD
H = 1.0                            # US local-linear kernel half-width (decades)
TAU = 0.25                         # how far a source may extrapolate past its outermost cell (decades)
MIN_YEARS = 5                      # a US class enters the pooled curve if present in >= 5 of the 10 years
D_SD = 0.15                        # bridge: OEWS direct share +- 0.15 (research round)
LAM_REST = (0.5, 1.5)              # restaurants: Table 12 labour share of COGS x0.5 .. x1.5
SRC_SINGLE = 0.04                  # source term where only one country has a cell (research round)
PROXY = {'vet': .04, 'land': .04, 'fit': .05, 'saas': .04, 'wash': .10}   # research round, judgment
TOL_CENSUS = 0.04
FX_YEAR = 2024
REF_X = {'1M': 6.0, '3M': float(np.log10(3e6))}
OWNER_MODE = 'soi'                 # CA owner pay removed with the SOI officers'-pay share (ISED's own reading -> SE)
OWNER_DROP = 0.5                   # CA cell dropped where that owner pay is >= half its labour bill (Q3 or Q4)
KAPPA_CAP = (0.5, 3.0)             # ISED calibration factor of the SOI officer-pay share, held to this range
PAYROLL_IND = ['medical', 'vet', 'msp', 'saas', 'prof', 'staff', 'land', 'fit', 'auto']   # round 3: payroll costs loaded on direct pay
PAYROLL_SE = 0.5                   # half the payroll-cost cut is carried as an SE term
MAT_SMALL_AC = 3                   # the SOI small cell for the materials size adjustment: asset class $1-500K
MEDIAN_RHO = 0.5                   # BACH median term: correlation between small-class knots and knots that blend in the medium class
RECENCY_IND = ['medical', 'prof']  # one-country industries resting on SOI 2004-13 alone
RECENCY_FRAC = 0.5                 # recency SE = this x |all-size SOI drift 2004-13 -> 2019-22| on the labour-loaded definition
SEED = 20261014

SEC_CONS = 'Construction'
SEC_HEALTH = 'Health care and social assistance'
SEC_PST = 'Professional, scientific, and technical services'
SEC_ADMIN = 'Administrative and support and waste management and remediation services'
# component cells: SOI Source Book minor industry, Table 12 sector (lambda), ISED NAICS, OEWS rule key
COMP = {
    'home':     dict(soi=238220, sec=SEC_CONS, ca='23822', oews='home', label='Plumbing, heating and air-conditioning contractors'),
    'constr':   dict(soi=238905, sec=SEC_CONS, ca='23899', oews='constr', label='Other specialty trade contractors'),
    'mfg':      dict(soi=31, sec='Manufacturing', ca='31-33', oews='mfg_all', label='Manufacturing, all'),
    'mfg_misc': dict(soi=339900, sec='Manufacturing', ca='33995', oews='mfg_misc', label='Other misc. manufacturing / sign manufacturing'),
    'auto':     dict(soi=811115, sec='Other services', ca='81111', oews='auto', label='Automotive repair and maintenance'),
    'dist':     dict(soi=423300, sec='Wholesale trade', ca='4163', oews=None, label='Lumber and construction materials wholesale'),
    'rest':     dict(soi=722, sec='Accommodation and food services', ca='72251', oews='rest', label='Food services / restaurants'),
    'phys':     dict(soi=621115, sec=SEC_HEALTH, ca='62111', oews='phys', label='Offices of physicians'),
    'dent':     dict(soi=621210, sec=SEC_HEALTH, ca='62121', oews='dent', label='Offices of dentists'),
    'vet':      dict(soi=541915, sec=SEC_PST, ca='54194', oews='vet', label='Veterinary (US: other professional services)'),
    'msp':      dict(soi=541515, sec=SEC_PST, ca='54151', oews='msp', label='Computer systems design'),
    'saas':     dict(soi=511210, sec='Information', ca=None, oews='saas', label='Software publishers'),
    'acct':     dict(soi=541215, sec=SEC_PST, ca='54121', oews='acct', label='Accounting, tax, bookkeeping, payroll'),
    'legal':    dict(soi=541115, sec=SEC_PST, ca='54111', oews='legal', label='Legal services'),
    'staff':    dict(soi=561300, sec=SEC_ADMIN, ca='56132', oews='staff', label='Employment services / temporary help'),
    'land':     dict(soi=561905, sec=SEC_ADMIN, ca='56173', oews='land', label='Landscaping (US: other admin and support)'),
    'fit':      dict(soi=713, sec='Arts, entertainment, and recreation', ca='71394', oews='fit', label='Fitness (US: amusement and recreation)'),
    'wash':     dict(soi=811115, sec='Other services', ca='81119', oews='wash', label='Car wash proxies (auto repair; other auto repair)'),
    'ecom':     dict(soi=454, sec='Retail trade', ca=None, oews=None, label='Nonstore retailers'),
}
OEWS_RULES = {
    'home': ('238220', ['47-', '49-', '51-', '53-']), 'constr': ('238900', ['47-', '49-', '51-', '53-']),
    'phys': ('621100', ['29-', '31-']), 'dent': ('621200', ['29-', '31-']), 'vet': ('541940', ['29-', '31-', '39-']),
    'msp': ('541500', ['15-']), 'saas': ('513200', ['15-1232', '15-1244', '43-4051', '15-1231']),
    'acct': ('541200', ['13-2', '43-3031', '43-3011']), 'legal': ('541100', ['23-']),
    'mfg_misc': ('339900', ['51-', '49-', '53-']), 'staff': ('561320', None),
    'land': ('561730', ['37-', '47-', '49-', '53-']), 'wash': ('811190', ['53-', '49-']),
    'auto': ('811110', ['49-', '51-', '53-']), 'fit': ('713940', ['39-']), 'rest': ('722511', ['35-']),
    'mfg_all': ('31-33', ['51-', '49-', '53-']),          # all manufacturing (natsector file): the shipped mfg cells (round 3)
}
STAFF_INTERNAL = ['13-1071', '13-1141', '43-4161', '41-3091', '13-1161']
IND_ORDER = ['home', 'medical', 'vet', 'msp', 'saas', 'prof', 'mfg', 'dist', 'ecom', 'staff', 'land', 'wash', 'auto',
             'constr', 'fit', 'rest']
BRIDGED = ['home', 'constr', 'mfg', 'auto', 'dist', 'rest', 'medical', 'vet', 'msp', 'saas', 'prof', 'staff', 'land', 'fit']
SHIPPED = BRIDGED + ['ecom']                                   # ecom: Empire Flippers sold listings (no bridge)
NOT_SHIPPED = ['wash']
US_STARTS_ABOVE_1M = ['staff']     # SOI employment-services cells start at $2.3M (2026 USD): 6.0 is carried from Canada's shape
DEFN = {'rest': 'exl', 'dist': 'rep', 'ecom': 'rep'}          # everything else: 'dir'
DEF_NAMES = {'rep': 'as reported (tax COGS)', 'dir': 'direct labour loaded, every industry (payroll costs for the services and auto; medical with Census materials)',
             'allw': 'all non-owner wages loaded', 'exl': 'materials only (COGS less its labour)',
             'conv': 'shipped convention (direct labour; restaurants food cost; distribution as reported)'}


def sha(path, n=16):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()[:n]


def r3(v):
    return None if v is None or not np.isfinite(v) else round(float(v), 3)


def r4(v):
    return None if v is None or not np.isfinite(v) else round(float(v), 4)


# ---------------------------------------------------------------- 0. price level and exchange rate
cpi = pd.read_csv(CPI_CSV)
cpi['year'] = cpi.observation_date.str[:4].astype(int)
cpi_full = cpi.groupby('year').CPIAUCNS.agg(['mean', 'count'])
CPI_ANN = {int(y): float(r['mean']) for y, r in cpi_full.iterrows() if r['count'] == 12}
CPI_AVAIL = {int(y): float(r['mean']) for y, r in cpi_full.iterrows() if r['count'] >= 6}   # 2025 lacks October (no release)
CPI_REF = float(cpi.CPIAUCNS.iloc[-12:].mean())
CPI_REF_SPAN = (cpi.observation_date.iloc[-12], cpi.observation_date.iloc[-1])
CPI_REF_MONTHS = int(cpi.CPIAUCNS.iloc[-12:].notna().sum())
fx = pd.read_csv(FX_CSV)
FX = float(fx.loc[fx.observation_date.str[:4].astype(int) == FX_YEAR, 'AEXCAUS'].iloc[0])     # CAD per USD
CA_TO_REAL = CPI_REF / CPI_ANN[FX_YEAR] / FX                                                  # CAD 2024 -> USD 2026
assert abs(CPI_ANN[2009] - 214.537) < 0.01 and abs(CPI_ANN[2013] - 232.957) < 0.01, 'CPI-U annual averages off'


def to_real(usd, year):
    return usd * CPI_REF / CPI_AVAIL[year]


# ---------------------------------------------------------------- 1a. SOI Source Book cells
SOI_ZIPS = sorted(glob.glob(os.path.join(SOI_SB_DIR, '*.zip')))
assert len(SOI_ZIPS) == 10, SOI_ZIPS
soi_codes = sorted({c['soi'] for c in COMP.values()})
rows = []
for zp in SOI_ZIPS:
    z = zipfile.ZipFile(zp)
    name = [n for n in z.namelist() if n.lower().endswith('sb1.csv')][0]
    t = pd.read_csv(io.BytesIO(z.read(name)), low_memory=False)
    t.columns = [c.upper() for c in t.columns]
    assert (t.TAB_NUM == 1).all()
    s = t[t.INDY_CD.isin(soi_codes) & t.AC.isin([1, 3, 4, 5, 6, 7])]
    for _, r in s.iterrows():
        flags = [r.get(f + '_IND', 0) for f in ('WT_CT', 'GRS_RCPTS', 'CST_OF_GDS', 'SLRS_WGS')]
        rows.append(dict(year=int(r.YEAR), code=int(r.INDY_CD), ac=int(r.AC), n=r.WT_CT, br=r.GRS_RCPTS,
                         cogs=r.CST_OF_GDS, sw=r.SLRS_WGS, oc=r.CMPNSTN_OFFCRS, pen=r.PNSN_PRFT_SHRNG_PLNS, ben=r.EMP_BNFT_PRG,
                         flag=int(np.nanmax([0 if pd.isna(f) else f for f in flags]))))
SOI = pd.DataFrame(rows)
for c in ['n', 'br', 'cogs', 'sw', 'oc', 'pen', 'ben']:
    SOI[c] = pd.to_numeric(SOI[c], errors='coerce')
SOI['ok'] = (SOI.n > 0) & (SOI.br > 0) & SOI.cogs.notna() & SOI.sw.notna()
SOI['rpr'] = SOI.br / SOI.n * 1000.0                                   # receipts per return, $ of the tax year
SOI['x_nom'] = np.log10(SOI.rpr.where(SOI.ok))
SOI['x_real'] = np.log10((SOI.rpr * CPI_REF / SOI.year.map(CPI_ANN)).where(SOI.ok))
SOI = SOI.sort_values(['code', 'year', 'ac'], kind='mergesort').reset_index(drop=True)
SOI_CELLS = SOI[SOI.ok & SOI.ac.between(3, 7)].copy()

# ---------------------------------------------------------------- 1b. SOI Table 12: labour share of COGS by sector, 2019-22
T12_FILES = sorted(glob.glob(os.path.join(SOI_CORP_DIR, '*co12ccr.xlsx')))
t12 = {}
for f in T12_FILES:
    yr = 2000 + int(os.path.basename(f)[:2])
    rws = list(openpyxl.load_workbook(f, read_only=True, data_only=True).worksheets[0].iter_rows(min_row=1, max_row=30, values_only=True))
    names = [re.sub(r'\s+', ' ', str(v)).strip() if v else '' for v in rws[4]]
    names[1] = 'All sectors'

    def get(lab):
        for r in rws[8:19]:
            if r[0] and str(r[0]).strip().lower().startswith(lab):
                return r
    cogs_r, lab_r = get('cost of goods sold'), get('cost of labor')
    for c in range(1, len(names)):
        if names[c]:
            try:
                t12.setdefault(names[c], {})[yr] = float(lab_r[c]) / float(cogs_r[c])
            except (TypeError, ValueError):
                pass
LAM = {k: float(np.mean([v[y] for y in sorted(v)])) for k, v in t12.items()}
for c in COMP.values():
    assert c['sec'] in LAM, c['sec']

# ---------------------------------------------------------------- 1b2. employer payroll costs per dollar of pay (the default definition loads them on direct pay)
STATUTORY = 0.0765 + 0.006                                               # employer FICA and FUTA, share of pay


def payroll_cost_factor(comp):
    """employer payroll costs per dollar of pay: SOI pension and employee-benefit deductions over all pay (all sizes,
    10-year mean) plus the statutory employer taxes"""
    g = SOI[(SOI.code == COMP[comp]['soi']) & (SOI.ac == 1) & SOI.ok]
    lam = LAM[COMP[comp]['sec']]
    return float(((g.pen.fillna(0) + g.ben.fillna(0)) / (g.oc + g.sw + lam * g.cogs)).mean()) + STATUTORY


PCF = {c: payroll_cost_factor(c) for c in COMP}


# ---------------------------------------------------------------- 1c. OEWS May 2025: direct-occupation wage shares
def load_oews():
    z = zipfile.ZipFile(OEWS_ZIP)
    fr = []
    for name in ['oesm25in4/nat4d_M2025_dl.xlsx', 'oesm25in4/nat5d_6d_M2025_dl.xlsx', 'oesm25in4/natsector_M2025_dl.xlsx']:
        ws = openpyxl.load_workbook(io.BytesIO(z.read(name)), read_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        hdr = next(it)
        d = pd.DataFrame(list(it), columns=hdr)
        d['file'] = name.split('/')[-1]
        fr.append(d)
    A = pd.concat(fr, ignore_index=True)
    A['NAICS'] = A.NAICS.astype(str).str.strip()
    A['OCC_CODE'] = A.OCC_CODE.astype(str).str.strip()
    for c in ['TOT_EMP', 'A_MEAN', 'H_MEAN']:
        A[c] = pd.to_numeric(A[c], errors='coerce')
    A['amean'] = A.A_MEAN.fillna(A.H_MEAN * 2080)
    A['wb'] = A.TOT_EMP * A.amean
    out = {}
    for k, (code, rule) in OEWS_RULES.items():
        s = A[A.NAICS == code]
        if k != 'mfg_all':
            s = s[s.file != 'natsector_M2025_dl.xlsx']
        assert len(s), code
        det = s[s.O_GROUP == 'detailed']
        assert not det.OCC_CODE.duplicated().any(), code
        tot = s[s.OCC_CODE == '00-0000']
        W, E = det.wb.sum(), det.TOT_EMP.sum()
        if rule is None:
            internal = det.OCC_CODE.str.startswith('11-') | det.OCC_CODE.str.startswith('41-') | det.OCC_CODE.isin(STAFF_INTERNAL)
            dW, dE = W - det.loc[internal, 'wb'].sum(), E - det.loc[internal, 'TOT_EMP'].sum()
        else:
            m = det.OCC_CODE.apply(lambda o: any(o.startswith(p) for p in rule))
            dW, dE = det.loc[m, 'wb'].sum(), det.loc[m, 'TOT_EMP'].sum()
        out[k] = dict(naics=code, title=str(tot.NAICS_TITLE.iloc[0]), emp_total=float(tot.TOT_EMP.iloc[0]),
                      direct_wage_share=float(dW / W), direct_emp_share=float(dE / E),
                      rule=rule if rule else 'all except 11-, 41-, ' + ', '.join(STAFF_INTERNAL))
    return out


OEWS = load_oews()


def delta(comp):
    k = COMP[comp]['oews']
    return OEWS[k]['direct_wage_share'] if k else 0.0


# ---------------------------------------------------------------- 1d. ISED / Statistics Canada pages
def page_lines(path):
    t = open(path, encoding='utf-8', errors='replace').read()
    t = re.sub(r'<script.*?</script>', '', t, flags=re.S)
    t = re.sub(r'<style.*?</style>', '', t, flags=re.S)
    t = html.unescape(re.sub(r'<[^>]+>', '\n', t))
    return [l.strip() for l in t.split('\n') if l.strip()]


def grab(lines, label, k=6, flags=r'[A-F]'):
    """the k numbers that follow a row label (reliability letters skipped); the longest run if the label repeats"""
    best = None
    for i, l in enumerate(lines):
        if l.lower() == label.lower():
            vals, j = [], i + 1
            while len(vals) < k and j < len(lines):
                v = lines[j].replace(',', '')
                if re.fullmatch(flags, v):
                    j += 1
                    continue
                m = re.match(r'^(-?[\d.]+|N/A)\s*([A-F])?$', v)
                if not m:
                    break
                vals.append(m.group(1))
                j += 1
            if len(vals) >= min(k, 2) and (best is None or len(vals) > len(best)):
                best = vals
            if best and len(best) >= k:
                return best
    return best


def parse_ised(path):
    t = open(path, encoding='utf-8', errors='replace').read()
    if len(t) < 5000:
        return None
    lines = page_lines(path)
    rec = {}
    yr = [lines[i + 1] for i, l in enumerate(lines) if l.startswith('Data year')]
    rec['year'] = int(yr[0]) if yr else None
    nb = grab(lines, 'Number of businesses', 1)
    rec['n_businesses'] = float(nb[0]) if nb else None
    for lab, key in [('Total revenue', 'rev'), ('Cost of sales (direct expenses)', 'cos'), ('Wages and benefits', 'dw'),
                     ('Labour and commissions', 'il'), ('Gross margin (%)', 'gm')]:
        v = grab(lines, lab, 6)
        if v:
            for qi, qn in enumerate(['all', 'q1', 'q2', 'q3', 'q4', 'pct_rep']):
                if qi < len(v):
                    rec[f'{key}_{qn}'] = float(v[qi]) if v[qi] != 'N/A' else np.nan
    return rec


ISED_FILES = sorted(glob.glob(os.path.join(ISED_DIR, 'rev_*.html')) + glob.glob(os.path.join(ISED_DIR, 'pm_*.html')))
ISED_REV, ISED_PM = {}, {}
for f in ISED_FILES:
    b = os.path.basename(f)
    rec = parse_ised(f)
    if rec is None:
        continue
    (ISED_REV if b.startswith('rev_') else ISED_PM)[b.split('_', 1)[1][:-5]] = rec
for c in COMP.values():
    if c['ca']:
        r = ISED_REV[c['ca']]
        assert r['year'] == 2024 and all(np.isfinite(r[f'{v}_{q}']) for v in ('rev', 'cos', 'dw', 'il') for q in ('q1', 'q2', 'q3', 'q4')), c['ca']


# ---------------------------------------------------------------- 1e. ISED Financial Performance Data by incorporation status (new download)
def parse_fpd(path):
    L = page_lines(path)
    rec = {}
    for lab, key in [('Incorporation status:', 'status'), ('Revenue range:', 'range'), ('Data year:', 'year')]:
        v = [L[i + 1] for i, l in enumerate(L) if l == lab]
        rec[key] = v[0] if v else None
    rec['year'] = int(rec['year'])
    nb = grab(L, 'Number of businesses', 1)
    rec['n'] = float(nb[0])
    rec['lo'] = [float(v) for v in grab(L, 'Low Value ($000)', 5)]
    rec['hi'] = [float(v) for v in grab(L, 'High Value ($000)', 5)]
    for lab, key in [('Total revenue', 'rev'), ('Cost of sales (direct expenses)', 'cos'), ('Wages and benefits', 'dw'),
                     ('Labour and commissions', 'il'), ('Net profit/loss', 'np')]:
        v = grab(L, lab, 6)
        rec[key] = [np.nan if x == 'N/A' else float(x) for x in v[:5]]
    return rec


FPD_FILES = sorted(glob.glob(os.path.join(FPD_DIR, 'fpd_*.html')))
FPD = {}
for f in FPD_FILES:
    m = re.match(r'fpd_(.+)_bt(\d)_rr(\d)\.html$', os.path.basename(f))
    FPD[(m.group(1), int(m.group(2)), int(m.group(3)))] = parse_fpd(f)
CA_CODES = sorted({c['ca'] for c in COMP.values() if c['ca']})
FPD_CONSIST = {}
for code in CA_CODES:
    A, I, U, B = FPD[(code, 3, 0)], FPD[(code, 2, 0)], FPD[(code, 1, 0)], FPD[(code, 2, 1)]
    assert A['status'] == 'All businesses' and I['status'] == 'Incorporated businesses' and U['status'] == 'Unincorporated businesses', code
    assert B['status'] == 'Incorporated businesses' and '20,000,000' in B['range'] and all(r['year'] == 2024 for r in (A, I, U, B)), code
    assert abs(I['n'] + U['n'] - A['n']) <= 1, (code, I['n'], U['n'], A['n'])
    r = ISED_REV[code]
    FPD_CONSIST[code] = max(abs(A[k][q] - r[f'{k}_q{q}']) for k in ('rev', 'cos', 'dw', 'il') for q in (1, 2, 3, 4))
assert max(FPD_CONSIST.values()) < 0.051, FPD_CONSIST           # the FPD 'all businesses' report is the page round 1 used


def quartile_cdf(rec):
    """share of the population below revenue r: piecewise linear in log revenue through the published quartile bounds"""
    px = np.log([rec['lo'][1], rec['hi'][1], rec['hi'][2], rec['hi'][3], rec['hi'][4]])
    return lambda r: float(np.interp(np.log(r), px, [0, .25, .5, .75, 1.0]))


def incorporated_share(code):
    """per all-business revenue quartile: incorporated firms / all firms, from the two status reports' quartile bounds"""
    A, I, U = FPD[(code, 3, 0)], FPD[(code, 2, 0)], FPD[(code, 1, 0)]
    Fi, Fu = quartile_cdf(I), quartile_cdf(U)
    out = []
    for q in (1, 2, 3, 4):
        a, b = A['lo'][q], A['hi'][q]
        ni, nu = I['n'] * (Fi(b) - Fi(a)), U['n'] * (Fu(b) - Fu(a))
        out.append(ni / (ni + nu))
    return out


PI_INC = {code: incorporated_share(code) for code in CA_CODES}

# ---------------------------------------------------------------- 1f. the page: hand-set IND values and BIZCOMPS NAICS-6 counts
PAGE = open(INDEX, encoding='utf-8').read()
IND_BLOCK = PAGE[PAGE.index('const IND = {'):]
IND_BLOCK = IND_BLOCK[:IND_BLOCK.index('};')]
_m = list(re.finditer(r'^\s*(\w+):\s*\{n:"([^"]*)",\s*anchor:[\d.]+,\s*buyers:\d+,\s*gm:([\d.]+)', IND_BLOCK, re.M))
HAND = {m.group(1): float(m.group(3)) for m in _m}
IND_LABEL = {m.group(1): m.group(2) for m in _m}
assert sorted(HAND) == sorted(IND_ORDER), HAND
MODELS = json.loads(re.search(r'const MODELS = (\{.*?\});\n', PAGE).group(1))
N6 = MODELS['value']['industry_n6']
W_MED = {'phys': N6['621111'], 'dent': N6['621210']}
W_PROF = {'acct': N6['541211'] + N6['541213'] + N6['541219'], 'legal': N6['541110']}
assert (W_MED['phys'], W_MED['dent'], W_PROF['acct'], W_PROF['legal']) == (176, 109, 235, 19), (W_MED, W_PROF)
Z_GM = float(re.search(r'\bgm\s*:\s*\(s\.gm-I\.gm\)/([\d.]+),', PAGE).group(1))      # model(): z.gm = (gm - I.gm)/.13
assert Z_GM == 0.13, Z_GM
INDUSTRIES = {k: [(k, 1.0)] for k in IND_ORDER}
INDUSTRIES['medical'] = [('phys', float(W_MED['phys'])), ('dent', float(W_MED['dent']))]
INDUSTRIES['prof'] = [('acct', float(W_PROF['acct'])), ('legal', float(W_PROF['legal']))]
PAYROLL_COMPS = sorted({c for i in PAYROLL_IND for c, _ in INDUSTRIES[i]})

# ---------------------------------------------------------------- 1g. ISED line definitions (glossary; new download, round 3)
GLOSSARY_TXT = re.sub(r'\s+', ' ', ' '.join(page_lines(ISED_GLOSSARY)))
GLOSSARY = dict(
    wages_and_benefits=bool(re.search(r'Wages and benefits includes wages and benefits paid to employees that are shown in the cost of sales\.', GLOSSARY_TXT)),
    labour_and_commissions=bool(re.search(r'Labour and commissions include remuneration paid to the employees of the business not shown in the cost of '
                                          r'sales and includes: salaries and wages\.', GLOSSARY_TXT)))
assert all(GLOSSARY.values()), GLOSSARY     # cost-of-sales wages carry benefits; labour and commissions are salaries and wages

# ---------------------------------------------------------------- 1h. medical materials (round 3): Census SAS 2022 Table 5 over Economic Census 2022 receipts
SAS_ITEMS = {'supplies': 'Medical supplies', 'other_materials': 'Expensed purchases of other materials, parts, and supplies',
             'purchased_services': 'Purchased professional and technical services', 'temp_staff': 'Temporary staff and leased employee expense',
             'payroll': 'Gross annual payroll', 'benefits': "Employer's cost for fringe benefits", 'expenses': 'Expenses'}
MAT_NAICS = {'phys': '6211', 'dent': '6212'}


def load_sas():
    ws = openpyxl.load_workbook(SAS5, read_only=True, data_only=True).worksheets[0]
    out = {}
    for r in ws.iter_rows(values_only=True):
        if r and str(r[0]).strip() in MAT_NAICS.values() and r[3] == 'All Establishments' and r[4] == 'Employer Firms':
            v = str(r[5]).replace(',', '')
            out.setdefault(str(r[0]).strip(), {})[str(r[2]).strip()] = float(v) if re.fullmatch(r'[\d.]+', v) else np.nan
    return out


SAS = load_sas()
EC_ALL = pd.read_csv(EC_US, sep='|', dtype=str, header=None)       # US rows of EC2200BASIC: 13 NAICS, 18 tax status, 21 firms, 25 receipts ($k), 27 payroll ($k)


def ec_row(code):
    r = EC_ALL[(EC_ALL[13] == code) & (EC_ALL[18] == '00')]
    assert len(r) == 1, (code, len(r))
    return r.iloc[0]


MAT, MAT_ALTS, MAT_AC4 = {}, ['no_purchased_services', 'pooled_small_classes', 'no_size_adjustment'], {}
for _c, _code in MAT_NAICS.items():
    s, e = SAS[_code], ec_row(_code)
    R_ = float(e[25]) / 1000.0                                          # receipts, $ million (SAS units)
    sh = {k: s[v] / R_ for k, v in SAS_ITEMS.items()}
    g_ = SOI[(SOI.code == COMP[_c]['soi']) & SOI.ok]
    c_all = float((g_[g_.ac == 1].cogs / g_[g_.ac == 1].br).mean())
    c_small = float((g_[g_.ac == MAT_SMALL_AC].cogs / g_[g_.ac == MAT_SMALL_AC].br).mean())
    _c4 = g_[g_.ac == MAT_SMALL_AC + 1].cogs / g_[g_.ac == MAT_SMALL_AC + 1].br         # the next class up, yearly (context)
    MAT_AC4[_c] = dict(mean=float(_c4.mean()), min=float(_c4.min()), max=float(_c4.max()))
    _cells = SOI_CELLS[(SOI_CELLS.code == COMP[_c]['soi'])]
    _ac_x = _cells.groupby('ac').x_real.median()
    _pool_ac = sorted(int(a) for a, x in _ac_x.items() if KNOTS[0] - TAU <= x <= KNOTS[-1] + TAU)
    _pool = _cells[_cells.ac.isin(_pool_ac)]
    c_pool = float(_pool.cogs.sum() / _pool.br.sum()) / float(g_[g_.ac == 1].cogs.sum() / g_[g_.ac == 1].br.sum())
    m_all = sh['supplies'] + sh['other_materials'] + sh['purchased_services']
    m_nops = sh['supplies'] + sh['other_materials']
    MAT[_c] = dict(naics=_code, receipts_musd=R_, firms=int(e[21]), shares={k: r4(v) for k, v in sh.items()},
                   census_share=m_all, soi_cogs_all=c_all, soi_cogs_small=c_small, small_ratio=c_small / c_all,
                   pooled_classes=_pool_ac, pooled_ratio=c_pool, share=m_all * c_small / c_all,
                   alternatives=dict(no_purchased_services=m_nops * c_small / c_all, pooled_small_classes=m_all * c_pool,
                                     no_size_adjustment=m_all))


# ---------------------------------------------------------------- 2. margins under each definition
def margin_us(defn, g, d, lam, pc=0.0, mat=None):
    """pc: employer payroll costs per dollar of the direct pay P = d(SW + L) (SOI salaries and wages carry none);
    mat: a materials share of receipts that replaces tax COGS less its labour where larger (physicians, dentists)"""
    L = lam * g.cogs
    if defn == 'rep':
        return 1 - g.cogs / g.br
    if defn == 'exl':
        return 1 - (g.cogs - L) / g.br
    if defn == 'allw':
        return 1 - (g.cogs + g.sw) / g.br
    P = d * (g.sw + L)
    if mat is None:
        return 1 - (g.cogs + np.maximum(0.0, P - L) + pc * P) / g.br                      # 'dir'
    return 1 - (np.maximum(g.cogs - L, mat * g.br) + np.maximum(L, P) + pc * P) / g.br


def margin_ca(defn, r, q, d, op=0.0, pc=0.0, mat=None):
    """op: owner-managers' salaries inside the cell's labour (CAD thousands per firm); never loaded. pc: payroll costs
    on the direct pay drawn from labour and commissions (salaries and wages only; cost-of-sales wages carry benefits,
    ISED glossary). mat: as margin_us, against purchases, materials and sub-contracts (cost of sales less its wages)"""
    rev, cos, dw, il = r['rev_' + q], r['cos_' + q], r['dw_' + q], r['il_' + q]
    if defn == 'rep':
        return 1 - cos / rev
    if defn == 'exl':
        return 1 - (cos - dw) / rev
    if defn == 'allw':
        return 1 - (cos + max(0.0, il - op)) / rev
    extra = max(0.0, d * (dw + il - op) - dw)
    if mat is not None:
        cos = dw + max(cos - dw, mat * rev)
    return 1 - (cos + extra * (1.0 + pc)) / rev


def adj(comp, defn, payroll=True, materials=True):
    """round-3 corrections on the default ('dir') definition: (payroll-cost factor, materials share). payroll True
    applies the factor to the PAYROLL_IND components, 'all' to every component (sensitivity), False to none;
    materials True uses the central Census share for physicians and dentists, a MAT_ALTS key an alternative, False none"""
    if defn != 'dir':
        return 0.0, None
    pc = PCF[comp] if (payroll == 'all' or (payroll is True and comp in PAYROLL_COMPS)) else 0.0
    if comp not in MAT or materials is False:
        return pc, None
    return pc, (MAT[comp]['share'] if materials is True else MAT[comp]['alternatives'][materials])


# ---------------------------------------------------------------- 3. size curves
class Curve:
    def __init__(self, x, y, w, kind):
        o = np.argsort(x, kind='mergesort')
        self.x, self.y, self.w, self.kind = np.asarray(x, float)[o], np.asarray(y, float)[o], np.asarray(w, float)[o], kind
        self.lo, self.hi = float(self.x.min()), float(self.x.max())

    def __call__(self, k):
        x, y = self.x, self.y
        if self.kind == 'interp':
            if k < x[0]:
                return float(y[0] + (y[1] - y[0]) / (x[1] - x[0]) * (k - x[0]))
            if k > x[-1]:
                return float(y[-1] + (y[-1] - y[-2]) / (x[-1] - x[-2]) * (k - x[-1]))
            return float(np.interp(k, x, y))
        if self.kind == 'line':                                         # global weighted line (research replication)
            X = np.c_[np.ones_like(x), x]
            sw = np.sqrt(self.w)
            b = np.linalg.lstsq(X * sw[:, None], y * sw, rcond=None)[0]
            return float(b[0] + b[1] * k)
        kw = self.w * np.clip(1 - np.abs(x - k) / H, 0, None)          # local-linear, triangular kernel
        if (kw > 0).sum() < 2:
            j = np.argsort(np.abs(x - k), kind='mergesort')[:2]
            kw = np.zeros_like(x)
            kw[j] = self.w[j]
        X = np.c_[np.ones_like(x), x - k]
        sw = np.sqrt(kw)
        b = np.linalg.lstsq(X * sw[:, None], y * sw, rcond=None)[0]
        return float(b[0])


def soi_class_means(comp, values, xcol='x_real', years=None, min_years=MIN_YEARS):
    g = SOI_CELLS[SOI_CELLS.code == COMP[comp]['soi']]
    if years is not None:
        g = g[g.year.isin(years)]
    v = values(g)
    t = pd.DataFrame(dict(ac=g.ac.values, x=g[xcol].values, m=np.asarray(v, float), n=g.n.values))
    c = t.groupby('ac').agg(x=('x', lambda s: np.log10(np.median(10 ** s))), m=('m', 'mean'), n=('n', 'mean'), yrs=('m', 'count'))
    return c[c.yrs >= min_years]


def us_curve(comp, defn, d, lam, years=None, xcol='x_real', kind='loclin', min_years=MIN_YEARS, pc=0.0, mat=None):
    g = SOI_CELLS[SOI_CELLS.code == COMP[comp]['soi']]
    if years is not None and len(years) == 1:
        g = g[g.year.isin(years)]
        m = margin_us(defn, g, d, lam, pc, mat)
        return Curve(g[xcol].values, m.values, np.sqrt(g.n.values), kind)
    c = soi_class_means(comp, lambda h: margin_us(defn, h, d, lam, pc, mat), xcol=xcol, years=years, min_years=min_years)
    return Curve(c.x.values, c.m.values, np.sqrt(c.n.values), kind)


# officers' compensation / business receipts by asset class (US corporations), smoothed the same way
OC_CURVES = {}
for _c in COMP:
    _t = soi_class_means(_c, lambda h: h.oc / h.br)
    OC_CURVES[_c] = Curve(_t.x.values, _t.m.values, np.sqrt(_t.n.values), 'loclin')


def oc_share(comp, x):
    """US officers' pay as a share of receipts at revenue 10**x (2026 USD); held flat beyond the cells +- TAU"""
    C = OC_CURVES[comp]
    return max(0.0, C(min(max(x, C.lo - TAU), C.hi + TAU)))


def ised_kappa(comp):
    """ISED's own reading of owner pay: incorporated minus unincorporated labour share at matched revenue (unincorporated
    owners' pay sits in profit), over the SOI officers'-pay share at the same revenue"""
    code = COMP[comp]['ca']
    I, U = FPD[(code, 2, 0)], FPD[(code, 1, 0)]
    ri, ru = np.array(I['rev'][1:5]), np.array(U['rev'][1:5])
    li = (np.array(I['dw'][1:5]) + np.array(I['il'][1:5])) / ri
    lu = (np.array(U['dw'][1:5]) + np.array(U['il'][1:5])) / ru
    pi_ = np.array(I['np'][1:5]) / ri
    pu_ = np.array(U['np'][1:5]) / ru
    pts = [r for r in ru if ri[0] <= r <= ri[-1]]
    rows = []
    for r in pts:
        lr = np.log(r)
        s_lab = float(np.interp(lr, np.log(ri), li) - np.interp(lr, np.log(ru), lu))
        s_np = float(np.interp(lr, np.log(ru), pu_) - np.interp(lr, np.log(ri), pi_))
        x = float(np.log10(r * 1000 * CA_TO_REAL))
        rows.append(dict(rev_cad_k=round(float(r), 1), x_real=round(x, 3), ised_labour_gap=r4(s_lab), ised_profit_gap=r4(s_np),
                         soi_officer_share=r4(oc_share(comp, x))))
    k_raw = sum(max(0.0, w['ised_labour_gap']) for w in rows) / sum(w['soi_officer_share'] for w in rows)
    return float(np.clip(k_raw, *KAPPA_CAP)), float(k_raw), rows


KAPPA, KAPPA_RAW, KAPPA_ROWS = {}, {}, {}
for _c in COMP:
    if COMP[_c]['ca']:
        KAPPA[_c], KAPPA_RAW[_c], KAPPA_ROWS[_c] = ised_kappa(_c)


def owner_factor(comp, owner):
    if owner == 'none':
        return 0.0
    return {'soi': 1.0, 'ised': KAPPA[comp], 'mid': (1.0 + KAPPA[comp]) / 2}[owner]


def ca_owner_pay(comp, q, owner):
    """owner-managers' salaries in the CA cell (CAD thousands per firm): officer-pay share x incorporated share x revenue"""
    r = ISED_REV[COMP[comp]['ca']]
    x = float(np.log10(r['rev_' + q] * 1000.0 * CA_TO_REAL))
    k = owner_factor(comp, owner)
    op = k * oc_share(comp, x) * PI_INC[COMP[comp]['ca']][int(q[1]) - 1] * r['rev_' + q]
    return min(op, r['dw_' + q] + r['il_' + q])


def owner_labour_share(comp, q, owner=OWNER_MODE):
    r = ISED_REV[COMP[comp]['ca']]
    return ca_owner_pay(comp, q, owner) / (r['dw_' + q] + r['il_' + q])


OWNER_LAB = {c: {q: owner_labour_share(c, q) for q in ('q1', 'q2', 'q3', 'q4')} for c in COMP if COMP[c]['ca']}
CA_DROPPED = sorted(c for c, v in OWNER_LAB.items() if max(v['q3'], v['q4']) >= OWNER_DROP)


def ca_curve(comp, defn, d, xcol='x_real', quartiles=('q1', 'q2', 'q3', 'q4'), kind='interp', owner=OWNER_MODE, pc=0.0, mat=None):
    code = COMP[comp]['ca']
    if code is None:
        return None
    r = ISED_REV[code]
    xs = [np.log10(r['rev_' + q] * 1000.0 * (CA_TO_REAL if xcol == 'x_real' else 1.0)) for q in quartiles]
    ys = [margin_ca(defn, r, q, d, ca_owner_pay(comp, q, owner) if defn in ('dir', 'allw') else 0.0, pc, mat) for q in quartiles]
    return Curve(xs, ys, np.ones(len(xs)), kind)


def combine(U, C, knots):
    """Mean of the two countries at each knot; a source past its cells (+TAU) follows the other's shape. With one
    country, its own local-linear trend is followed past its cells (dist records how far)."""
    us, ca, mode, dist = [], [], [], []
    for k in knots:
        if C is None:
            e = min(max(k, U.lo - TAU), U.hi + TAU)
            us.append(U(k)); ca.append(np.nan); dist.append(abs(k - e))
            mode.append('us_only' if e == k else 'us_only_extrapolated')
            continue
        uin, cin = U.lo - TAU <= k <= U.hi + TAU, C.lo - TAU <= k <= C.hi + TAU
        if uin and cin:
            u, c, md, e = U(k), C(k), 'both', k
        elif uin:
            e = min(max(k, C.lo - TAU), C.hi + TAU)
            u, c, md = U(k), C(e) + U(k) - U(e), 'ca_carried'
        elif cin:
            e = min(max(k, U.lo - TAU), U.hi + TAU)
            u, c, md = U(e) + C(k) - C(e), C(k), 'us_carried'
        else:
            u, c, md, e = U(k), C(k), 'neither_extrapolated', k
        us.append(u); ca.append(c); mode.append(md); dist.append(abs(k - e))
    us, ca = np.array(us), np.array(ca)
    est = us if C is None else (us + ca) / 2
    return est, us, ca, mode, np.array(dist)


def comp_curve(comp, defn, dshift=0.0, lamscale=1.0, years=None, xcol='x_real', knots=KNOTS, us_kind='loclin',
               ca_q=('q1', 'q2', 'q3', 'q4'), owner=OWNER_MODE, allow_drop=True, payroll=True, materials=True, d_override=None):
    d = float(np.clip((delta(comp) if d_override is None else d_override) + dshift, 0, 1))
    lam = LAM[COMP[comp]['sec']] * lamscale
    pc, mat = adj(comp, defn, payroll, materials)
    U = us_curve(comp, defn, d, lam, years=years, xcol=xcol, kind=us_kind, pc=pc, mat=mat)
    drop = allow_drop and owner != 'none' and comp in CA_DROPPED and defn in ('dir', 'allw')
    C = None if drop else ca_curve(comp, defn, d, xcol=xcol, quartiles=ca_q, kind='interp' if us_kind == 'loclin' else 'line', owner=owner, pc=pc, mat=mat)
    est, us, ca, mode, dist = combine(U, C, knots)
    return dict(est=est, us=us, ca=ca, mode=mode, dist=dist, U=U, C=C)


def ind_defn(ind, defn):
    return DEFN.get(ind, 'dir') if defn == 'conv' else defn


def ind_curve(ind, defn='conv', **kw):
    parts = [(c, w, comp_curve(c, ind_defn(ind, defn), **kw)) for c, w in INDUSTRIES[ind]]
    W = sum(w for _, w, _ in parts)
    est = sum(w * p['est'] for _, w, p in parts) / W
    us = sum(w * p['us'] for _, w, p in parts) / W
    ca = sum(w * p['ca'] for _, w, p in parts) / W
    return dict(est=est, us=us, ca=ca, parts=parts, W=W)


def ols_weights(knots):
    k = np.asarray(knots, float)
    return (k - k.mean()) / np.sum((k - k.mean()) ** 2)


def ols_slope(knots, v):
    return float(ols_weights(knots) @ np.asarray(v, float))


def page_value(knots, gm, x):
    """the page's rule: linear in log10 revenue between knots, flat beyond the outer knots"""
    return float(np.interp(x, knots, gm))


YEARS = sorted(SOI_CELLS.year.unique())


# ---------------------------------------------------------------- 4. size slope across the two countries (computed first: it gives the shape-transfer term of the source SE)
SLOPE_PAIRS = [('home', 'home'), ('constr', 'constr'), ('mfg', 'mfg'), ('mfg_misc', 'mfg_misc'), ('dist', 'dist'),
               ('vet', 'vet'), ('msp', 'msp'), ('acct', 'prof'), ('staff', 'staff'), ('land', 'land'), ('auto', 'auto'),
               ('wash', 'wash'), ('fit', 'fit'), ('rest', 'rest'), ('phys', 'medical'), ('dent', 'medical'), ('legal', 'prof')]


def slope_test(defn_mode, owner=OWNER_MODE, corrections=True):
    """corrections: the round-3 payroll-cost and materials corrections (False: the round-1/2 estimator)"""
    rows = []
    for comp, ind in SLOPE_PAIRS:
        d_ = ind_defn(ind, 'conv') if defn_mode == 'conv' else 'rep'
        dd = delta(comp)
        pc, mat = adj(comp, d_, corrections, corrections)
        U = us_curve(comp, d_, dd, LAM[COMP[comp]['sec']], pc=pc, mat=mat)
        r = ISED_REV[COMP[comp]['ca']]
        x3, x4 = (np.log10(r['rev_' + q] * 1000 * CA_TO_REAL) for q in ('q3', 'q4'))
        op = {q: (ca_owner_pay(comp, q, owner) if d_ in ('dir', 'allw') else 0.0) for q in ('q3', 'q4')}
        y3, y4 = margin_ca(d_, r, 'q3', dd, op['q3'], pc, mat), margin_ca(d_, r, 'q4', dd, op['q4'], pc, mat)
        rows.append(dict(component=comp, industry=ind, ca_dropped=comp in CA_DROPPED, x_q3=round(x3, 3), x_q4=round(x4, 3),
                         us_slope_per_decade=r4((U(x4) - U(x3)) / (x4 - x3)), ca_slope_per_decade=r4((y4 - y3) / (x4 - x3)),
                         pred_change=r4(U(x4) - U(x3)), actual_change=r4(y4 - y3), pred_level_q4=r4(U(x4)), actual_level_q4=r4(y4),
                         us_extrapolated_below=round(max(0.0, U.lo - x3), 3)))
    use = [r for r in rows if not r['ca_dropped']] if (defn_mode == 'conv' and owner != 'none') else rows
    pc = np.array([r['pred_change'] for r in use]); ac = np.array([r['actual_change'] for r in use])
    return dict(rows=rows, n=len(use), r=r4(np.corrcoef(pc, ac)[0, 1]), mean_pred=r4(pc.mean()), mean_actual=r4(ac.mean()),
                same_sign=int(np.sum(np.sign(pc) == np.sign(ac))),
                level_mad_q4=r4(np.mean([abs(r['pred_level_q4'] - r['actual_level_q4']) for r in use])),
                pairs_used='all %d' % len(use) if len(use) == len(rows) else 'the %d pairs whose Canadian cell is kept' % len(use))


SLOPE = dict(convention=slope_test('conv'), as_reported=slope_test('rep'), convention_round1=slope_test('conv', owner='none', corrections=False),
             convention_round2=slope_test('conv', corrections=False), convention_no_owner_fix_same_pairs=None,
             note='US curve change between the CA Q3 and Q4 revenues (2026 USD) against the CA change; the US curve is '
                  'extrapolated below its smallest class where the CA Q3 revenue lies below it. convention: the shipped definition '
                  '(owner-pay removal, payroll costs, medical materials); convention_round2: without the round-3 corrections; '
                  'convention_round1: no owner-pay removal either')
_r1 = [r for r in SLOPE['convention_round1']['rows'] if r['component'] not in CA_DROPPED]
_pc, _ac = np.array([r['pred_change'] for r in _r1]), np.array([r['actual_change'] for r in _r1])
SLOPE['convention_no_owner_fix_same_pairs'] = dict(n=len(_r1), r=r4(np.corrcoef(_pc, _ac)[0, 1]), mean_pred=r4(_pc.mean()), mean_actual=r4(_ac.mean()),
                                                   slope_gap_rms=r4(np.sqrt(np.mean([(r['us_slope_per_decade'] - r['ca_slope_per_decade']) ** 2 for r in _r1]))))
# where both countries have cells their size slopes differ by DB; past one country's cells its slope is taken from the
# other (or, with one country, its own trend is extended), so its level drifts by about RMS(DB) x distance
DB = {r['component']: r['us_slope_per_decade'] - r['ca_slope_per_decade'] for r in SLOPE['convention']['rows'] if not r['ca_dropped']}
DB_RMS = float(np.sqrt(np.mean([v ** 2 for v in DB.values()])))
SLOPE['convention']['slope_gap_rms'] = r4(DB_RMS)
SLOPE['convention_round2']['slope_gap_rms'] = r4(np.sqrt(np.mean([(r['us_slope_per_decade'] - r['ca_slope_per_decade']) ** 2
                                                                  for r in SLOPE['convention_round2']['rows'] if not r['ca_dropped']])))


# ---------------------------------------------------------------- 5. BACH: within-division spread and the median-minus-weighted-mean gap (derived statistics only)
BACH_SEC = {'home': 'F43', 'medical': 'Q86', 'vet': 'M75', 'msp': 'J62', 'saas': 'J58', 'prof': 'M69', 'mfg': 'C', 'dist': 'G46',
            'ecom': 'G47', 'staff': 'N78', 'land': 'N81', 'wash': 'G45', 'auto': 'G45', 'constr': 'F43', 'fit': 'R93', 'rest': 'I56'}
BACH_CONV = {k: ('mat' if k in ('dist', 'ecom', 'rest', 'saas') else 'mid') for k in BACH_SEC}


def bach():
    z = zipfile.ZipFile(BACH_ZIP)
    name = z.namelist()[0]
    f = io.TextIOWrapper(z.open(name), encoding='utf-8', errors='replace')
    f.readline()
    rd = csv.reader(f, delimiter=';')
    hdr = next(rd)
    ix = {h: i for i, h in enumerate(hdr)}
    secs = set(BACH_SEC.values())
    vals = [v + '_' + s for v in ('I5', 'r31', 'r32') for s in ('wm', 'q1', 'q2', 'q3')]
    rows, nrows = [], 0
    for row in rd:
        nrows += 1
        if row[ix['sector']] in secs and row[ix['size']] in ('1a', '1b') and row[ix['sample']] in ('0', '-1') \
                and row[ix['year']].isdigit() and 2019 <= int(row[ix['year']]) <= 2023:
            rec = dict(country=row[ix['country']], year=int(row[ix['year']]), sector=row[ix['sector']], size=row[ix['size']],
                       sample=row[ix['sample']])
            for v in vals:
                s = row[ix[v]]
                rec[v] = float(s) if s not in ('', None) else np.nan
            rows.append(rec)
    D = pd.DataFrame(rows)
    # margins: materials 1 - I5; value added R31; gross operating profit R32. median minus weighted mean, per cell:
    D['g_mat'] = (D.I5_wm - D.I5_q2) / 100
    D['g_r31'] = (D.r31_q2 - D.r31_wm) / 100
    D['g_r32'] = (D.r32_q2 - D.r32_wm) / 100
    D['g_mid'] = (D.g_r31 + D.g_r32) / 2               # the round-1 convention bracket for labour-loaded margins
    D['g_matgop'] = (D.g_mat + D.g_r32) / 2            # the other bracket: materials-only and all-labour margins
    res = {}
    for (sec, size, smp), g in D.groupby(['sector', 'size', 'sample']):
        g2 = g.dropna(subset=['I5_q1', 'I5_q3'])
        if not len(g2):
            continue
        rec = dict(cells=int(len(g2)), countries=int(g2.country.nunique()))
        for v, lab in [('I5', 'mat'), ('r31', 'gva'), ('r32', 'gop')]:
            rec[lab + '_sd'] = float(np.nanmedian((g2[v + '_q3'] - g2[v + '_q1']) / 100 / 1.349))
        for v in ('g_mat', 'g_r31', 'g_r32', 'g_mid', 'g_matgop'):
            x = g[v].dropna()
            rec[v] = float(np.median(x)) if len(x) else np.nan
            rec[v + '_pos_share'] = float((x > 0).mean()) if len(x) else np.nan
            rec[v + '_n'] = int(len(x))
        res[(sec, size, smp)] = rec
    return res, nrows, sorted(D.country.unique())


BACH_RES, BACH_ROWS, BACH_COUNTRIES = bach()


def bach_gap(ind, size):
    """convention-matched median-minus-weighted-mean gap and its bracket RMS (sample 0, 2019-23)"""
    b = BACH_RES[(BACH_SEC[ind], size, '0')]
    if BACH_CONV[ind] == 'mat':
        return b['g_mat'], abs(b['g_mat'])
    return b['g_mid'], float(np.sqrt((b['g_mid'] ** 2 + b['g_matgop'] ** 2) / 2))


GAP_DIVS = sorted({(BACH_SEC[i], BACH_CONV[i]) for i in BRIDGED})
GAP_SIGNS = {}
for size in ('1a', '1b'):
    g = [bach_gap(next(i for i in BRIDGED if (BACH_SEC[i], BACH_CONV[i]) == dv), size)[0] for dv in GAP_DIVS]
    GAP_SIGNS[size] = dict(n=len(g), positive=int(sum(v > 0 for v in g)), negative=int(sum(v < 0 for v in g)),
                           median=r4(np.median(g)), min=r4(min(g)), max=r4(max(g)))
MEDIAN_RULE = 'shift' if all(GAP_SIGNS[s]['positive'] in (0, GAP_SIGNS[s]['n']) for s in ('1a', '1b')) else 'se'


def medium_weight(x):
    """weight on BACH's medium class at revenue 10**x: 0 up to $3.2M, 1/2 at $10M (the EUR 10M class boundary)"""
    return 0.5 * float(np.clip((x - 6.5) / 0.5, 0, 1))


def median_gap(ind, x):
    """(signed gap, size) at 10**x: small-class gap, blended with the medium class toward $10M"""
    w = medium_weight(x)
    gs, ss = bach_gap(ind, '1a')
    gm, sm_ = bach_gap(ind, '1b')
    return (1 - w) * gs + w * gm, float(np.sqrt((1 - w) * ss ** 2 + w * sm_ ** 2))


# ---------------------------------------------------------------- 5b. recency (round 3): SOI Table 5.1 2019-22 against the Source Book 2004-13
def soi_t51():
    rows = []
    for f in sorted(glob.glob(os.path.join(SOI_CORP_DIR, '*co51ccr.xlsx'))):
        yr = 2000 + int(os.path.basename(f)[:2])
        data = list(openpyxl.load_workbook(f, read_only=True, data_only=True).worksheets[0].iter_rows(min_row=1, max_row=120, values_only=True))
        cn = next(i for i, r in enumerate(data) if len(r) > 3 and r[1] == 1 and r[2] == 2 and r[3] == 3)
        names = {c: ' || '.join(re.sub(r'\s+', ' ', str(data[i][c])).strip() for i in range(4, cn) if c < len(data[i]) and data[i][c] not in (None, ''))
                 for c in range(1, len(data[cn]))}

        def getrow(label):
            for i in range(cn + 1, len(data)):
                v = data[i][0]
                if v and re.sub(r'\s+', ' ', str(v)).strip().lower().startswith(label):
                    return data[i]
        br, cg, sw = getrow('business receipts'), getrow('cost of goods sold'), getrow('salaries and wages')
        assert br is not None and cg is not None and sw is not None, f
        for c in range(1, len(data[cn])):
            if data[cn][c] is None:
                continue
            rows.append(dict(year=yr, name=names[c], br=pd.to_numeric(br[c] if c < len(br) else None, errors='coerce'),
                             cogs=pd.to_numeric(cg[c] if c < len(cg) else None, errors='coerce'),
                             sw=pd.to_numeric(sw[c] if c < len(sw) else None, errors='coerce')))
    return pd.DataFrame(rows)


T51 = soi_t51()
RECENCY_CELLS = {'home': (238220, 'Plumbing, heating'), 'constr': (238905, 'Other specialty trade'), 'mfg_misc': (339900, 'Other miscellaneous manufacturing'),
                 'dist': (423300, 'Lumber and other construction'), 'saas': (511210, 'Software publishers'), 'legal': (541115, 'Legal services'),
                 'acct': (541215, 'Accounting, tax'), 'msp': (541515, 'Computer systems design'), 'vet': (541915, 'Other professional, scientific'),
                 'staff': (561300, 'Employment services'), 'land': (561905, 'Other administrative and support'), 'phys': (621115, 'Offices of physicians'),
                 'dent': (621210, 'Offices of dentists'), 'fit': (713, 'Amusement, gambling'), 'rest': (722, 'Food services and drinking'),
                 'auto': (811115, 'Automotive repair')}


def soi_drift(comp):
    """all-size margins, Source Book 2004-13 (AC 1) and Table 5.1 2019-22: as reported, and on the labour-loaded
    definition without the round-3 corrections (the reviewer's reading; d and lambda as the bridge)"""
    code, pref = RECENCY_CELLS[comp]
    old = SOI[(SOI.code == code) & (SOI.ac == 1) & SOI.ok]
    new = T51[T51.name.str.startswith(pref)].groupby('year').first()
    d, lam = delta(comp), LAM[COMP[comp]['sec']]
    return dict(rep_old=float((1 - old.cogs / old.br).mean()), rep_new=float((1 - new.cogs / new.br).mean()),
                dir_old=float(margin_us('dir', old, d, lam).mean()), dir_new=float(margin_us('dir', new, d, lam).mean()), years_new=int(len(new)))


def small_class_trend(comp, horizon=2024 - 2008.5):
    """2004-13 trend of the small asset classes' labour-loaded margin (classes $1-500K to $1-5M; class effects, sqrt-returns
    weights), carried from the sample midpoint to 2024: a cross-check of the recency term"""
    g = SOI_CELLS[(SOI_CELLS.code == COMP[comp]['soi']) & SOI_CELLS.ac.isin([3, 4, 5])]
    m = margin_us('dir', g, delta(comp), LAM[COMP[comp]['sec']]).values
    X = np.c_[pd.get_dummies(g.ac.astype(str)).astype(float).values, g.year.values - 2008.5]
    w = np.sqrt(g.n.values)
    b = np.linalg.lstsq(X * w[:, None], m * w, rcond=None)[0]
    res = m - X @ b
    s2 = float(np.sum(w ** 2 * res ** 2) / (len(m) - X.shape[1]))
    se = float(np.sqrt(s2 * np.linalg.inv((X * (w ** 2)[:, None]).T @ X)[-1, -1]))
    return dict(per_year=r4(b[-1]), se_per_year=r4(se), to_2024=r4(b[-1] * horizon), to_2024_se=r4(se * horizon))


RECENCY, RECENCY_DETAIL, RECENCY_TREND = {}, {}, {}
for _i in RECENCY_IND:
    _W = sum(w for _, w in INDUSTRIES[_i])
    _dr = {c: soi_drift(c) for c, _ in INDUSTRIES[_i]}
    for c, _ in INDUSTRIES[_i]:
        RECENCY_TREND[c] = small_class_trend(c)
    _wd = sum(w / _W * (_dr[c]['dir_new'] - _dr[c]['dir_old']) for c, w in INDUSTRIES[_i])
    _wt = sum(w / _W * RECENCY_TREND[c]['to_2024'] for c, w in INDUSTRIES[_i])
    RECENCY[_i] = RECENCY_FRAC * abs(_wd)
    RECENCY_DETAIL[_i] = dict(se_term=r4(RECENCY[_i]), all_size_drift_labour_loaded=r4(_wd), small_class_trend_to_2024=r4(_wt),
                              components={c: dict(weight=r4(w / _W), drift=r4(_dr[c]['dir_new'] - _dr[c]['dir_old']), trend=RECENCY_TREND[c])
                                          for c, w in INDUSTRIES[_i]},
                              covers_trend=bool(abs(_wt) <= RECENCY[_i]))


# ---------------------------------------------------------------- 6. the bridge: knots, error covariance and SE terms
SE_TERMS = ['time', 'source', 'shape', 'bridge', 'proxy', 'mix', 'owner', 'median', 'payroll', 'materials', 'recency']


KNOT_MID = float(np.mean(KNOTS))


def fit_industry(ind, defn='conv', xcol='x_real', xs=KNOTS, tilt=0.0):
    xs = [float(v) for v in xs]
    n = len(xs)
    base = ind_curve(ind, defn, xcol=xcol, knots=xs)
    est = np.asarray(base['est'], float)
    d_ = ind_defn(ind, defn)
    per_year = np.array([ind_curve(ind, defn, years=[y], xcol=xcol, knots=xs)['est'] for y in YEARS])
    cov = {'time': np.atleast_2d(np.cov(per_year, rowvar=False, ddof=1)) / len(YEARS)}
    vecs = {k: [] for k in SE_TERMS if k != 'time'}
    single = False
    for c, w, p in base['parts']:
        f = w / base['W']
        if p['C'] is None:
            single = True
            vecs['source'].append(f * SRC_SINGLE * np.ones(n))
            vecs['shape'].append(f * DB_RMS * p['dist'])                  # one country: its own trend past its cells
            vecs['source'].append(f * tilt * (np.array(xs) - KNOT_MID))   # and the size slope it cannot check against a second country
        else:
            vecs['source'].append(f * (p['us'] - p['ca']) / 2)            # half the US-CA gap
            vecs['shape'].append(f * DB_RMS * p['dist'] / 2)              # a carried country takes the other's slope
    if d_ == 'dir':
        hi, lo = ind_curve(ind, defn, dshift=D_SD, xcol=xcol, knots=xs)['est'], ind_curve(ind, defn, dshift=-D_SD, xcol=xcol, knots=xs)['est']
    elif d_ == 'exl':
        hi, lo = ind_curve(ind, defn, lamscale=LAM_REST[1], xcol=xcol, knots=xs)['est'], ind_curve(ind, defn, lamscale=LAM_REST[0], xcol=xcol, knots=xs)['est']
    else:
        hi = lo = est
    vecs['bridge'].append((np.asarray(hi) - np.asarray(lo)) / 2)
    vecs['proxy'].append(PROXY.get(ind, 0.0) * np.ones(n))
    if len(base['parts']) == 2:
        (c1, w1, p1), (c2, w2, p2) = base['parts']
        pw = w1 / (w1 + w2)
        sp = np.sqrt(pw * (1 - pw) / (w1 + w2))
        vecs['mix'].append(sp * (p1['est'] - p2['est']))
    if d_ in ('dir', 'allw') and xcol == 'x_real':
        e_ised = ind_curve(ind, defn, xcol=xcol, knots=xs, owner='ised')['est']
        vecs['owner'].append((np.asarray(e_ised) - est) / 2)             # ISED's own reading of owner pay, half the change
    gaps = [median_gap(ind, x) for x in xs]
    shift = np.array([g[0] for g in gaps]) if MEDIAN_RULE == 'shift' else np.zeros(n)
    if MEDIAN_RULE == 'se':
        # one level within the knots read off BACH's small class and within those that blend in the medium class;
        # MEDIAN_RHO between the two groups (round 3: 1 before)
        m = np.array([g[1] for g in gaps])
        grp = np.array([medium_weight(x) > 0 for x in xs])
        vecs['median'] += [np.sqrt(MEDIAN_RHO) * m, np.sqrt(1 - MEDIAN_RHO) * m * ~grp, np.sqrt(1 - MEDIAN_RHO) * m * grp]
    pay_vec = np.zeros(n)
    if d_ == 'dir' and any(c in PAYROLL_COMPS for c, _, _ in base['parts']):
        e_np = ind_curve(ind, defn, xcol=xcol, knots=xs, payroll=False)['est']
        pay_vec = PAYROLL_SE * (np.asarray(e_np) - est)                 # half the payroll-cost cut, one sign across sizes
        vecs['payroll'].append(pay_vec)
    for c, w, p in base['parts']:
        if d_ == 'dir' and c in MAT:                                    # materials: RMS of the alternative readings, per component
            dev = [np.asarray(comp_curve(c, d_, xcol=xcol, knots=xs, materials=a)['est']) - p['est'] for a in MAT_ALTS]
            vecs['materials'].append(w / base['W'] * np.sqrt(np.mean(np.square(dev), axis=0)))
    if ind in RECENCY and defn == 'conv':
        vecs['recency'].append(RECENCY[ind] * np.ones(n))                # a level term
    for k, vs in vecs.items():
        cov[k] = sum(np.outer(v, v) for v in vs) if vs else np.zeros((n, n))
    Sigma = sum(cov[k] for k in SE_TERMS)
    return dict(xs=xs, est=est + shift, est_ratio_of_sums=est, shift=shift, cov=Sigma, cov_terms=cov, payroll_vec=pay_vec,
                se=np.sqrt(np.diag(Sigma)), se_terms={k: np.sqrt(np.clip(np.diag(cov[k]), 0, None)) for k in SE_TERMS},
                us=base['us'], ca=base['ca'], single_source=single, per_year=per_year,
                modes={c: p['mode'] for c, _, p in base['parts']},
                carry_distance={c: [round(float(v), 3) for v in p['dist']] for c, _, p in base['parts']},
                components={c: dict(weight=w, est=p['est'].tolist(), us=p['us'].tolist(), ca=p['ca'].tolist(), ca_dropped=p['C'] is None and COMP[c]['ca'] is not None,
                                    us_cells_x=[round(v, 4) for v in p['U'].x], us_cells_gm=[round(v, 4) for v in p['U'].y],
                                    ca_cells_x=None if p['C'] is None else [round(v, 4) for v in p['C'].x],
                                    ca_cells_gm=None if p['C'] is None else [round(v, 4) for v in p['C'].y])
                            for c, w, p in base['parts']})


def knot_view(F, idx=range(len(KNOTS))):
    """restrict a fit to the knots: est, SE, slope and the slope's SE from the covariance"""
    idx = list(idx)
    S = F['cov'][np.ix_(idx, idx)]
    c = ols_weights(KNOTS)
    return dict(est=F['est'][idx], se=np.sqrt(np.diag(S)), cov=S, slope=float(c @ F['est'][idx]), slope_var=float(c @ S @ c))


FITS = {ind: fit_industry(ind) for ind in IND_ORDER if ind != 'ecom'}
# one-country industries: the cross-country slope uncertainty they cannot measure is the median two-country source term


def slope_source_term(F):
    W = sum(v['weight'] for v in F['components'].values())
    t = 0.0
    for v in F['components'].values():
        if v['ca_cells_x'] is not None:
            t += (v['weight'] / W * abs(ols_slope(KNOTS, (np.array(v['us'][:len(KNOTS)]) - np.array(v['ca'][:len(KNOTS)])) / 2))) ** 2
    return float(np.sqrt(t))


SLOPE_SRC_MED = float(np.median([slope_source_term(FITS[i]) for i in BRIDGED if not FITS[i]['single_source']]))
SINGLE = [i for i in FITS if FITS[i]['single_source']]
for _i in SINGLE:
    FITS[_i] = fit_industry(_i, tilt=SLOPE_SRC_MED)


# ---------------------------------------------------------------- 7. direct surveys
PDFTOTEXT = shutil.which('pdftotext')


def pdf_text(path, page, layout=True):
    if not PDFTOTEXT:
        return None
    cmd = [PDFTOTEXT] + (['-layout'] if layout else []) + ['-f', str(page), '-l', str(page), path, '-']
    return subprocess.run(cmd, capture_output=True, text=True, check=True).stdout


def bins_from_text(txt, bins, col):
    """re-read 'label   NN%' entries of one column of a two-column layout page: col 'L' (label before column 40), 'R' (after)"""
    if txt is None:
        return None
    out = []
    for lab, *_ in bins:
        found = None
        for line in txt.split('\n'):
            for m in re.finditer(re.escape(lab) + r'\s+(\d+)%', line):
                if (col == 'L' and m.start() < 40) or (col == 'R' and m.start() >= 40):
                    found = int(m.group(1))
                    break
            if found is not None:
                break
        out.append(found)
    return out


def grouped_median(bins, counts):
    tot, cum = float(sum(counts)), 0.0
    for (lab, a, b, _), n in zip(bins, counts):
        if cum + n >= tot / 2:
            assert a is not None and b is not None, 'median in an open bin'
            return a + (b - a) * (tot / 2 - cum) / n
        cum += n


def icnormal(bins, counts):
    """normal fitted to interval-censored counts by maximum likelihood (open end bins allowed)"""
    def nll(p):
        mu, s = p[0], np.exp(p[1])
        ll = 0.0
        for (lab, a, b, _), n in zip(bins, counts):
            pa = 0.0 if a is None else stats.norm.cdf((a - mu) / s)
            pb = 1.0 if b is None else stats.norm.cdf((b - mu) / s)
            ll += n * np.log(max(pb - pa, 1e-300))
        return -ll
    r = optimize.minimize(nll, x0=[grouped_median(bins, counts), np.log(0.1)], method='Nelder-Mead',
                          options=dict(xatol=1e-9, fatol=1e-12, maxiter=20000))
    return float(r.x[0]), float(np.exp(r.x[1]))


def grouped_logmedian(bins, counts, open_lo, open_hi):
    """median of a revenue distribution given in bins, log-uniform within bins (open bins closed at open_lo / open_hi)"""
    tot, cum = float(sum(counts)), 0.0
    for (lab, a, b, _), n in zip(bins, counts):
        a, b = (open_lo if a is None else a), (open_hi if b is None else b)
        if cum + n >= tot / 2:
            return float(np.log10(a) + (np.log10(b) - np.log10(a)) * (tot / 2 - cum) / n)
        cum += n


RNG = np.random.default_rng(SEED)

# ---- 7a. Ratchet+Wrench 2026 Industry Survey (auto repair). Constants transcribed from the PDF text layer, re-read below.
RW = dict(
    gp_page=9, gp_bins=[('Less than 30%', None, .30, 10), ('30-39%', .30, .40, 18), ('40-49%', .40, .50, 22), ('50-59%', .50, .60, 28),
                        ('60% or more', .60, None, 13)], gp_not_tracked=9,
    rev_page=5, rev_bins=[('Under $250,000', None, 250e3, 14), ('$250,000-$499,999', 250e3, 500e3, 16), ('$500,000-$749,999', 500e3, 750e3, 14),
                          ('$750,000-$999,999', 750e3, 1e6, 12), ('$1,000,000-$2,499,999', 1e6, 2.5e6, 30), ('Over $2,500,000', 2.5e6, None, 10)],
    rev_unknown=4, rev_open=(100e3, 10e6),
    dept_page=10, parts_bins=[('Less than 30%', None, .30, 10), ('30-39%', .30, .40, 22), ('40-49%', .40, .50, 24), ('50-59%', .50, .60, 27),
                              ('60% or more', .60, None, 10)], parts_not_tracked=7,
    labor_bins=[('Less than 50%', None, .50, 13), ('50-59%', .50, .60, 22), ('60-69%', .60, .70, 23), ('70-79%', .70, .80, 21),
                ('80% or more', .80, None, 7)], labor_not_tracked=14,
    resp_page=2, n_responses=430, survey='first quarter of 2026', revenue_year=2025)
rw_txt = {p: pdf_text(RW_PDF, p) for p in (RW['gp_page'], RW['rev_page'], RW['dept_page'])}
rw_raw2 = pdf_text(RW_PDF, RW['resp_page'], layout=False)
RW_QUOTES = {
    'gp bins (p. 9, right column, under OVERALL GROSS PROFIT MARGIN)': None if rw_txt[9] is None else bins_from_text(rw_txt[9][rw_txt[9].index('OVERALL GROSS PROFIT MARGIN'):], RW['gp_bins'], 'R') == [b[3] for b in RW['gp_bins']],
    'gp does not track 9% (p. 9)': None if rw_txt[9] is None else bool(re.search(r'OVERALL GROSS PROFIT MARGIN[\s\S]*Does not track\s+9%', rw_txt[9])),
    'revenue bins (p. 5)': None if rw_txt[5] is None else bins_from_text(rw_txt[5], RW['rev_bins'], 'L') == [b[3] for b in RW['rev_bins']],
    'parts gp bins (p. 10, left)': None if rw_txt[10] is None else bins_from_text(rw_txt[10], RW['parts_bins'], 'L') == [b[3] for b in RW['parts_bins']],
    'labor gp bins (p. 10, right)': None if rw_txt[10] is None else bins_from_text(rw_txt[10], RW['labor_bins'], 'R') == [b[3] for b in RW['labor_bins']],
    'more than 430 responses (p. 2)': None if rw_raw2 is None else bool(re.search(r'more\s+than\s+430\s+survey\s+responses', rw_raw2)),
}
assert all(v is not False for v in RW_QUOTES.values()), RW_QUOTES
rw_counts = [b[3] for b in RW['gp_bins']]
RW_TRACK = sum(rw_counts) / 100.0
RW_N = int(round(RW['n_responses'] * RW_TRACK))                      # respondents who track it (430 is a lower bound)
RW_MED = grouped_median(RW['gp_bins'], rw_counts)
RW_MU, RW_SD = icnormal(RW['gp_bins'], rw_counts)
_bs = [grouped_median(RW['gp_bins'], list(RNG.multinomial(RW_N, np.array(rw_counts) / sum(rw_counts)))) for _ in range(4000)]
RW_SE_SAMPLING = float(np.std(_bs, ddof=1))
RW_SE_BINNING = abs(RW_MED - RW_MU)                                  # linear-in-bin reading against the fitted normal's median
RW_SE = float(np.hypot(RW_SE_SAMPLING, RW_SE_BINNING))
RW_SE_POPULATION = 0.02                                               # judgment, as KBCM's: self-selected trade-magazine respondents, an undefined term
rw_rev_counts = [b[3] for b in RW['rev_bins']]
RW_X = float(grouped_logmedian(RW['rev_bins'], rw_rev_counts, *RW['rev_open']) + np.log10(CPI_REF / CPI_AVAIL[RW['revenue_year']]))
RW_PARTS_MED = grouped_median(RW['parts_bins'], [b[3] for b in RW['parts_bins']])
RW_LABOR_MED = grouped_median(RW['labor_bins'], [b[3] for b in RW['labor_bins']])
# what the survey's term leaves out when a shop counts technician pay without its payroll costs (the job-costing habit):
# technician pay ~ the direct-labour load of the US bridge at the survey's size; payroll costs = SOI employee-benefit
# and pension deductions over pay (811115, 10-year mean) + statutory employer taxes (FICA 7.65%, FUTA 0.6%). The US
# half of the bridge loads SOI salaries without them too, so after the round-3 payroll cut both are put on the shipped
# basis: the bridge by its own cut, the survey by this gap (the full US-return cut at its size)
def direct_pay_curve(comp):
    """the direct staff's pay as a share of receipts on US returns (d x (salaries + labour in COGS)), by size"""
    lam = LAM[COMP[comp]['sec']]
    c = soi_class_means(comp, lambda h: delta(comp) * (h.sw + lam * h.cogs) / h.br)
    return Curve(c.x.values, c.m.values, np.sqrt(c.n.values), 'loclin')


RW_PAYROLL_FACTOR = PCF['auto']
RW_TECH_PAY = direct_pay_curve('auto')(RW_X)
RW_DEF_GAP = RW_TECH_PAY * RW_PAYROLL_FACTOR                            # 0 if payroll costs are counted, this much if not

# ---- 7b. 2021 KBCM Private SaaS Company Survey (firms with $5M+ of 2020 revenue). Constants re-read from the PDF text.
KBCM = dict(sub_page=52, sub_bins=[('Below 55%', None, .55, 4), ('55% - 60%', .55, .60, 5), ('60% - 65%', .60, .65, 8),
                                   ('65% - 70%', .65, .70, 21), ('70% - 75%', .70, .75, 22), ('75% - 80%', .75, .80, 22),
                                   ('80% - 85%', .80, .85, 51), ('85% - 90%', .85, .90, 23), ('Over 90%', .90, None, 16)],
            sub_n=172, sub_median_printed=0.80,
            total_page=56, total_median=0.73, total_n=169,
            size_page=57, size_total={'all': 0.73, '5-25M': 0.75, '25-50M': 0.75, '50-100M': 0.68, '>100M': 0.75},
            size_n={'5-25M': 101, '25-50M': 33, '50-100M': 22, '>100M': 17},
            arr_page=51, arr_bins=[('$5MM-10MM', 5e6, 10e6, 59), ('$10MM-25MM', 10e6, 25e6, 74)], revenue_year=2020,
            cs_page=53, cs_total={'all or most in COGS': (43, 0.72), 'all or most in S&M': (65, 0.76), 'split': (40, 0.75)})
kb_txt = {p: pdf_text(KBCM_PDF, p) for p in (KBCM['sub_page'], KBCM['total_page'], KBCM['size_page'], KBCM['arr_page'], KBCM['cs_page'])}


def _kb_sub_counts(txt):
    if txt is None:
        return None
    out = []
    for lab, *_ in KBCM['sub_bins']:
        m = re.search(re.escape(lab) + r'\s+(\d+)\s*$', txt, re.M)
        out.append(int(m.group(1)) if m else None)
    return out


KBCM_QUOTES = {
    'subscription gm bins (p. 52)': None if kb_txt[52] is None else _kb_sub_counts(kb_txt[52]) == [b[3] for b in KBCM['sub_bins']],
    '172 respondents (p. 52)': None if kb_txt[52] is None else bool(re.search(r'172 respondents', kb_txt[52])),
    'total gm median 73%, support in COGS (p. 56)': None if kb_txt[56] is None else bool(re.search(r'Total Gross Margin\d?\s+73%', kb_txt[56]) and re.search(r'including customer support in COGS', kb_txt[56])),
    'gross margin: 169 respondents (p. 56)': None if kb_txt[56] is None else bool(re.search(r'Gross Margin: 169', kb_txt[56])),
    'total gm by size 73/75/75/68/75 (p. 57)': None if kb_txt[57] is None else bool(re.search(r'Total Gross Margin\s+73%\s+75%\s+75%\s+68%\s+75%', kb_txt[57])),
    '$5MM-$25MM: 101 respondents (p. 57)': None if kb_txt[57] is None else bool(re.search(r'\$5MM-\$25MM: 101', kb_txt[57])),
    'ARR bins $5-10M 59, $10-25M 74 (p. 51)': None if kb_txt[51] is None else bool(re.search(r'\$5MM\s*.10MM: 59, \$10MM - 25MM: 74', kb_txt[51])),
    'SBC added back (p. 56)': None if kb_txt[56] is None else bool(re.search(r'stock-based compensation add-back', kb_txt[56])),
}
assert all(v is not False for v in KBCM_QUOTES.values()), KBCM_QUOTES
kb_counts = [b[3] for b in KBCM['sub_bins']]
assert sum(kb_counts) == KBCM['sub_n']
KB_SUB_MED = grouped_median(KBCM['sub_bins'], kb_counts)
KB_MU, KB_SD = icnormal(KBCM['sub_bins'], kb_counts)
KB_LEVEL = KBCM['size_total']['5-25M']
KB_N = KBCM['size_n']['5-25M']
KB_X = float(grouped_logmedian(KBCM['arr_bins'], [b[3] for b in KBCM['arr_bins']], None, None) + np.log10(CPI_REF / CPI_AVAIL[KBCM['revenue_year']]))
KB_SE_SAMPLING = float(np.sqrt(np.pi / 2) * KB_SD / np.sqrt(KB_N))      # median of n draws; spread from the subscription histogram
KB_SE_ROUNDING = 0.005 / np.sqrt(3)                                     # medians printed to whole percent
KB_SE_POPULATION = 0.02                                                 # judgment: 2020, venture- and PE-backed respondents
KB_SE = float(np.sqrt(KB_SE_SAMPLING ** 2 + KB_SE_ROUNDING ** 2 + KB_SE_POPULATION ** 2))

# ---- 7c. Empire Flippers sold listings: cogs_percent (EF: 'number of units sold X landed cost') and tacos_percent
EF_PROD = ['Amazon FBA', 'eCommerce', 'DropShipping', 'Amazon FBM', 'Subscription-Box']     # businesses that sell merchandise
EF_COLS = ['listing_number', 'listing_status', 'monetization_primary', 'cogs_percent', 'tacos_percent', 'profit_margin',
           'average_annual_gross_revenue', 'first_listed_date']
ef_all = pd.read_csv(EF_CSV, usecols=EF_COLS, low_memory=False)
ef = ef_all[(ef_all.listing_status == 'Sold') & ef_all.cogs_percent.notna() & ef_all.monetization_primary.isin(EF_PROD)].copy()
ef = ef.sort_values('listing_number', kind='mergesort').reset_index(drop=True)
ef['gm'] = 1 - ef.cogs_percent / 100.0
ef['year'] = ef.first_listed_date.str[:4].astype(int)
ef['x'] = np.log10([to_real(r, y) for r, y in zip(ef.average_annual_gross_revenue, ef.year)])
assert ef.gm.between(0, 1).all() and len(ef) >= 250, len(ef)
EF_N = int(len(ef))
EF_MED = float(np.median(ef.gm))
EF_SE_SAMPLING = float(np.std([np.median(RNG.choice(ef.gm.values, EF_N)) for _ in range(4000)], ddof=1))
EF_Q1, EF_Q3 = (float(v) for v in np.percentile(ef.gm, [25, 75]))
EF_SD_IQR, EF_SD_RAW = (EF_Q3 - EF_Q1) / 1.349, float(ef.gm.std(ddof=1))
_qr = sm.QuantReg(ef.gm.values, sm.add_constant(ef.x.values)).fit(q=0.5)
EF_SLOPE, EF_SLOPE_SE = float(_qr.params[1]), float(_qr.bse[1])


def _wmedian(v, w):
    o = np.argsort(v, kind='mergesort')
    cw = np.cumsum(w[o])
    return float(v[o][np.searchsorted(cw, 0.5 * cw[-1])])


_fba = (ef.monetization_primary == 'Amazon FBA').values
EF_FBA_SHARE = float(_fba.mean())
EF_MED_FBA, EF_MED_NONFBA = float(np.median(ef.gm[_fba])), float(np.median(ef.gm[~_fba]))
EF_MED_HALF = _wmedian(ef.gm.values, np.where(_fba, 0.5 / _fba.sum(), 0.5 / (~_fba).sum()))     # half Amazon, half other sellers
EF_SE_MIX = max(abs(EF_MED_HALF - EF_MED), abs(EF_MED_FBA - EF_MED))
_yr = ef.groupby('year').gm.agg(['count', 'median'])
_yr = _yr[_yr['count'] >= 30]
EF_SE_TIME = float(_yr['median'].std(ddof=1) / np.sqrt(len(_yr)))
EF_SE_DEFN = 0.01                                                       # judgment: sellers' compliance with 'landed cost'
EF_X90 = float(np.percentile(ef.x, 90))


def ef_se(x):
    beyond = max(0.0, x - EF_X90) * EF_SLOPE_SE                          # flat curve carried past 90% of the listings
    return float(np.sqrt(EF_SE_SAMPLING ** 2 + EF_SE_MIX ** 2 + EF_SE_TIME ** 2 + EF_SE_DEFN ** 2 + beyond ** 2))


EF_TACOS_MED = float(np.median(ef.tacos_percent.dropna())) / 100
EF_OTHER = {s: dict(n=int(((ef_all.listing_status == s) & ef_all.cogs_percent.notna() & ef_all.monetization_primary.isin(EF_PROD)).sum()),
                    median=r4(float(np.median(1 - ef_all.loc[(ef_all.listing_status == s) & ef_all.cogs_percent.notna() & ef_all.monetization_primary.isin(EF_PROD), 'cogs_percent'] / 100))))
            for s in ('For Sale', 'Pending Sold')}
def page_text(path):
    return re.sub(r'\s+', ' ', ' '.join(page_lines(path)))


EF_DOC_QUOTE = bool(re.search(r'Cost of Goods Sold \(COGS\) in accrual format \(number of units sold X landed cost\)', page_text(EF_DOC)))
assert EF_DOC_QUOTE
FB_QUOTE = bool(re.search(r'Labor . 60 percent \(attained by paying techs on a 60/40 split\)', page_text(FB_DOC)))
assert FB_QUOTE


# ---------------------------------------------------------------- 8. the bridge with the surveys
def blup(est, cov, s, S, var_s, c=None):
    """inverse-variance update at the survey's size; the other points move through the bridge's own error covariance.
    c: covariance of the bridge's errors with the survey's (a shared adjustment, e.g. the payroll-cost cut); None = 0"""
    k = cov[:, s] if c is None else cov[:, s] - c
    den = cov[s, s] + var_s - (0.0 if c is None else 2.0 * c[s])
    g = k / den
    return est + g * (S - est[s]), cov - np.outer(k, k) / den, g


NK = len(KNOTS)
CW = ols_weights(KNOTS)
# saas: KBCM total gross margin ($5-25M of 2020 revenue): same definition (support and hosting in COGS, payroll costs
# included); the population matches at the top knot only, so it enters at its own size
F_SAAS = fit_industry('saas', xs=KNOTS + [KB_X], tilt=SLOPE_SRC_MED)


def mixture_offset(ind, bins, open_lo, open_hi, year, sd, x_at, tilt=0.0):
    """the bridge's own median over a survey's revenue mix (within-industry spread sd) minus the bridge at x_at"""
    xs_, w_ = [], []
    for (lab, a, b, w) in bins:
        a, b = (open_lo if a is None else a), (open_hi if b is None else b)
        xs_.append(0.5 * (np.log10(a) + np.log10(b)) + np.log10(CPI_REF / CPI_AVAIL[year]))
        w_.append(w)
    bx = ind_curve(ind, knots=xs_ + [x_at])['est']
    w_ = np.array(w_, float)
    m = optimize.brentq(lambda v: float(np.sum(w_ * stats.norm.cdf((v - bx[:-1]) / sd))) / w_.sum() - 0.5, 0.0, 1.0)
    return float(m - bx[-1]), [round(float(v), 3) for v in xs_]


KB_MIX, KB_MIX_X = mixture_offset('saas', KBCM['arr_bins'], None, None, KBCM['revenue_year'], KB_SD, KB_X)
KB_LEVEL_AT_X = KB_LEVEL - KB_MIX                     # the survey median moved to the bridge's point
SAAS_EST, SAAS_COV, SAAS_GAIN = blup(F_SAAS['est'], F_SAAS['cov'], NK, KB_LEVEL_AT_X, KB_SE ** 2)
SAAS_W = float(F_SAAS['cov'][NK, NK] / (F_SAAS['cov'][NK, NK] + KB_SE ** 2))
SAAS_TOP_ONLY = dict(gm=[r4(v) for v in list(F_SAAS['est'][:NK - 1]) + [SAAS_EST[NK - 1]]],
                     se=[r4(v) for v in list(np.sqrt(np.diag(F_SAAS['cov']))[:NK - 1]) + [np.sqrt(SAAS_COV[NK - 1, NK - 1])]])
# the cross-size correlation of the BACH median term (round 3 ships MEDIAN_RHO; round 2 used 1)
SAAS_RHO_ALT = {}
_mt = np.sqrt(np.clip(np.diag(F_SAAS['cov_terms']['median']), 0, None))
_grp = np.array([medium_weight(x) > 0 for x in F_SAAS['xs']])
for _rho in (1.0, MEDIAN_RHO, 0.0):
    _R = np.where(_grp[:, None] == _grp[None, :], 1.0, _rho)
    _e, _c, _ = blup(F_SAAS['est'], F_SAAS['cov'] - F_SAAS['cov_terms']['median'] + np.outer(_mt, _mt) * _R, NK, KB_LEVEL_AT_X, KB_SE ** 2)
    SAAS_RHO_ALT['rho_%.1f' % _rho] = dict(gm=[r4(v) for v in _e[:NK]], se=[r4(v) for v in np.sqrt(np.diag(_c))[:NK]])
assert SAAS_RHO_ALT['rho_%.1f' % MEDIAN_RHO]['gm'] == [r4(v) for v in SAAS_EST[:NK]], 'the shipped saas update is the rho reading'
# auto (round 3): the bridge's US half and Ratchet+Wrench both counted direct pay without its payroll costs; after the
# payroll cut both are on the shipped basis (the survey moved by the full US-return cut at its size), and they are
# combined by inverse variance. The cut's doubt (half the cut) is shared, so it enters as a correlated error.
F_AUTO = fit_industry('auto', xs=KNOTS + [RW_X])
RW_MIX, RW_MIX_X = mixture_offset('auto', RW['rev_bins'], RW['rev_open'][0], RW['rev_open'][1], RW['revenue_year'], RW_SD, RW_X)
RW_LEVEL = RW_MED - RW_MIX - RW_DEF_GAP                     # the survey median moved to the bridge's point, on the shipped basis
RW_PAY_SE = PAYROLL_SE * RW_DEF_GAP
RW_VAR = RW_SE ** 2 + RW_SE_POPULATION ** 2 + RW_PAY_SE ** 2
RW_C = F_AUTO['payroll_vec'] * RW_PAY_SE                     # the shared payroll-cost doubt: perfectly correlated
AUTO_EST, AUTO_COV, AUTO_GAIN = blup(F_AUTO['est'], F_AUTO['cov'], NK, RW_LEVEL, RW_VAR, RW_C)
AUTO_W = float((F_AUTO['cov'][NK, NK] - RW_C[NK]) / (F_AUTO['cov'][NK, NK] + RW_VAR - 2 * RW_C[NK]))
AUTO_ALT = {}
for tag, S, v, cc in [('survey_read_as_payroll_loaded', RW_MED - RW_MIX, RW_SE ** 2 + RW_SE_POPULATION ** 2, None),
                      ('payroll_doubt_independent', RW_LEVEL, RW_VAR, None),
                      ('no_population_term', RW_LEVEL, RW_SE ** 2 + RW_PAY_SE ** 2, RW_C),
                      ('bridge_only', None, None, None)]:
    if S is None:
        e2, c2, wt = F_AUTO['est'], F_AUTO['cov'], 0.0
    else:
        e2, c2, _ = blup(F_AUTO['est'], F_AUTO['cov'], NK, S, v, cc)
        c0 = 0.0 if cc is None else float(cc[NK])
        wt = (F_AUTO['cov'][NK, NK] - c0) / (F_AUTO['cov'][NK, NK] + v - 2 * c0)
    AUTO_ALT[tag] = dict(survey_value=None if S is None else r4(S), survey_se=None if v is None else r4(np.sqrt(v)), weight=r3(wt),
                         gm=[r3(x) for x in e2[:NK]], se=[r3(x) for x in np.sqrt(np.diag(c2))[:NK]])
FINAL = {}
for ind in BRIDGED:
    if ind == 'saas':
        est, cov = SAAS_EST[:NK], SAAS_COV[:NK, :NK]
    elif ind == 'auto':
        est, cov = AUTO_EST[:NK], AUTO_COV[:NK, :NK]
    else:
        est, cov = FITS[ind]['est'][:NK], FITS[ind]['cov'][:NK, :NK]
    slope_var = float(CW @ cov @ CW)
    FINAL[ind] = dict(est=np.asarray(est, float), se=np.sqrt(np.diag(cov)), cov=np.asarray(cov, float), slope=float(CW @ est), slope_se=float(np.sqrt(slope_var)))
_ef_lvl = float(np.sqrt(EF_SE_SAMPLING ** 2 + EF_SE_MIX ** 2 + EF_SE_TIME ** 2 + EF_SE_DEFN ** 2))
_ef_beyond = np.array([max(0.0, k - EF_X90) * EF_SLOPE_SE for k in KNOTS])
FINAL['ecom'] = dict(est=np.full(NK, EF_MED), se=np.array([ef_se(k) for k in KNOTS]), cov=_ef_lvl ** 2 * np.ones((NK, NK)) + np.outer(_ef_beyond, _ef_beyond),
                     slope=0.0, slope_se=EF_SLOPE_SE)
SHIP = {ind: dict(gm=[r3(v) for v in FINAL[ind]['est']], se=[r3(v) for v in FINAL[ind]['se']], slope=r3(FINAL[ind]['slope']),
                  slope_se=r3(FINAL[ind]['slope_se'])) for ind in IND_ORDER if ind in FINAL}
for ind in SHIPPED:
    assert all(0.05 < v < 0.98 for v in SHIP[ind]['gm']), (ind, SHIP[ind])
for ind in BRIDGED:
    for c, ms in FITS[ind]['modes'].items():
        assert all(m != 'neither_extrapolated' for m in ms), (ind, c, ms)
        assert all(m != 'us_only_extrapolated' or c in CA_DROPPED for m in ms), (ind, c, ms)
        assert ms[1] in ('both', 'us_only') or c in US_STARTS_ABOVE_1M, (ind, c, ms)


# ================================================================ 9. checks
# ---------------------------------------------------------------- 9a. Census 2022 direct-cost decomposition, against the decided wording
def read_ec(zpath, member):
    with zipfile.ZipFile(zpath) as z:
        d = pd.read_csv(io.BytesIO(z.read(member)), sep='|', dtype=str)
    d.columns = [c.lstrip('#') for c in d.columns]
    return d[d.GEOTYPE == '01']


ec_c = read_ec(EC_CONSTR, 'EC2223BASIC.dat')
ec_m = read_ec(EC_MFG, 'EC2231BASIC.dat')
CENSUS = {}
for code, d, kind in [('238220', ec_c, 'c'), ('238990', ec_c, 'c'), ('31-33', ec_m, 'm'), ('339950', ec_m, 'm'), ('332710', ec_m, 'm')]:
    r = d[d.NAICS2022 == code]
    assert len(r) == 1, (code, len(r))
    r = r.iloc[0]
    f = lambda c: float(r[c])
    if kind == 'c':
        direct_mat, wages = f('CSTMPRT') + f('CSTEFT') + f('CSTSCNT'), f('PAYANCW')
    else:
        direct_mat, wages = f('CSTMTOT'), f('PAYANPW')
    ben = f('BENEFIT') * wages / f('PAYANN')
    rcp = f('RCPTOT')
    rpf = rcp / f('FIRM') * 1000.0
    CENSUS[code] = dict(firms=int(f('FIRM')), receipts_per_firm_2022usd=round(rpf), x_real=float(np.log10(rpf * CPI_REF / CPI_ANN[2022])),
                        materials_subcontracts=direct_mat / rcp, direct_wages=wages / rcp, direct_benefits=ben / rcp,
                        gm_wages_only=1 - (direct_mat + wages) / rcp, gm_with_benefits=1 - (direct_mat + wages + ben) / rcp)
census_rows = []
CENSUS_ROLE = {'238220': ('home', 'shipped', 'home (shipped curve)'), '238990': ('constr', 'shipped', 'constr (shipped curve)'),
               '31-33': ('mfg', 'shipped', 'mfg (shipped curve: all-manufacturing cells, page rule beyond $10M)'),
               '339950': ('mfg_misc', 'component', 'misc-manufacturing component (SOI 339900 / ISED 33995, OEWS 3399 direct share), not the shipped mfg curve'),
               '332710': ('mfg', 'context', 'mfg (shipped curve) against machine shops, a small-firm manufacturing industry (context)')}
for code, (ind, role, label) in CENSUS_ROLE.items():
    c = CENSUS[code]
    x = c['x_real']
    if role != 'component':
        shipped = page_value(KNOTS, FINAL[ind]['est'], x)
        own = float(ind_curve(ind, 'conv', knots=[x])['est'][0])                  # the bridge's own curve at x (no page clamp)
        r1 = float(ind_curve(ind, 'conv', knots=[x], owner='none', payroll=False, materials=False,
                             d_override=OEWS['mfg_misc']['direct_wage_share'] if ind == 'mfg' else None)['est'][0])   # round 1's estimator
        asrep = float(ind_curve(ind, 'rep', knots=[x])['est'][0])
    else:
        shipped = float(comp_curve(ind, 'dir', knots=[x])['est'][0])            # same estimator on the 3399/33995 cells
        own = shipped
        r1 = float(comp_curve(ind, 'dir', knots=[x], owner='none')['est'][0])
        asrep = float(comp_curve(ind, 'rep', knots=[x])['est'][0])
    rec = dict(naics=code, tested=ind, role=role, label=label, census_x=round(x, 3), census_with_payroll_costs=r4(c['gm_with_benefits']),
               census_wages_only=r4(c['gm_wages_only']), census_benefit_share=r4(c['direct_benefits']),
               fitted=r4(shipped), fitted_curve_at_x=r4(own), fitted_round1=r4(r1), fitted_as_reported=r4(asrep),
               diff_with_payroll_costs=r4(shipped - c['gm_with_benefits']), diff_wages_only=r4(shipped - c['gm_wages_only']),
               pass_with_payroll_costs=bool(abs(shipped - c['gm_with_benefits']) <= TOL_CENSUS),
               pass_wages_only=bool(abs(shipped - c['gm_wages_only']) <= TOL_CENSUS))
    if code == '339950':
        rec['shipped_mfg_at_this_size'] = r4(page_value(KNOTS, FINAL['mfg']['est'], x))
        rec['shipped_mfg_diff_with_payroll_costs'] = r4(page_value(KNOTS, FINAL['mfg']['est'], x) - c['gm_with_benefits'])
    census_rows.append(rec)
_cs = [r for r in census_rows if r['role'] == 'shipped']
CENSUS_CHECK = dict(rows=census_rows, tolerance=TOL_CENSUS, wording='the decided default: wages and payroll costs (benefits, payroll taxes)',
                    pass_with_payroll_costs=sum(r['pass_with_payroll_costs'] for r in _cs), pass_wages_only=sum(r['pass_wages_only'] for r in _cs),
                    n=len(_cs), shipped_tests=[r['naics'] for r in _cs],
                    component_339950=dict(pass_with_payroll_costs=next(r['pass_with_payroll_costs'] for r in census_rows if r['naics'] == '339950')),
                    context_332710=dict(pass_with_payroll_costs=next(r['pass_with_payroll_costs'] for r in census_rows if r['naics'] == '332710')),
                    note='Census 2022 Economic Census, all employer firms, ratio of sums at the Census receipts per firm; the '
                         'fitted value is the shipped curve read by the page rule (flat beyond the $10M knot; no median shift is '
                         'applied, so it is also a ratio of sums). 339950 tests the misc-manufacturing component with the same '
                         'estimator; the shipped mfg curve (all manufacturing) is tested against 31-33.')

# ---------------------------------------------------------------- 9a2. the decided wording and the loaded wages: payroll costs (applied in round 3)
# The bridge loads the direct staff's wages as the returns report them: SOI salaries and wages (no benefits or payroll
# taxes) and ISED labour and commissions (salaries and wages; the cost-of-sales wages already carry benefits, ISED
# glossary). For trades and manufacturing the Census test above shows the result already sits at the with-payroll-costs
# margin (tax COGS carries job costs), so nothing is loaded there. For the services and auto the factor is loaded on the
# US direct pay and on the Canadian direct pay drawn from labour and commissions; half the cut is an SE term.
PAYROLL_SENS = []
for ind in BRIDGED:
    if ind_defn(ind, 'conv') != 'dir':
        continue
    W_ = sum(w for _, w in INDUSTRIES[ind])
    us_eff, betas, pay = np.zeros(NK), {}, {}
    for c, w in INDUSTRIES[ind]:
        betas[c] = PCF[c]
        Cp = direct_pay_curve(c)
        share = np.array([max(0.0, Cp(min(max(k, Cp.lo - TAU), Cp.hi + TAU))) for k in KNOTS])
        pay[c] = [r4(v) for v in share]
        us_eff += w / W_ * (-betas[c] * share)
    applied = ind in PAYROLL_IND
    on, off = ind_curve(ind, payroll=True if applied else 'all'), ind_curve(ind, payroll=False)
    cut = np.asarray(off['est']) - np.asarray(on['est'])
    cut_us = np.asarray(off['us']) - np.asarray(on['us'])
    cut_ca = np.asarray(off['ca']) - np.asarray(on['ca'])
    ca_extra = {}
    for c, w in INDUSTRIES[ind]:
        code = COMP[c]['ca']
        if code is None or (c in CA_DROPPED):
            continue
        r = ISED_REV[code]
        ca_extra[c] = [r4(max(0.0, delta(c) * (r['dw_' + q] + r['il_' + q] - ca_owner_pay(c, q, OWNER_MODE)) - r['dw_' + q]) / r['rev_' + q])
                       for q in ('q1', 'q2', 'q3', 'q4')]
    PAYROLL_SENS.append(dict(industry=ind, applied=applied, tested_by_census=ind in ('home', 'constr', 'mfg'),
                             payroll_cost_factor={c: r4(v) for c, v in betas.items()}, direct_pay_share_us=pay,
                             ca_direct_pay_from_labour_and_commissions=ca_extra,
                             effect_us_returns_only=[r4(v) for v in us_eff],
                             cut=[r4(v) for v in cut], cut_us_half=[r4(v) for v in cut_us], cut_ca_half=[r4(v) for v in cut_ca],
                             se_term=[r4(PAYROLL_SE * v) for v in cut],
                             note='applied' if applied else 'not applied: the Census 2022 test shows payroll costs already in the margin'))
PAYROLL_SENS_JS = {r['industry']: dict(cut_at_1M=r3(page_value(KNOTS, r['cut'], 6.0)), applied=r['applied']) for r in PAYROLL_SENS}


# ---------------------------------------------------------------- 9b. US against CA at matched size (x = 6.0), before and after the owner-pay fix
X_MATCH = 6.0
usca = []
for ind in BRIDGED + NOT_SHIPPED:
    parts = INDUSTRIES[ind]
    if any(COMP[c]['ca'] is None for c, _ in parts):
        continue
    row = dict(industry=ind, ca_dropped=[c for c, _ in parts if c in CA_DROPPED])
    for tag, defn, owner, corr in [('round1', 'conv', 'none', False), ('round2', 'conv', OWNER_MODE, False), ('fixed', 'conv', OWNER_MODE, True),
                                   ('rep', 'rep', 'none', False)]:
        r = ind_curve(ind, defn, knots=[X_MATCH], owner=owner, allow_drop=False, payroll=corr, materials=corr)
        row[tag + '_us'], row[tag + '_ca'] = r4(r['us'][0]), r4(r['ca'][0])
        row[tag + '_diff'] = r4(r['us'][0] - r['ca'][0])
        row[tag + '_modes'] = {c: p['mode'][0] for c, _, p in r['parts']}
    row['widening_round1'] = r4(row['round1_diff'] - row['rep_diff'])
    row['widening_round2'] = r4(row['round2_diff'] - row['rep_diff'])
    row['widening_fixed'] = r4(row['fixed_diff'] - row['rep_diff'])
    usca.append(row)
U_KEEP = [r for r in usca if r['industry'] in BRIDGED and not r['ca_dropped']]
U_ALL = [r for r in usca if r['industry'] in BRIDGED]


def _usca_summary(rows, tag):
    g = [r[tag + '_diff'] for r in rows]
    return dict(n=len(rows), within_004=int(sum(abs(v) <= .04 for v in g)), within_007=int(sum(abs(v) <= .07 for v in g)),
                median_abs=r4(np.median(np.abs(g))), mean_us_minus_ca=r4(np.mean(g)))


USCA = dict(x=X_MATCH, revenue_usd=1e6, rows=usca,
            fixed_kept=_usca_summary(U_KEEP, 'fixed'), round2_same=_usca_summary(U_KEEP, 'round2'), round1_same=_usca_summary(U_KEEP, 'round1'),
            as_reported_same=_usca_summary(U_KEEP, 'rep'),
            round1_all=_usca_summary(U_ALL, 'round1'), fixed_all_incl_dropped=_usca_summary(U_ALL, 'fixed'),
            widening_round1_mean=r4(np.mean([r['widening_round1'] for r in U_KEEP])), widening_round2_mean=r4(np.mean([r['widening_round2'] for r in U_KEEP])),
            widening_fixed_mean=r4(np.mean([r['widening_fixed'] for r in U_KEEP])),
            note='fixed: the shipped definition (Canadian owner-managers\' pay removed, SOI officers\' share x incorporated share; payroll '
                 'costs loaded on direct pay for the services and auto, on the Canadian side only on the pay drawn from labour and '
                 'commissions; medical materials). round2: without the payroll and materials corrections. CA kept for the comparison '
                 'even where the shipped bridge drops it (ca_dropped). widening = (gap on the definition) - (gap as reported)')

# ---------------------------------------------------------------- 6d. DealLedger stated margins (selected)
DL_IND = {
    'rest': r'restaurant|pizza|cafe\b|coffee shop|bakery|\bbar\b|\bpub\b|tavern|grill\b|diner|sushi|taco|food truck|catering|brewery|ice cream|yogurt|\bdeli\b|bistro',
    'home': r'hvac|plumb|electrical contractor|electrician|heating|air condition|roofing|pest control|garage door|handyman|home service|septic|water heater',
    'medical': r'dental|dentist|medical practice|clinic\b|physician|chiropract|physical therap|optometr|med ?spa|orthodont|urgent care',
    'vet': r'veterinar|animal hospital|vet clinic|pet clinic',
    'msp': r'managed service|it services|it support|\bmsp\b|cybersecurity services|network services',
    'saas': r'\bsaas\b|software|subscription platform|\bapp\b',
    'prof': r'accounting|\bcpa\b|bookkeeping|tax practice|law firm|consulting firm|engineering firm|architect',
    'mfg': r'manufactur|machine shop|fabricat|\bcnc\b|plastics|printing company',
    'dist': r'distribut|wholesale',
    'ecom': r'e-?commerce|amazon|shopify|online store|\bfba\b|dtc brand|direct-to-consumer',
    'staff': r'staffing|recruit|employment agency|temp agency',
    'land': r'landscap|lawn|tree service|irrigation|snow removal',
    'wash': r'car wash|carwash',
    'auto': r'auto repair|automotive repair|collision|body shop|\btire|oil change|transmission|mechanic',
    'constr': r'construction|contractor|concrete|excavat|paving|drywall|framing|remodel',
    'fit': r'fitness|\bgym\b|yoga|pilates|martial arts|crossfit|dance studio',
}


def dealledger():
    snaps = sorted(glob.glob(os.path.join(DL_SNAPS, '*', 'listings.csv')))
    keep = {'id', 'title', 'description', 'asking_price', 'cash_flow', 'revenue', 'vertical', 'broker_domain'}
    hsh, nbytes, fr = hashlib.sha256(), 0, []
    for f in snaps:
        with open(f, 'rb') as fh:
            b = fh.read()
        hsh.update(b); nbytes += len(b)
        d = pd.read_csv(io.BytesIO(b), low_memory=False, dtype=str, usecols=lambda c: c in keep)
        d['snap'] = os.path.basename(os.path.dirname(f))
        fr.append(d)
    A = pd.concat(fr, ignore_index=True)
    for c in ['asking_price', 'cash_flow', 'revenue']:
        A[c] = pd.to_numeric(A[c], errors='coerce')
    A['dlen'] = A['description'].fillna('').str.len()
    A = A.sort_values(['id', 'dlen', 'snap'], kind='mergesort')
    g = A.groupby('id')
    P = g.last()                                         # last non-null per column: longest description, latest snapshot
    P['n_snaps'] = g['snap'].count()
    P = P.reset_index()
    n_ids = len(P)
    P['text'] = (P['title'].fillna('') + ' . ' + P['description'].fillna('')).str.lower().str.replace(r'\s+', ' ', regex=True)
    P['tkey'] = P['title'].fillna('').str.lower().str.replace(r'[^a-z0-9]', '', regex=True).str[:60]
    P = P.sort_values(['n_snaps', 'id'], ascending=[False, True], kind='mergesort').drop_duplicates(subset=['tkey', 'asking_price', 'cash_flow'])
    P['dkey'] = P['description'].fillna('').str.lower().str.replace(r'[^a-z0-9]', '', regex=True).str[:300]
    P = P[(P.dkey == '') | ~P.duplicated(subset=['dkey'])]
    NUM = r'(\d{1,2}(?:\.\d{1,2})?)\s?%'
    Q = r'(?:of|at|is|are|was|were|near|around|approximately|approx\.?|about|roughly|averaging|average|averaged|exceeding|over|above|in the|consistently|typically|between|from|~|:|-|–|—|\s)*'
    pat_after = re.compile(r'gross (?:profit )?margins?\s*(?:\(gm\))?\s*' + Q + r'\s*' + NUM)
    pat_before = re.compile(NUM + r'\s*(?:\+\s*)?(?:average |avg\.? |blended |overall |plus |\+ )?gross (?:profit )?margins?')
    pat_gp_pct = re.compile(r'gross profit\s*(?:%|percent(?:age)?)\s*' + Q + r'\s*' + NUM)
    pat_cogs = re.compile(r'(?:\bcogs\b|cost of goods(?: sold)?|cost of sales)\s*' + Q + r'\s*' + NUM + r'(?!\s*(?:gross|margin))')
    pat_food = re.compile(r'(?:food (?:and|&) (?:beverage|bev) costs?|food costs?|cost of food)\s*' + Q + r'\s*' + NUM)
    pat_food2 = re.compile(NUM + r'\s*food costs?')
    pat_gp_amt = re.compile(r'gross profit\s*(?:of|was|is|:|-|–)?\s*(?:approximately |about |~)?\$\s?([\d,]+(?:\.\d+)?)\s*(k|m|million|mm)?')

    def money(x, suf):
        v = float(x.replace(',', ''))
        return v * (1e3 if suf == 'k' else 1e6 if suf in ('m', 'million', 'mm') else 1)

    def cls(t):
        for k, p in DL_IND.items():
            if re.search(p, t[:200]):
                return k
        for k, p in DL_IND.items():
            if re.search(p, t):
                return k
        return 'other'
    out = []
    for r in P.itertuples(index=False):
        t, rec = r.text, {}
        v = [float(m.group(1)) for m in pat_after.finditer(t)] + [float(m.group(1)) for m in pat_before.finditer(t)] + \
            [float(m.group(1)) for m in pat_gp_pct.finditer(t)]
        v = [x for x in v if 5 <= x <= 97]
        if v:
            rec['gm_stated'] = float(np.median(v)) / 100
        c = [x for x in (float(m.group(1)) for m in pat_cogs.finditer(t)) if 3 <= x <= 95]
        if c:
            rec['gm_from_cogs'] = 1 - float(np.median(c)) / 100
        fo = [x for x in [float(m.group(1)) for m in pat_food.finditer(t)] + [float(m.group(1)) for m in pat_food2.finditer(t)] if 10 <= x <= 60]
        if fo:
            rec['gm_food_only'] = 1 - float(np.median(fo)) / 100
        if pd.notna(r.revenue) and r.revenue > 0:
            a = [money(m.group(1), (m.group(2) or '').strip()) for m in pat_gp_amt.finditer(t)]
            a = [x for x in a if 0.03 * r.revenue <= x <= 0.97 * r.revenue]
            if a:
                rec['gm_from_gp_amt'] = float(np.median(a)) / r.revenue
        if rec:
            rec.update(id=r.id, ind=cls(t), rev=r.revenue)
            out.append(rec)
    D = pd.DataFrame(out)
    D['gm'] = D['gm_stated'].fillna(D['gm_from_cogs']).fillna(D['gm_from_gp_amt'])
    rng = np.random.default_rng(20261014)
    rows = []
    for ind in IND_ORDER:
        for col, lab in [('gm', 'gross margin stated (or 1 - COGS share, or gross profit / revenue)'), ('gm_food_only', '1 - food cost share')]:
            x = D[(D.ind == ind) & D[col].notna()]
            if len(x) < 5 or (col == 'gm_food_only' and ind != 'rest'):
                continue
            med = float(np.median(x[col]))
            bs = [np.median(rng.choice(x[col].values, len(x))) for _ in range(2000)]
            q1, q3 = np.percentile(x[col], [25, 75])
            rec = dict(industry=ind, measure=lab, n=int(len(x)), median=r3(med), se_median=r3(np.std(bs)), iqr=[r3(q1), r3(q3)],
                       median_revenue=None if x.rev.notna().sum() == 0 else float(np.median(x.rev.dropna())))
            if ind in SHIP:
                xs = np.log10(x.rev.where(x.rev > 0).fillna(1e6).clip(lower=1e4))   # listings without revenue: read at $1M
                fitted = np.array([page_value(KNOTS, SHIP[ind]['gm'], v) for v in xs])
                rec['shipped_at_listing_revenue'] = r3(np.median(fitted))
                rec['diff_stated_minus_shipped'] = r3(med - np.median(fitted))
                rec['n_without_revenue'] = int(x.rev.isna().sum() + (x.rev <= 0).sum())
            rows.append(rec)
    D.round(4).to_csv(os.path.join(OUT, 'dealledger_stated_margins.csv'), index=False)
    return dict(rows=rows, n_listing_ids=n_ids, n_after_dedupe=int(len(P)), n_with_figure=int(D.gm.notna().sum()),
                n_food_cost=int(D.gm_food_only.notna().sum()), snapshots=len(snaps), bytes=nbytes, sha256=hsh.hexdigest()[:16],
                note='selected: margins are stated when favourable; a check on definitions, not a level')


DLC = dealledger()

# ---------------------------------------------------------------- 9d. definition test against the 16 hand-set values (context only)
deftest, deftab = {}, []
CURVES = {}
for ind in IND_ORDER:
    for defn in ['rep', 'dir', 'allw', 'exl']:
        CURVES[(ind, defn)] = ind_curve(ind, defn)['est']
    CURVES[(ind, 'conv')] = FINAL[ind]['est'] if ind in FINAL else ind_curve(ind, 'conv')['est']
for tag, x in REF_X.items():
    for defn in ['rep', 'dir', 'allw', 'exl', 'conv']:
        v = np.array([page_value(KNOTS, CURVES[(ind, defn)], x) for ind in IND_ORDER])
        h = np.array([HAND[ind] for ind in IND_ORDER])
        deftest[f'{defn}_{tag}'] = dict(definition=DEF_NAMES[defn], at=tag, n=len(v), mad=r3(np.mean(np.abs(h - v))),
                                        rmse=r3(np.sqrt(np.mean((h - v) ** 2))), r=r3(np.corrcoef(h, v)[0, 1]), bias_hand_minus_data=r3(np.mean(h - v)))
        for ind, vv in zip(IND_ORDER, v):
            deftab.append(dict(industry=ind, at=tag, definition=defn, hand_set=HAND[ind], data=r4(vv)))

# ---------------------------------------------------------------- 9e. spread: BACH bracket and the surveys' within-industry SDs against the .13 z-scale (kept)
spread_rows = []
for ind in IND_ORDER:
    rec = dict(industry=ind, bach_division=BACH_SEC[ind], convention=BACH_CONV[ind])
    for size in ('1a', '1b'):
        for smp in ('0', '-1'):
            b = BACH_RES.get((BACH_SEC[ind], size, smp))
            if b is None:
                continue
            sd = b['mat_sd'] if BACH_CONV[ind] == 'mat' else 0.5 * (b['gva_sd'] + b['gop_sd'])
            rec[f'sd_{size}_s{smp}'] = r4(sd)
    pm = ISED_PM.get({'medical': '62111', 'prof': '54121', 'saas': '51821', 'ecom': '44-45'}.get(ind, COMP[ind]['ca'] if ind in COMP else None))
    if pm:
        g4 = np.array([pm[f'gm_q{i}'] for i in range(1, 5)]) / 100
        rec['ca_between_profit_quartile_sd'] = r4(np.sqrt(np.mean((g4 - g4.mean()) ** 2)))
    spread_rows.append(rec)


def med_col(c, inds=IND_ORDER):
    v = [r[c] for r in spread_rows if r['industry'] in inds and r.get(c) is not None]
    return dict(median=r3(np.median(v)), min=r3(min(v)), max=r3(max(v)), n=len(v))


DL_SAAS = next((r for r in DLC['rows'] if r['industry'] == 'saas' and r['measure'].startswith('gross')), None)
SURVEY_SD = {'auto_ratchetwrench': dict(sd=r3(RW_SD), n=RW_N, how='normal fitted to the 5 overall-GP bins (interval-censored ML)'),
             'saas_kbcm_subscription': dict(sd=r3(KB_SD), n=KBCM['sub_n'], how='normal fitted to the 9 subscription-GM bins (total GM has no histogram)'),
             'ecom_empireflippers': dict(sd=r3(EF_SD_IQR), sd_raw=r3(EF_SD_RAW), n=EF_N, how='IQR/1.349 of 1 - cogs_percent, sold listings'),
             'saas_dealledger': None if DL_SAAS is None else dict(sd=r3((DL_SAAS['iqr'][1] - DL_SAAS['iqr'][0]) / 1.349), n=DL_SAAS['n'],
                                                                  how='IQR/1.349 of stated margins (selected: stated when favourable)')}
SPREAD = dict(small_s0=med_col('sd_1a_s0'), medium_s0=med_col('sd_1b_s0'), small_research_sample=med_col('sd_1a_s-1'),
              medium_research_sample=med_col('sd_1b_s-1'), ca_profit_quartile_floor=med_col('ca_between_profit_quartile_sd'),
              rows=spread_rows, bach_rows_scanned=BACH_ROWS, surveys=SURVEY_SD, z_scale=Z_GM)
SPREAD['bracket'] = [SPREAD['medium_s0']['median'], SPREAD['small_s0']['median']]
SPREAD['z_scale_inside_bracket'] = bool(SPREAD['bracket'][0] <= Z_GM <= SPREAD['bracket'][1])
_sd = [v['sd'] for v in SURVEY_SD.values() if v is not None]
SPREAD['survey_sd_range'] = [min(_sd), max(_sd)]
SPREAD['survey_sd_median'] = r3(np.median(_sd))

# ---------------------------------------------------------------- 9f. replication of the research round's convention column (its estimator, no owner fix)
RESEARCH_MAP = {'medical': 'phys', 'prof': 'acct'}


def research_value(ind, R):
    comp = RESEARCH_MAP.get(ind, ind)
    d_ = DEFN.get(ind, 'dir')
    dd = OEWS['mfg_misc']['direct_wage_share'] if comp == 'mfg' else delta(comp)      # the research round used the 3399 share for mfg
    us = us_curve(comp, d_, float(np.clip(dd, 0, 1)), LAM[COMP[comp]['sec']], xcol='x_nom', kind='line', min_years=1)
    v_us = us(np.log10(R))
    C = ca_curve(comp, d_, dd, xcol='x_nom', quartiles=('q3', 'q4'), kind='line', owner='none')
    return v_us if C is None else (v_us + C(np.log10(R))) / 2


rep_rows = []
research_tab = pd.read_csv(RESEARCH_TABLE).set_index('key') if os.path.exists(RESEARCH_TABLE) else None
for ind in IND_ORDER:
    rec = dict(industry=ind)
    for tag, R in [('1M', 1e6), ('3M', 3e6)]:
        rec[f'replicated_{tag}'] = r4(research_value(ind, R))
        if research_tab is not None:
            rec[f'research_{tag}'] = float(research_tab.loc[ind, f'conv_{tag}'])
            rec[f'diff_{tag}'] = r4(rec[f'replicated_{tag}'] - rec[f'research_{tag}'])
        if ind in FINAL:
            rec[f'shipped_{tag}'] = r4(page_value(KNOTS, FINAL[ind]['est'], REF_X[tag]))
    rep_rows.append(rec)
REPLICATION = dict(rows=rep_rows, max_abs_diff=r4(max(max(abs(r['diff_1M']), abs(r['diff_3M'])) for r in rep_rows)) if research_tab is not None else None,
                   note='the research estimator (global weighted line through the SOI classes, CA Q3-Q4 line, nominal 2004-13 USD and 2024 CAD; '
                        'medical = physicians, prof = accounting; no owner-pay fix) reproduces its final_table to rounding; shipped_* is this fit')
NOMINAL = {ind: [r4(v) for v in ind_curve(ind, 'conv', xcol='x_nom')['est']] for ind in BRIDGED}


# ---------------------------------------------------------------- 9g. recency: SOI Table 5.1 2019-22 vs Source Book 2004-13 (all sizes): as reported and labour-loaded
recency = []
for comp, (code, pref) in RECENCY_CELLS.items():
    dr = soi_drift(comp)
    recency.append(dict(component=comp, soi_code=code, gm_2004_13=r4(dr['rep_old']), gm_2019_22=r4(dr['rep_new']), years_new=dr['years_new'],
                        drift=r4(dr['rep_new'] - dr['rep_old']), dir_2004_13=r4(dr['dir_old']), dir_2019_22=r4(dr['dir_new']),
                        drift_dir=r4(dr['dir_new'] - dr['dir_old']),
                        small_class_trend=None if comp not in RECENCY_TREND else RECENCY_TREND[comp]))


# ---------------------------------------------------------------- 9h. Census AWTS 4233 / ARTS 4541 (all employer firms, all sizes)
def census_table(path, code, label_col, label_text, val_from):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    hdr, out = None, None
    for r in ws.iter_rows(values_only=True):
        if r and r[0] in ('2017 NAICS Code', 'NAICS Code'):
            hdr = [str(v).rstrip('r') if v is not None else None for v in r]
        if r and r[0] == code and (label_text is None or (r[label_col] and str(r[label_col]).strip().startswith(label_text))):
            out = {int(h): float(v) for h, v in zip(hdr[val_from:], r[val_from:]) if h and h.isdigit() and isinstance(v, (int, float))}
    return out


awts = census_table(AWTS, 4233, 1, 'Gross margins as a percent of sales', 4)
arts = census_table(ARTS, 4541, 1, None, 2)
CENSUS_ALLFIRM = dict(awts_4233_2018_22=r4(np.mean([awts[y] for y in range(2018, 2023)]) / 100),
                      arts_4541_2018_22=r4(np.mean([arts[y] for y in range(2018, 2023)]) / 100),
                      dist_shipped_at_10M=SHIP['dist']['gm'][-1], ecom_shipped=SHIP['ecom']['gm'][0],
                      ecom_soi_nonstore_as_reported=[r3(v) for v in ind_curve('ecom', 'rep')['est']],
                      note='ARTS 4541 is all employer electronic shopping and mail-order firms, revenue-weighted, on the Census cost of goods '
                           'sold; it is not the console\'s population or definition (context only)')


# ---------------------------------------------------------------- 9i. Canada above $5M: the incorporated $5-20M report against the carried Canadian curve
def fpd_as_rev(rec):
    out = {}
    for k in ('rev', 'cos', 'dw', 'il'):
        for q in (1, 2, 3, 4):
            out[f'{k}_q{q}'] = rec[k][q]
    return out


ca_big = []
for comp in COMP:
    code = COMP[comp]['ca']
    if code is None or comp in ('ecom',):
        continue
    ind = {'phys': 'medical', 'dent': 'medical', 'acct': 'prof', 'legal': 'prof', 'mfg_misc': 'mfg'}.get(comp, comp)
    defn = ind_defn(ind, 'conv')
    d = delta(comp)
    B = FPD[(code, 2, 1)]
    rb = fpd_as_rev(B)
    xs = [float(np.log10(B['rev'][q] * 1000 * CA_TO_REAL)) for q in (1, 2, 3, 4)]
    act = []
    pc_, mat_ = adj(comp, defn)
    for qi, q in enumerate(('q1', 'q2', 'q3', 'q4')):
        op = oc_share(comp, xs[qi]) * B['rev'][qi + 1] if defn in ('dir', 'allw') else 0.0   # all incorporated
        act.append(margin_ca(defn, rb, q, d, min(op, rb['dw_' + q] + rb['il_' + q]), pc_, mat_))
    cc = comp_curve(comp, defn, knots=xs, allow_drop=False)
    carried = list(cc['ca'])
    diff = np.array(act) - np.array(carried)
    k65, k70 = (float(np.interp(k, xs, diff, left=diff[0], right=diff[-1])) for k in (6.5, 7.0))
    ca_big.append(dict(component=comp, industry=ind, ca_dropped=comp in CA_DROPPED, n_businesses=int(B['n']),
                       x=[round(v, 3) for v in xs], ca_5_20M=[r4(v) for v in act], ca_carried=[r4(v) for v in carried],
                       us=[r4(v) for v in cc['us']], mean_diff=r4(diff.mean()), knot_effect_if_used={'6.5': r4(k65 / 2), '7.0': r4(k70 / 2)}))
_kept_big = [r for r in ca_big if not r['ca_dropped'] and r['industry'] in BRIDGED]
CA_BIG = dict(rows=ca_big, n=len(_kept_big), mean_diff=r4(np.mean([r['mean_diff'] for r in _kept_big])),
              median_abs_diff=r4(np.median([abs(r['mean_diff']) for r in _kept_big])),
              within_004=int(sum(abs(r['mean_diff']) <= .04 for r in _kept_big)),
              note='ISED/Statistics Canada FPD 2024, incorporated businesses with revenue CAD 5-20M (quartile averages; reliability '
                   'mostly E), same definition and owner-pay removal (all incorporated). knot_effect_if_used = half the difference at '
                   'the knot (the knot is the mean of the two countries). A check: the estimator is unchanged')

# ---------------------------------------------------------------- 9j. dentists: ADA Health Policy Institute, Survey of Dental Practice 2025 (free tables)
ADA = {}
_wb = openpyxl.load_workbook(ADA_XLSX, read_only=True, data_only=True)
for tab in ('Table 5', 'Table 13'):
    rows_ = [r for r in _wb[tab].iter_rows(values_only=True) if r and r[0] is not None]
    gp = next(i for i, r in enumerate(rows_) if str(r[0]).strip() == 'General Practitioners')
    r = rows_[gp + 1]
    assert str(r[0]).strip() == 'All owners', r
    ADA[tab] = dict(average=float(r[1]), q1=float(r[2]), median=float(r[3]), q3=float(r[4]), n=int(r[6]))
_cats = ' '.join(str(c) for ws in _wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None)
ADA_HAS_SHARES = bool(re.search(r'(laborator|dental suppl|staff salar)\w*[^.]{0,80}(percent|%)', _cats, re.I))
ADA_EXP_SHARE_MEAN = ADA['Table 13']['average'] / ADA['Table 5']['average']
ADA_EXP_SHARE_MED = ADA['Table 13']['median'] / ADA['Table 5']['median']
_dent_x = float(np.log10(to_real(ADA['Table 5']['average'], 2025)))
_dent_est = float(comp_curve('dent', 'dir', knots=[_dent_x])['est'][0])
_dent_est_r2 = float(comp_curve('dent', 'dir', knots=[_dent_x], payroll=False, materials=False)['est'][0])
_md = MAT['dent']
ADA_CHECK = dict(expense_shares_published=ADA_HAS_SHARES, gross_billings_per_owner_gp=ADA['Table 5'], expenses_per_owner_gp_excl_owner_pay=ADA['Table 13'],
                 total_expense_share_of_billings_means=r4(ADA_EXP_SHARE_MEAN), total_expense_share_of_billings_medians=r4(ADA_EXP_SHARE_MED),
                 gm_lower_bound=r4(1 - max(ADA_EXP_SHARE_MEAN, ADA_EXP_SHARE_MED)), dent_component_at_mean_billings=r4(_dent_est),
                 dent_component_round2=r4(_dent_est_r2),
                 x=round(_dent_x, 3), passes=bool(_dent_est >= 1 - max(ADA_EXP_SHARE_MEAN, ADA_EXP_SHARE_MED)),
                 sas_2022_materials=dict(census_share_all_sizes=r4(_md['census_share']), supplies=_md['shares']['supplies'],
                                         other_materials=_md['shares']['other_materials'], purchased_services=_md['shares']['purchased_services'],
                                         small_cell_ratio=r4(_md['small_ratio']), share_used=r4(_md['share']),
                                         soi_cogs_small_cell=r4(_md['soi_cogs_small']), bridge_materials_round2=r4((1 - LAM[SEC_HEALTH]) * _md['soi_cogs_small'])),
                 note='the free 2025 tables publish total practice expenses (excluding owner pay) and gross billings per owner '
                      'dentist, not staff, laboratory and supply shares; total expenses include rent, front-office staff and other '
                      'overhead that the shipped margin does not deduct, so 1 - total share is only a lower bound. The direct test of '
                      'the materials is Census SAS 2022 (checks.medical_materials): supplies, other materials and purchased professional '
                      'and technical services are %.1f%% of dentists\' receipts (all employer firms) against the %.1f%% of tax cost of '
                      'goods (less its labour) the round-2 bridge deducted in the small cell' % (100 * _md['census_share'], 100 * (1 - LAM[SEC_HEALTH]) * _md['soi_cogs_small']))

# ---------------------------------------------------------------- 9j2. medical materials (round 3): the correction, its alternative readings and its SE term
mat_rows = []
for c in MAT:
    m_ = MAT[c]
    on, off = comp_curve(c, 'dir'), comp_curve(c, 'dir', materials=False)
    alts = {a: np.asarray(comp_curve(c, 'dir', materials=a)['est']) for a in MAT_ALTS}
    se_ = np.sqrt(np.mean([np.square(alts[a] - on['est']) for a in MAT_ALTS], axis=0))
    mat_rows.append(dict(component=c, naics=m_['naics'], firms_2022=m_['firms'], receipts_2022_musd=round(m_['receipts_musd'], 1), sas_shares=m_['shares'],
                         census_share_all_sizes=r4(m_['census_share']), soi_cogs_all_2004_13=r4(m_['soi_cogs_all']),
                         soi_cogs_small_cell_2004_13=r4(m_['soi_cogs_small']), small_cell_ratio=r4(m_['small_ratio']),
                         share_used=r4(m_['share']), bridge_materials_small_cell_round2=r4((1 - LAM[COMP[c]['sec']]) * m_['soi_cogs_small']),
                         pooled_classes=m_['pooled_classes'], pooled_ratio=r4(m_['pooled_ratio']),
                         soi_cogs_next_class_yearly={k: r4(v) for k, v in MAT_AC4[c].items()},
                         alternative_shares={a: r4(v) for a, v in m_['alternatives'].items()},
                         gm_round2_definition=[r4(v) for v in off['est']], gm=[r4(v) for v in on['est']],
                         correction=[r4(a - b) for a, b in zip(on['est'], off['est'])],
                         alternatives={a: [r4(v - b) for v, b in zip(alts[a], off['est'])] for a in MAT_ALTS},
                         se_term=[r4(v) for v in se_]))
_med_on, _med_off = ind_curve('medical'), ind_curve('medical', materials=False)
MATERIALS = dict(rows=mat_rows, medical_correction=[r4(a - b) for a, b in zip(_med_on['est'], _med_off['est'])],
                 medical_se_term=[r4(v) for v in FITS['medical']['se_terms']['materials'][:NK]],
                 temp_staff_not_deducted={c: MAT[c]['shares']['temp_staff'] for c in MAT},
                 note='Census SAS 2022 Table 5 (employer firms; medical supplies, other materials and supplies, purchased professional '
                      'and technical services) over Economic Census 2022 receipts (the same population: 621 matches SAS Table 1 exactly), '
                      'size-adjusted by the SOI $1-500K asset class\'s cost-of-goods share over the all-size share (2004-13). At every '
                      'SOI cell the materials deducted are the larger of tax COGS less its labour and that share of receipts. SE term: '
                      'RMS over three alternative readings (purchased services left out; the pooled small classes at the knots as the '
                      'size adjustment; no size adjustment). Temporary and leased staff (%s of receipts) are not deducted.'
                      % ', '.join('%s %.1f%%' % (c, 100 * MAT[c]['shares']['temp_staff']) for c in MAT))


# ---------------------------------------------------------------- 9k. effect of the owner-pay fix, by industry and knot
owner_rows = []
for ind in BRIDGED:
    if any(COMP[c]['ca'] is None for c, _ in INDUSTRIES[ind]):
        continue
    v = {tag: ind_curve(ind, owner=o, allow_drop=ad)['est'] for tag, o, ad in
         [('no_owner_fix', 'none', True), ('shipped_bridge', OWNER_MODE, True), ('soi_ca_kept', OWNER_MODE, False), ('ised_reading', 'ised', True),
          ('ised_ca_kept', 'ised', False)]}
    v['round1_estimator'] = ind_curve(ind, owner='none', payroll=False, materials=False,
                                      d_override=OEWS['mfg_misc']['direct_wage_share'] if ind == 'mfg' else None)['est']
    owner_rows.append(dict(industry=ind, ca_dropped=[c for c, _ in INDUSTRIES[ind] if c in CA_DROPPED],
                           **{k: [r4(x) for x in vv] for k, vv in v.items()},
                           effect=[r4(a - b) for a, b in zip(v['shipped_bridge'], v['no_owner_fix'])],
                           note='every column on the round-3 definition (payroll costs, materials) except round1_estimator; effect = the owner-pay fix'))
OWNER = dict(rows=owner_rows, kappa={c: r3(KAPPA[c]) for c in KAPPA}, kappa_raw={c: r3(KAPPA_RAW[c]) for c in KAPPA_RAW},
             kappa_points=KAPPA_ROWS, incorporated_share=PI_INC, owner_share_of_ca_labour={c: {q: r3(v) for q, v in d.items()} for c, d in OWNER_LAB.items()},
             ca_dropped=CA_DROPPED, rule=OWNER_DROP, fpd_matches_round1_pages=bool(max(FPD_CONSIST.values()) < 0.051))

# ---------------------------------------------------------------- 9l. the BACH median-minus-weighted-mean gap, by industry and knot
median_rows = []
for ind in BRIDGED:
    b_s, b_m = BACH_RES[(BACH_SEC[ind], '1a', '0')], BACH_RES[(BACH_SEC[ind], '1b', '0')]
    median_rows.append(dict(industry=ind, division=BACH_SEC[ind], convention=BACH_CONV[ind],
                            gap_by_knot=[r4(median_gap(ind, k)[0]) for k in KNOTS], se_term_by_knot=[r4(median_gap(ind, k)[1]) for k in KNOTS],
                            small=dict(g_mid=r4(b_s['g_mid']), g_matgop=r4(b_s['g_matgop']), g_mat=r4(b_s['g_mat']), pos_share=r3(b_s['g_mid_pos_share' if BACH_CONV[ind] == 'mid' else 'g_mat_pos_share']), n=b_s['g_mid_n']),
                            medium=dict(g_mid=r4(b_m['g_mid']), g_matgop=r4(b_m['g_matgop']), g_mat=r4(b_m['g_mat']), pos_share=r3(b_m['g_mid_pos_share' if BACH_CONV[ind] == 'mid' else 'g_mat_pos_share']), n=b_m['g_mid_n'])))
MEDIAN = dict(rule=MEDIAN_RULE, signs=GAP_SIGNS, rows=median_rows,
              materials_margin_small=r4(np.median([BACH_RES[(dv, '1a', '0')]['g_mat'] for dv, _ in GAP_DIVS])),
              note='median firm minus the ratio of sums (BACH weighted mean), sample 0, 2019-23, median over country-years. '
                   'Convention-matched: materials margin (1 - I5) for distribution, restaurants and software; for the labour-loaded '
                   'margins the midpoint of the value-added and gross-operating-profit gaps, with the materials/GOP midpoint as the '
                   'other bracket (SE = RMS of the two). Knots up to $3.2M use the small class (turnover < EUR 10M); $10M blends in '
                   'the medium class at weight 1/2. Derived statistics only (BACH licence).')

# ---------------------------------------------------------------- 9m. car wash: evidence for keeping the hand-set value
_cw = ec_row('811192')
CW_CENSUS = dict(firms=int(_cw[21]), receipts_k=float(_cw[25]), payroll_k=float(_cw[27]), payroll_share=r4(float(_cw[27]) / float(_cw[25])),
                 receipts_per_firm=round(float(_cw[25]) / float(_cw[21]) * 1000))
PCD_TXT = page_text(PCD_DOC)
PCD = dict(chemicals_share_of_operating_costs=0.13 if re.search(r'carwash chemicals account for 13% of operators. total operating costs', PCD_TXT) else None,
           labour_share_of_operating_costs=0.24 if re.search(r'labor . at 24% . ranked as the highest', PCD_TXT) else None,
           maintenance_share=0.11 if re.search(r'11% of operating costs across wash platforms are for maintenance', PCD_TXT) else None)
assert all(v is not None for v in PCD.values()), PCD
WASH_DEF = 'Revenue less chemicals, water, utilities and supplies (car washes: wages are not in it).'
WASH = dict(definition=WASH_DEF, proxy_bridge_direct_labour=dict(gm=[r3(v) for v in FITS['wash']['est']], se=[r3(v) for v in FITS['wash']['se']]),
            materials_only_at_1M=r3(page_value(KNOTS, ind_curve('wash', 'exl')['est'], 6.0)),
            ca_materials_only_at_1M=r3(page_value(KNOTS, ind_curve('wash', 'exl')['ca'], 6.0)),
            census_2022_car_washes=CW_CENSUS, pcd_2019_survey=PCD,
            verdict='kept at the hand-set value (judgment) on its own definition, definition.wash: revenue less chemicals, water, '
                    'utilities and supplies, wages not in it. No car-wash cell or survey measures that margin. The closest proxy reading '
                    'is materials only: auto repair (US 811115) and other auto repair including car washes (CA 81119) give %.2f at $1M; '
                    'with direct labour loaded (the default definition) they give %.2f. The 2019 PC&D operator survey gives cost shares '
                    'of operating costs, not of revenue (chemicals 13%%, labour 24%%, maintenance 11%%; paid report, figures from PC&D\'s '
                    'own article). Census 2022 car washes: payroll %.1f%% of receipts, all employer firms, with no materials line.'
                    % (page_value(KNOTS, ind_curve('wash', 'exl')['est'], 6.0), page_value(KNOTS, FITS['wash']['est'], 6.0), 100 * CW_CENSUS['payroll_share']))


# ---------------------------------------------------------------- 9n. materiality: a business that enters today's hand-set median, before and after
MAT_JS = r"""
const H=require(%s);const fs=require('fs');const SNAP=fs.readFileSync(H.SNAP,'utf8');const A=H.load();
const spec=JSON.parse(process.argv[process.argv.length-1]);
function biz(API,ind,ern){const I=API.IND[ind],RF=API.sizeRefs(ern),rev=ern/API.typMargin(ind,ern);
 return {ind,ern,rev,ask:3*ern,yrs:6,buyers:I.buyers,franchise:false,st:null,rec:I.rec,cc:.16,own:RF.own,mgt:RF.mgt,
  grw:I.grw,gm:I.gm,ret:I.ret,fin:RF.fin,plt:.40,budget:1e5,lamT:900,theta:.75};}
function ov(ind,f,v){const m=SNAP.match(new RegExp(ind+':\\s*\\{[^}]*\\}'));if(!m)throw new Error('no IND row '+ind);
 const re=new RegExp('(\\b'+f+':)([-.0-9]+)');if(!re.test(m[0]))throw new Error('no field');return [m[0],m[0].replace(re,'$1'+String(v))];}
function pv(k,g,x){if(x<=k[0])return g[0];if(x>=k[k.length-1])return g[g.length-1];
 for(let i=1;i<k.length;i++){if(x<=k[i]){const t=(x-k[i-1])/(k[i]-k[i-1]);return g[i-1]+t*(g[i]-g[i-1]);}}}
const out=[];
for(const g of spec.groups){const s=biz(A,g.ind,spec.ern);const b=A.model(s);const x=Math.log10(s.rev);
 const v=Math.round(pv(spec.knots,g.gm,x)*1000)/1000;const B=H.load([ov(g.ind,'gm',v)]);const a=B.model(s);
 out.push({ind:g.ind,revenue:s.rev,x:x,gm_hand:A.IND[g.ind].gm,gm_new:v,peer_after:B.IND[g.ind].gm,value_before:b.V,value_after:a.V,
  pct:a.V/b.V-1,z_gm_before:b.z.gm,z_gm_after:a.z.gm});}
process.stdout.write(JSON.stringify(out));
""" % json.dumps(HARNESS)
NODE = shutil.which('node')
if NODE and os.path.exists(HARNESS):
    _spec = json.dumps(dict(ern=SDE_REF, knots=KNOTS, groups=[dict(ind=i, gm=SHIP[i]['gm']) for i in IND_ORDER if i in SHIPPED]))
    MAT_ROWS = json.loads(subprocess.run([NODE, '-e', MAT_JS, _spec], capture_output=True, text=True, check=True).stdout)
    for r in MAT_ROWS:
        r.update(revenue=round(r['revenue']), x=r4(r['x']), value_before=round(r['value_before']), value_after=round(r['value_after']),
                 pct=r4(r['pct']), z_gm_before=r3(r['z_gm_before']), z_gm_after=r3(r['z_gm_after']))
    MATERIALITY = dict(sde=SDE_REF, rows=MAT_ROWS, note="model value (harness, snapshot HEAD 56ba230) of a business with $708k SDE that enters "
                       "today's hand-set gm, before and after IND.gm is set to the shipped median at its revenue; nothing else moves",
                       direction_ok=all((r['pct'] < 0) == (r['gm_new'] > r['gm_hand']) for r in MAT_ROWS if abs(r['gm_new'] - r['gm_hand']) > 0.005),
                       peer_set=all(abs(r['peer_after'] - r['gm_new']) < 1e-12 for r in MAT_ROWS))
else:
    MATERIALITY = dict(skipped='node or the harness is not available')


# ================================================================ 10. outputs
DEF_TEXT = {
    'default': ('Revenue less materials, merchandise, subcontractors and the wages and payroll costs (benefits, payroll taxes) of the '
                'non-owner staff who produce or deliver the work; the owner\'s own pay is not a cost here, because SDE adds it back.'),
    'rest': 'Revenue less food, beverage and supplies.',
    'dist': 'Revenue less the cost of the merchandise sold.',
    'ecom': 'Revenue less the landed cost of the merchandise sold (marketplace fees, fulfilment and advertising are not in it).',
    'wash': WASH_DEF,
}
DEF_KEY = {ind: (ind if ind in ('rest', 'dist', 'ecom', 'wash') else 'default') for ind in SHIPPED + NOT_SHIPPED}
TIER = {ind: ('listing-level' if ind == 'ecom' else 'survey-combined' if ind in ('saas', 'auto') else 'aggregate bridge') for ind in SHIPPED}
TIER_TEXT = {'aggregate bridge': 'tax-return aggregates (IRS SOI, Statistics Canada) bridged to the definition (Aggregate)',
             'survey-combined': 'the aggregate bridge combined by inverse variance with a direct industry survey on the same basis (Aggregate)',
             'listing-level': 'the median of listing-level data, Empire Flippers sold listings (Fitted)'}
NOT_ID = {'wash': ('Hand-set (judgment): %.2f on definition.wash (revenue less chemicals, water, utilities and supplies; car washes: wages '
                   'are not in it). No car-wash cell or survey measures that margin. The closest proxy reading, materials only (US auto '
                   'repair, CA other auto repair incl. car washes), gives %.2f at $1M; with direct labour loaded (the default definition) '
                   'the proxies give %.2f (SE %.2f). PC&D\'s 2019 operator survey gives cost shares, not margins (chemicals 13%%, labour '
                   '24%% of operating costs); Census 2022 car-wash payroll is %.1f%% of receipts.')
                  % (HAND['wash'], WASH['materials_only_at_1M'], page_value(KNOTS, FITS['wash']['est'], 6.0),
                     float(np.interp(6.0, KNOTS, FITS['wash']['se'][:NK])), 100 * CW_CENSUS['payroll_share'])}
_saas_bridge_at = float(F_SAAS['est'][NK])
_rw_bridge_mix = float(F_AUTO['est'][NK] + RW_MIX)                       # the bridge's median over the respondents' revenue mix, shipped basis
_rw_z = (RW_MED - RW_DEF_GAP - _rw_bridge_mix) / np.sqrt(F_AUTO['cov'][NK, NK] + RW_VAR - 2 * RW_C[NK])
SURVEYS_JS = {
    'auto': dict(source='Ratchet+Wrench 2026 Industry Survey, p. 9', measure='overall shop gross profit margin (parts, sublet, technician pay)',
                 median=r3(RW_MED), n=RW_N, sd=r3(RW_SD), revenue=int(round(10 ** RW_X, -3)), payroll_cost_adjustment=r3(-RW_DEF_GAP),
                 median_shipped_basis=r3(RW_MED - RW_DEF_GAP), bridge=r3(_rw_bridge_mix), definition_match=True, weight=r3(AUTO_W),
                 verdict='combined by inverse variance at its size: shops count technician pay without its payroll costs and so did '
                         'the US half of the bridge; after the payroll cut both are on the shipped basis (the survey moved by the full '
                         'US-return cut, %.3f), and the doubt about that cut is shared; the survey carries a %.2f judgment term for its '
                         'population, as KBCM does. Before combining it sat %+.1f SE from the bridge'
                         % (RW_DEF_GAP, RW_SE_POPULATION, _rw_z)),
    'saas': dict(source='KBCM 2021 Private SaaS Company Survey, pp. 52, 56-57', measure='total gross margin, support and hosting in COGS',
                 median=KB_LEVEL, n=KB_N, sd=r3(KB_SD), revenue=int(round(10 ** KB_X, -5)), bridge=r3(_saas_bridge_at + KB_MIX),
                 definition_match=True, weight=r3(SAAS_W),
                 verdict='combined by inverse variance at its size ($5-25M of revenue); the lower knots move with the bridge\'s own error '
                         'correlation, with the BACH median term %.1f correlated between them and the top' % MEDIAN_RHO),
    'saas_listings': None if DL_SAAS is None else dict(source='DealLedger broker listings that state a margin', measure='stated gross margin (undefined)',
                                                       median=DL_SAAS['median'], n=DL_SAAS['n'], sd=SURVEY_SD['saas_dealledger']['sd'],
                                                       definition_match=False, verdict='check only: undefined and stated when favourable'),
    'ecom': dict(source='Empire Flippers sold listings, cogs_percent', measure='1 - COGS share (units sold x landed cost)', median=r3(EF_MED),
                 n=EF_N, sd=r3(EF_SD_IQR), revenue=int(round(10 ** float(np.median(ef.x)), -3)), definition_match=True,
                 verdict='sets the level: no bridge on this definition; flat (size slope %.3f, SE %.3f)' % (EF_SLOPE, EF_SLOPE_SE)),
}
census_js = {}
for r in census_rows:
    census_js[r['naics']] = dict(tests=r['label'], with_payroll_costs=r3(r['census_with_payroll_costs']), wages_only=r3(r['census_wages_only']),
                                 fitted=r3(r['fitted']), pass_with_payroll_costs=r['pass_with_payroll_costs'], pass_wages_only=r['pass_wages_only'])
    if r['naics'] == '339950':
        census_js[r['naics']].update(shipped_mfg=r3(r['shipped_mfg_at_this_size']), shipped_mfg_diff_with_payroll_costs=r3(r['shipped_mfg_diff_with_payroll_costs']))
census_js['tolerance'] = TOL_CENSUS
census_js['shipped_pass_with_payroll_costs'] = '%d of %d (%s)' % (CENSUS_CHECK['pass_with_payroll_costs'], CENSUS_CHECK['n'], ', '.join(CENSUS_CHECK['shipped_tests']))
_mm = {r['component']: r for r in mat_rows}
js = {
    'definition': DEF_TEXT,
    'knots_log10_rev': KNOTS,
    'by_industry': {ind: dict(SHIP[ind], tier=TIER[ind]) for ind in IND_ORDER if ind in SHIPPED},
    'tiers': TIER_TEXT,
    'not_identified': NOT_ID,
    'surveys': SURVEYS_JS,
    'spread_check': dict(z_scale=Z_GM, kept=True, bach_small_sd=SPREAD['small_s0']['median'], bach_medium_sd=SPREAD['medium_s0']['median'],
                         bach_small_range=[SPREAD['small_s0']['min'], SPREAD['small_s0']['max']], ca_floor_sd=SPREAD['ca_profit_quartile_floor']['median'],
                         survey_sd={k: v['sd'] for k, v in SURVEY_SD.items() if v is not None}, inside=SPREAD['z_scale_inside_bracket'],
                         note='BACH database (ECCBSO), within NACE division, IQR/1.349, 2019-23 (derived statistics only); surveys: '
                              'Ratchet+Wrench 2026, KBCM 2021 subscription margin, Empire Flippers, DealLedger'),
    'checks': dict(census_2022=census_js,
                   payroll_costs=dict(applied_to=PAYROLL_IND, cut_at_1M={k: v['cut_at_1M'] for k, v in PAYROLL_SENS_JS.items() if v['applied']},
                                      not_applied_census_tested={k: v['cut_at_1M'] for k, v in PAYROLL_SENS_JS.items() if not v['applied']},
                                      se='half the cut', factor={c: r3(PCF[c]) for c in PAYROLL_COMPS},
                                      note='factor = SOI pension and employee-benefit deductions over pay plus employer FICA and FUTA, loaded '
                                           'on the US direct pay and on the Canadian direct pay drawn from labour and commissions (salaries and '
                                           'wages); Canadian cost-of-sales wages already include benefits (ISED glossary)'),
                   medical_materials=dict(**{c: dict(census_share=_mm[c]['census_share_all_sizes'], small_cell_ratio=_mm[c]['small_cell_ratio'],
                                                     share_used=_mm[c]['share_used'], round2_deducted=_mm[c]['bridge_materials_small_cell_round2'])
                                             for c in MAT},
                                          correction_at_1M=r3(page_value(KNOTS, MATERIALS['medical_correction'], 6.0)),
                                          se_at_1M=r3(float(np.interp(6.0, KNOTS, MATERIALS['medical_se_term']))),
                                          note='Census SAS 2022 supplies, other materials and purchased professional and technical services '
                                               '(6211, 6212) over Economic Census 2022 receipts, times the SOI $1-500K to all-size cost-of-goods '
                                               'ratio; replaces tax COGS less its labour where larger'),
                   recency=dict(**{i: r3(RECENCY[i]) for i in RECENCY_IND},
                                note='half the all-size SOI drift 2004-13 to 2019-22 on the labour-loaded definition, a level SE term for the '
                                     'industries on US 2004-13 cells alone; the small asset classes\' own 2004-13 trends carried to 2024 '
                                     'fall inside it'),
                   us_vs_ca_at_1M=dict(n=USCA['fixed_kept']['n'], within_004=USCA['fixed_kept']['within_004'], within_007=USCA['fixed_kept']['within_007'],
                                       median_abs_gap=r3(USCA['fixed_kept']['median_abs']), round2_within_004=USCA['round2_same']['within_004'],
                                       round2_median_abs_gap=r3(USCA['round2_same']['median_abs']), round1_within_004=USCA['round1_same']['within_004'],
                                       round1_median_abs_gap=r3(USCA['round1_same']['median_abs']), as_reported_within_004=USCA['as_reported_same']['within_004']),
                   size_slope=dict(r=r3(SLOPE['convention']['r']), mean_us_pred=r3(SLOPE['convention']['mean_pred']),
                                   mean_ca=r3(SLOPE['convention']['mean_actual']), n=SLOPE['convention']['n']),
                   ca_above_5M=dict(n=CA_BIG['n'], mean_diff=r3(CA_BIG['mean_diff']), within_004=CA_BIG['within_004']),
                   owner_pay=dict(ca_dropped=CA_DROPPED, effect_at_1M={r['industry']: r3(page_value(KNOTS, r['effect'], 6.0)) for r in owner_rows}),
                   median_gap=dict(rule=MEDIAN_RULE, small_positive=GAP_SIGNS['1a']['positive'], small_n=GAP_SIGNS['1a']['n'],
                                   medium_positive=GAP_SIGNS['1b']['positive'], medium_n=GAP_SIGNS['1b']['n'], cross_size_correlation=MEDIAN_RHO),
                   ada_dental=dict(gm_lower_bound=ADA_CHECK['gm_lower_bound'], dent=ADA_CHECK['dent_component_at_mean_billings'], passes=ADA_CHECK['passes'],
                                   shares_published=ADA_CHECK['expense_shares_published'],
                                   sas_2022_materials_share=ADA_CHECK['sas_2022_materials']['census_share_all_sizes'],
                                   sas_share_used=ADA_CHECK['sas_2022_materials']['share_used'],
                                   round2_deducted=ADA_CHECK['sas_2022_materials']['bridge_materials_round2']),
                   dealledger={(r['industry'] + ('_food_cost' if r['measure'].startswith('1 - food') else '')): dict(n=r['n'], stated=r['median'],
                               shipped=r.get('shipped_at_listing_revenue')) for r in DLC['rows'] if r['industry'] in SHIPPED},
                   definition_test_1M=dict(caveat='context only: the hand-set values mix definitions, so agreement with them is not '
                                                  'evidence for any definition and is not a test of the fit',
                                           by_definition={k.split('_')[0]: dict(mad=v['mad'], r=v['r']) for k, v in deftest.items() if k.endswith('_1M')})),
    'source': ('IRS SOI Corporation Source Book 2004-13 (asset classes $1-25M); Statistics Canada / ISED Financial Performance Data 2024 '
               '(revenue quartiles CAD 30k-5M, all and by incorporation status); BLS OEWS May 2025 direct-occupation wage shares; SOI '
               'Table 12 labour share of COGS; SOI pension and benefit deductions (payroll costs); SOI officers\' pay for the Canadian '
               'owner-pay removal; Census SAS 2022 and Economic Census 2022 (medical materials); KBCM 2021 SaaS survey; Ratchet+Wrench '
               '2026 (auto); Empire Flippers sold listings (ecom); revenue in mid-2026 USD. Checks: Census 2022, DealLedger, BACH (ECCBSO), '
               'ADA HPI, SOI Table 5.1 2019-22.'),
    'method': ('Per country, the industry margin on the stated definition by size cell: US asset classes smoothed by a local-linear fit in '
               'log10 revenue, CA revenue quartiles joined linearly; past its cells (+0.25 decade) a country follows the other\'s shape. '
               'Canadian owner-managers\' salaries are not loaded (SOI officers\' share of receipts x the incorporated share); where that '
               'pay is half or more of the Canadian labour bill (physicians, dentists, lawyers, accountants) the US cells are used alone. '
               'Direct pay carries its payroll costs (benefits and employer taxes) for the services and auto; physicians and dentists '
               'deduct the Census 2022 supplies and purchased services, size-adjusted. Knot = mean of the two countries (ratio of sums; '
               'BACH shows no consistent median-minus-mean gap, so it is carried in the SE). saas: KBCM total margin added by inverse '
               'variance at $14M; auto: Ratchet+Wrench at $845k on the same basis; the other knots move with the bridge\'s error '
               'covariance. ecom: median of Empire Flippers sold listings, flat. SE = time, source, shape, bridge, proxy, blend, owner-pay, '
               'median, payroll-cost, materials and recency terms in quadrature. The page interpolates linearly in log10 revenue and holds '
               'the end values beyond 5.5 and 7.0.'),
}
JS_STR = json.dumps(js, separators=(',', ':'), ensure_ascii=False, allow_nan=False)
with open(os.path.join(OUT, 'peer_gm.json'), 'w', encoding='utf-8') as f:
    f.write(JS_STR + '\n')

# ---- tables
krows = []
for ind in IND_ORDER:
    if ind == 'ecom':
        for i, k in enumerate(KNOTS):
            krows.append(dict(industry=ind, shipped=True, definition='ecom', knot=k, gm=r4(FINAL['ecom']['est'][i]), se=r4(FINAL['ecom']['se'][i]),
                              hand_set=HAND[ind], note='Empire Flippers median, flat'))
        continue
    F = FITS[ind]
    fin = FINAL.get(ind)
    for i, k in enumerate(KNOTS):
        rec = dict(industry=ind, shipped=ind in SHIPPED, definition=DEFN.get(ind, 'dir'), knot=k,
                   gm=r4(fin['est'][i] if fin else F['est'][i]), se=r4(fin['se'][i] if fin else F['se'][i]),
                   bridge_gm=r4(F['est'][i]), bridge_se=r4(F['se'][i]), us=r4(F['us'][i]), ca=r4(F['ca'][i]), hand_set=HAND[ind],
                   modes=';'.join(f'{c}:{m[i]}' for c, m in F['modes'].items()), carry=';'.join(f'{c}:{d[i]}' for c, d in F['carry_distance'].items()))
        for t in SE_TERMS:
            rec['se_' + t] = r4(F['se_terms'][t][i])
        krows.append(rec)
pd.DataFrame(krows).to_csv(os.path.join(OUT, 'knots.csv'), index=False)
cell_rows = []
for comp in COMP:
    for defn in ['conv', 'rep', 'dir', 'exl', 'allw']:
        d_ = {'rest': 'exl', 'dist': 'rep', 'ecom': 'rep'}.get(comp, 'dir') if defn == 'conv' else defn
        pc_, mat_ = adj(comp, d_)                                            # the round-3 corrections where they apply
        U = us_curve(comp, d_, delta(comp), LAM[COMP[comp]['sec']], pc=pc_, mat=mat_)
        for x, y, w in zip(U.x, U.y, U.w):
            cell_rows.append(dict(component=comp, source='US SOI', definition=defn, x_real=round(x, 4), gm=round(y, 4), sqrt_returns=round(w, 1)))
        for owner in ('none', OWNER_MODE):
            C = ca_curve(comp, d_, delta(comp), owner=owner, pc=pc_, mat=mat_)
            if C is not None:
                for x, y in zip(C.x, C.y):
                    cell_rows.append(dict(component=comp, source='CA ISED' + (' owner pay removed' if owner != 'none' else ''), definition=defn,
                                          x_real=round(x, 4), gm=round(y, 4), sqrt_returns=None, ca_dropped=comp in CA_DROPPED))
pd.DataFrame(cell_rows).to_csv(os.path.join(OUT, 'cells.csv'), index=False)
pd.DataFrame([dict(key=k, **v) for k, v in OEWS.items()]).to_csv(os.path.join(OUT, 'oews_direct_share.csv'), index=False)
pd.DataFrame(census_rows).to_csv(os.path.join(OUT, 'census_check.csv'), index=False)
pd.DataFrame([{k: (json.dumps(v) if isinstance(v, (dict, list)) else v) for k, v in r.items()} for r in usca]).to_csv(os.path.join(OUT, 'us_vs_ca.csv'), index=False)
pd.DataFrame(SLOPE['convention']['rows']).to_csv(os.path.join(OUT, 'slope_test.csv'), index=False)
pd.DataFrame(DLC['rows']).to_csv(os.path.join(OUT, 'dealledger_check.csv'), index=False)
pd.DataFrame(deftab).to_csv(os.path.join(OUT, 'definition_test.csv'), index=False)
pd.DataFrame(spread_rows).to_csv(os.path.join(OUT, 'spread_bach_derived.csv'), index=False)
pd.DataFrame(rep_rows).to_csv(os.path.join(OUT, 'research_replication.csv'), index=False)
pd.DataFrame(recency).to_csv(os.path.join(OUT, 'recency.csv'), index=False)
pd.DataFrame([{k: (json.dumps(v) if isinstance(v, (dict, list)) else v) for k, v in r.items()} for r in owner_rows]).to_csv(os.path.join(OUT, 'owner_pay.csv'), index=False)
pd.DataFrame([dict(component=c, kappa=KAPPA[c], kappa_raw=KAPPA_RAW[c], **{f'pi_q{i + 1}': PI_INC[COMP[c]['ca']][i] for i in range(4)},
                   **{f'owner_share_labour_{q}': OWNER_LAB[c][q] for q in ('q1', 'q2', 'q3', 'q4')}, ca_dropped=c in CA_DROPPED) for c in KAPPA]).round(4).to_csv(os.path.join(OUT, 'owner_pay_inputs.csv'), index=False)
pd.DataFrame([{k: (json.dumps(v) if isinstance(v, (dict, list)) else v) for k, v in r.items()} for r in median_rows]).to_csv(os.path.join(OUT, 'median_gap_bach_derived.csv'), index=False)
pd.DataFrame([{k: (json.dumps(v) if isinstance(v, (dict, list)) else v) for k, v in r.items()} for r in ca_big]).to_csv(os.path.join(OUT, 'ca_5_20M_check.csv'), index=False)
ef[['listing_number', 'monetization_primary', 'cogs_percent', 'tacos_percent', 'average_annual_gross_revenue', 'year', 'gm', 'x']].round(4).to_csv(os.path.join(OUT, 'ef_listings_used.csv'), index=False)
pd.DataFrame([{k: (json.dumps(v) if isinstance(v, (dict, list)) else v) for k, v in r.items()} for r in PAYROLL_SENS]).to_csv(os.path.join(OUT, 'payroll_cut.csv'), index=False)
pd.DataFrame([{k: (json.dumps(v) if isinstance(v, (dict, list)) else v) for k, v in r.items()} for r in mat_rows]).to_csv(os.path.join(OUT, 'medical_materials.csv'), index=False)

# ---- the change from round 2, by industry and knot, with its parts (round 2's shipped table read from its RESULT)
PREV = json.load(open(PREV_RESULT, encoding='utf-8')) if os.path.exists(PREV_RESULT) else None
_z = lambda v: None if v is None else v + 0.0                 # no signed zeros in the change table
change_rows, CHANGE_1M = [], {}
_mfg_d_effect = np.asarray(ind_curve('mfg')['est']) - np.asarray(ind_curve('mfg', d_override=OEWS['mfg_misc']['direct_wage_share'])['est'])
for ind in SHIPPED:
    if PREV is None:
        break
    pvs = PREV['shipped_values'][ind]
    r2_gm, r2_se = np.array(pvs['gm_unrounded'], float), np.array(pvs['se_unrounded'], float)
    r3_gm, r3_se = FINAL[ind]['est'], FINAL[ind]['se']
    parts = dict(payroll=np.zeros(NK), materials=np.zeros(NK), mfg_direct_share=np.zeros(NK), survey=np.zeros(NK))
    ps = next((r for r in PAYROLL_SENS if r['industry'] == ind), None)
    if ps is not None and ps['applied']:
        parts['payroll'] = -np.array(ps['cut'], float)
    if ind == 'medical':
        parts['materials'] = np.array(MATERIALS['medical_correction'], float)
    if ind == 'mfg':
        parts['mfg_direct_share'] = _mfg_d_effect
    if ind in ('saas', 'auto'):
        s3 = FINAL[ind]['est'] - FITS[ind]['est'][:NK]
        s2 = r2_gm - np.array(pvs['bridge_gm'], float) if 'bridge_gm' in pvs else np.zeros(NK)
        parts['survey'] = s3 - s2
    other = (r3_gm - r2_gm) - sum(parts.values())
    for i, k in enumerate(KNOTS):
        change_rows.append(dict(industry=ind, knot=k, round2_gm=r4(r2_gm[i]), round2_se=r4(r2_se[i]), round3_gm=r4(r3_gm[i]), round3_se=r4(r3_se[i]),
                                change=_z(r4(r3_gm[i] - r2_gm[i])), **{p: _z(r4(v[i])) for p, v in parts.items()}, other=_z(r4(other[i]))))
    CHANGE_1M[ind] = dict(round2=r3(page_value(KNOTS, r2_gm, 6.0)), round3=r3(page_value(KNOTS, r3_gm, 6.0)),
                          change=_z(r3(page_value(KNOTS, r3_gm - r2_gm, 6.0))), **{p: _z(r3(page_value(KNOTS, v, 6.0))) for p, v in parts.items()},
                          other=_z(r3(page_value(KNOTS, other, 6.0))),
                          se_round2=r3(float(np.interp(6.0, KNOTS, r2_se))), se_round3=r3(float(np.interp(6.0, KNOTS, r3_se))))
if change_rows:
    pd.DataFrame(change_rows).to_csv(os.path.join(OUT, 'change_from_round2.csv'), index=False)

# ---- sources
SOURCES = [
    dict(name='IRS SOI Corporation Source Book flat files, Table 1 (sb1.csv)', url='https://www.irs.gov/statistics/soi-tax-stats-corporation-source-book-data-file',
         files=[dict(file=p, sha256=sha(p)) for p in SOI_ZIPS], years='2004-2013', n=dict(cells_used=int(len(SOI_CELLS)), codes=soi_codes)),
    dict(name='IRS SOI Corporation Complete Report Table 12 (cost of goods sold components) and Table 5.1', url='https://www.irs.gov/statistics/soi-tax-stats-corporation-complete-report',
         files=[dict(file=p, sha256=sha(p)) for p in T12_FILES + sorted(glob.glob(os.path.join(SOI_CORP_DIR, '*co51ccr.xlsx')))], years='2019-2022', n=dict(sectors=len(LAM))),
    dict(name='Statistics Canada / ISED Financial Performance Data (revenue-quartile pages, all businesses, CAD 30k-5M)',
         url='https://ised-isde.canada.ca/app/ixb/cis/performance/rev/<NAICS>', files=[dict(file=p, sha256=sha(p)) for p in ISED_FILES],
         years='2024', n=dict(revenue_quartile_pages=len(ISED_REV), profit_quartile_pages=len(ISED_PM),
                              businesses={COMP[c]['ca']: ISED_REV[COMP[c]['ca']]['n_businesses'] for c in COMP if COMP[c]['ca']})),
    dict(name='Statistics Canada / ISED Financial Performance Data reports by incorporation status (all, incorporated, unincorporated; '
              'CAD 30k-5M) and incorporated CAD 5-20M (new download, round 2)',
         url='https://ised-isde.canada.ca/app/ixb/fpd-dpf/profilestart?lang=eng (form POST to /app/ixb/fpd-dpf/start: reportCriteria.businessType 3/2/1, '
             'revenueRange 0/1, distributedBy 0, expenseUnit 1, area 00, searchResultsSelectCode <NAICS>)',
         files=[dict(file=p, sha256=sha(p)) for p in FPD_FILES], years='2024',
         n=dict(reports=len(FPD_FILES), codes=len(CA_CODES), incorporated={c: int(FPD[(c, 2, 0)]['n']) for c in CA_CODES},
                unincorporated={c: int(FPD[(c, 1, 0)]['n']) for c in CA_CODES}, over_5M={c: int(FPD[(c, 2, 1)]['n']) for c in CA_CODES})),
    dict(name='BLS OEWS May 2025 national industry-specific estimates (nat4d, nat5d_6d, natsector)', url='https://www.bls.gov/oes/special-requests/oesm25in4.zip',
         files=[dict(file=OEWS_ZIP, sha256=sha(OEWS_ZIP))], years='May 2025', n=dict(industries=len(OEWS))),
    dict(name='Census 2022 Economic Census, construction (EC2223BASIC) and manufacturing (EC2231BASIC); US rows of EC2200BASIC (car washes; '
              'physicians and dentists for the materials shares)',
         url='https://www2.census.gov/programs-surveys/economic-census/data/2022/',
         files=[dict(file=EC_CONSTR, sha256=sha(EC_CONSTR)), dict(file=EC_MFG, sha256=sha(EC_MFG)),
                dict(file=EC_US, sha256=sha(EC_US), note='US rows extracted from sector00/EC2200BASIC.zip by the research round')], years='2022',
         n=dict(firms={**{k: v['firms'] for k, v in CENSUS.items()}, '811192': CW_CENSUS['firms'], **{MAT[c]['naics']: MAT[c]['firms'] for c in MAT}})),
    dict(name='Census Service Annual Survey 2022, Table 5: estimated selected expenses for employer firms (6211, 6212 used)',
         url='https://www2.census.gov/programs-surveys/sas/data/2022/ (Table5.xlsx, downloaded by the research round)',
         files=[dict(file=SAS5, sha256=sha(SAS5))], years='2022 estimates', n=dict(items_used=sorted(SAS_ITEMS.values()))),
    dict(name='ISED Financial Performance Data glossary (line definitions: cost-of-sales wages include benefits; labour and commissions are '
              'salaries and wages) (new download, round 3)',
         url='https://ised-isde.canada.ca/site/financial-performance-data/en/glossary', files=[dict(file=ISED_GLOSSARY, sha256=sha(ISED_GLOSSARY))],
         years='page modified 2026-07-27', n=dict(definitions_confirmed=GLOSSARY)),
    dict(name='Round 2 outputs of this fit (read for the change table only)', url='local', files=[dict(file=PREV_RESULT, sha256=sha(PREV_RESULT))] if PREV else [],
         years='2026-10-14', n=None),
    dict(name='Census AWTS 2022 Table 4 and ARTS 2022 gross margin % of sales', url='https://www2.census.gov/programs-surveys/awts/ and /arts/',
         files=[dict(file=AWTS, sha256=sha(AWTS)), dict(file=ARTS, sha256=sha(ARTS))], years='2018-2022', n=dict(cells=2)),
    dict(name='BACH database (ECCBSO; countries: %s) - derived statistics only, no redistribution' % ' '.join(BACH_COUNTRIES),
         url='https://www.bach.banque-france.fr/bach-ws-fo/ws/log-documentation/downloadocs?file_name=bach.zip',
         files=[dict(file=BACH_ZIP, sha256=sha(BACH_ZIP))], years='2019-2023', n=dict(rows_scanned=BACH_ROWS, countries=len(BACH_COUNTRIES))),
    dict(name='Ratchet+Wrench 2026 Industry Survey (auto repair; pp. 2, 5, 9, 10)',
         url='https://img.ratchetandwrench.com/files/base/ebm/ratchetandwrench/document/2026/06/6a4277e469691969c387beca-2026_rwindustry_survey_fnl.pdf',
         files=[dict(file=RW_PDF, sha256=sha(RW_PDF))], years='survey Q1 2026 (2025 revenue)', n=dict(responses='more than 430', tracking_gp=RW_N)),
    dict(name='2021 KBCM (KeyBanc) Private SaaS Company Survey (pp. 51-53, 56-57)', url='https://www.forentrepreneurs.com/wp-content/uploads/2021/10/2021-KBCM-SaaS-Survey.pdf',
         files=[dict(file=KBCM_PDF, sha256=sha(KBCM_PDF))], years='2020 data', n=dict(total_gm=KBCM['total_n'], total_gm_5_25M=KB_N, subscription_gm=KBCM['sub_n'])),
    dict(name='Empire Flippers marketplace API, sold listings (cogs_percent, tacos_percent)', url='https://empireflippers.com/marketplace-api/',
         files=[dict(file=EF_CSV, sha256=sha(EF_CSV)), dict(file=EF_DOC, sha256=sha(EF_DOC), note="EF listing checklist: COGS = units sold x landed cost (https://empireflippers.com/amazing/)")],
         years='listed 2021-2026', n=dict(sold_with_cogs_product_sellers=EF_N)),
    dict(name='DealLedger broker listings (CC0)', url='https://github.com/jeffsosville/dealledger',
         files=[dict(file=DL_SNAPS + '/*/listings.csv', sha256=DLC['sha256'], note='sha256 of the snapshot files concatenated in date order')],
         years='2026-02 to 2026-10', n=dict(snapshots=DLC['snapshots'], listing_ids=DLC['n_listing_ids'], after_dedupe=DLC['n_after_dedupe'],
                                           with_margin_figure=DLC['n_with_figure'])),
    dict(name='ADA Health Policy Institute, Survey of Dental Practice 2025, free tables (Tables 5, 13)',
         url='https://www.ada.org/-/media/project/ada-organization/ada/ada-org/files/resources/research/hpi/hpidata_sdpi_2025.xlsx',
         files=[dict(file=ADA_XLSX, sha256=sha(ADA_XLSX))], years='2025', n=dict(gross_billings=ADA['Table 5']['n'], expenses=ADA['Table 13']['n'])),
    dict(name='Professional Carwashing & Detailing, article quoting its 2019 Industry Report (car-wash cost shares)',
         url='https://www.carwash.com/winning-combinations-chemicals-equipment/', files=[dict(file=PCD_DOC, sha256=sha(PCD_DOC))], years='2019 survey', n=None),
    dict(name='FenderBender (Ratchet+Wrench\'s sister title): shop gross profit by job costing', url='https://www.fenderbender.com/running-a-shop/operations/article/33026361/basic-financial-knowledge-allows-shop-owners-to-measure-improve-performance',
         files=[dict(file=FB_DOC, sha256=sha(FB_DOC))], years='2020', n=None),
    dict(name='FRED CPI-U (CPIAUCNS, monthly) and Canada/US exchange rate (AEXCAUS, annual)', url='https://fred.stlouisfed.org/series/CPIAUCNS ; /AEXCAUS',
         files=[dict(file=CPI_CSV, sha256=sha(CPI_CSV)), dict(file=FX_CSV, sha256=sha(FX_CSV))], years='2004-2026',
         n=dict(cpi_ref=round(CPI_REF, 3), cpi_ref_months=list(CPI_REF_SPAN), months_with_data=CPI_REF_MONTHS, cad_per_usd_2024=FX)),
    dict(name='Console snapshot index.html (hand-set IND values; BIZCOMPS NAICS-6 counts for the blends; the z.gm scale)', url='local, read only',
         files=[dict(file=INDEX, sha256=sha(INDEX))], years='HEAD 56ba230', n=dict(industries=len(HAND))),
]

# ---- shipped values against the console's current values
shipped_values = {}
for ind in SHIPPED:
    fin = FINAL[ind]
    rec = dict(label=IND_LABEL[ind], definition=DEF_TEXT[DEF_KEY[ind]], definition_key=DEF_KEY[ind], hand_set=HAND[ind], tier=TIER[ind],
               gm=SHIP[ind]['gm'], se=SHIP[ind]['se'], slope_per_decade=SHIP[ind]['slope'], slope_se=SHIP[ind]['slope_se'],
               gm_unrounded=[r4(v) for v in fin['est']], se_unrounded=[r4(v) for v in fin['se']], cov_knots=[[float('%.6g' % v) for v in row] for row in fin['cov']],
               z_hand_vs_knot=[r3((HAND[ind] - g) / s_) for g, s_ in zip(fin['est'], fin['se'])],
               at_1M=r3(page_value(KNOTS, fin['est'], REF_X['1M'])), at_3M=r3(page_value(KNOTS, fin['est'], REF_X['3M'])),
               z_hand_vs_1M=r3((HAND[ind] - page_value(KNOTS, fin['est'], REF_X['1M'])) / float(np.interp(REF_X['1M'], KNOTS, fin['se']))))
    if ind == 'ecom':
        rec.update(source='Empire Flippers sold listings', se_terms=dict(sampling=r4(EF_SE_SAMPLING), mix=r4(EF_SE_MIX), time=r4(EF_SE_TIME),
                                                                        definition=EF_SE_DEFN, beyond_data=[r4(max(0.0, k - EF_X90) * EF_SLOPE_SE) for k in KNOTS]),
                   round2_shipped=None if PREV is None else dict(gm=PREV['shipped_values'][ind]['gm_unrounded'], se=PREV['shipped_values'][ind]['se_unrounded']))
    else:
        F = FITS[ind]
        rec.update(source='bridge' + {'saas': ' + KBCM 2021', 'auto': ' + Ratchet+Wrench 2026'}.get(ind, ''),
                   bridge_gm=[r4(v) for v in F['est'][:NK]], bridge_se=[r4(v) for v in F['se'][:NK]],
                   round1_estimator_gm=[r4(v) for v in ind_curve(ind, owner='none', payroll=False, materials=False,
                                                                 d_override=OEWS['mfg_misc']['direct_wage_share'] if ind == 'mfg' else None)['est']],
                   round2_shipped=None if PREV is None else dict(gm=PREV['shipped_values'][ind]['gm_unrounded'], se=PREV['shipped_values'][ind]['se_unrounded']),
                   se_terms={t: [r4(v) for v in F['se_terms'][t][:NK]] for t in SE_TERMS},
                   se_terms_note='the bridge\'s terms' + (' before the survey update (se is after it)' if ind in ('saas', 'auto') else ''), us=[r4(v) for v in F['us'][:NK]], ca=[r4(v) for v in F['ca'][:NK]],
                   modes=F['modes'], carry_distance=F['carry_distance'],
                   components={c: dict(cells=COMP[c], weight=v['weight'], direct_share=r4(delta(c)), lambda_sector=r4(LAM[COMP[c]['sec']]), ca_dropped=v['ca_dropped'],
                                       est=[r4(e) for e in v['est'][:NK]], us=[r4(e) for e in v['us'][:NK]], ca=[r4(e) for e in v['ca'][:NK]],
                                       us_cells_x=v['us_cells_x'], us_cells_gm=v['us_cells_gm'], ca_cells_x=v['ca_cells_x'], ca_cells_gm=v['ca_cells_gm'])
                               for c, v in F['components'].items()},
                   nominal_units_sensitivity=NOMINAL[ind])
    shipped_values[ind] = rec

SURVEYS = dict(
    ratchetwrench_auto=dict(SURVEYS_JS['auto'], quotes=RW_QUOTES, grouped_median=r4(RW_MED), normal_fit=dict(mu=r4(RW_MU), sd=r4(RW_SD)),
                            se=dict(sampling=r4(RW_SE_SAMPLING), binning=r4(RW_SE_BINNING), total=r4(RW_SE), population_judgment=RW_SE_POPULATION,
                                    payroll_adjustment=r4(RW_PAY_SE), used_in_combination=r4(np.sqrt(RW_VAR))), tracking_share=RW_TRACK,
                            respondents_median_revenue_2026usd=round(10 ** RW_X), bins=RW['gp_bins'], revenue_bins=RW['rev_bins'], revenue_open_bins_closed_at=list(RW['rev_open']),
                            parts_gp_median=r4(RW_PARTS_MED), labour_gp_median=r4(RW_LABOR_MED),
                            bridge_at_median_revenue=r4(F_AUTO['est'][NK]), bridge_mixture_median=r4(F_AUTO['est'][NK] + RW_MIX), bridge_se=r4(np.sqrt(F_AUTO['cov'][NK, NK])),
                            payroll_cost_adjustment=dict(technician_pay_share=r4(RW_TECH_PAY), payroll_cost_factor=r4(RW_PAYROLL_FACTOR), gap=r4(RW_DEF_GAP),
                                                         se=r4(RW_PAY_SE), correlation_with_bridge_cut=1.0),
                            survey_value_at_x=r4(RW_LEVEL), survey_se_total=r4(np.sqrt(RW_VAR)), weight_on_survey=r4(AUTO_W),
                            gain_by_knot=[r3(v) for v in AUTO_GAIN[:NK]], z_before_combining=r3(_rw_z),
                            auto_before=dict(gm=[r4(v) for v in F_AUTO['est'][:NK]], se=[r4(v) for v in np.sqrt(np.diag(F_AUTO['cov']))[:NK]]),
                            auto_after=dict(gm=[r4(v) for v in AUTO_EST[:NK]], se=[r4(v) for v in np.sqrt(np.diag(AUTO_COV))[:NK]]),
                            prior=dict(xs=[r4(v) for v in F_AUTO['xs']], est=[float('%.8g' % v) for v in F_AUTO['est']],
                                       cov=[[float('%.8g' % v) for v in row] for row in F_AUTO['cov']], survey_value_at_x=float('%.8g' % RW_LEVEL),
                                       survey_var=float('%.8g' % RW_VAR), shared_cov=[float('%.8g' % v) for v in RW_C]),
                            job_costing_quote=FB_QUOTE, alternatives=AUTO_ALT,
                            definition_note='The survey prints only "overall gross profit margin". The trade\'s job costing (FenderBender, the '
                                            'sister title: "Labor - 60 percent (attained by paying techs on a 60/40 split)") counts parts, '
                                            'paint and materials, sublet and technician pay; payroll taxes and benefits of that pay are usually '
                                            'left in overhead. The US half of the bridge loaded SOI salaries without them as well, so both stood '
                                            'on the wages-only basis; round 3 moves both to the shipped basis (the bridge by its own payroll cut, '
                                            'the survey by technician pay x the payroll-cost factor at its size) and combines them by inverse '
                                            'variance, with the doubt about that cut shared (correlation 1).'),
    kbcm_saas=dict(SURVEYS_JS['saas'], quotes=KBCM_QUOTES, subscription=dict(grouped_median=r4(KB_SUB_MED), printed=KBCM['sub_median_printed'],
                                                                             normal_fit=dict(mu=r4(KB_MU), sd=r4(KB_SD)), n=KBCM['sub_n']),
                   total_all=dict(median=KBCM['total_median'], n=KBCM['total_n']), total_by_size=KBCM['size_total'], size_n=KBCM['size_n'],
                   cs_allocation=KBCM['cs_total'], se=dict(sampling=r4(KB_SE_SAMPLING), rounding=r4(KB_SE_ROUNDING), population_judgment=KB_SE_POPULATION, total=r4(KB_SE)),
                   median_revenue_2026usd=round(10 ** KB_X), bridge_at=r4(_saas_bridge_at), bridge_se_at=r4(np.sqrt(F_SAAS['cov'][NK, NK])), mixture_offset=r4(KB_MIX),
                   weight_on_survey=r4(SAAS_W), gain_by_knot=[r3(v) for v in SAAS_GAIN[:NK]], top_knot_only_alternative=SAAS_TOP_ONLY,
                   median_term_cross_size_correlation=MEDIAN_RHO, correlation_alternatives=SAAS_RHO_ALT,
                   prior=dict(xs=[r4(v) for v in F_SAAS['xs']], est=[float('%.8g' % v) for v in F_SAAS['est']], cov=[[float('%.8g' % v) for v in row] for row in F_SAAS['cov']],
                              survey_value_at_x=float('%.8g' % KB_LEVEL_AT_X), survey_var=float('%.8g' % KB_SE ** 2)),
                   saas_before=dict(gm=[r4(v) for v in F_SAAS['est'][:NK]], se=[r4(v) for v in np.sqrt(np.diag(F_SAAS['cov']))[:NK]]),
                   saas_after=dict(gm=[r4(v) for v in SAAS_EST[:NK]], se=[r4(v) for v in np.sqrt(np.diag(SAAS_COV))[:NK]]),
                   definition_note='GAAP total gross margin with customer support (and hosting) in COGS, payroll costs included, stock-based '
                                   'pay added back: the shipped definition for software. Population: private SaaS firms with $5M+ of 2020 '
                                   'revenue; the $5-25M group (median about $14M in 2026 dollars) overlaps only the top knot.'),
    dealledger_saas=SURVEYS_JS['saas_listings'],
    empireflippers_ecom=dict(SURVEYS_JS['ecom'], median=r4(EF_MED), se=dict(sampling=r4(EF_SE_SAMPLING), mix=r4(EF_SE_MIX), time=r4(EF_SE_TIME), definition=EF_SE_DEFN),
                             iqr=[r4(EF_Q1), r4(EF_Q3)], sd_raw=r4(EF_SD_RAW), slope=dict(per_decade=r4(EF_SLOPE), se=r4(EF_SLOPE_SE), method='median regression on log10 revenue (2026 USD)'),
                             by_type={m: dict(n=int((ef.monetization_primary == m).sum()), median=r4(float(np.median(ef.gm[ef.monetization_primary == m]))))
                                      for m in EF_PROD if (ef.monetization_primary == m).any()},
                             by_year={int(y): dict(n=int(r['count']), median=r4(r['median'])) for y, r in ef.groupby('year').gm.agg(['count', 'median']).iterrows()},
                             amazon_share=r3(EF_FBA_SHARE), median_half_amazon=r4(EF_MED_HALF), tacos_median=r4(EF_TACOS_MED),
                             unsold_listings=EF_OTHER, revenue_p10_p90=[round(10 ** float(np.percentile(ef.x, 10))), round(10 ** EF_X90)],
                             definition_quote=EF_DOC_QUOTE,
                             definition_note='EF asks sellers for a P&L "with Cost of Goods Sold (COGS) in accrual format (number of units sold X '
                                             'landed cost)"; marketplace fees, fulfilment and advertising (TACoS) are separate lines. Product '
                                             'sellers only (Amazon FBA/FBM, e-commerce, dropshipping, subscription box); sold listings.'),
)

# ---- what stays judgment (or is not used), with its evidence
_z_sd = {k: v['sd'] for k, v in SURVEY_SD.items() if v is not None}
NOT_SHIPPED = dict(
    wash=dict(parameter='IND.gm wash', current=HAND['wash'], kept=True, **WASH),
    z_scale=dict(parameter='z.gm scale', current=Z_GM, kept=True,
                 evidence=dict(bach_small=SPREAD['small_s0'], bach_medium=SPREAD['medium_s0'], ca_floor=SPREAD['ca_profit_quartile_floor'],
                               surveys=SURVEY_SD),
                 verdict='kept at %.2f: BACH within-division SD %.3f (small firms) to %.3f (medium); the direct surveys give %s'
                         % (Z_GM, SPREAD['small_s0']['median'], SPREAD['medium_s0']['median'],
                            ', '.join('%s %.3f' % (k, v) for k, v in _z_sd.items()))),
    median_shift=dict(parameter='median-minus-ratio-of-sums shift', applied=MEDIAN_RULE == 'shift', rule=MEDIAN_RULE, signs=GAP_SIGNS,
                      verdict='not applied as a shift: the convention-matched gap has no consistent sign across the %d divisions (small firms '
                              '%d positive, %d negative; medium %d positive, %d negative); carried as an SE term (per industry and knot in '
                              'checks.median_gap). The pure materials margin does lean positive for small firms (median %+.3f).'
                              % (GAP_SIGNS['1a']['n'], GAP_SIGNS['1a']['positive'], GAP_SIGNS['1a']['negative'], GAP_SIGNS['1b']['positive'],
                                 GAP_SIGNS['1b']['negative'], MEDIAN['materials_margin_small'])),
    auto_survey_alternatives=dict(parameter='IND.gm auto: other readings of Ratchet+Wrench', applied=False, alternatives=AUTO_ALT,
                                  verdict='shipped: bridge and survey combined on the shipped basis with the payroll-cost doubt shared; '
                                          'the alternatives read the survey as already payroll-loaded, treat that doubt as independent, '
                                          'or drop the survey'),
    saas_recency=dict(parameter='a recency SE term for saas', applied=False,
                      verdict='not added (the analyst named medical and prof): the all-size software drift on the labour-loaded definition is '
                              '%+.3f, but KBCM 2020 pins the top knot and the bridge sat %.1f SE from it there' % (
                                  next(r['drift_dir'] for r in recency if r['component'] == 'saas'),
                                  (KB_LEVEL - (_saas_bridge_at + KB_MIX)) / np.sqrt(F_SAAS['cov'][NK, NK] + KB_SE ** 2))),
    dental_expense_shares=dict(parameter='dental half of medical, direct test', available=ADA_CHECK['expense_shares_published'], check=ADA_CHECK,
                               verdict='the ADA free tables still give only a bound; the materials are now tested and corrected with Census SAS 2022 '
                                       '(checks.medical_materials)'),
    canada_above_5M=dict(parameter='Canadian cells above CAD 5M', used=False, check=dict(n=CA_BIG['n'], mean_diff=CA_BIG['mean_diff'],
                                                                                         median_abs_diff=CA_BIG['median_abs_diff'])),
)

# ---- check outcomes (one line each) and open issues, from this run's numbers
_c = CENSUS_CHECK
_owner1M = {r['industry']: page_value(KNOTS, r['effect'], 6.0) for r in owner_rows}
_cut1M = {r['industry']: page_value(KNOTS, r['cut'], 6.0) for r in PAYROLL_SENS}
_kb_z = (KB_LEVEL - (_saas_bridge_at + KB_MIX)) / np.sqrt(F_SAAS['cov'][NK, NK] + KB_SE ** 2)
_cr = {r['naics']: r for r in census_rows}
_at1 = lambda v: page_value(KNOTS, v, 6.0)
_se1 = lambda v: float(np.interp(6.0, KNOTS, v))
CHECK_OUTCOMES = [
    dict(check='Census 2022 direct-cost test of the shipped curves on the decided wording (with payroll costs), tolerance .04',
         outcome='pass %d of %d (%s)' % (_c['pass_with_payroll_costs'], _c['n'], ', '.join(_c['shipped_tests'])),
         detail='diffs %s; wages-only passes %d of %d (diffs %s). The misc-manufacturing component against 339950: %+.3f (%s); the shipped mfg '
                'curve at 339950\'s size: %+.3f; against machine shops 332710 (context): %+.3f'
                % ([r['diff_with_payroll_costs'] for r in census_rows if r['role'] == 'shipped'], _c['pass_wages_only'], _c['n'],
                   [r['diff_wages_only'] for r in census_rows if r['role'] == 'shipped'], _cr['339950']['diff_with_payroll_costs'],
                   'pass' if _cr['339950']['pass_with_payroll_costs'] else 'fail', _cr['339950']['shipped_mfg_diff_with_payroll_costs'],
                   _cr['332710']['diff_with_payroll_costs'])),
    dict(check='payroll costs loaded on direct pay (decision 1: medical, vet, msp, saas, prof, staff, land, fit, auto)',
         outcome='applied; cut at $1M %s' % ', '.join('%s %.3f' % (k, v) for k, v in _cut1M.items() if k in PAYROLL_IND),
         detail='factor (SOI benefits and pension over pay + FICA/FUTA) %s; the Canadian half is cut only on the direct pay drawn from labour and '
                'commissions (ISED glossary: cost-of-sales wages include benefits); half of each cut is an SE term. Not applied (Census-tested): %s'
                % ({c: round(PCF[c], 3) for c in PAYROLL_COMPS}, ', '.join('%s %.3f' % (k, v) for k, v in _cut1M.items() if k not in PAYROLL_IND))),
    dict(check='medical materials against Census SAS 2022 (decision 2)',
         outcome='corrected: medical %+.3f at $1M (materials SE %.3f)' % (_at1(MATERIALS['medical_correction']), _se1(MATERIALS['medical_se_term'])),
         detail='; '.join('%s: SAS supplies + materials + purchased services %.3f of receipts, x small-cell ratio %.2f = %.3f against %.3f deducted in round 2; '
                          'component %+.3f at $1M' % (r['component'], r['census_share_all_sizes'], r['small_cell_ratio'], r['share_used'],
                                                      r['bridge_materials_small_cell_round2'], _at1(r['correction'])) for r in mat_rows)),
    dict(check='recency SE term (decision 5)', outcome=', '.join('%s %.3f' % (i, RECENCY[i]) for i in RECENCY_IND),
         detail='; '.join('%s: all-size drift %+.3f (labour-loaded, 2004-13 to 2019-22), small classes\' own trend carried to 2024 %+.3f, inside the '
                          'term: %s' % (i, d['all_size_drift_labour_loaded'], d['small_class_trend_to_2024'], d['covers_trend'])
                          for i, d in RECENCY_DETAIL.items())),
    dict(check='US against CA at $1M on the shipped definition (industries whose Canadian cells are kept)',
         outcome='%d of %d within .04, %d within .07' % (USCA['fixed_kept']['within_004'], USCA['fixed_kept']['n'], USCA['fixed_kept']['within_007']),
         detail='median |gap| %.3f, mean US - CA %+.3f; round 2 %d and %d, median %.3f, mean %+.3f; before the owner-pay fix %d and %d; as reported %d within .04'
                % (USCA['fixed_kept']['median_abs'], USCA['fixed_kept']['mean_us_minus_ca'], USCA['round2_same']['within_004'], USCA['round2_same']['within_007'],
                   USCA['round2_same']['median_abs'], USCA['round2_same']['mean_us_minus_ca'], USCA['round1_same']['within_004'], USCA['round1_same']['within_007'],
                   USCA['as_reported_same']['within_004'])),
    dict(check='size slope across the two countries (CA Q3 to Q4)', outcome='r = %.2f over %d pairs' % (SLOPE['convention']['r'], SLOPE['convention']['n']),
         detail='mean change US %.3f against CA %.3f; RMS slope gap %.3f per decade (round 2 %.3f; same pairs before the owner-pay fix %.3f)'
                % (SLOPE['convention']['mean_pred'], SLOPE['convention']['mean_actual'], DB_RMS, SLOPE['convention_round2']['slope_gap_rms'],
                   SLOPE['convention_no_owner_fix_same_pairs']['slope_gap_rms'])),
    dict(check='Ratchet+Wrench 2026 (auto), combined on the shipped basis (decision 6)', outcome='consistent: %+.1f SE before combining; weight on the survey %.2f' % (_rw_z, AUTO_W),
         detail='survey median %.3f (n %d) less the payroll cut at its size %.3f = %.3f, against the bridge %.3f over the respondents\' revenue mix; '
                'auto at $1M %.3f (SE %.3f) -> %.3f (SE %.3f)' % (RW_MED, RW_N, RW_DEF_GAP, RW_MED - RW_DEF_GAP, _rw_bridge_mix,
                                                               _at1(F_AUTO['est'][:NK]), _se1(np.sqrt(np.diag(F_AUTO['cov']))[:NK]),
                                                               _at1(AUTO_EST[:NK]), _se1(np.sqrt(np.diag(AUTO_COV))[:NK]))),
    dict(check='KBCM 2021 (saas), combined; BACH median term %.1f correlated across the size groups (decision 3)' % MEDIAN_RHO,
         outcome='consistent: %+.1f SE before combining; weight on the survey %.2f' % (_kb_z, SAAS_W),
         detail='$5-25M median %.2f (n %d) against the bridge %.3f at $%.1fM; saas SEs %s (perfect correlation %s, none %s)'
                % (KB_LEVEL, KB_N, _saas_bridge_at + KB_MIX, 10 ** KB_X / 1e6, SHIP['saas']['se'], SAAS_RHO_ALT['rho_1.0']['se'], SAAS_RHO_ALT['rho_0.0']['se'])),
    dict(check='DealLedger stated margins (selected), saas and ecom', outcome='saas %.2f (n %d) vs %.3f shipped; ecom %.3f (n %d) vs %.3f' % (
        DL_SAAS['median'], DL_SAAS['n'], DL_SAAS['shipped_at_listing_revenue'], *(lambda r: (r['median'], r['n'], r['shipped_at_listing_revenue']))(next(r for r in DLC['rows'] if r['industry'] == 'ecom' and r['measure'].startswith('gross')))), detail='stated when favourable; a check on definitions'),
    dict(check='Empire Flippers robustness (ecom)', outcome='median %.3f, flat (slope %.3f, SE %.3f)' % (EF_MED, EF_SLOPE, EF_SLOPE_SE),
         detail='Amazon FBA %.3f, other product sellers %.3f, half-and-half %.3f; yearly medians %s; unsold listings %s' % (EF_MED_FBA, EF_MED_NONFBA, EF_MED_HALF, {int(y): r4(v) for y, v in _yr['median'].items()}, EF_OTHER)),
    dict(check='BACH median-minus-weighted-mean sign test', outcome='inconsistent: carried as an SE term',
         detail='small firms %d+ / %d-, medium %d+ / %d- over %d divisions' % (GAP_SIGNS['1a']['positive'], GAP_SIGNS['1a']['negative'], GAP_SIGNS['1b']['positive'], GAP_SIGNS['1b']['negative'], GAP_SIGNS['1a']['n'])),
    dict(check='z-scale .13 against measured within-industry SDs', outcome='kept (inside the BACH bracket %.3f-%.3f)' % tuple(SPREAD['bracket']),
         detail='surveys: %s' % {k: v for k, v in _z_sd.items()}),
    dict(check='Canadian cells above CAD 5M (incorporated) against the carried Canadian curve', outcome='mean diff %+.3f, %d of %d within .04' % (CA_BIG['mean_diff'], CA_BIG['within_004'], CA_BIG['n']),
         detail='largest: %s' % sorted(((r['component'], r['mean_diff']) for r in ca_big if not r['ca_dropped']), key=lambda t: -abs(t[1]))[:4]),
    dict(check='ADA dental lower bound (no free expense shares)', outcome='pass' if ADA_CHECK['passes'] else 'fail',
         detail='1 - total expenses/billings = %.3f; dentists component %.3f (round 2 %.3f)' % (ADA_CHECK['gm_lower_bound'], ADA_CHECK['dent_component_at_mean_billings'], ADA_CHECK['dent_component_round2'])),
    dict(check='materiality at $708k SDE (harness): value of a business entering today\'s hand-set gm', outcome=('direction %s' % ('as expected' if MATERIALITY.get('direction_ok') else 'NOT as expected')) if 'rows' in MATERIALITY else 'skipped',
         detail=('; '.join('%s %+.1f%%' % (r['ind'], 100 * r['pct']) for r in MATERIALITY['rows'])) if 'rows' in MATERIALITY else MATERIALITY.get('skipped')),
    dict(check='replication of the research round', outcome='max |diff| %.4f' % REPLICATION['max_abs_diff'], detail='its estimator and units, no owner-pay fix'),
    dict(check='quotes re-read from the PDFs and pages', outcome='all %s' % ('confirmed' if PDFTOTEXT else 'not run (no pdftotext)'),
         detail='Ratchet+Wrench %d, KBCM %d, EF landed cost %s, FenderBender job costing %s, PC&D %s, ISED glossary %s' % (
             sum(bool(v) for v in RW_QUOTES.values()), sum(bool(v) for v in KBCM_QUOTES.values()), EF_DOC_QUOTE, FB_QUOTE,
             all(v is not None for v in PCD.values()), all(GLOSSARY.values()))),
]
_pay_rng = [v for k, v in _cut1M.items() if k in PAYROLL_IND]
_mm_ph, _mm_de = _mm['phys'], _mm['dent']
OPEN_ISSUES = [
    'Payroll costs (decision 1): the cut at $1M runs %.3f-%.3f (%s), half of it carried as an SE term. The factor is the US one (SOI benefit and pension '
    'deductions over pay plus FICA and FUTA, %.2f-%.2f); it is applied to the Canadian direct pay drawn from labour and commissions as well, since '
    'ISED defines that line as salaries and wages, while its cost-of-sales wages already include benefits. Canadian statutory rates (CPP, EI) are '
    'close to FICA and FUTA; Canadian benefit levels were not measured.'
    % (min(_pay_rng), max(_pay_rng), ', '.join('%s %.3f' % (k, v) for k, v in _cut1M.items() if k in PAYROLL_IND),
       min(PCF[c] for c in PAYROLL_COMPS), max(PCF[c] for c in PAYROLL_COMPS)),
    'Medical materials (decision 2): physicians and dentists now deduct the Census SAS 2022 supplies, other materials and purchased professional '
    'and technical services (%.3f and %.3f of receipts, all employer firms), scaled by the SOI $1-500K-to-all-size cost-of-goods ratio (%.2f and '
    '%.2f) to %.3f and %.3f, in place of tax COGS less its labour where larger (round 2 deducted %.3f and %.3f in that cell). medical moves %+.3f at '
    '$1M; the materials SE term (%.3f at $1M) is the RMS of three alternative readings (purchased services left out: the Census item covers '
    'professional and technical services in general, such as legal, accounting and consulting, much of it overhead for physicians; the pooled '
    'small classes at the knots as the size adjustment; no size adjustment). '
    'Temporary and leased staff (%.1f%% and %.1f%% of receipts) are not deducted. The size ratio is from 2004-13 and the shares from 2022. One '
    'share is held across sizes, so the medical size slope flattens (%.3f per decade, round 2 %.3f): the SOI cost-of-goods share does rise '
    'from the $1-500K class to the next (physicians %.3f to %.3f) but swings %.3f-%.3f across years in that class, so it is not used as a gradient.'
    % (_mm_ph['census_share_all_sizes'], _mm_de['census_share_all_sizes'], _mm_ph['small_cell_ratio'], _mm_de['small_cell_ratio'],
       _mm_ph['share_used'], _mm_de['share_used'], _mm_ph['bridge_materials_small_cell_round2'], _mm_de['bridge_materials_small_cell_round2'],
       _at1(MATERIALS['medical_correction']), _se1(MATERIALS['medical_se_term']), 100 * MAT['phys']['shares']['temp_staff'], 100 * MAT['dent']['shares']['temp_staff'],
       SHIP['medical']['slope'], PREV['js_table']['by_industry']['medical']['slope'] if PREV else float('nan'),
       MAT['phys']['soi_cogs_small'], MAT_AC4['phys']['mean'], MAT_AC4['phys']['min'], MAT_AC4['phys']['max']),
    'saas (decision 3): the BACH median term is now %.1f correlated between the knots read off BACH\'s small class (5.5-6.5) and those that blend '
    'in the medium class (7.0 and KBCM\'s $%.1fM), the reviewer\'s reading; the same error model is used for every industry, which changes only '
    'the two survey updates and the slope SEs. saas SEs %s (perfect correlation %s; none %s); the top knot keeps the KBCM level (%.3f). '
    'No recency term is added for saas (the software drift on the labour-loaded definition is %+.3f, but KBCM 2020 pins the top knot).'
    % (MEDIAN_RHO, 10 ** KB_X / 1e6, SHIP['saas']['se'], SAAS_RHO_ALT['rho_1.0']['se'], SAAS_RHO_ALT['rho_0.0']['se'], SHIP['saas']['gm'][-1],
       next(r['drift_dir'] for r in recency if r['component'] == 'saas')),
    'Auto (decision 6): Ratchet+Wrench and the bridge now share the shipped basis and are combined by inverse variance (weight %.2f at $%.0fk): '
    'auto %s (SE %s), bridge alone %s (SE %s). The survey carries a %.2f judgment term for its population (self-selected trade-magazine '
    'respondents, an undefined margin read through the trade\'s job costing), as KBCM does; without it auto would be %s (SE %s). Reading the '
    'survey as already payroll-loaded would give %s; treating the payroll doubt as independent %s.'
    % (AUTO_W, 10 ** RW_X / 1e3, SHIP['auto']['gm'], SHIP['auto']['se'], AUTO_ALT['bridge_only']['gm'], AUTO_ALT['bridge_only']['se'], RW_SE_POPULATION,
       AUTO_ALT['no_population_term']['gm'], AUTO_ALT['no_population_term']['se'], AUTO_ALT['survey_read_as_payroll_loaded']['gm'],
       AUTO_ALT['payroll_doubt_independent']['gm']),
    'Recency (decision 5): medical %.3f and prof %.3f, half the all-size SOI drift on the labour-loaded definition (%s); the small classes\' own '
    '2004-13 trends carried to 2024 (%s) fall inside. For medical part of that drift is the rising cost-of-goods share that the Census materials '
    'now cover, so the term is conservative there.'
    % (RECENCY['medical'], RECENCY['prof'], {i: d['all_size_drift_labour_loaded'] for i, d in RECENCY_DETAIL.items()},
       {i: d['small_class_trend_to_2024'] for i, d in RECENCY_DETAIL.items()}),
    'mfg (decision 7): the all-manufacturing cells now load the all-manufacturing OEWS direct share (%.3f; round 2 used 3399\'s %.3f), which moves '
    'mfg %+.3f at $1M. Against Census 31-33 the shipped curve passes (%+.3f), but that test is weak: all manufacturing averages $%.0fM of receipts '
    'per firm, past the $10M knot the page holds flat. Machine shops 332710 (context, $%.1fM): %+.3f.'
    % (OEWS['mfg_all']['direct_wage_share'], OEWS['mfg_misc']['direct_wage_share'], _at1(_mfg_d_effect), _cr['31-33']['diff_with_payroll_costs'],
       CENSUS['31-33']['receipts_per_firm_2022usd'] / 1e6, CENSUS['332710']['receipts_per_firm_2022usd'] / 1e6, _cr['332710']['diff_with_payroll_costs']),
    'Car wash (decision 4): definition.wash is "%s" and the hand-set %.2f stays judgment on it; its closest proxy reading, materials only, is %.2f at $1M.'
    % (WASH_DEF, HAND['wash'], WASH['materials_only_at_1M']),
    'Owner pay: removing Canadian owner-managers\' pay (SOI officers\' share x incorporated share) moves the medians at $1M by %s. For physicians, '
    'dentists, lawyers and accountants that pay is %s of the Canadian labour bill at Q3/Q4, so their Canadian cells are dropped (US cells alone). '
    'ISED\'s own reading is a median %.1fx the SOI share (%.1f-%.1fx); half its change is carried as an SE term.'
    % (', '.join('%s %+.3f' % (k, v) for k, v in _owner1M.items()), {c: '%.2f/%.2f' % (OWNER_LAB[c]['q3'], OWNER_LAB[c]['q4']) for c in CA_DROPPED},
       float(np.median(list(KAPPA_RAW.values()))), min(KAPPA_RAW.values()), max(KAPPA_RAW.values())),
    'ecom moves from %.2f to %.3f only because the definition now counts the landed cost of goods alone; a typical e-commerce seller entered on the '
    'old basis would score z = %.1f. The slider must show definition.ecom. Empire Flippers is %.0f%% Amazon FBA; a half-Amazon mix gives %.3f (inside its SE).'
    % (HAND['ecom'], EF_MED, (HAND['ecom'] - EF_MED) / ef_se(6.0), 100 * EF_FBA_SHARE, EF_MED_HALF),
    'Canadian incorporated cells above CAD 5M (not used) sit %+.3f above the carried Canadian curve on average; msp, constr and fit would gain %s at 6.5/7.0 if they were used.'
    % (CA_BIG['mean_diff'], {r['component']: r['knot_effect_if_used'] for r in ca_big if r['component'] in ('msp', 'constr', 'fit')}),
    'Integration (page work, not done here): by_industry carries a tier per industry (%s); the slider text is definition[ind] || definition.default, '
    'and wash now has its own; checks.definition_test_1M is context only. Sites that read I.gm need the size-varying median at the business\'s revenue.'
    % ', '.join('%s %s' % (t, sorted(i for i in SHIPPED if TIER[i] == t)) for t in ('aggregate bridge', 'survey-combined', 'listing-level')),
]

RESULT = dict(
    name='peer_gross_margin', js_key='peer_gm', asof='2026-10-14', script='/tmp/claude-0/fit6/fitscripts/peer_gross_margin.py', revision=3,
    tier='mixed', tier_note='per industry in js_table.by_industry[ind].tier: %s; wash stays judgment' % '; '.join(
        '%s (%s): %s' % (t, TIER_TEXT[t], ', '.join(i for i in IND_ORDER if TIER.get(i) == t)) for t in ('aggregate bridge', 'survey-combined', 'listing-level')),
    js_table=js, js_table_compact=JS_STR, change_from_round2=dict(at_1M=CHANGE_1M, by_knot=change_rows),
    parameters=dict(knots=KNOTS, kernel_half_width=H, tau=TAU, min_years=MIN_YEARS, direct_share_sd=D_SD, lambda_rest_range=LAM_REST,
                    source_single=SRC_SINGLE, proxy=PROXY, owner_mode=OWNER_MODE, owner_drop=OWNER_DROP, kappa_cap=KAPPA_CAP,
                    slope_gap_rms=r4(DB_RMS), single_country_tilt=r4(SLOPE_SRC_MED), median_rule=MEDIAN_RULE, seed=SEED,
                    payroll_industries=PAYROLL_IND, payroll_cost_factor={c: r4(PCF[c]) for c in PAYROLL_COMPS}, payroll_se_share=PAYROLL_SE,
                    materials_small_asset_class=MAT_SMALL_AC, materials_share={c: r4(MAT[c]['share']) for c in MAT}, median_cross_size_rho=MEDIAN_RHO,
                    recency_industries=RECENCY_IND, recency_fraction=RECENCY_FRAC, recency_se={i: r4(v) for i, v in RECENCY.items()},
                    mfg_direct_share='31-33 all manufacturing (OEWS natsector)',
                    cpi_ref=round(CPI_REF, 3), cpi_ref_months=list(CPI_REF_SPAN), cpi_ref_months_with_data=CPI_REF_MONTHS,
                    cpi_annual={str(y): round(CPI_AVAIL[y], 3) for y in list(range(2004, 2014)) + [2020, 2022, 2024, 2025]}, cad_per_usd_2024=FX,
                    cad2024_to_usd2026=round(CA_TO_REAL, 4), blend_weights=dict(medical=W_MED, prof=W_PROF), lambda_by_sector={k: r4(v) for k, v in LAM.items()},
                    rounding='gm, se, slope and slope_se rounded to 3 decimals with Python round(); unrounded values in shipped_values'),
    shipped_values=shipped_values, surveys=SURVEYS, not_shipped=NOT_SHIPPED, check_outcomes=CHECK_OUTCOMES, open_issues=OPEN_ISSUES,
    oews_direct_share={k: dict(naics=v['naics'], title=v['title'], direct_wage_share=r4(v['direct_wage_share']), emp=v['emp_total'], rule=v['rule'])
                       for k, v in OEWS.items()},
    checks=dict(census_2022=CENSUS_CHECK, census_cells={k: {kk: (r4(vv) if isinstance(vv, float) else vv) for kk, vv in v.items()} for k, v in CENSUS.items()},
                materiality=MATERIALITY, payroll_costs_of_loaded_wages=PAYROLL_SENS, medical_materials=MATERIALS, owner_pay=OWNER, median_gap=MEDIAN,
                us_vs_ca=USCA, size_slope=SLOPE, ca_above_5M=CA_BIG, dealledger=DLC,
                definition_test=dict(results=deftest, hand_set=HAND, note='context only: the hand-set values mix definitions, so agreement with '
                                                                          'them is not evidence for any definition and not a test of the fit'),
                spread=SPREAD, ada_dental=ADA_CHECK, research_replication=REPLICATION, recency=dict(rows=recency, se_terms=RECENCY_DETAIL),
                census_all_firm=CENSUS_ALLFIRM,
                mfg_direct_share=dict(misc_3399=r4(OEWS['mfg_misc']['direct_wage_share']), all_31_33=r4(OEWS['mfg_all']['direct_wage_share']),
                                      shipped_uses='31-33 for the all-manufacturing cells (round 3); 3399 for the misc-manufacturing component',
                                      effect_by_knot=[r4(v) for v in _mfg_d_effect]),
                quotes=dict(ratchetwrench=RW_QUOTES, kbcm=KBCM_QUOTES, empireflippers_landed_cost=EF_DOC_QUOTE, fenderbender_job_costing=FB_QUOTE,
                            pcd=PCD, ised_glossary=GLOSSARY, pdftotext=bool(PDFTOTEXT))),
    sources=SOURCES,
)
with open(os.path.join(OUT, 'RESULT.json'), 'w', encoding='utf-8') as f:
    json.dump(RESULT, f, indent=1, ensure_ascii=False, allow_nan=False, default=lambda o: o.item() if hasattr(o, 'item') else str(o))

# ---------------------------------------------------------------- print
print('CPI ref %.3f (%s..%s, %d months), CAD/USD %.4f, CAD2024->USD2026 x%.4f' % (CPI_REF, CPI_REF_SPAN[0], CPI_REF_SPAN[1], CPI_REF_MONTHS, FX, CA_TO_REAL))
print('owner pay: CA dropped %s; kappa %s' % (CA_DROPPED, {c: round(v, 2) for c, v in KAPPA.items()}))
print('median gap rule: %s %s' % (MEDIAN_RULE, GAP_SIGNS))
print('%-8s %5s | %s | slope (se) | z(hand vs 1M)' % ('ind', 'hand', '  '.join('gm@%.1f (se)' % k for k in KNOTS)))
for ind in IND_ORDER:
    fin = FINAL.get(ind) or dict(est=FITS[ind]['est'], se=FITS[ind]['se'], slope=ols_slope(KNOTS, FITS[ind]['est'][:NK]), slope_se=np.nan)
    tag = '' if ind in SHIPPED else '  [not shipped]'
    print('%-8s %5.2f | %s | %+.3f (%.3f) | %+5.2f%s' % (ind, HAND[ind], '  '.join('%.3f (%.3f)' % (a, b) for a, b in zip(fin['est'][:NK], fin['se'][:NK])),
                                                       fin['slope'], fin['slope_se'], (HAND[ind] - page_value(KNOTS, fin['est'][:NK], 6.0)) / float(np.interp(6.0, KNOTS, fin['se'][:NK])), tag))
print('Census 2022 (with payroll costs / wages only):', [(r['naics'], r['fitted'], r['census_with_payroll_costs'], r['census_wages_only'], r['pass_with_payroll_costs'], r['pass_wages_only']) for r in census_rows])
print('US vs CA at $1M, CA kept (n %d): shipped %d within .04, %d within .07, median |gap| %.3f; round 2 %d, %d, %.3f; round 1 %d, %d, %.3f; widening %.3f / %.3f -> %.3f' % (
    USCA['fixed_kept']['n'], USCA['fixed_kept']['within_004'], USCA['fixed_kept']['within_007'], USCA['fixed_kept']['median_abs'],
    USCA['round2_same']['within_004'], USCA['round2_same']['within_007'], USCA['round2_same']['median_abs'],
    USCA['round1_same']['within_004'], USCA['round1_same']['within_007'], USCA['round1_same']['median_abs'], USCA['widening_round1_mean'],
    USCA['widening_round2_mean'], USCA['widening_fixed_mean']))
print('size slope r %.2f (n %d), US %.3f vs CA %.3f; RMS slope gap %.3f (round 2 %.3f)' % (SLOPE['convention']['r'], SLOPE['convention']['n'], SLOPE['convention']['mean_pred'], SLOPE['convention']['mean_actual'], DB_RMS, SLOPE['convention_round2']['slope_gap_rms']))
print('payroll cut at $1M: %s' % {k: round(v, 3) for k, v in _cut1M.items()})
print('medical materials: %s; medical correction %s, SE %s' % ({c: (round(MAT[c]['census_share'], 4), round(MAT[c]['small_ratio'], 3), round(MAT[c]['share'], 4)) for c in MAT},
                                                             MATERIALS['medical_correction'], MATERIALS['medical_se_term']))
print('recency: %s' % RECENCY_DETAIL)
print('R+W: median %.3f (n %d, sd %.3f, se %.3f) at $%.0fk; payroll gap %.3f -> %.3f vs bridge mixture %.3f (se %.3f); z %+.2f, weight %.2f' % (
    RW_MED, RW_N, RW_SD, RW_SE, 10 ** RW_X / 1e3, RW_DEF_GAP, RW_MED - RW_DEF_GAP, _rw_bridge_mix, np.sqrt(F_AUTO['cov'][NK, NK]), _rw_z, AUTO_W))
print('KBCM: %.2f (n %d, se %.3f) at $%.1fM vs bridge %.3f (se %.3f), weight %.2f; rho alternatives %s' % (KB_LEVEL, KB_N, KB_SE, 10 ** KB_X / 1e6, _saas_bridge_at + KB_MIX, np.sqrt(F_SAAS['cov'][NK, NK]), SAAS_W, SAAS_RHO_ALT))
print('change from round 2 at $1M: %s' % CHANGE_1M)
print('EF ecom: median %.3f (n %d, se %.3f, iqr-sd %.3f), slope %.3f (%.3f)' % (EF_MED, EF_N, EF_SE_SAMPLING, EF_SD_IQR, EF_SLOPE, EF_SLOPE_SE))
print('spread: BACH small %.3f medium %.3f; surveys %s' % (SPREAD['small_s0']['median'], SPREAD['medium_s0']['median'], _z_sd))
print('CA above $5M vs carried: mean diff %+.3f (n %d, %d within .04)' % (CA_BIG['mean_diff'], CA_BIG['n'], CA_BIG['within_004']))
print('ADA dental: total expense share %.3f -> gm >= %.3f; dent component %.3f (%s)' % (ADA_EXP_SHARE_MEAN, ADA_CHECK['gm_lower_bound'], _dent_est, ADA_CHECK['passes']))
print('research replication max |diff| %s' % REPLICATION['max_abs_diff'])
print('DealLedger:', [(r['industry'], r['n'], r['median'], r.get('shipped_at_listing_revenue')) for r in DLC['rows']])
for o in CHECK_OUTCOMES:
    print('check: %s -> %s' % (o['check'], o['outcome']))
print('materiality:', MATERIALITY.get('rows') and '; '.join('%s %+.1f%%' % (r['ind'], 100 * r['pct']) for r in MATERIALITY['rows']))
print('wrote', OUT + '/RESULT.json', 'and', OUT + '/peer_gm.json', '(%d bytes)' % len(JS_STR))
